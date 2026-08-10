# 注意：本脚本为历史/辅助生成器，权威提交文件以 01_技术文件 与 02_样题 中现有 docx 为准。
# 请勿用本脚本输出直接覆盖提交集权威文件（尤其 V1 技术文件与完整版样题）。
# 默认输出到 03_脚本与方案/_generated/（可复跑草稿目录）。
"""
江西省"振兴杯"职业技能竞赛 - 印制电路制作工样题生成脚本
"""
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml, OxmlElement

doc = Document()

# ============================================================
# 全局样式
# ============================================================
style = doc.styles['Normal']
style.font.name = '宋体'
style.font.size = Pt(12)
style.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
pf = style.paragraph_format
pf.line_spacing = Pt(28)
pf.space_before = Pt(0)
pf.space_after = Pt(0)

for section in doc.sections:
    section.top_margin = Cm(2.8)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(2.6)

# ============================================================
# 辅助函数
# ============================================================
def set_cell_shading(cell, color="D9D9D9"):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def add_heading1(text):
    p = doc.add_heading(text, level=1)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in p.runs:
        run.font.name = '黑体'
        run.font.size = Pt(16)
        run.element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = Pt(28)
    return p

def add_heading2(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    run.bold = True
    run.font.name = '黑体'
    run.font.size = Pt(14)
    run.element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = Pt(28)
    return p

def add_body(text, indent=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    run = p.add_run(text)
    run.font.name = '宋体'
    run.font.size = Pt(12)
    run.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    p.paragraph_format.line_spacing = Pt(28)
    if indent:
        p.paragraph_format.first_line_indent = Cm(0.74)
    return p

def add_body_bold_prefix(prefix, text, indent=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    run1 = p.add_run(prefix)
    run1.bold = True
    run1.font.name = '宋体'
    run1.font.size = Pt(12)
    run1.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    run2 = p.add_run(text)
    run2.font.name = '宋体'
    run2.font.size = Pt(12)
    run2.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    p.paragraph_format.line_spacing = Pt(28)
    if indent:
        p.paragraph_format.first_line_indent = Cm(0.74)
    return p

def add_table(headers, rows, col_widths=None):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ''
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(header)
        run.bold = True
        run.font.name = '黑体'
        run.font.size = Pt(10.5)
        run.element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
        set_cell_shading(cell, "D9D9D9")
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    for r_idx, row in enumerate(rows):
        for c_idx, val in enumerate(row):
            cell = table.rows[r_idx + 1].cells[c_idx]
            cell.text = ''
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(str(val))
            run.font.name = '宋体'
            run.font.size = Pt(10.5)
            run.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    if col_widths:
        for row in table.rows:
            for i, width in enumerate(col_widths):
                row.cells[i].width = Cm(width)
    return table

def add_page_break():
    doc.add_page_break()

def add_attachment_title(text):
    add_page_break()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.bold = True
    run.font.name = '黑体'
    run.font.size = Pt(16)
    run.element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
    p.paragraph_format.space_after = Pt(12)
    return p

# ============================================================
# 封面
# ============================================================
for _ in range(6):
    doc.add_paragraph()

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run('江西省"振兴杯"职业技能竞赛')
run.font.name = '黑体'
run.font.size = Pt(22)
run.element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
p.paragraph_format.line_spacing = Pt(36)

doc.add_paragraph()

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run("印制电路制作工")
run.bold = True
run.font.name = '黑体'
run.font.size = Pt(36)
run.element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
p.paragraph_format.line_spacing = Pt(48)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run("竞赛样题")
run.bold = True
run.font.name = '黑体'
run.font.size = Pt(28)
run.element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
p.paragraph_format.line_spacing = Pt(42)

doc.add_paragraph()

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run("工程资料处理与质量检测方向")
run.font.name = '黑体'
run.font.size = Pt(16)
run.element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
p.paragraph_format.line_spacing = Pt(28)

for _ in range(4):
    doc.add_paragraph()

info_lines = [
    ("主办单位：", "XXXX"),
    ("承办单位：", "XXXX"),
    ("技术支持：", "南云信息科技有限公司"),
    ("发布日期：", "2026年X月X日"),
]
for label, value in info_lines:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run1 = p.add_run(label)
    run1.font.name = '宋体'
    run1.font.size = Pt(14)
    run1.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    run2 = p.add_run(value)
    run2.font.name = '宋体'
    run2.font.size = Pt(14)
    run2.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    run2.underline = True
    p.paragraph_format.line_spacing = Pt(28)

# ============================================================
# 选手须知
# ============================================================
add_page_break()

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run("选手须知")
run.bold = True
run.font.name = '黑体'
run.font.size = Pt(22)
run.element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
p.paragraph_format.space_after = Pt(12)

add_body("1. 本样题仅供选手了解竞赛形式和内容范围，实际竞赛题目以现场发放为准。")
add_body("2. 竞赛总时长为4小时，总分100分，分为理论知识、PCB工程资料处理、DFM与工艺文件、裸板缺陷检测四个模块。")
add_body("3. 竞赛使用嘉立创EDA专业版客户端，由组委会统一安装，选手不得自行安装软件或连接外部设备。")
add_body("4. 缺陷检测环节所需工具（放大镜、万用表、游标卡尺等）由选手自带，组委会不统一提供。")
add_body("5. 选手不得携带手机、平板电脑、智能手表等具有通信、存储、联网功能的电子设备进入赛场。")
add_body("6. 作品须按规定目录结构提交，未按要求提交的相关模块按零分处理。")
add_body("7. 竞赛期间不得抄袭、协助他人或接受他人协助，违反者取消竞赛资格。")

# ============================================================
# 模块一：理论知识
# ============================================================
add_page_break()
add_heading1("模块一：理论知识（30分）")

add_body("本模块为闭卷考试，共30题，满分30分。考试时间由组委会统一安排。", indent=True)

add_heading2("一、单项选择题（每题1分，共15分）")

questions_single = [
    ("1. 根据《印制电路制作工国家职业技能标准》，印制电路制作工共设（    ）个等级。", "A. 3", "B. 4", "C. 5", "D. 6", "C"),
    ("2. PCB基材FR-4中的"FR"代表（    ）。", "A. 柔性", "B. 阻燃", "C. 高频", "D. 耐高温", "B"),
    ("3. 以下哪种文件格式是PCB Gerber文件的标准格式？（    ）", "A. DXF", "B. RS-274X", "C. STEP", "D. PDF", "B"),
    ("4. PCB制造中，化学沉铜的目的是（    ）。", "A. 提高导电性", "B. 在非金属化孔壁沉积铜层", "C. 防止氧化", "D. 提高焊接性", "B"),
    ("5. 以下哪种缺陷属于PCB裸板的致命缺陷？（    ）", "A. 丝印偏位", "B. 阻焊气泡", "C. 孔壁铜层断裂", "D. 板面划痕", "C"),
    ("6. DFM审查中，'环宽'指的是（    ）。", "A. 线宽", "B. 焊盘外径与孔径之差的一半", "C. 线距", "D. 板边距", "B"),
    ("7. NC Drill钻孔文件通常采用（    ）格式。", "A. Gerber", "B. DXF", "C. Excellon", "D. ODB++", "C"),
    ("8. PCB阻焊层的主要作用是（    ）。", "A. 美观", "B. 防止焊接短路和保护线路", "C. 提高导电性", "D. 增加板厚", "B"),
    ("9. 以下哪种工艺属于PCB外层图形转移？（    ）", "A. 内层蚀刻", "B. 外层贴膜曝光显影", "C. 钻孔", "D. 压合", "B"),
    ("10. 安全生产中，使用化学药品时应佩戴的基本防护用品不包括（    ）。", "A. 防护手套", "B. 护目镜", "C. 防毒面具", "D. 安全鞋", "D"),
    ("11. IPC-A-600标准是关于（    ）的验收标准。", "A. PCB设计", "B. PCB裸板可接受性", "C. 焊接质量", "D. 组装工艺", "B"),
    ("12. PCB制造中，蚀刻工序的主要目的是（    ）。", "A. 去除多余铜箔，形成线路图形", "B. 在孔壁沉积铜层", "C. 涂覆阻焊油墨", "D. 印刷丝印", "A"),
    ("13. 以下哪种不是PCB常见的基材类型？（    ）", "A. FR-4", "B. CEM-3", "C. 铝基板", "D. 硅基板", "D"),
    ("14. PCB板翘曲度超标可能导致的问题不包括（    ）。", "A. 贴片偏位", "B. 焊接不良", "C. 插件困难", "D. 线路开路", "D"),
    ("15. 以下关于PCB定位孔的说法，错误的是（    ）。", "A. 用于制造过程中的定位", "B. 一般为非金属化孔", "C. 直径通常为3.2mm", "D. 可以作为装配孔使用", "D"),
]

for q in questions_single:
    add_body(q[0])
    add_body(f"    {q[1]}    {q[2]}    {q[3]}    {q[4]}")
    doc.add_paragraph()

add_heading2("二、判断题（每题1分，共10分）")

questions_judge = [
    ("1. PCB制造中，最小线宽和线距只取决于设计能力，与制造能力无关。（    ）", "×"),
    ("2. Gerber文件输出时，钻孔图层（Drill Drawing）可以省略。（    ）", "×"),
    ("3. DFM审查的目的是在制造前发现设计中可能导致制造困难或良率降低的问题。（    ）", "√"),
    ("4. PCB制造中，阻焊开窗应比焊盘略大，以确保焊接可靠性。（    ）", "√"),
    ("5. 竞赛期间，选手可以使用自带的U盘传输文件。（    ）", "×"),
    ("6. NC Drill文件中的孔径信息必须与钻孔表一致。（    ）", "√"),
    ("7. PCB丝印可以印在焊盘上，不影响焊接质量。（    ）", "×"),
    ("8. 缺陷检测中，开路指的是线路中间断开，导致电气连接中断。（    ）", "√"),
    ("9. PCB制造中，工艺边的主要作用是方便制造过程中的夹持和传输。（    ）", "√"),
    ("10. DRC检查通过即可保证PCB完全可制造，无需进行DFM审查。（    ）", "×"),
]

for q in questions_judge:
    add_body(f"{q[0]}")

add_heading2("三、简答题（每题2.5分，共5分）")

add_body("1. 请简述PCB制造中开路和短路两种缺陷的定义、产生原因及检测方法。")
doc.add_paragraph()
add_body("    答题区：")
for _ in range(5):
    doc.add_paragraph()

add_body("2. 请简述DFM审查中，对线宽、线距、孔径三项参数进行检查的意义和标准。")
doc.add_paragraph()
add_body("    答题区：")
for _ in range(5):
    doc.add_paragraph()

# ============================================================
# 模块二：PCB工程资料处理
# ============================================================
add_page_break()
add_heading1("模块二：PCB工程资料处理（30分）")

add_body("选手根据本任务书和题目包中提供的文件，在嘉立创EDA中完成PCB设计并输出工程文件。", indent=True)

add_heading2("一、任务背景")
add_body("某电子产品需设计一块双层PCB板，用于STM32最小系统的核心板。选手需根据提供的原理图和封装库，完成PCB布局布线，并输出完整的制造文件。", indent=True)

add_heading2("二、设计要求")

add_body_bold_prefix("1. 板框尺寸：", "50mm × 40mm，圆角R=1mm。板框文件已包含在题目包中，选手不得修改。")
add_body_bold_prefix("2. 层数：", "双面板（Top Layer + Bottom Layer）。")
add_body_bold_prefix("3. 板厚：", "1.6mm。")
add_body_bold_prefix("4. 表面处理：", "HASL（喷锡）。")
add_body_bold_prefix("5. 阻焊颜色：", "绿色。")
add_body_bold_prefix("6. 丝印颜色：", "白色。")

add_heading2("三、制造规则")
add_body("选手须按以下规则设置设计参数：")

add_table(
    ["序号", "规则项目", "参数要求"],
    [
        ["1", "最小线宽", "≥8mil（0.2mm）"],
        ["2", "最小线距", "≥8mil（0.2mm）"],
        ["3", "最小孔径", "≥0.3mm"],
        ["4", "最小焊盘环宽", "单边≥0.15mm"],
        ["5", "板边距", "走线和焊盘距板边≥0.3mm"],
        ["6", "定位孔", "3.2mm非金属化孔，4个，位于四角"],
        ["7", "工艺边", "四边各留≥5mm"],
        ["8", "过孔", "0.5mm孔径/0.9mm外径"],
        ["9", "阻焊开窗", "比焊盘单边大0.05mm"],
        ["10", "丝印线宽", "≥0.15mm，丝印距焊盘≥0.2mm"],
    ],
    col_widths=[1.2, 4.0, 10.3]
)

add_heading2("四、设计约束")
add_body("1. 选手不得更改原理图电气连接，不得删除指定元件和接口。")
add_body("2. J1（电源接口）、J2（调试接口）、SW1（复位按键）必须放置在板边，便于外部连接。")
add_body("3. 晶振Y1（8MHz）应靠近MCU放置，走线尽量短。")
add_body("4. 电源滤波电容C1-C4应靠近MCU电源引脚放置。")
add_body("5. 所有元件丝印方向应统一，不得倒置或侧放。")

add_heading2("五、提交要求")
add_body("选手须输出以下文件：")
add_body("1. Gerber文件：包含Top Layer、Bottom Layer、Top Solder、Bottom Solder、Top Silkscreen、Bottom Silkscreen、Mechanical层。")
add_body("2. NC Drill钻孔文件及钻孔表。")
add_body("3. DRC报告：不得存在严重错误（Error）。")
add_body("4. PDF预览图：包含各层预览。")

add_heading2("六、评分标准")

add_table(
    ["序号", "评分项目", "分值", "评分说明"],
    [
        ["1", "规则设置正确", "5分", "线宽、线距、过孔等规则设置符合制造要求"],
        ["2", "布局合理", "5分", "元器件布局紧凑合理，信号流向清晰，约束元件位置正确"],
        ["3", "布线符合制造要求", "8分", "布线满足DRC，无严重违规，走线规范，电源/信号走线合理"],
        ["4", "DRC检查通过", "5分", "DRC报告无严重错误"],
        ["5", "Gerber/钻孔文件完整", "7分", "层文件完整，命名规范，钻孔文件正确"],
    ],
    col_widths=[1.5, 5.0, 1.5, 7.5]
)

# ============================================================
# 模块三：DFM与工艺文件
# ============================================================
add_page_break()
add_heading1("模块三：DFM与工艺文件（20分）")

add_body("选手根据模块二的设计成果和制造规则，对设计进行DFM（可制造性设计）审查，并填写DFM审查表。", indent=True)

add_heading2("一、审查内容")
add_body("选手需检查以下项目，并在DFM审查表中逐项填写检查结果：")

add_table(
    ["序号", "检查项目", "检查要点"],
    [
        ["1", "线宽线距", "所有走线线宽≥8mil，线距≥8mil，重点关注BGA区域和密集走线区域"],
        ["2", "孔径与环宽", "通孔孔径≥0.3mm，焊盘环宽单边≥0.15mm，过孔尺寸符合规则"],
        ["3", "板边距", "所有走线和焊盘距板边≥0.3mm"],
        ["4", "定位孔", "四角定位孔直径3.2mm，非金属化，位置正确"],
        ["5", "工艺边", "四边工艺边宽度≥5mm"],
        ["6", "阻焊开窗", "所有焊盘阻焊开窗正确，无漏开窗或过度开窗"],
        ["7", "丝印规范", "丝印线宽≥0.15mm，丝印不覆盖焊盘，方向统一"],
        ["8", "元件位号", "所有元件位号清晰可辨，无遗漏"],
    ],
    col_widths=[1.2, 3.5, 10.8]
)

add_heading2("二、提交要求")
add_body('选手须填写附件中的《DFM审查表》，逐项记录检查结果。对于不合格项，须在问题描述栏说明具体问题和位置。')

add_heading2("三、评分标准")

add_table(
    ["序号", "评分项目", "分值", "评分说明"],
    [
        ["1", "线宽线距检查", "4分", "正确检查并记录线宽线距是否满足制造能力"],
        ["2", "孔径与环宽检查", "4分", "正确检查并记录孔径和焊盘环宽"],
        ["3", "板边距、定位孔检查", "4分", "正确检查并记录板边距和定位孔设置"],
        ["4", "阻焊、丝印检查", "4分", "正确检查并记录阻焊开窗和丝印规范"],
        ["5", "工艺记录完整", "4分", "DFM审查表填写完整、规范，问题描述准确"],
    ],
    col_widths=[1.5, 5.0, 1.5, 7.5]
)

# ============================================================
# 模块四：裸板缺陷检测
# ============================================================
add_page_break()
add_heading1("模块四：裸板缺陷检测（20分）")

add_body("选手对裁判提供的裸板样板进行缺陷检测，识别缺陷类型、标注位置、分析原因并提出处理建议。", indent=True)

add_heading2("一、检测对象")
add_body("组委会提供1块含有多种缺陷的PCB裸板样板（或缺陷图册），样板上预设了8处缺陷，涵盖以下类型：")

add_table(
    ["序号", "缺陷类型", "典型表现"],
    [
        ["1", "开路", "线路中间断开"],
        ["2", "短路", "不应连接的线路之间形成导通"],
        ["3", "缺口/缺口", "线路边缘缺损"],
        ["4", "针孔", "阻焊层或铜层上的微小孔洞"],
        ["5", "偏位", "焊盘或孔相对于设计位置偏移"],
        ["6", "阻焊不良", "阻焊层脱落、气泡、覆盖不全"],
        ["7", "丝印不良", "丝印模糊、偏位、缺字"],
        ["8", "表面污染", "板面残留异物或化学污染"],
    ],
    col_widths=[1.2, 3.5, 10.8]
)

add_heading2("二、检测要求")
add_body("1. 选手使用自带的放大镜、万用表、游标卡尺等工具对样板进行逐项检测。")
add_body("2. 识别每处缺陷的类型，在样板上或图册上标注缺陷位置。")
add_body("3. 分析每处缺陷可能产生的原因（至少从制造工艺角度分析）。")
add_body("4. 对每处缺陷提出合理的处理建议（返修、报废、让步接收等）。")

add_heading2("三、提交要求")
add_body("选手须填写附件中的《缺陷检测记录表》，每处缺陷填写以下信息：")
add_body("1. 缺陷编号。")
add_body("2. 缺陷类型（参照IPC-A-600标准分类）。")
add_body("3. 缺陷位置（用坐标或区域描述）。")
add_body("4. 严重程度（致命/严重/一般/轻微）。")
add_body("5. 原因分析。")
add_body("6. 处理建议。")

add_heading2("四、评分标准")

add_table(
    ["序号", "评分项目", "分值", "评分说明"],
    [
        ["1", "缺陷识别准确", "8分", "正确识别8处缺陷的类型，无漏检误检（每处1分）"],
        ["2", "缺陷位置标注准确", "4分", "缺陷位置标注清晰准确"],
        ["3", "原因分析合理", "4分", "缺陷原因分析符合制造工艺逻辑"],
        ["4", "处理建议合理", "4分", "提出的处理建议具有可操作性"],
    ],
    col_widths=[1.5, 5.0, 1.5, 7.5]
)

# ============================================================
# 附件
# ============================================================
add_attachment_title("附件1：DFM审查表")

add_body("选手编号：____________    审查日期：____________")
doc.add_paragraph()

add_table(
    ["序号", "检查项目", "标准要求", "检查结果", "问题描述"],
    [
        ["1", "最小线宽", "≥8mil", "□合格 □不合格", ""],
        ["2", "最小线距", "≥8mil", "□合格 □不合格", ""],
        ["3", "最小孔径", "≥0.3mm", "□合格 □不合格", ""],
        ["4", "焊盘环宽", "单边≥0.15mm", "□合格 □不合格", ""],
        ["5", "板边距", "≥0.3mm", "□合格 □不合格", ""],
        ["6", "定位孔", "3.2mm非金属化", "□合格 □不合格", ""],
        ["7", "工艺边", "≥5mm", "□合格 □不合格", ""],
        ["8", "阻焊开窗", "比焊盘单边大0.05mm", "□合格 □不合格", ""],
        ["9", "丝印规范", "线宽≥0.15mm", "□合格 □不合格", ""],
        ["10", "丝印距焊盘", "≥0.2mm", "□合格 □不合格", ""],
    ],
    col_widths=[1.2, 3.0, 4.0, 3.5, 3.8]
)

doc.add_paragraph()
add_body("审查结论：□ 全部合格    □ 存在不合格项（需整改）")
add_body("审查人签名：____________")

add_attachment_title("附件2：缺陷检测记录表")

add_body("选手编号：____________    检测日期：____________")
doc.add_paragraph()

add_table(
    ["缺陷编号", "缺陷类型", "缺陷位置", "严重程度", "原因分析", "处理建议"],
    [
        ["1", "", "", "□致命 □严重 □一般 □轻微", "", ""],
        ["2", "", "", "□致命 □严重 □一般 □轻微", "", ""],
        ["3", "", "", "□致命 □严重 □一般 □轻微", "", ""],
        ["4", "", "", "□致命 □严重 □一般 □轻微", "", ""],
        ["5", "", "", "□致命 □严重 □一般 □轻微", "", ""],
        ["6", "", "", "□致命 □严重 □一般 □轻微", "", ""],
        ["7", "", "", "□致命 □严重 □一般 □轻微", "", ""],
        ["8", "", "", "□致命 □严重 □一般 □轻微", "", ""],
    ],
    col_widths=[1.8, 2.5, 2.5, 3.2, 2.8, 2.7]
)

doc.add_paragraph()
add_body("检测人签名：____________    日期：____________")

# ============================================================
# 页眉页脚
# ============================================================
for section in doc.sections:
    header = section.header
    header.is_linked_to_previous = False
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    hr = hp.add_run('江西省"振兴杯"职业技能竞赛印制电路制作工样题')
    hr.font.name = '宋体'
    hr.font.size = Pt(9)
    hr.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    hp.paragraph_format.space_after = Pt(6)

    footer = section.footer
    footer.is_linked_to_previous = False
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER

    def add_footer_text(paragraph, text):
        r = paragraph.add_run(text)
        r.font.name = '宋体'
        r.font.size = Pt(9)
        r.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')

    def add_page_field(paragraph, field_name):
        run_begin = paragraph.add_run()
        run_begin._r.append(parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="begin"/>'))
        run_instr = paragraph.add_run()
        run_instr._r.append(parse_xml(f'<w:instrText {nsdecls("w")} xml:space="preserve"> {field_name} </w:instrText>'))
        run_end = paragraph.add_run()
        run_end._r.append(parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="end"/>'))

    add_footer_text(fp, "第 ")
    add_page_field(fp, "PAGE")
    add_footer_text(fp, " 页 / 共 ")
    add_page_field(fp, "NUMPAGES")
    add_footer_text(fp, " 页")

# ============================================================
# 保存
# ============================================================
# 草稿输出（勿覆盖 02_样题 完整版权威样题）
_SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
_GEN_DIR = os.path.join(_SCRIPT_DIR, "_generated")
os.makedirs(_GEN_DIR, exist_ok=True)
output_path = os.path.join(_GEN_DIR, "印制电路制作工竞赛样题_draft.docx")
doc.save(output_path)
print(f"样题已生成: {output_path}")
