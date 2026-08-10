# -*- coding: utf-8 -*-
"""Build Module D referee answer sheet; sync exam M04 hole to 3.50 mm."""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[2]
PACK = ROOT / "02_样题" / "第一套"
OUT = PACK / "模块D_检测记录表_参考答案.docx"
EXAM = PACK / "印制电路制作工赛项_竞赛样题（完整版）.docx"
# keep user's copy; also leave blank form as is


def set_run_font(run, name="宋体", size=10.5, bold=False):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold


def add_p(doc, text, *, size=10.5, bold=False, center=False, after=4, color=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.1
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold)
    if color is not None:
        run.font.color.rgb = color
    return p


def shade_cell(cell, fill="FFF2CC"):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    for old in tcPr.findall(qn("w:shd")):
        tcPr.remove(old)
    shd = tcPr.makeelement(
        qn("w:shd"),
        {qn("w:val"): "clear", qn("w:color"): "auto", qn("w:fill"): fill},
    )
    tcPr.append(shd)


def set_cell(cell, text, *, bold=False, size=9.5, fill=None, center=False):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold)
    if fill:
        shade_cell(cell, fill)


def build_answer() -> Document:
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.left_margin = Cm(1.5)
    sec.right_margin = Cm(1.5)
    sec.top_margin = Cm(1.2)
    sec.bottom_margin = Cm(1.2)

    add_p(doc, '江西省"振兴杯"职业技能大赛', size=12, bold=True, center=True, after=2)
    add_p(doc, '"印制电路制作工"赛项', size=12, bold=True, center=True, after=2)
    add_p(doc, "模块D：成品PCB裸板质量检测与判定", size=14, bold=True, center=True, after=2)
    add_p(doc, "检测记录表（参考答案·裁判用）", size=12, bold=True, center=True, after=6)
    add_p(
        doc,
        "【黄色底纹为参考答案】正式评分允许位置同义表述与合理测量误差；"
        "以实物三测标定为准，下列数值为设计/标定骨架（孔径按本套板 3.50 mm）。"
        "本文件不得发给选手。",
        size=9,
        after=6,
        color=RGBColor(0x66, 0x33, 0x00),
    )

    # header
    add_p(doc, "表头", size=11, bold=True, after=3)
    t0 = doc.add_table(rows=2, cols=4)
    t0.style = "Table Grid"
    set_cell(t0.rows[0].cells[0], "赛位号", bold=True, fill="D9E2F3")
    set_cell(t0.rows[0].cells[1], "（裁判用·空白）", fill="FFF2CC")
    set_cell(t0.rows[0].cells[2], "日期", bold=True, fill="D9E2F3")
    set_cell(t0.rows[0].cells[3], "赛日", fill="FFF2CC")
    set_cell(t0.rows[1].cells[0], "样件编号", bold=True, fill="D9E2F3")
    set_cell(t0.rows[1].cells[1], "标准缺陷样件", fill="FFF2CC")
    set_cell(t0.rows[1].cells[2], "版本（丝印）", bold=True, fill="D9E2F3")
    set_cell(t0.rows[1].cells[3], "☑ MOD-D-S-A　□ MOD-D-S-B\n（B 版类型相同、位置见发放表）", fill="FFF2CC")

    add_p(doc, "", after=4)
    add_p(doc, "一、D-1 检验准备（示例勾选）", size=11, bold=True, after=3)
    add_p(
        doc,
        "☑ 已阅读赛题模块D第（三）节检测依据　☑ 顶面/底面方向已确认　"
        "☑ 卡尺已对零　☑ 厚度规可用　☑ 显微镜/放大镜可用　☑ 万用表可用　"
        "☑ 版本丝印已抄录　☑ 样件外观运输情况：正常",
        size=9,
        after=6,
    )

    add_p(doc, "二、D-2 外观缺陷检测（V01～V05 必检）", size=11, bold=True, after=3)
    add_p(
        doc,
        "说明：类型名以样题标准名为准；括号内为同义/现象。"
        "底面至少含 V02、V03。不设 V06 必检（额外丝印问题可作干扰项，不计必检分）。",
        size=9,
        after=3,
    )

    tv = doc.add_table(rows=6, cols=6)
    tv.style = "Table Grid"
    headers = ["编号", "所在面", "网格/特征", "缺陷类型（标准名）", "现象简述", "判定"]
    for i, h in enumerate(headers):
        set_cell(tv.rows[0].cells[i], h, bold=True, size=9, fill="D9E2F3", center=True)

    visual = [
        (
            "V01",
            "顶层",
            "D3 / LED1 限流支路",
            "顶层线路缺口",
            "LED1 限流电阻相关铜线断开（线路缺口/断线）；不断开整网叙述时以实物为准",
            "不合格",
        ),
        (
            "V02",
            "底层",
            "B2 / 走线区",
            "底层线宽变窄（颈缩）",
            "电源相关网络线宽局部突然变细，细段仍连通（非开路）；与 M05 不同点",
            "不合格",
        ),
        (
            "V03",
            "底层",
            "D2 / 孔阵",
            "孔破盘",
            "钻孔偏出焊盘，焊环一侧不足/破盘（非 3.5 mm 定位孔）",
            "不合格",
        ),
        (
            "V04",
            "顶层",
            "C3 / L1",
            "阻焊未开窗或开窗偏移",
            "L1 电感有一引脚阻焊未开窗（或严重盖盘）",
            "不合格",
        ),
        (
            "V05",
            "顶层",
            "C2 / 电容丝印",
            "丝印缺失或压焊盘",
            "电容位号丝印不完整/缺失",
            "不合格",
        ),
    ]
    for r, row in enumerate(visual, start=1):
        for c, val in enumerate(row):
            set_cell(tv.rows[r].cells[c], val, size=9, fill="FFF2CC")

    add_p(doc, "", after=4)
    add_p(doc, "三、D-3 尺寸与指定点测量（M01～M05）", size=11, bold=True, after=3)
    add_p(
        doc,
        "单位 mm。M04 本套板定位孔名义 3.50 mm（±0.10）。"
        "M05 须用测量显微镜；禁止用普通卡尺测细线宽。"
        "实测值投板三测后可微调，下表为设计/标定骨架。",
        size=9,
        after=3,
    )
    add_p(
        doc,
        "M05 指定点位置/标识：板面丝印「M05」（或等价标识）　测量对象：☑ 线宽　□ 线距"
        "　【不得与 V02 颈缩为同一处】",
        size=9,
        after=3,
    )

    tm = doc.add_table(rows=6, cols=5)
    tm.style = "Table Grid"
    for i, h in enumerate(["编号", "项目", "工具", "实测值 (mm)", "合格/不合格"]):
        set_cell(tm.rows[0].cells[i], h, bold=True, size=9, fill="D9E2F3", center=True)

    measures = [
        ("M01", "板长 L", "游标卡尺", "80.00（允许 79.80～80.20）", "合格*"),
        ("M02", "板宽 W", "游标卡尺", "60.00（允许 59.80～60.20）", "合格*"),
        ("M03", "板厚 T", "厚度规/千分尺", "1.60（允许 1.45～1.75）", "合格*"),
        ("M04", "定位孔径", "卡尺/孔规", "3.50（允许 3.40～3.60）", "合格*"),
        (
            "M05",
            "指定点线宽",
            "测量显微镜",
            "以板面 M05 标识为准；设计骨架约 0.25～0.30；允许 ±0.05",
            "合格*",
        ),
    ]
    for r, row in enumerate(measures, start=1):
        for c, val in enumerate(row):
            set_cell(tm.rows[r].cells[c], val, size=9, fill="FFF2CC")

    add_p(
        doc,
        "*「合格」指相对本卷尺寸基准表；投板后若超差，以三测真值为准改判定，并同步评分。",
        size=8,
        after=6,
    )

    add_p(doc, "四、D-4 基础电气检测（E01～E02）", size=11, bold=True, after=3)
    te = doc.add_table(rows=3, cols=7)
    te.style = "Table Grid"
    for i, h in enumerate(
        ["编号", "类型", "测试点", "档位", "读数", "判定", "与预期一致"]
    ):
        set_cell(te.rows[0].cells[i], h, bold=True, size=9, fill="D9E2F3", center=True)

    elec = [
        (
            "E01",
            "开路",
            "TP1–TP2",
            "电阻档（高阻/200MΩ 等）",
            "≥1 MΩ 或 OL；示例 >10 MΩ",
            "开路",
            "☑ 是",
        ),
        (
            "E02",
            "短路",
            "TP3–TP4",
            "电阻档（低阻/100Ω/通断）",
            "≤1 Ω；示例 ≈0 Ω",
            "短路",
            "☑ 是",
        ),
    ]
    for r, row in enumerate(elec, start=1):
        for c, val in enumerate(row):
            set_cell(te.rows[r].cells[c], val, size=9, fill="FFF2CC")

    add_p(doc, "", after=4)
    add_p(doc, "五、D-5 分类判定与提交", size=11, bold=True, after=3)
    add_p(
        doc,
        "1. 外观必检：V01 不合格　V02 不合格　V03 不合格　V04 不合格　V05 不合格　（已检查☑）",
        size=9,
        after=2,
    )
    add_p(
        doc,
        "2. 分类汇总：线路 2（V01 缺口 + V02 颈缩）　孔/盘 1（V03）　阻焊 1（V04）　丝印 1（V05）　其它 0",
        size=9,
        after=2,
    )
    add_p(
        doc,
        "3. 尺寸：☑ M01～M05 全部合格（按上表骨架；投板超差则改）　□ 有不合格：—",
        size=9,
        after=2,
    )
    add_p(
        doc,
        "4. 电气：☑ E01、E02 均与预期一致　□ 不符合：—",
        size=9,
        after=2,
    )
    add_p(doc, "5. 质量处置：□ 接收　□ 返工　☑ 报废", size=9, after=2)
    add_p(
        doc,
        "6. 依据：存在顶层线路缺口（V01）、孔破盘（V03）等严重影响使用的缺陷；"
        "另有底层颈缩、阻焊未开窗、丝印缺失。虽 E01/E02 符合本卷预设预期，"
        "但按样题报废规则应判定报废，不得接收。",
        size=9,
        after=2,
    )
    add_p(
        doc,
        "7. 其他说明：选手多写干扰项丝印不强制另设 V06 得分；"
        "类型同义（断线/缺口、颈缩/变窄、盖盘/未开窗）可给分。",
        size=9,
        after=6,
    )

    add_p(doc, "六、裁判给分口径（摘要·40 分骨架可折算）", size=11, bold=True, after=3)
    for line in [
        "外观 14：V01～V05 各约 2.8；面+类型+位置大致正确即可；V02 须体现底层+变窄且不断开。",
        "尺寸 12：M01～M04 各约 2，M05 约 4；读数在公差内且工具正确。",
        "电气 8：E01 开路判定正确 4；E02 短路判定正确 4。",
        "综合 4：分类合理 + 处置选报废 + 依据提到致命缺陷。",
        "规范 2：版本/单位/记录完整。",
        "误把正常线当 V02、或 M05 与 V02 同点、或定位孔写成 2.00 与本套 3.50 冲突 → 按错扣。",
    ]:
        add_p(doc, "• " + line, size=9, after=2)

    add_p(doc, "", after=4)
    add_p(
        doc,
        "签注：命题/裁判长 ________　　日期 ________　　本答案随投板三测修订后锁版",
        size=9,
        after=2,
    )
    return doc


def patch_exam_m04():
    if not EXAM.exists():
        print("exam missing")
        return
    doc = Document(str(EXAM))
    n = 0
    for table in doc.tables:
        body = "\n".join(c.text for r in table.rows for c in r.cells)
        if "M04" not in body or "定位孔" not in body:
            continue
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if len(cells) >= 4 and cells[0].replace(" ", "") == "M04":
                # 编号 | 项目 | 工具 | 名义值 | 允许误差
                if "2.00" in row.cells[3].text or "2.0" in row.cells[3].text:
                    row.cells[3].text = "3.50"
                    n += 1
                # also fix tool note if needed
        break
    # also patch paragraph mentions of 2.00 mm locating hole if any
    for p in doc.paragraphs:
        if "定位孔" in p.text and "2.00" in p.text:
            for run in p.runs:
                if "2.00" in run.text:
                    run.text = run.text.replace("2.00", "3.50")
                    n += 1
    doc.save(str(EXAM))
    print(f"patched exam M04 sites: {n}")


def main():
    ans = build_answer()
    ans.save(OUT)
    print("wrote", OUT)
    patch_exam_m04()
    # rename note: user copy left as-is
    print("done")


if __name__ == "__main__":
    main()
