# -*- coding: utf-8 -*-
"""
完整版模块D回写：选手-only 题干
- 原则 B + S1 四表 + T1 + J1 + K1 + N1
- 不写总时长、不写各任务建议分钟
- 删除命题口吻（分册/骨架/F1/标定/评分细则等）
"""
from __future__ import annotations

import shutil
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

ROOT = Path(__file__).resolve().parents[2]
SAMPLE = ROOT / "02_样题"
PACK = SAMPLE / "学生组_模块D"
FULL = SAMPLE / "印制电路制作工赛项_竞赛样题（完整版）.docx"
BACKUP = ROOT / "05_归档备份" / f"module-d-player-only-{datetime.now().strftime('%Y%m%d_%H%M%S')}"
C_BLUE = "2E75B6"


def set_run_font(run, name_cn="仿宋", name_en="Times New Roman", size_pt=12, bold=False, color_hex=None):
    run.bold = bold
    run.font.size = Pt(size_pt) if size_pt else None
    run.font.name = name_en
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:ascii"), name_en)
    rFonts.set(qn("w:hAnsi"), name_en)
    rFonts.set(qn("w:eastAsia"), name_cn)
    rFonts.set(qn("w:cs"), name_en)
    if color_hex:
        run.font.color.rgb = RGBColor.from_string(color_hex)
    else:
        try:
            run.font.color.rgb = RGBColor(0, 0, 0)
        except Exception:
            pass


def add_para(doc, text, *, cn="仿宋", size=12, bold=False, space_after=4, space_before=0):
    p = doc.add_paragraph()
    try:
        p.style = doc.styles["Normal"]
    except Exception:
        pass
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(space_before)
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    run = p.add_run(text)
    en = "黑体" if cn == "黑体" else "Times New Roman"
    set_run_font(run, name_cn=cn, name_en=en, size_pt=size, bold=bold)
    return p


def shade_cell(cell, fill_hex):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill_hex)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def set_cell(cell, text, *, header=False, size=9.5):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if header else WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(str(text))
    set_run_font(
        run,
        name_cn="黑体" if header else "仿宋",
        name_en="Times New Roman",
        size_pt=9 if header else size,
        bold=header,
        color_hex="FFFFFF" if header else None,
    )
    if header:
        shade_cell(cell, C_BLUE)


def add_table(doc, rows):
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            set_cell(table.cell(ri, ci), val, header=(ri == 0))
    add_para(doc, "", space_after=2)
    return table


def el_text(el) -> str:
    return "".join(t.text or "" for t in el.iter(qn("w:t"))).strip()


def find_module_d_body_start(body) -> int:
    kids = list(body.iterchildren())
    title_idx = None
    for i, child in enumerate(kids):
        if child.tag.split("}")[-1] != "p":
            continue
        t = el_text(child)
        if t.startswith("模块D：") and ("成品" in t or "质量检测" in t or "裸板" in t):
            title_idx = i
            break
    if title_idx is None:
        raise RuntimeError("未找到模块D标题")
    start = title_idx
    j = title_idx - 1
    while j >= 0:
        child = kids[j]
        if child.tag.split("}")[-1] == "tbl":
            tx = el_text(child)
            if any(k in tx for k in ("检测记录表", "尺寸基准", "验收", "基础电气", "板面分区", "成品裸板", "双面成品")):
                start = j
                j -= 1
                continue
        break
    return start


def remove_body_from(body, start_idx: int):
    kids = list(body.iterchildren())
    to_remove = []
    for i, child in enumerate(kids):
        if i < start_idx:
            continue
        if child.tag.split("}")[-1] == "sectPr":
            continue
        to_remove.append(child)
    for el in to_remove:
        body.remove(el)
    print(f"[FULL] removed {len(to_remove)} elements from {start_idx}")


def write_module_d(doc: Document):
    H = lambda t, sz=14: add_para(doc, t, cn="黑体", size=sz, bold=True, space_before=8, space_after=4)
    L = lambda t: add_para(doc, t, cn="黑体", size=12, bold=True, space_before=6, space_after=2)
    B = lambda t, **kw: add_para(doc, t, cn="仿宋", size=12, bold=False, **kw)

    H("模块D：成品PCB裸板质量检测与判定（学生组）", 16)

    H("（一）考核要求")
    B("对赛场提供的双面成品PCB裸板（无元器件，已完成线路、阻焊、丝印和表面处理）完成：")
    B("1. 外观缺陷检测（顶面、底面）；")
    B("2. 关键尺寸与指定点线宽/线距测量；")
    B("3. 指定测试点的开路、短路检测；")
    B("4. 缺陷分类与质量处置结论（接收 / 返工 / 报废）及简要依据。")
    B("仅检验本模块样件；只做指定测试点的开短路；不上电；不测量 Gerber。")

    H("（二）任务说明")
    L("【检验对象】")
    B("双面成品裸板 1 块。")
    B("外形约 80 mm × 60 mm，标称板厚 1.6 mm；FR-4 双面，绿阻焊，白丝印，无铅喷锡。")
    B("丝印含版本标识（MOD-D-S-A 或 MOD-D-S-B）、分区参考及测试点 TP1～TP4。")
    L("【提供材料】")
    B("1. 双面成品裸板样件 1 块；")
    B("2. 《模块D 检测记录表》（空白，填写后提交）。")
    L("【使用工具】")
    B("游标卡尺、厚度规（或千分尺）、测量显微镜（或带刻度放大镜）、放大镜、数字万用表、照明及防静电用品等（以赛场实际提供为准）。")

    H("（三）检测依据")
    B("下列说明与表格为本模块判定与填写依据。请对照填写《模块D 检测记录表》。")

    L("1. 板面分区与测试点")
    B("版本以板面丝印为准（MOD-D-S-A / MOD-D-S-B），记录表版本栏须与实物一致。")
    B("分区网格：顶层参考 A1～D4（列 1～4，行 A～D）。记录缺陷时填写：所在面 + 网格或特征名 + 方位。")
    B("特征区名称以板面丝印为准，例如：走线区、孔阵、阻焊对比区、丝印条、板框、测试点区。")
    B("测试点用途：")
    add_table(
        doc,
        [
            ["测试点", "用途"],
            ["TP1、TP2", "E01 开路检测两端"],
            ["TP3、TP4", "E02 短路检测两端"],
        ],
    )

    L("2. 尺寸基准（M01～M05）")
    B("单位：mm。M01～M04 使用卡尺或厚度规；M05 使用测量显微镜，只测指定点（不得与 V02 为同一处）。")
    B("禁止用普通卡尺测量约 0.30 mm 及以下小孔；禁止对全板任意扫测微距。")
    add_table(
        doc,
        [
            ["编号", "项目", "工具", "名义值", "允许误差", "判定"],
            ["M01", "板长 L", "游标卡尺", "80.00", "±0.20", "合格 / 不合格"],
            ["M02", "板宽 W", "游标卡尺", "60.00", "±0.20", "合格 / 不合格"],
            ["M03", "板厚 T", "厚度规/千分尺", "1.60", "±0.15", "合格 / 不合格"],
            ["M04", "定位孔径", "卡尺/孔规", "2.00", "±0.10", "合格 / 不合格"],
            ["M05", "指定点线宽或线距", "测量显微镜", "以板面指定点标识为准", "±0.05", "合格 / 不合格"],
        ],
    )
    B("M05：按板面指定点标识选择测线宽或线距其一；读数一般保留到 0.01 mm。")

    L("3. 外观缺陷判定（V01～V05）")
    B("须检查顶面与底面。下列 5 类为必检内容；记录类型、位置与现象。")
    add_table(
        doc,
        [
            ["编号", "类型", "判定为不合格的情形"],
            ["V01", "顶层线路缺口", "顶层导体可见断开或缺口"],
            ["V02", "底层线宽变窄", "底层局部线宽明显变窄（目视/放大观察即可，此处不要求读数）"],
            ["V03", "孔破盘", "孔与焊盘关系异常，出现破盘或焊盘缺损"],
            ["V04", "阻焊未开窗或开窗偏移", "应开窗处被阻焊覆盖，或开窗相对焊盘明显偏移"],
            ["V05", "丝印缺失或压焊盘", "字符缺失/残缺，或丝印压在焊盘上"],
        ],
    )
    B("板面可能存在其他现象，是否记入由你判断；上述 5 类应完整检查并记录。")

    L("4. 基础电气检测（E01～E02）")
    B("仅在指定测试点之间测量。使用数字万用表电阻档或通断档。")
    add_table(
        doc,
        [
            ["编号", "类型", "测试点", "预期结果", "判定参考"],
            ["E01", "开路", "TP1–TP2", "开路", "电阻很大（约 ≥1 MΩ）或显示开路/OL"],
            ["E02", "短路", "TP3–TP4", "短路", "电阻很小（约 ≤1 Ω）或通断档导通"],
        ],
    )
    B("记录：测试点、档位、电阻读数（或通断结果）、判定（开路/短路/正常）、是否与预期一致。")

    H("（四）模块任务")
    B("请完成下列任务，结果填入《模块D 检测记录表》后提交。判定以本卷第（三）节为准。")

    L("D-1 检验准备")
    B("1. 核对样件编号与版本丝印，填入记录表；")
    B("2. 确认顶面/底面方向及分区约定；")
    B("3. 检查所用工具可用（卡尺对零、万用表档位等）；")
    B("4. 如发现明显运输损伤，先在记录表注明。")

    L("D-2 外观缺陷检测")
    B("1. 检查顶面与底面；")
    B("2. 完成 V01～V05 必检内容，对照第（三）节第 3 款判定；")
    B("3. 每处缺陷在记录表中填写：所在面、类型、网格/特征、现象等；")
    B("4. 位置描述清楚即可（例如“顶层 B2 / 走线区 / 中部”），不要求坐标；")
    B("5. 类型写具体名称（如“线路缺口”“孔破盘”）。")

    L("D-3 尺寸与指定点测量")
    B("1. 按第（三）节第 2 款测量 M01～M05；")
    B("2. 填写实测值，判定合格/不合格；")
    B("3. 以 mm 记录。")

    L("D-4 基础电气检测")
    B("1. 按第（三）节第 4 款，在 TP1～TP4 完成 E01、E02；")
    B("2. 记录读数与判定。")

    L("D-5 分类判定与提交")
    B("1. 对外观缺陷分类汇总；")
    B("2. 汇总尺寸、电气是否符合第（三）节要求；")
    B("3. 给出质量处置结论，并写简要依据：")
    B("   （1）接收：必检外观、尺寸、电气均符合本卷要求；")
    B("   （2）返工：存在可返工缺陷（如丝印、部分阻焊问题），且尺寸与电气总体可接受；")
    B("   （3）报废：存在严重影响使用的缺陷（如线路缺口、孔破盘），或电气结果与预期严重不符；")
    B("4. 检查记录完整后提交《模块D 检测记录表》。")
    B("   电子提交目录示例：D:\\提交资料\\模块D\\ ；命名示例：赛位号_模块D。")

    L("注意事项")
    B("1. 请合理安排检查、测量、点测与填写时间。")
    B("2. 只检验本模块样件；只做指定点开短路；不上电；不测量 Gerber。")
    B("3. 判定以本卷第（三）节为准；提交物为《模块D 检测记录表》。")
    B("4. 爱护样件与仪器；注意防静电；正确使用万用表档位。")
    B("5. 成果物不得标注姓名等身份信息；设备异常请举手示意。")


def sync_front_tables(doc: Document):
    for table in doc.tables:
        if len(table.rows) >= 5 and "模块编号" in table.rows[0].cells[0].text:
            for row in table.rows:
                if row.cells[0].text.strip() == "模块D":
                    want = "成品PCB裸板质量检测与判定（学生组）"
                    if want not in row.cells[1].text:
                        row.cells[1].text = ""
                        run = row.cells[1].paragraphs[0].add_run(want)
                        set_run_font(run, name_cn="仿宋", size_pt=12)
        if len(table.columns) >= 3 and "提交内容" in table.rows[0].cells[1].text:
            for row in table.rows:
                if row.cells[0].text.strip() == "模块D":
                    want = "模块D 检测记录表（外观/尺寸/电气/综合判定）"
                    if row.cells[1].text.strip() != want:
                        row.cells[1].text = ""
                        run = row.cells[1].paragraphs[0].add_run(want)
                        set_run_font(run, name_cn="仿宋", size_pt=12)


def update_record_form_title():
    """检测记录表提示与样题用词对齐，去掉时长。"""
    path = PACK / "选手" / "学生组_模块D_检测记录表.docx"
    if not path.exists():
        print("[SKIP] record form missing")
        return
    doc = Document(str(path))
    changed = 0
    for p in doc.paragraphs:
        t = p.text
        if not t:
            continue
        nt = t
        nt = nt.replace("正式时长45分钟（D-1～D-5）。", "")
        nt = nt.replace("正式时长 45 分钟（D-1～D-5）。", "")
        nt = nt.replace("时长建议 45 分钟（D-1～D-5）。", "")
        nt = nt.replace("学生组_模块D_检测记录表", "模块D 检测记录表")
        if "完整版" in nt and "第（三）节" in nt:
            nt = (
                "⚠ 注意：判定与填写依据见《竞赛样题》模块D第（三）节检测依据"
                "（分区与测试点、尺寸基准、外观判定、电气判据）。"
                "本表填写后提交。命名示例：赛位号_模块D。不得标注姓名（按赛场规则）。"
                "电气仅为指定测试点开短路。"
            )
        if nt != t:
            for r in list(p.runs):
                r._element.getparent().remove(r._element)
            run = p.add_run(nt)
            run.bold = True
            run.font.size = Pt(10)
            run.font.name = "微软雅黑"
            rPr = run._element.get_or_add_rPr()
            rFonts = rPr.get_or_add_rFonts()
            for k in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
                rFonts.set(qn(k), "微软雅黑")
            run.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
            changed += 1
            print("[RECORD]", nt[:80])
    # title para if any
    for p in doc.paragraphs[:3]:
        if "检测记录表" in p.text and "PCB" in p.text:
            nt = "模块D 检测记录表（学生组）"
            if p.text.strip() != nt:
                for r in list(p.runs):
                    r._element.getparent().remove(r._element)
                run = p.add_run(nt)
                run.bold = True
                run.font.size = Pt(14)
                run.font.name = "黑体"
                rPr = run._element.get_or_add_rPr()
                rFonts = rPr.get_or_add_rFonts()
                for k in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
                    rFonts.set(qn(k), "黑体")
                changed += 1
    doc.save(str(path))
    print("[OK] record form", changed)


def write_docs():
    (PACK / "完整版学生组模块D正文稿_选手卷_审阅.md").write_text(
        Path(__file__).read_text(encoding="utf-8")  # placeholder replaced below
        if False
        else "",
        encoding="utf-8",
    )
    # write clean mirror of player text
    (PACK / "完整版学生组模块D正文稿_选手卷_审阅.md").write_text(
        """# 模块D：成品PCB裸板质量检测与判定（学生组）

> 已回写完整版 docx。原则：选手-only；**不写时长**；第（三）节仅四表；处置在 D-5。

（内容与完整版模块D一致，以 docx 为准。）

## 决策摘要

| 项 | 内容 |
|----|------|
| 原则 | 选手-only + 考场纪律 |
| 时长 | **卷面不写**（总时长与分段建议均不写） |
| 第（三）节 | 检测依据：分区TP / 尺寸 / 外观 / 电气 |
| 处置 | D-5 |
| 安全 | 注意事项 |
| 编号 | 保留 V/M/E/TP |
""",
        encoding="utf-8",
    )
    (SAMPLE / "00_学生组模块D现行包说明.md").write_text(
        """# 学生组模块D现行包说明（选手卷口径）

## 样题

- 文件：`02_样题/印制电路制作工赛项_竞赛样题（完整版）.docx` **模块D**
- 面向：**选手**；仅解题必备 + 必要考场纪律
- **卷面不写时长**（总时长/分段建议均不写；时长以竞赛通知或技术工作文件为准）
- 第（三）节「检测依据」仅四块：分区与TP、尺寸、外观、电气
- 提交：仅《模块D 检测记录表》

## 下发

| 材料 | 形式 |
|------|------|
| 赛题模块D | 完整版内 |
| 检测记录表 | `学生组_模块D/选手/学生组_模块D_检测记录表.docx` |
| 样件与量具 | 赛场提供 |

命题/裁判材料仍在 `学生组_模块D/` 包内，**不印入选手样题**。
""",
        encoding="utf-8",
    )
    print("[OK] package notes")


def verify(doc: Document):
    full = "\n".join(p.text for p in doc.paragraphs)
    checks = {
        "标题": "模块D：成品PCB裸板质量检测与判定" in full,
        "检测依据": "检测依据" in full,
        "V01": "V01" in full,
        "M05": "M05" in full,
        "E01": "E01" in full,
        "TP1": "TP1" in full,
        "无45分钟": "45分钟" not in full and "45 分钟" not in full,
        "无建议分钟": "建议3分钟" not in full and "建议12分钟" not in full,
        "无分册": "分册" not in full,
        "无骨架": "骨架" not in full,
        "无F1": "F1" not in full,
        "无标定": "标定" not in full,
        "模块B": "模块B：" in full,
        "模块C": "模块C：" in full,
    }
    # 模块D段内不应出现 45
    in_d = False
    d_text = []
    for p in doc.paragraphs:
        t = p.text.strip()
        if t.startswith("模块D："):
            in_d = True
        if in_d:
            d_text.append(t)
    d_join = "\n".join(d_text)
    checks["D段无时长"] = "45" not in d_join and "分钟" not in d_join.replace("命名", "")
    # allow 命名 without 分钟 - check 分钟 specifically in D
    checks["D段无分钟"] = "分钟" not in d_join

    for k, v in checks.items():
        print(("OK" if v else "FAIL"), k)
        if not v and k in ("D段无分钟", "D段无时长"):
            # show lines with 分钟
            for line in d_text:
                if "分钟" in line or "45" in line:
                    print("   >>", line[:100])
    if not all(checks.values()):
        # soft fail on 分钟 if only false positive - recheck
        bad = [k for k, v in checks.items() if not v]
        if bad == ["D段无分钟"] or bad == ["D段无时长", "D段无分钟"]:
            # if only empty issues
            pass
        if any(k for k in bad if k not in ("D段无分钟", "D段无时长")):
            raise SystemExit(2)
        if "分钟" in d_join:
            raise SystemExit(3)
    print("VERIFY PASS")


def main():
    BACKUP.mkdir(parents=True, exist_ok=True)
    shutil.copy2(FULL, BACKUP / FULL.name)
    print("[BACKUP]", BACKUP)

    doc = Document(str(FULL))
    body = doc.element.body
    start = find_module_d_body_start(body)
    print("[START]", start, el_text(list(body.iterchildren())[start])[:50])
    remove_body_from(body, start)
    write_module_d(doc)
    sync_front_tables(doc)
    doc.save(str(FULL))
    print("[OK] saved", FULL)

    update_record_form_title()
    write_docs()

    doc2 = Document(str(FULL))
    verify(doc2)
    # print outline
    for p in doc2.paragraphs:
        t = p.text.strip()
        if t.startswith("模块D") or t.startswith("（") or t.startswith("D-") or t == "注意事项":
            if len(t) < 40:
                print(" ", t)
    print("DONE")


if __name__ == "__main__":
    main()
