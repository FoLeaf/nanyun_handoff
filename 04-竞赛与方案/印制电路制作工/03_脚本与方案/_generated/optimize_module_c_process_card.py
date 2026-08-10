# -*- coding: utf-8 -*-
"""Optimize Module C process card: only Gerber-measurable fields + sync sample exam C-2."""
from __future__ import annotations

from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[2]
PACK = ROOT / "02_样题" / "第一套"
OUT_PLAYER = PACK / "模块C_PCB制程工艺卡.docx"
OUT_ANSWER = PACK / "模块C_PCB制程工艺卡（参考答案版）.docx"
OUT_DUP_REMOVE = PACK / "模块C_PCB制程工艺卡 - 副本.docx"
FULL_EXAM = PACK / "印制电路制作工赛项_竞赛样题（完整版）.docx"


def set_run_font(run, name="宋体", size=11, bold=False, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def add_para(doc, text, *, size=11, bold=False, center=False, space_after=6, color=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color)
    return p


def set_cell_text(cell, text, *, bold=False, size=10.5, fill=None, center=False):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold)
    if fill:
        tc = cell._tc
        tcPr = tc.get_or_add_tcPr()
        shd = tcPr.makeelement(
            qn("w:shd"),
            {
                qn("w:val"): "clear",
                qn("w:color"): "auto",
                qn("w:fill"): fill,
            },
        )
        # remove old shd
        for old in tcPr.findall(qn("w:shd")):
            tcPr.remove(old)
        tcPr.append(shd)


def build_card(answer: bool) -> Document:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)
    section.top_margin = Cm(1.5)
    section.bottom_margin = Cm(1.5)

    title = "印制电路板（PCB）简版制程工艺卡"
    if answer:
        title += "（参考答案版·模块C）"
    add_para(doc, title, size=16, bold=True, center=True, space_after=4)
    add_para(
        doc,
        "模块C · CAM审核配套 · 仅依据 Gerber/钻孔实测填写（单位优先 mil）",
        size=10.5,
        center=True,
        space_after=8,
        color=RGBColor(0x33, 0x33, 0x33),
    )

    add_para(doc, "一、填写说明", size=12, bold=True, space_after=4)
    rules = [
        "本卡只考「能量/能数清」的文件参数，不要求填写板厚、铜厚、板材、表面处理、阻焊颜色等订单约定项。",
        "最小线宽/线距：量正常布线区，不得把故意缺陷段（如极细线、贴线）当作规格填入。",
        "最小孔径：取钻孔文件中的最小工具直径（过孔优先）。",
        "外形尺寸：量板框层（Board Outline），长×宽，mil 与 mm 可双写。",
        "文件完整性：对照应有层（顶/底线路、顶/底阻焊、顶/底丝印、板框、钻孔）勾选或简述。",
        "编制栏只写赛位号与日期，不得填写姓名、单位。",
    ]
    if answer:
        rules.append("【黄色底纹为裁判参考答案；正式评分允许合理测量误差。】")
    for r in rules:
        add_para(doc, "• " + r, size=10, space_after=2)

    add_para(doc, "二、基本信息（选手填写）", size=12, bold=True, space_after=4)

    # Field definitions: (label, player_blank_hint, answer_value)
    fields = [
        ("赛位号", "________", "（裁判用·空白）"),
        ("日期", "____年__月__日", "（赛日）"),
        ("产品名称", "（可由任务背景填写，如 LoRa 核心板）", "LoRa 核心板"),
        ("层数", "（据顶/底层线路判断）", "2 层（双面板）"),
        ("外形尺寸（长×宽）", "____ mil 或 ____ mm", "80.00 × 45.00 mm（约 3150 × 1772 mil）"),
        ("最小线宽（正常区）", "____ mil（可附 mm）", "6 mil（0.15 mm）\n【正常信号线；不含故意缺陷段】"),
        ("最小线距（正常区）", "____ mil（可附 mm）", "6 mil（0.15 mm）\n【正常间距；不含故意贴线段】"),
        ("最小孔径", "____ mil（可附 mm）", "12 mil（约 0.30～0.31 mm）\n【钻孔最小工具】"),
        ("阻焊开窗方式（据阻焊/铜层对照）", "如：1:1 开窗 / 其他", "1:1 开窗（正常焊盘）"),
        (
            "文件完整性",
            "勾选缺层或写「齐全/缺××层」",
            "不齐全：缺少底层丝印层（.GBO）\n应有：GTL/GBL/GTS/GBS/GTO/GBO/GKO/钻孔",
        ),
        ("备注（可选）", "测量单位、工具说明等", "单位：mil 为主；Gerbv 测量"),
    ]

    table = doc.add_table(rows=1 + len(fields), cols=2)
    table.style = "Table Grid"
    table.autofit = True

    set_cell_text(table.rows[0].cells[0], "填写项", bold=True, size=10.5, fill="D9E2F3", center=True)
    set_cell_text(
        table.rows[0].cells[1],
        "填写内容" if not answer else "参考答案（裁判）",
        bold=True,
        size=10.5,
        fill="D9E2F3",
        center=True,
    )

    ans_fill = "FFF2CC" if answer else None
    for i, (label, blank, ans) in enumerate(fields, start=1):
        set_cell_text(table.rows[i].cells[0], label, bold=True, size=10)
        val = ans if answer else blank
        set_cell_text(
            table.rows[i].cells[1],
            val,
            size=10,
            fill=ans_fill if answer and i >= 3 else None,  # don't highlight seat/date
        )

    add_para(doc, "", space_after=6)
    add_para(doc, "三、与缺陷记录表的边界", size=12, bold=True, space_after=4)
    add_para(
        doc,
        "缺陷记录表记录「违规/缺陷」；本工艺卡记录「文件体现的正常工艺参数 + 文件齐全性」。"
        "二者不得混填（例如不得把 2 mil 缺陷线宽填为本卡最小线宽）。",
        size=10,
        space_after=8,
    )

    add_para(doc, "四、签注", size=12, bold=True, space_after=4)
    sign = doc.add_table(rows=2, cols=4)
    sign.style = "Table Grid"
    set_cell_text(sign.rows[0].cells[0], "编制（赛位号）", bold=True, size=10, fill="F2F2F2")
    set_cell_text(sign.rows[0].cells[1], "" if not answer else "—", size=10)
    set_cell_text(sign.rows[0].cells[2], "复核（裁判）", bold=True, size=10, fill="F2F2F2")
    set_cell_text(sign.rows[0].cells[3], "" if not answer else "—", size=10)
    set_cell_text(sign.rows[1].cells[0], "日期", bold=True, size=10, fill="F2F2F2")
    set_cell_text(sign.rows[1].cells[1], "", size=10)
    set_cell_text(sign.rows[1].cells[2], "得分（工艺卡）", bold=True, size=10, fill="F2F2F2")
    set_cell_text(sign.rows[1].cells[3], "" if not answer else "以评分表为准", size=10)

    if answer:
        add_para(doc, "", space_after=6)
        add_para(doc, "五、裁判给分口径（摘要）", size=12, bold=True, space_after=4)
        for line in [
            "层数：2 / 双面 → 给分",
            "外形：80×45 mm（±0.2 mm）或等价 mil → 给分",
            "最小孔径：12 mil 或 0.30～0.31 mm → 给分",
            "最小线宽/线距：正常区约 4～6 mil 量级；写成缺陷 2 mil 当规格 → 不得分或扣分",
            "阻焊开窗：1:1 / 焊盘开窗 意思对即可",
            "文件完整性：指出缺底层丝印（GBO）→ 给分；写「齐全」→ 不得分",
            "板厚/铜厚/表面处理/颜色等：本卡不考核，选手不填不扣分",
        ]:
            add_para(doc, "• " + line, size=10, space_after=2)

    return doc


def replace_table_12_fill_items(doc: Document) -> bool:
    """Replace exam TABLE that lists 12 fill items for process card."""
    target_headers = {"序号", "填写项"}
    for table in doc.tables:
        if len(table.rows) < 3 or len(table.columns) < 2:
            continue
        h0 = " ".join(table.rows[0].cells[0].text.split())
        h1 = " ".join(table.rows[0].cells[1].text.split())
        if h0 != "序号" or "填写" not in h1:
            continue
        # confirm it's the process card list (has 产品名称)
        body = "\n".join(" ".join(c.text.split()) for r in table.rows for c in r.cells)
        if "产品名称" not in body and "最小线宽" not in body:
            continue

        new_items = [
            "赛位号 / 日期",
            "产品名称（任务背景即可）",
            "层数（据顶/底线路）",
            "外形尺寸长×宽（量板框，mil 或 mm）",
            "最小线宽（正常区，mil）",
            "最小线距（正常区，mil）",
            "最小孔径（钻孔最小工具，mil）",
            "阻焊开窗方式（据阻焊/铜层对照）",
            "文件完整性（齐全或缺哪一层）",
            "备注（可选：单位/测量说明）",
        ]
        # rebuild rows: keep header, replace data rows
        # python-docx can't easily delete rows; clear and rewrite existing, add if needed
        while len(table.rows) < 1 + len(new_items):
            table.add_row()
        # clear extra rows content if any
        for i, item in enumerate(new_items, start=1):
            table.rows[i].cells[0].text = str(i)
            table.rows[i].cells[1].text = item
        for i in range(1 + len(new_items), len(table.rows)):
            for c in table.rows[i].cells:
                c.text = ""
        return True
    return False


def patch_c2_paragraphs(doc: Document) -> int:
    """Update C-2 wording in full sample exam."""
    n = 0
    replacements = [
        (
            "根据Gerber文件中的信息，结合板厂工艺能力，填写制程工艺卡的基本信息。",
            "根据Gerber与钻孔文件实测，填写简版制程工艺卡。"
            "本卡仅填写能量/能数清的参数（层数、外形、正常区最小线宽与线距、最小孔径、阻焊开窗方式、文件完整性等）；"
            "不要求填写板厚、铜厚、板材、表面处理、阻焊颜色等订单约定项。",
        ),
        (
            "根据 Gerber 文件中的信息，结合板厂工艺能力，填写制程工艺卡的基本信息。",
            "根据Gerber与钻孔文件实测，填写简版制程工艺卡。"
            "本卡仅填写能量/能数清的参数（层数、外形、正常区最小线宽与线距、最小孔径、阻焊开窗方式、文件完整性等）；"
            "不要求填写板厚、铜厚、板材、表面处理、阻焊颜色等订单约定项。",
        ),
    ]
    for p in doc.paragraphs:
        t = p.text
        for old, new in replacements:
            if old in t:
                # replace full paragraph text preserving first run style if possible
                if p.runs:
                    p.runs[0].text = t.replace(old, new)
                    for r in p.runs[1:]:
                        r.text = ""
                else:
                    p.add_run(t.replace(old, new))
                n += 1
                break
    return n


def main():
    PACK.mkdir(parents=True, exist_ok=True)

    player = build_card(answer=False)
    player.save(OUT_PLAYER)
    print("wrote", OUT_PLAYER)

    ans = build_card(answer=True)
    ans.save(OUT_ANSWER)
    print("wrote", OUT_ANSWER)

    if OUT_DUP_REMOVE.exists():
        OUT_DUP_REMOVE.unlink()
        print("removed", OUT_DUP_REMOVE)

    if FULL_EXAM.exists():
        exam = Document(str(FULL_EXAM))
        n = patch_c2_paragraphs(exam)
        ok = replace_table_12_fill_items(exam)
        exam.save(str(FULL_EXAM))
        print(f"patched full exam: para={n}, table12={ok}")
    else:
        print("full exam missing, skip")


if __name__ == "__main__":
    main()
