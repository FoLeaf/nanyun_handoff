# -*- coding: utf-8 -*-
"""
模块D 题干嵌入：除《检测记录表》外，文字/表格素材并入完整版赛题。
- 保留独立 docx：学生组_模块D_检测记录表.docx（选手提交）
- 其余：验收/尺寸/电气规则/分区/安全/材料清单 → 嵌入完整版模块D
"""
from __future__ import annotations

import re
import shutil
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[2]
SAMPLE = ROOT / "02_样题"
PACK = SAMPLE / "学生组_模块D"
PLAYER = PACK / "选手"
FULL = SAMPLE / "印制电路制作工赛项_竞赛样题（完整版）.docx"
BACKUP = ROOT / "05_归档备份" / f"module-d-embed-stem-{datetime.now().strftime('%Y%m%d_%H%M%S')}"

C_BLUE = "2E75B6"


def set_run_font(run, name_cn="仿宋", name_en="Times New Roman", size_pt=12, bold=False, color_hex=None):
    run.bold = bold
    run.font.size = Pt(size_pt) if size_pt else None
    run.font.name = name_en
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:ascii"), name_en)
    rFonts.set(qn("w:hAnsi"), name_en)
    rFonts.set(qn("w:eastAsia"), name_cn)
    rFonts.set(qn("w:cs"), name_en)
    if color_hex:
        run.font.color.rgb = RGBColor.from_string(color_hex)


def add_para(doc, text, *, cn="仿宋", en="Times New Roman", size=12, bold=False, space_after=4, space_before=0):
    p = doc.add_paragraph()
    try:
        p.style = doc.styles["Normal"]
    except Exception:
        pass
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(space_before)
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    run = p.add_run(text)
    set_run_font(run, name_cn=cn, name_en=en if cn != "黑体" else "黑体", size_pt=size, bold=bold)
    return p


def shade_cell(cell, fill_hex):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill_hex)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def set_cell(cell, text, *, header=False, size=9.5):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if header else WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(str(text))
    set_run_font(
        run,
        name_cn="黑体" if header else "仿宋",
        name_en="Times New Roman",
        size_pt=size if not header else 9,
        bold=header,
        color_hex="FFFFFF" if header else None,
    )
    if header:
        shade_cell(cell, C_BLUE)


def add_table(doc, rows):
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            set_cell(table.cell(ri, ci), val, header=(ri == 0))
    return table


def delete_paragraph(paragraph):
    el = paragraph._element
    el.getparent().remove(el)


def find_module_d_range(doc):
    start = None
    for i, p in enumerate(doc.paragraphs):
        t = p.text.strip()
        if t.startswith("模块D：") and ("成品" in t or "质量检测" in t):
            start = i
            break
    if start is None:
        raise RuntimeError("未找到模块D标题")
    # 到文末（当前文件模块D为最后模块）
    end = len(doc.paragraphs)
    return start, end


def clear_range(doc, start, end):
    # 从后往前删，避免索引错位
    for i in range(end - 1, start - 1, -1):
        delete_paragraph(doc.paragraphs[i])


def insert_module_d(doc):
    """在文档末尾追加完整模块D（清除旧段落后调用）。"""
    # 标题
    add_para(doc, "模块D：成品PCB裸板质量检测与判定（学生组）", cn="黑体", size=16, bold=True, space_before=8, space_after=6)

    add_para(doc, "（一）模块考核点", cn="黑体", size=14, bold=True, space_before=6)
    add_para(
        doc,
        "本模块（学生组）考核选手对双面成品PCB裸板（无元器件，已完成线路、阻焊、丝印和表面处理）进行"
        "外观缺陷检测、关键尺寸与指定点微几何测量、基础电气开短路点测、缺陷分类与单件质量处置（接收/返工/报废）的能力。"
        "正式考核时长45分钟（D-1～D-5）；模块内原始分40分（外观14／尺寸·微几何12／电气8／综合4／规范2；折算见技术工作文件）。"
        "不考核复杂网络分段、故障定位树、应通应断全表、上电功能调试、实际返修、贴装/焊接质量及Gerber量测。",
    )

    add_para(doc, "（二）模块简介", cn="黑体", size=14, bold=True, space_before=6)
    add_para(doc, "【任务背景】", cn="黑体", size=12, bold=True)
    add_para(
        doc,
        "本模块为独立质检任务。检验对象为赛场提供的双面成品裸板质检样件（专用检测板），"
        "与模块B（EDA工程设计）、模块C（CAM审核与工艺文件编制）相互独立；"
        "仅做基础开短路点测，不做故障定位与网络分段，不上电，不量测Gerber。"
        "某批次双面成品裸板完成后进入抽检环节，你作为质检人员，需完成规范检验、准确测量、基础点测，"
        "对照本卷题干中的验收要求、尺寸基准与电气判据进行分类判定，并给出接收／返工／报废结论及简要依据。",
    )
    add_para(doc, "【检验对象】", cn="黑体", size=12, bold=True)
    add_para(
        doc,
        "双面成品裸板质检样件1块（无元器件）。外形约80×60 mm，标称板厚1.6 mm；FR-4双面、绿阻焊白丝印、无铅喷锡。"
        "板面丝印含版本标识（MOD-D-S-A或MOD-D-S-B）、A1–D4分区网格、特征名及专用测试点（TP1–TP4）。"
        "须完成：外观必检5处（V01–V05）、尺寸M01–M04与指定点线宽/线距M05、基础电气E01开路与E02短路；"
        "可能含1～2处临界合格干扰特征（不计必检5）。须检查顶面与底面（底面缺陷不少于2处）。",
    )
    add_para(doc, "【提供材料】", cn="黑体", size=12, bold=True)
    add_para(doc, "1. 双面成品裸板质检样件1块；")
    add_para(doc, "2. 《学生组_模块D_检测记录表》（空白，选手填写并提交的唯一表格附件）；")
    add_para(doc, "3. 本卷模块D题干内嵌素材：板面分区与测试点、尺寸基准、外观验收要求、基础电气判据、仪器与安全、材料与工具清单。")
    add_para(doc, "【使用工具】", cn="黑体", size=12, bold=True)
    add_para(
        doc,
        "游标卡尺、厚度规（或千分尺）、测量显微镜（10～20×或带刻度放大镜）、放大镜、数字万用表、侧光/照明、防静电用品等（以赛场清单为准）。",
    )

    # —— 嵌入素材 ——
    add_para(doc, "（三）题干素材（直接使用，无需另附分册）", cn="黑体", size=14, bold=True, space_before=8)
    add_para(
        doc,
        "下列表格与说明为本模块试题组成部分。选手依据本卷填写《检测记录表》并提交；"
        "除检测记录表外，不再另发验收表/尺寸表/电气表/分区说明等文字表格附件。",
        space_after=6,
    )

    # 3.1 分区与测试点
    add_para(doc, "1. 板面分区与测试点说明", cn="黑体", size=12, bold=True, space_before=6)
    add_para(doc, "版本丝印：MOD-D-S-A / MOD-D-S-B（以实物为准）。答题版本栏须与丝印一致。")
    add_para(doc, "分区网格：顶层参考A1–D4（列1–4，行A–D）。记录缺陷时填写所在面＋网格/特征名＋方位。")
    add_para(doc, "特征区名称（示例，以板面丝印为准）：走线区、孔阵、阻焊对比区、丝印条、板框、测试点区。")
    add_para(doc, "专用测试点仅服务基础电气E01/E02，不得理解为多网络故障定位表：")
    add_table(
        doc,
        [
            ["丝印", "用途", "说明"],
            ["TP1", "E01一端", "开路网络一端"],
            ["TP2", "E01另一端", "开路网络另一端"],
            ["TP3", "E02一端", "短路网络一端"],
            ["TP4", "E02另一端", "短路网络另一端"],
        ],
    )
    add_para(doc, "", space_after=2)

    # 3.2 尺寸基准
    add_para(doc, "2. 尺寸基准要求（M01–M05）", cn="黑体", size=12, bold=True, space_before=6)
    add_para(
        doc,
        "单位mm。M01–M04用卡尺/厚度规；M05用测量显微镜测指定点（与V02不同点）。"
        "禁止普通卡尺测0.30 mm级小孔；禁止无指定点全板扫微距。"
        "名义与公差为竞赛骨架，实物以赛场样件为准；判定时实测与下表比对。",
    )
    add_table(
        doc,
        [
            ["编号", "项目", "工具", "名义值", "允许误差", "判定"],
            ["M01", "板长L", "游标卡尺", "80.00", "±0.20", "合格/不合格"],
            ["M02", "板宽W", "游标卡尺", "60.00", "±0.20", "合格/不合格"],
            ["M03", "板厚T", "厚度规/千分尺", "1.60", "±0.15", "合格/不合格"],
            ["M04", "定位孔径", "卡尺/孔规", "2.00", "±0.10", "合格/不合格"],
            ["M05", "指定点线宽或线距", "测量显微镜", "见板面指定点丝印/特征名", "±0.05或对照阈值", "合格/不合格"],
        ],
    )
    add_para(
        doc,
        "M05指定点：以板面丝印或特征名唯一标识的测量位为准（与V02底层线宽变窄外观点不得为同一采分点）。"
        "测量对象勾选线宽或线距其一；读数保留至0.01 mm（或赛场规定）。",
        space_before=4,
    )

    # 3.3 外观验收
    add_para(doc, "3. 外观缺陷验收要求（F1 · V01–V05）", cn="黑体", size=12, bold=True, space_before=6)
    add_para(doc, "下列为竞赛用不合格判定表述（可与同义表述对应，见裁判培训材料）。须检顶底两面；底面缺陷不少于2处。")
    add_table(
        doc,
        [
            ["编号", "类型", "不合格判定（竞赛用表述）", "备注"],
            ["V01", "顶层线路缺口", "顶层导体出现可见断开/缺口，存在连通风险", "须定位顶面"],
            ["V02", "底层线宽变窄", "底层局部线宽明显变窄（外观定性，不要求在此处读数）", "与M05不同点"],
            ["V03", "孔破盘", "孔与焊盘关系异常导致破盘/焊盘缺损", "顶或底以实物为准"],
            ["V04", "阻焊未开窗或开窗偏移", "应开窗处被阻焊覆盖，或开窗相对焊盘明显偏移", "二选一现象"],
            ["V05", "丝印缺失或压焊盘", "字符缺失/残缺，或丝印压覆焊盘", "二选一现象"],
        ],
    )
    add_para(doc, "干扰项：可出现1～2处接近合格临界的特征，不计入必检5；误报不强制重扣外观分（评分细则另定）。", space_before=4)

    # 3.4 电气
    add_para(doc, "4. 基础电气检测要求（E01–E02）", cn="黑体", size=12, bold=True, space_before=6)
    add_para(
        doc,
        "仅对指定测试点做点对点开路/短路判定。不做网络分段、故障定位树、应通应断全表、高阻分析、上电功能测试。"
        "电阻判据骨架如下（赛场书面说明或实物标定优先）：",
    )
    add_table(
        doc,
        [
            ["编号", "类型", "测试点对", "预期", "默认电阻判据"],
            ["E01", "开路", "TP1–TP2", "开路", "≥1 MΩ 或表显OL/开路"],
            ["E02", "短路", "TP3–TP4", "短路", "≤1 Ω 或蜂鸣导通"],
        ],
    )
    add_para(
        doc,
        "记录要求：测试点对、档位、电阻读数、单位、判定（开路/短路/正常）、与预期是否符合。",
        space_before=4,
    )

    # 3.5 综合处置
    add_para(doc, "5. 综合质量处置规则（接收／返工／报废）", cn="黑体", size=12, bold=True, space_before=6)
    add_para(doc, "接收：无必检致命缺陷，尺寸M01–M05与电气E01/E02均符合本卷要求。")
    add_para(doc, "返工：存在可返工类缺陷（如丝印、部分阻焊类）且尺寸/电气主体可接受——按本卷验收表述与裁判细则。")
    add_para(doc, "报废：存在线路缺口、孔破盘等严重影响功能/可靠性，或电气与预期严重不符且不可接受。")
    add_para(doc, "选手须在检测记录表给出结论＋简要依据（对照本卷第3节条款/现象名）。不要求深层工艺根因分析。")

    # 3.6 仪器安全
    add_para(doc, "6. 仪器使用与安全要点", cn="黑体", size=12, bold=True, space_before=6)
    add_para(doc, "（1）卡尺使用前对零；测量板厚避开铜瘤/丝印厚堆；定位孔测量方法全卷统一。")
    add_para(doc, "（2）显微镜仅用于M05指定点及必要时外观确认；爱护光学部件与样件。")
    add_para(doc, "（3）万用表仅用于指定TP点测；注意档位，禁止对样件进行上电功能测试。")
    add_para(doc, "（4）防静电、轻拿轻放；如遇设备故障举手示意裁判。")

    # 3.7 材料工具清单（题干内）
    add_para(doc, "7. 本模块材料与工具清单（摘要）", cn="黑体", size=12, bold=True, space_before=6)
    add_table(
        doc,
        [
            ["类别", "名称", "备注"],
            ["样件", "双面成品裸板质检样件", "A或B版；无元器件"],
            ["提交表", "学生组_模块D_检测记录表", "唯一须提交的表格附件"],
            ["题干素材", "本卷第（三）节全部表格与说明", "已嵌入，不另附分册"],
            ["工具", "卡尺/厚度规/显微镜/放大镜/万用表等", "以赛场清单为准"],
        ],
    )

    # （四）任务
    add_para(doc, "（四）模块任务", cn="黑体", size=14, bold=True, space_before=8)
    add_para(
        doc,
        "请在45分钟内完成D-1～D-5，并将结果填写在《学生组_模块D_检测记录表》中提交。"
        "判定一律以本卷第（三）节题干素材为准。",
        space_after=6,
    )

    add_para(doc, "D-1  检验准备（建议3分钟）", cn="黑体", size=12, bold=True, space_before=4)
    add_para(doc, "核对样件编号与版本丝印、确认顶/底方向与分区约定；检查卡尺零位、厚度规、放大镜/显微镜、照明及万用表；在检测记录表填写基本信息并勾选准备项。")
    add_para(doc, "（1）核对样件编号、版本丝印（MOD-D-S-A / MOD-D-S-B）与记录表；")
    add_para(doc, "（2）确认顶/底方向与本卷分区说明一致；")
    add_para(doc, "（3）检查量具与万用表可用；")
    add_para(doc, "（4）勾选准备项；如有运输损伤先记录。")

    add_para(doc, "D-2  外观缺陷检测（建议12分钟）", cn="黑体", size=12, bold=True, space_before=6)
    add_para(
        doc,
        "对样件线路、焊盘/孔、阻焊、丝印等进行外观检查（须检顶面与底面），完成5处必检（V01–V05），"
        "按本卷第（三）节第3款判定，将结果记入检测记录表。可结合分区网格定位；干扰项可不计入必检5。",
    )
    add_para(doc, "（1）每发现一处缺陷填写一行：所在面、类型、网格/特征、现象、严重程度等；")
    add_para(doc, "（2）位置描述应便于复核（如“顶层B2/走线区/中部”），无需精确坐标；")
    add_para(doc, "（3）类型应具体（如“线路缺口”“孔破盘”），避免仅写“线路有问题”。")

    add_para(doc, "D-3  尺寸与指定点微几何测量（建议10分钟）", cn="黑体", size=12, bold=True, space_before=6)
    add_para(doc, "按本卷第（三）节第2款测量M01–M05，填写实测值并判定合格/不合格。")
    add_para(doc, "（1）M01板长；（2）M02板宽；（3）M03板厚；（4）M04定位孔径；（5）M05指定点线宽或线距（显微镜）。")
    add_para(doc, "不要求测量普通卡尺无法可靠完成的项目（如0.30 mm级微孔）。记录以mm为主。")

    add_para(doc, "D-4  基础电气检测（建议8分钟）", cn="黑体", size=12, bold=True, space_before=6)
    add_para(doc, "按本卷第（三）节第4款，在TP1–TP4完成E01开路、E02短路点测，记录电阻与判定。")
    add_para(doc, "仅做指定测试点基础开短路；不做故障定位、网络分段、应通应断全表、高阻分析与上电功能测试。")

    add_para(doc, "D-5  分类判定与提交（建议12分钟）", cn="黑体", size=12, bold=True, space_before=6)
    add_para(doc, "（1）对已记录外观缺陷按类型分类计数；")
    add_para(doc, "（2）汇总尺寸M01–M05是否全部满足基准；")
    add_para(doc, "（3）汇总电气E01/E02结果；")
    add_para(doc, "（4）对照本卷第（三）节第3、5款，给出接收/返工/报废结论及简要依据；")
    add_para(doc, "（5）不要求深层制造工艺根因分析；")
    add_para(
        doc,
        "（6）检查记录完整性后提交《学生组_模块D_检测记录表》；如需电子版，示例目录D:\\提交资料\\模块D\\，命名示例：赛位号_模块D。",
    )

    add_para(doc, "注意事项", cn="黑体", size=12, bold=True, space_before=8)
    add_para(doc, "1. 请合理分配D-1～D-5时间（建议合计45分钟）；时长与评分以技术工作文件/竞赛平台为准；")
    add_para(doc, "2. 本模块与模块B/C独立；仅做基础开短路点测，不做故障定位与网络分段，不上电，不量测Gerber；")
    add_para(doc, "3. 判定依据以本卷题干素材为准；检测记录表为唯一提交表格附件；")
    add_para(doc, "4. 爱护样件与仪器；按安全规范使用防静电措施与万用表；")
    add_para(doc, "5. 成果上不得标注姓名等身份信息；如遇设备故障请举手示意裁判。")


def update_package_docs():
    # 现行包说明
    note = SAMPLE / "00_学生组模块D现行包说明.md"
    if note.exists():
        text = note.read_text(encoding="utf-8")
    else:
        text = ""
    embed_note = """
## 下发结构（2026-07-21 调整）

| 类型 | 文件 | 说明 |
|------|------|------|
| **赛题题干** | `印制电路制作工赛项_竞赛样题（完整版）.docx` 模块D | 内嵌分区/尺寸基准/验收/电气判据/安全/清单等**全部文字表格素材** |
| **选手提交** | `学生组_模块D/选手/学生组_模块D_检测记录表.docx` | **唯一**单独下发的填写/提交表 |
| 裁判内部 | `学生组_模块D/裁判/*` | 不向选手下发 |
| 命题规格 | `学生组_模块D/检测板/*`、边界确认表等 | 不向选手下发 |

> 原单独印发的验收要求表、尺寸基准表、基础电气检测表、板面分区说明、选手材料清单、仪器安全等，**内容已并入完整版模块D第（三）节**；包内对应 docx 可作命题底稿/备份，**正式赛不对选手另附**（除非赛场另有通知）。
"""
    # append or replace section
    if "下发结构（2026-07-21 调整）" not in text:
        text = text.rstrip() + "\n" + embed_note
        note.write_text(text, encoding="utf-8")
        print("[OK] updated", note.name)

    # 选手目录 README
    readme = PLAYER / "00_下发说明.md"
    readme.write_text(
        """# 选手侧下发说明（学生组·模块D）

> 2026-07-21

## 正式对选手下发

1. **赛题**：完整版样题中的**模块D全文**（题干已内嵌分区、尺寸基准、验收、电气判据、仪器安全、清单等）
2. **提交表**：`学生组_模块D_检测记录表.docx`（唯一单独表格附件）
3. **实物**：双面成品裸板质检样件 + 赛场量具

## 本目录其他 docx / md

| 文件 | 状态 |
|------|------|
| 检测记录表.docx | **现行提交表** |
| 验收要求表 / 尺寸基准表 / 基础电气 / 分区说明 / 材料清单 | 内容已嵌入完整版题干；保留作命题底稿，**默认不另发** |
| 仪器使用与安全要求.md | 要点已嵌入题干；可作培训印发 |
| 任务书要点.md | 命题/培训摘要，非选手必发 |

## 提交

仅提交填写完整的《检测记录表》（纸质/电子按赛场要求）。命名示例：`赛位号_模块D`。
""",
        encoding="utf-8",
    )
    print("[OK]", readme)

    # 更新镜像 md
    mirror = PACK / "完整版学生组模块D正文稿_R1.md"
    # 简短指针，避免与 docx 双源长期漂移——写清以 docx 为准
    mirror.write_text(
        """# 完整版样题 · 学生组模块D（R1 · 题干嵌入版）

> **权威正文**：`02_样题/印制电路制作工赛项_竞赛样题（完整版）.docx` 模块D 段
> 本文件为结构说明，不替代 docx。

## 结构

1. （一）模块考核点
2. （二）模块简介（任务背景／检验对象／提供材料／工具）
3. **（三）题干素材（嵌入）**
   - 板面分区与测试点
   - 尺寸基准 M01–M05
   - 外观验收 V01–V05
   - 基础电气 E01–E02
   - 综合处置规则
   - 仪器与安全
   - 材料工具清单摘要
4. （四）模块任务 D-1～D-5
5. 注意事项

## 下发原则

- **唯一单独表格附件**：《检测记录表》
- 其余文字/表格素材：**全部在题干第（三）节**，不另附分册

## 时长与口径

- 45 分钟；R1；基础开短路；非故障定位
""",
        encoding="utf-8",
    )
    print("[OK] mirror note")

    # 任务书要点同步
    tp = PLAYER / "任务书要点.md"
    if tp.exists():
        body = tp.read_text(encoding="utf-8")
        body = re.sub(
            r"## 4\. 成果物.*?(?=## 5\.)",
            """## 4. 成果物（须提交）

1. **《学生组_模块D_检测记录表》**（唯一单独表格附件；含 D-1～D-5）
2. 判定依据为**完整版样题模块D第（三）节题干素材**（分区/尺寸/验收/电气/安全已嵌入，不另附分册）
3. 纸质/电子按赛场要求；命名示例 `赛位号_模块D`

""",
            body,
            count=1,
            flags=re.S,
        )
        body = re.sub(
            r"## 7\. 随卷材料索引.*",
            """## 7. 随卷材料索引

| 材料 | 形式 |
|------|------|
| 赛题模块D全文（含题干素材表） | 完整版样题内嵌 |
| 检测记录表 | **单独 docx，提交用** |
| 样件与量具 | 赛场提供 |

""",
            body,
            count=1,
            flags=re.S,
        )
        tp.write_text(body, encoding="utf-8")
        print("[OK] 任务书要点")


def main():
    BACKUP.mkdir(parents=True, exist_ok=True)
    shutil.copy2(FULL, BACKUP / FULL.name)
    print("[BACKUP]", BACKUP)

    doc = Document(str(FULL))
    start, end = find_module_d_range(doc)
    print(f"[FULL] clear Module D paras [{start}, {end})")
    clear_range(doc, start, end)
    insert_module_d(doc)
    doc.save(str(FULL))
    print("[OK] full exam Module D embedded stem")

    update_package_docs()

    # 轻量校验
    d2 = Document(str(FULL))
    text = "\n".join(p.text for p in d2.paragraphs)
    for key in ["题干素材", "M01", "V01", "E01", "TP1", "检测记录表", "不另发", "45"]:
        # 45 may appear as 45分钟
        ok = key in text or (key == "45" and "45分钟" in text)
        print(("OK" if ok else "MISS"), key)
    print("tables", len(d2.tables))
    print("DONE")


if __name__ == "__main__":
    main()
