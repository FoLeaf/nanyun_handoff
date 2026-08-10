# -*- coding: utf-8 -*-
"""
样题版式交叉对齐（2026-07-21）

目标：
1) 完整版样题 · 学生组模块D 段：对齐同文件模块 B/C 的「黑体标题 + 正文继承/仿宋」体系
2) 学生组_模块D 选手/裁判附件 docx：对齐模块 B/C 独立样题的「微软雅黑 + 蓝灰标题色」体系

不改正文口径（R1 内容保持）。
"""
from __future__ import annotations

import re
import shutil
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[2]
SAMPLE = ROOT / "02_样题"
PACK = SAMPLE / "学生组_模块D"
PLAYER = PACK / "选手"
JUDGE = PACK / "裁判"
FULL = SAMPLE / "印制电路制作工赛项_竞赛样题（完整版）.docx"
BACKUP = ROOT / "05_归档备份" / f"module-d-format-align-{datetime.now().strftime('%Y%m%d_%H%M%S')}"

# B/C 独立附件色板
C_NAVY = "1F3864"
C_BLUE = "2E75B6"
C_GRAY = "595959"
C_RED = "C00000"


def set_run_font(
    run,
    *,
    name_cn="微软雅黑",
    name_en="微软雅黑",
    size_pt=10.5,
    bold=False,
    color_hex=None,
):
    run.bold = bold
    run.font.size = Pt(size_pt) if size_pt is not None else None
    run.font.name = name_en
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:ascii"), name_en)
    rFonts.set(qn("w:hAnsi"), name_en)
    rFonts.set(qn("w:eastAsia"), name_cn)
    rFonts.set(qn("w:cs"), name_en)
    if color_hex:
        run.font.color.rgb = RGBColor.from_string(color_hex)
    else:
        # 清主题色/直设色，交给自动黑
        try:
            run.font.color.rgb = None
        except Exception:
            pass


def clear_and_set_para(p, text, **font_kw):
    """Replace段落全部 runs 为单一格式。"""
    # 保留段落样式名，但清直接格式 runs
    for r in list(p.runs):
        r._element.getparent().remove(r._element)
    # 若无 runs，add_run
    run = p.add_run(text)
    set_run_font(run, **font_kw)
    return p


def classify_full_d_para(text: str) -> str:
    t = text.strip()
    if not t:
        return "empty"
    if re.match(r"^模块D[：:]", t):
        return "module_title"
    if re.match(r"^（[一二三四五六七八九十]+）", t):
        return "section"
    if re.match(r"^【.+】$", t):
        return "label"
    if re.match(r"^D-\d+", t):
        return "task_title"
    if t in ("准备要求", "填写要求", "测量要求", "判定要求", "提交要求") or t.endswith("要求") and len(t) <= 8:
        return "sub_label"
    if re.match(r"^（\d+）", t) or re.match(r"^\d+[\.、]", t):
        return "list"
    if t.startswith("注意事项") or t.startswith("注意："):
        return "note_title"
    return "body"


def fix_full_exam_module_d():
    """完整版：模块D 段对齐模块C（黑体16/14 + 正文不强制微软雅黑，用仿宋/继承）。"""
    doc = Document(str(FULL))
    # 定位模块D起止
    start = None
    end = None
    for i, p in enumerate(doc.paragraphs):
        t = p.text.strip()
        if start is None and t.startswith("模块D") and ("成品" in t or "质量检测" in t or "学生组" in t):
            start = i
            continue
        if start is not None and i > start:
            # 下一顶级模块或附录
            if re.match(r"^模块[A-E][：:]", t) and not t.startswith("模块D"):
                end = i
                break
            if t.startswith("附录") or t.startswith("评分标准") and "模块D" not in t:
                # 保守：仅当明显离开 D
                pass
    if start is None:
        raise RuntimeError("完整版中未找到学生组模块D标题")
    if end is None:
        end = len(doc.paragraphs)

    print(f"[FULL] Module D paragraphs [{start}, {end})")
    for i in range(start, end):
        p = doc.paragraphs[i]
        t = p.text
        if not t.strip():
            continue
        kind = classify_full_d_para(t)
        # 统一用 Normal，避免 Heading 主题蓝
        try:
            p.style = doc.styles["Normal"]
        except Exception:
            pass
        pf = p.paragraph_format
        pf.line_spacing_rule = WD_LINE_SPACING.SINGLE

        if kind == "module_title":
            clear_and_set_para(
                p, t.strip(), name_cn="黑体", name_en="黑体", size_pt=16, bold=True
            )
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        elif kind == "section":
            clear_and_set_para(
                p, t.strip(), name_cn="黑体", name_en="黑体", size_pt=14, bold=True
            )
        elif kind in ("label", "task_title", "sub_label", "note_title"):
            clear_and_set_para(
                p, t.strip(), name_cn="黑体", name_en="黑体", size_pt=12, bold=True
            )
        elif kind == "list":
            # 列表：仿宋 12，与完整版其它模块正文一致（不强制 10.5 微软雅黑）
            clear_and_set_para(
                p, t.strip(), name_cn="仿宋", name_en="Times New Roman", size_pt=12, bold=False
            )
        else:
            # body：仿宋 12
            clear_and_set_para(
                p, t.strip(), name_cn="仿宋", name_en="Times New Roman", size_pt=12, bold=False
            )

    doc.save(str(FULL))
    print("[OK] full exam Module D format aligned to B/C-in-full style")


def add_para(doc, text, *, cn="微软雅黑", en=None, size=10.5, bold=False, color=None, align=None, space_after=6, space_before=0):
    en = en or cn
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(space_before)
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    run = p.add_run(text)
    set_run_font(run, name_cn=cn, name_en=en, size_pt=size, bold=bold, color_hex=color)
    return p


def set_cell(cell, text, *, cn="微软雅黑", size=9.5, bold=False, color=None, center=False):
    cell.text = ""
    p = cell.paragraphs[0]
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(str(text))
    set_run_font(run, name_cn=cn, name_en=cn, size_pt=size, bold=bold, color_hex=color)


def fill_table(table, rows, header=True):
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            is_h = header and ri == 0
            set_cell(
                table.cell(ri, ci),
                val,
                cn="微软雅黑",
                size=9 if is_h else 9.5,
                bold=is_h,
                color="FFFFFF" if is_h else None,
                center=is_h,
            )
            if is_h:
                # 表头底色
                from docx.oxml import OxmlElement

                tc = table.cell(ri, ci)._tc
                tcPr = tc.get_or_add_tcPr()
                shd = OxmlElement("w:shd")
                shd.set(qn("w:fill"), C_BLUE)
                shd.set(qn("w:val"), "clear")
                tcPr.append(shd)


def setup_page(doc):
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.left_margin = Cm(2.0)
    sec.right_margin = Cm(2.0)
    sec.top_margin = Cm(1.8)
    sec.bottom_margin = Cm(1.8)


def new_doc():
    doc = Document()
    setup_page(doc)
    return doc


def add_bc_header(doc, module_title: str, subtitle: str = "（竞赛样题 · 学生组）"):
    add_para(doc, '江西省"振兴杯"职业技能大赛', cn="黑体", size=20, bold=True, color=C_NAVY, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    add_para(doc, '"印制电路制作工"赛项', cn="黑体", size=18, bold=True, color=C_NAVY, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    add_para(doc, "━━━━━━━━━━━━━━━━━━━━", cn="微软雅黑", size=12, color=C_BLUE, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
    add_para(doc, module_title, cn="黑体", size=28, bold=True, color=C_RED, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    add_para(doc, subtitle, cn="微软雅黑", size=14, bold=True, color=C_GRAY, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
    add_para(doc, "━━━━━━━━━━━━━━━━━━━━", cn="微软雅黑", size=12, color=C_BLUE, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=8)


def rebuild_player_docs():
    """按 B/C 附件视觉规范重写选手侧全部 docx（内容保持 R1）。"""
    # 1 检测记录表
    doc = new_doc()
    add_bc_header(doc, "模块D：成品PCB裸板质量检测与判定", "检测记录表（学生组 · R1 / 45分钟）")
    add_para(doc, "模块说明", cn="黑体", size=14, bold=True, color=C_NAVY, space_before=4)
    add_para(
        doc,
        "⚠ 注意：结合《验收要求表》《尺寸基准表》《基础电气检测表》《板面分区说明》填写。"
        "正式时长 45 分钟（D-1～D-5）。电子命名示例：赛位号_模块D。不得标注姓名（按赛场规则）。"
        "电气仅为点对点开短路，非故障定位。",
        size=10,
        bold=True,
        color=C_RED,
        space_after=8,
    )

    t0 = doc.add_table(rows=3, cols=4)
    t0.style = "Table Grid"
    fill_table(
        t0,
        [
            ["赛位号", "", "日期", ""],
            ["样件编号", "", "版本（丝印）", "□A(MOD-D-S-A)  □B(MOD-D-S-B)"],
            ["板面方向确认", "顶/底已确认□", "记录格式提示", "版本-编号，如 A-V01"],
        ],
        header=False,
    )
    for row in t0.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for r in p.runs:
                    set_run_font(r, name_cn="微软雅黑", name_en="微软雅黑", size_pt=9.5, bold=False)

    add_para(doc, "一、D-1 检验准备（建议 3 min）", cn="黑体", size=14, bold=True, color=C_NAVY, space_before=10)
    add_para(
        doc,
        "□ 文件齐全（记录表/验收表/尺寸基准/基础电气/分区说明）  "
        "□ 卡尺零位已检  □ 厚度规可用  □ 测量显微镜/放大镜可用  "
        "□ 万用表电阻档可用  □ 样件无异常运输损伤（如有：________）  "
        "□ 版本丝印已抄录",
        size=10.5,
    )

    add_para(doc, "二、D-2 外观缺陷检测记录（V01–V05 必检，建议 12 min）", cn="黑体", size=14, bold=True, color=C_NAVY, space_before=8)
    add_para(
        doc,
        "须检顶底两面。必检类型：V01 顶层线路缺口；V02 底层线宽变窄（定性）；V03 孔破盘；"
        "V04 阻焊未开窗或开窗偏移；V05 丝印缺失或压焊盘。干扰项可不计必检。记录格式：版本-编号。",
        size=10.5,
    )
    t1 = doc.add_table(rows=9, cols=8)
    t1.style = "Table Grid"
    rows = [[
        "编号", "所在面", "网格区", "特征名", "缺陷类型（标准名/同义）", "现象简述", "严重程度", "是否必检",
    ]]
    for r in [
        ("V01", "顶", "", "走线区", "顶层线路缺口", "", "轻/中/重", "是"),
        ("V02", "底", "", "走线区", "底层线宽变窄", "", "轻/中/重", "是"),
        ("V03", "底", "", "孔阵", "孔破盘", "", "轻/中/重", "是"),
        ("V04", "顶", "", "阻焊对比区", "阻焊未开窗/开窗偏移", "", "轻/中/重", "是"),
        ("V05", "", "", "丝印条", "丝印缺失/压焊盘", "", "轻/中/重", "是"),
    ]:
        rows.append(list(r))
    for _ in range(3):
        rows.append(["", "", "", "", "", "", "", "否/干扰"])
    fill_table(t1, rows)

    add_para(doc, "三、D-3 尺寸与指定点微几何（M01–M05，建议 10 min）", cn="黑体", size=14, bold=True, color=C_NAVY, space_before=8)
    add_para(
        doc,
        "单位 mm。M01–M04 用卡尺/厚度规；M05 用测量显微镜测指定点（与 V02 不同点）。"
        "禁止普通卡尺测 0.30 mm 级小孔；禁止无指定点全板扫微距。",
        size=10.5,
    )
    t2 = doc.add_table(rows=6, cols=6)
    t2.style = "Table Grid"
    fill_table(
        t2,
        [
            ["编号", "测量项目", "工具", "基准/公差（见尺寸基准表）", "实测值(mm)", "合格/不合格"],
            ["M01", "板长 L", "游标卡尺", "80.00±0.20", "", ""],
            ["M02", "板宽 W", "游标卡尺", "60.00±0.20", "", ""],
            ["M03", "板厚 T", "厚度规/千分尺", "1.60±0.15", "", ""],
            ["M04", "定位孔径", "卡尺/孔规", "2.00±0.10", "", ""],
            ["M05", "指定点线宽或线距", "测量显微镜", "见尺寸基准表指定点", "", ""],
        ],
    )
    add_para(
        doc,
        "M05 指定点特征名/位置：____________________  测量对象：□线宽  □线距    读数方法：____________________",
        size=10.5,
        space_before=4,
    )

    add_para(doc, "四、D-4 基础电气（E01–E02，建议 8 min）", cn="黑体", size=14, bold=True, color=C_NAVY, space_before=8)
    add_para(
        doc,
        "仅测指定测试点对。默认判据：开路 ≥1 MΩ 或 OL；短路 ≤1 Ω 或蜂鸣导通。"
        "详见《基础电气检测表》。不做故障定位、网络分段、上电。",
        size=10.5,
    )
    t3 = doc.add_table(rows=3, cols=8)
    t3.style = "Table Grid"
    fill_table(
        t3,
        [
            ["编号", "类型", "测试点对", "档位", "电阻读数", "单位", "判定(开路/短路/正常)", "与预期符合"],
            ["E01", "开路", "TP1 – TP2", "", "", "", "", "□是 □否"],
            ["E02", "短路", "TP3 – TP4", "", "", "", "", "□是 □否"],
        ],
    )

    add_para(doc, "五、D-5 分类、综合判定与复核（建议 12 min）", cn="黑体", size=14, bold=True, color=C_NAVY, space_before=8)
    for line in [
        "1. 外观必检计数：V01____ V02____ V03____ V04____ V05____；已发现必检合计____/5；干扰项____。",
        "2. 分类汇总：线路类____；孔/盘类____；阻焊类____；丝印类____；其它____。",
        "3. 尺寸：□ M01–M05 全部合格    □ 有不合格项：________________",
        "4. 电气：□ E01/E02 均与预期符合    □ 不符合项：________________",
        "5. 综合质量处置：□接收    □返工    □报废",
        "6. 依据（对照验收要求表）：________________________________________________",
        "7. 简要说明：________________________________________________",
        "复核：□ 单位齐全  □ 编号与版本一致  □ 分类与明细一致  □ 处置结论已填  签注：________",
    ]:
        add_para(doc, line, size=10.5, space_after=3)

    out = PLAYER / "学生组_模块D_检测记录表.docx"
    doc.save(str(out))
    print("[OK]", out.name)

    # 2 尺寸基准表
    doc = new_doc()
    add_bc_header(doc, "模块D：尺寸基准要求表", "学生组 · 随卷材料")
    add_para(doc, "一、说明", cn="黑体", size=14, bold=True, color=C_NAVY)
    add_para(
        doc,
        "本表给出 M01–M05 的工具、单位、名义与公差骨架。投板标定后以实物三测为准可微调公差。"
        "禁止普通卡尺测量 0.30 mm 级小孔。M05 为指定点，与外观 V02 不同点。",
        size=10.5,
    )
    t = doc.add_table(rows=6, cols=7)
    t.style = "Table Grid"
    fill_table(
        t,
        [
            ["编号", "项目", "工具", "单位", "名义值", "允许误差", "判定"],
            ["M01", "板长 L", "游标卡尺", "mm", "80.00", "±0.20", "合格/不合格"],
            ["M02", "板宽 W", "游标卡尺", "mm", "60.00", "±0.20", "合格/不合格"],
            ["M03", "板厚 T", "厚度规/千分尺", "mm", "1.60", "±0.15", "合格/不合格"],
            ["M04", "定位孔径", "卡尺/孔规", "mm", "2.00", "±0.10", "合格/不合格"],
            ["M05", "指定点线宽或线距", "测量显微镜", "mm", "见指定点", "±0.05（或对照阈值）", "合格/不合格"],
        ],
    )
    add_para(doc, "二、M05 指定点（卷面填写/丝印对照）", cn="黑体", size=14, bold=True, color=C_NAVY, space_before=8)
    add_para(doc, "特征名/位置：____________________    测量对象：□线宽  □线距", size=10.5)
    add_para(doc, "读数方法：显微镜估读/刻度读数；记录保留至 0.01 mm（或赛场规定）。", size=10.5)
    doc.save(str(PLAYER / "学生组_模块D_尺寸基准表.docx"))
    print("[OK] 尺寸基准表")

    # 3 验收要求表
    doc = new_doc()
    add_bc_header(doc, "模块D：验收要求 / 外观缺陷判定表", "学生组 · 随卷材料")
    add_para(doc, "一、外观必检（F1）判定语句", cn="黑体", size=14, bold=True, color=C_NAVY)
    t = doc.add_table(rows=6, cols=4)
    t.style = "Table Grid"
    fill_table(
        t,
        [
            ["编号", "类型", "不合格判定（竞赛用表述）", "备注"],
            ["V01", "顶层线路缺口", "顶层导体出现可见断开/缺口，影响连通风险", "须定位顶面"],
            ["V02", "底层线宽变窄", "底层局部线宽明显变窄（定性）；不要求此处读数", "与 M05 不同点"],
            ["V03", "孔破盘", "孔与焊盘关系异常导致破盘/焊盘缺损", "顶或底以实物为准"],
            ["V04", "阻焊未开窗或开窗偏移", "应开窗处被阻焊覆盖，或开窗相对焊盘明显偏移", "二选一现象"],
            ["V05", "丝印缺失或压焊盘", "字符缺失/残缺，或丝印压覆焊盘", "二选一现象"],
        ],
    )
    add_para(doc, "二、尺寸与电气", cn="黑体", size=14, bold=True, color=C_NAVY, space_before=8)
    add_para(doc, "尺寸：对照尺寸基准表，任一项超差即记不合格项。", size=10.5)
    add_para(doc, "电气：E01 预期开路；E02 预期短路。默认判据开路≥1 MΩ/OL，短路≤1 Ω/导通（以赛场书面或实物标定为准）。", size=10.5)
    add_para(doc, "三、综合处置（接收 / 返工 / 报废）", cn="黑体", size=14, bold=True, color=C_NAVY, space_before=8)
    add_para(doc, "接收：无必检致命缺陷，尺寸与电气均符合。", size=10.5)
    add_para(doc, "返工：存在可返工缺陷（如丝印、阻焊类）且尺寸/电气主体可接受——按验收表与裁判细则。", size=10.5)
    add_para(doc, "报废：存在线路开路缺口、孔破盘等严重影响功能/可靠性，或电气与预期严重不符且不可接受。", size=10.5)
    add_para(doc, "选手须给出结论 + 简要依据（对照本表条款/现象名）。不要求深层工艺根因分析。", size=10.5)
    doc.save(str(PLAYER / "学生组_模块D_验收要求表.docx"))
    print("[OK] 验收要求表")

    # 4 基础电气
    doc = new_doc()
    add_bc_header(doc, "模块D：基础电气检测表", "学生组 · 开短路点测（非故障定位）")
    add_para(doc, "一、检测范围", cn="黑体", size=14, bold=True, color=C_NAVY)
    add_para(doc, "仅对指定测试点做基础开路/短路判定。不做网络分段、故障定位树、高阻分析、上电功能测试。", size=10.5)
    add_para(doc, "二、测试点与预期", cn="黑体", size=14, bold=True, color=C_NAVY, space_before=6)
    t = doc.add_table(rows=3, cols=5)
    t.style = "Table Grid"
    fill_table(
        t,
        [
            ["编号", "类型", "测试点对", "预期", "默认电阻判据（可微调）"],
            ["E01", "开路", "TP1 – TP2", "开路", "≥1 MΩ 或表显 OL"],
            ["E02", "短路", "TP3 – TP4", "短路", "≤1 Ω 或蜂鸣导通"],
        ],
    )
    add_para(doc, "三、操作与记录", cn="黑体", size=14, bold=True, color=C_NAVY, space_before=6)
    for line in [
        "1. 确认万用表电阻档/通断档可用；表笔接触 TP 焊盘可靠。",
        "2. 记录：测试点对、档位、电阻读数、单位、判定、与预期是否符合。",
        "3. 裁判以实物标定为准；本表为选手填写与判据说明。",
    ]:
        add_para(doc, line, size=10.5, space_after=3)
    add_para(doc, "四、选手填写区", cn="黑体", size=14, bold=True, color=C_NAVY, space_before=6)
    t = doc.add_table(rows=3, cols=7)
    t.style = "Table Grid"
    fill_table(
        t,
        [
            ["编号", "测试点对", "档位", "电阻读数", "单位", "判定", "与预期符合"],
            ["E01", "TP1–TP2", "", "", "", "", "□是 □否"],
            ["E02", "TP3–TP4", "", "", "", "", "□是 □否"],
        ],
    )
    doc.save(str(PLAYER / "学生组_模块D_基础电气检测表.docx"))
    print("[OK] 基础电气检测表")

    # 5 分区说明
    doc = new_doc()
    add_bc_header(doc, "模块D：板面分区与测试点说明", "学生组 · 随卷材料")
    add_para(doc, "一、版本丝印", cn="黑体", size=14, bold=True, color=C_NAVY)
    add_para(doc, "MOD-D-S-A / MOD-D-S-B。答题须与实物丝印版本一致。", size=10.5)
    add_para(doc, "二、分区网格（建议）", cn="黑体", size=14, bold=True, color=C_NAVY, space_before=6)
    add_para(doc, "顶层参考网格 A1–D4（列 1–4，行 A–D）。记录缺陷时填写所在面 + 网格/特征名 + 方位。", size=10.5)
    add_para(doc, "三、特征区名称", cn="黑体", size=14, bold=True, color=C_NAVY, space_before=6)
    add_para(doc, "走线区、孔阵、阻焊对比区、丝印条、板框、测试点区等（以板面丝印为准）。", size=10.5)
    add_para(doc, "四、测试点 TP（仅基础电气）", cn="黑体", size=14, bold=True, color=C_NAVY, space_before=6)
    t = doc.add_table(rows=5, cols=3)
    t.style = "Table Grid"
    fill_table(
        t,
        [
            ["丝印", "用途", "说明"],
            ["TP1", "E01 一端", "开路网络"],
            ["TP2", "E01 另一端", "开路网络"],
            ["TP3", "E02 一端", "短路网络"],
            ["TP4", "E02 另一端", "短路网络"],
        ],
    )
    add_para(doc, "禁止将本说明理解为多网络故障定位表。仅 E01/E02 两点对。", size=10.5, color=C_RED, bold=True, space_before=6)
    doc.save(str(PLAYER / "学生组_模块D_板面分区说明.docx"))
    print("[OK] 板面分区说明")

    # 6 材料清单
    doc = new_doc()
    add_bc_header(doc, "模块D：选手材料清单", "学生组 · 随卷材料")
    add_para(doc, "一、赛场提供（以现场清单为准）", cn="黑体", size=14, bold=True, color=C_NAVY)
    t = doc.add_table(rows=9, cols=3)
    t.style = "Table Grid"
    fill_table(
        t,
        [
            ["序号", "名称", "备注"],
            ["1", "双面成品裸板质检样件", "A 或 B 版；无元器件"],
            ["2", "检测记录表", "本包"],
            ["3", "尺寸基准表", "本包"],
            ["4", "验收要求表", "本包"],
            ["5", "基础电气检测表", "本包"],
            ["6", "板面分区与测试点说明", "本包"],
            ["7", "仪器使用与安全要求", "本包 md/印发"],
            ["8", "量具与仪表", "见下表"],
        ],
    )
    add_para(doc, "二、推荐工具", cn="黑体", size=14, bold=True, color=C_NAVY, space_before=8)
    t = doc.add_table(rows=7, cols=3)
    t.style = "Table Grid"
    fill_table(
        t,
        [
            ["工具", "用途", "精度/说明"],
            ["游标卡尺", "L/W/孔径/槽宽", "0.02 mm 常用"],
            ["厚度规/千分尺", "板厚 T", "避开厚堆"],
            ["测量显微镜", "M05 指定点线宽/线距", "10～20× 或带刻度"],
            ["放大镜", "外观检查", "—"],
            ["数字万用表", "E01/E02 开短路", "电阻/通断档"],
            ["照明/侧光、防静电", "辅助", "不上电功能测试"],
        ],
    )
    doc.save(str(PLAYER / "学生组_模块D_选手材料清单.docx"))
    print("[OK] 选手材料清单")


def rebuild_judge_docs():
    """裁判侧 docx 视觉对齐 B/C；真值仍为待投板回填。"""
    # 缺陷参考答案
    doc = new_doc()
    add_bc_header(doc, "模块D：缺陷参考答案（A/B）", "裁判用 · 骨架 · 待投板回填")
    add_para(doc, "说明：类型/面/网格为命题骨架；照片编号与实物坐标投板后回填。同义见《同义答案与采点说明》。", size=10.5, color=C_GRAY)
    for ver, face_v05, grid in [
        ("A", "顶", [("V01", "顶", "B2", "走线区", "顶层线路缺口", "严重"),
                     ("V02", "底", "C2", "走线区", "底层线宽变窄", "主要"),
                     ("V03", "底", "B4", "孔阵", "孔破盘", "严重"),
                     ("V04", "顶", "C1", "阻焊对比区", "阻焊未开窗或开窗偏移", "主要"),
                     ("V05", "顶", "A3", "丝印条", "丝印缺失或压焊盘", "次要")]),
        ("B", "底", [("V01", "顶", "C3", "走线区", "顶层线路缺口", "严重"),
                     ("V02", "底", "B3", "走线区", "底层线宽变窄", "主要"),
                     ("V03", "底", "A3", "孔阵", "孔破盘", "严重"),
                     ("V04", "顶", "B1", "阻焊对比区", "阻焊未开窗或开窗偏移", "主要"),
                     ("V05", "底", "D1", "丝印条", "丝印缺失或压焊盘", "次要")]),
    ]:
        add_para(doc, f"版本 {ver}", cn="黑体", size=14, bold=True, color=C_NAVY, space_before=8)
        t = doc.add_table(rows=6, cols=7)
        t.style = "Table Grid"
        rows = [["编号", "面", "网格", "特征区", "类型", "严重程度", "照片#"]]
        for r in grid:
            rows.append([r[0], r[1], r[2], r[3], r[4], r[5], "待投板回填"])
        fill_table(t, rows)
    doc.save(str(JUDGE / "学生组_模块D_缺陷参考答案_AB.docx"))
    print("[OK] 缺陷参考答案")

    # 尺寸与线宽
    doc = new_doc()
    add_bc_header(doc, "模块D：尺寸与线宽参考答案（A/B）", "裁判用 · 待投板回填")
    add_para(doc, "名义/公差为骨架；A/B 标定列不得伪造已测值。", size=10.5, color=C_GRAY)
    t = doc.add_table(rows=6, cols=7)
    t.style = "Table Grid"
    fill_table(
        t,
        [
            ["编号", "项目", "工具", "名义/公差骨架", "A标定", "B标定", "备注"],
            ["M01", "板长L", "卡尺", "80.00±0.20", "待投板回填", "待投板回填", "mm"],
            ["M02", "板宽W", "卡尺", "60.00±0.20", "待投板回填", "待投板回填", "mm"],
            ["M03", "板厚T", "厚度规", "1.60±0.15", "待投板回填", "待投板回填", "mm"],
            ["M04", "定位孔径", "卡尺/孔规", "2.00±0.10", "待投板回填", "待投板回填", "全卷统一"],
            ["M05", "指定点线宽/线距", "显微镜", "±0.05或阈值", "待投板回填", "待投板回填", "≠V02点"],
        ],
    )
    doc.save(str(JUDGE / "学生组_模块D_尺寸与线宽参考答案_AB.docx"))
    print("[OK] 尺寸与线宽参考答案")

    # 电气
    doc = new_doc()
    add_bc_header(doc, "模块D：电气参考答案（A/B）", "裁判用 · 基础开短路 · 待投板回填")
    add_para(doc, "禁止使用故障定位旧表口径。电阻以实物标定为准。", size=10.5, bold=True, color=C_RED)
    t = doc.add_table(rows=3, cols=7)
    t.style = "Table Grid"
    fill_table(
        t,
        [
            ["编号", "类型", "测试点", "预期", "判据骨架", "A电阻标定", "B电阻标定"],
            ["E01", "开路", "TP1–TP2", "开路", "≥1MΩ/OL", "待投板回填", "待投板回填"],
            ["E02", "短路", "TP3–TP4", "短路", "≤1Ω/导通", "待投板回填", "待投板回填"],
        ],
    )
    doc.save(str(JUDGE / "学生组_模块D_电气参考答案_AB.docx"))
    print("[OK] 电气参考答案")


def write_report():
    report = SAMPLE / "学生组_模块D" / "00_版式交叉检查与对齐说明.md"
    report.write_text(
        f"""# 样题版式交叉检查与对齐说明（2026-07-21）

## 1. 审计发现（对齐前）

| 文件族 | 中文主字体 | 正文字号 | 标题色 | 问题 |
|--------|------------|----------|--------|------|
| 完整版 · 模块B/C 段 | 黑体标题 + 正文多继承/仿宋 | 标题16/14，正文约12 | 黑（无主题蓝强制） | 基准 |
| 完整版 · 模块D 段（R1回写后） | 混用黑体/仿宋/微软雅黑；标签与正文对调 | 10.5/12/14 混乱 | 部分 Heading 主题蓝 | **已修** |
| 模块B/C 独立样题 | **微软雅黑**正文 + **黑体**标题 | 正文10.5，标题14/18/20/28 | 深蓝1F3864 / 蓝2E75B6 / 红C00000 | 附件基准 |
| 模块D 选手/裁判 docx（R1脚本） | 仿宋+Times New Roman | 9～12 混用 | 无色板 | **已按B/C重排** |

## 2. 冻结的两套规范（勿混用）

### 规范 F · 完整版内嵌模块正文
- 模块标题：黑体 16pt 加粗
- （一）（二）…：黑体 14pt 加粗
- 【标签】/ D-x 任务名：黑体 12pt 加粗
- 正文与列表：仿宋 + Times New Roman，12pt，黑色自动
- 段落样式：优先 Normal，避免 Heading 主题色蓝

### 规范 A · 独立模块附件（B/C/D 选手裁判表）
- 页眉竞赛名：黑体 20/18，色 `#1F3864`
- 模块大红题：黑体 **28**（与 B/C 一致），色 `#C00000`
- 分隔线：微软雅黑 12，色 `#2E75B6`
- 一级节标题：黑体 14，`#1F3864`
- 正文：微软雅黑 10.5
- 警告：微软雅黑 10 加粗，`#C00000`
- 表头：微软雅黑 9 加粗白字 + 底 `#2E75B6`
- 表体：微软雅黑 9.5

## 3. 本轮已执行

1. 备份 → `{BACKUP.name}`
2. 完整版模块D 段按规范 F 重刷直接格式
3. 选手 6 份 docx + 裁判 3 份 docx 按规范 A 重建（R1 内容不变）
4. 脚本：`03_脚本与方案/_generated/align_module_d_visual_format_20260721.py`

## 4. 未纳入本轮（可选后续）

- 完整版全文（模块A/B/C/封面）统一主题色（工作量大，B/C段已基本稳定）
- 模块C缺陷记录表等历史小表（非D包）
- md 文件的打印版式（任务书要点、评分骨架等）

## 5. 复检命令

```bash
python 03_脚本与方案/_generated/align_module_d_visual_format_20260721.py
```
""",
        encoding="utf-8",
    )
    print("[OK] report", report)


def main():
    BACKUP.mkdir(parents=True, exist_ok=True)
    # backup targets
    targets = [FULL]
    targets += list(PLAYER.glob("*.docx"))
    targets += list(JUDGE.glob("*.docx"))
    for t in targets:
        if t.exists():
            shutil.copy2(t, BACKUP / t.name)
    print("[BACKUP]", BACKUP)

    fix_full_exam_module_d()
    rebuild_player_docs()
    rebuild_judge_docs()
    write_report()
    print("DONE")


if __name__ == "__main__":
    main()
