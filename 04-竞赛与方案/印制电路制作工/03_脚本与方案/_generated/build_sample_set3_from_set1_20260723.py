# -*- coding: utf-8 -*-
"""
第三套选手材料生成（2026-07-25 · PCB6 换锚）：

- 版式：整份复制第一套 docx 母版
- 内容：无线开发底板实质换题（源 Netlist_PCB6；非 CAN）
- B-0：WM2 无线模组封装（脱敏）；扩展座题面位号 J1
- 套次：仅标题「竞赛样题（第三套）」
"""
from __future__ import annotations

import shutil
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parents[2]
SET1 = ROOT / "02_样题" / "第一套"
SET3 = ROOT / "02_样题" / "第三套"
S1_EXAM = SET1 / "01_样题"
S1_ANS = SET1 / "03_结果呈现_参考答案"
S1_FORM = SET1 / "04_需填写内容"
S3_EXAM = SET3 / "01_样题"
S3_ANS = SET3 / "03_结果呈现_参考答案"
S3_FORM = SET3 / "04_需填写内容"
S3_SPEC = SET3 / "99_命题规格"
for d in (S3_EXAM, S3_ANS, S3_FORM, SET3 / "02_评分标准", SET3 / "05_U盘资料", S3_SPEC):
    d.mkdir(parents=True, exist_ok=True)


def set_paragraph_text(paragraph, text: str) -> None:
    if not paragraph.runs:
        paragraph.add_run(text)
        return
    paragraph.runs[0].text = text
    for r in paragraph.runs[1:]:
        r.text = ""


def set_cell_text(cell, text: str) -> None:
    if not cell.paragraphs:
        cell.text = text
        return
    set_paragraph_text(cell.paragraphs[0], text)
    for p in cell.paragraphs[1:]:
        set_paragraph_text(p, "")


def set_table_row(table, row_idx: int, values: list[str]) -> None:
    row = table.rows[row_idx]
    for ci, val in enumerate(values):
        if ci < len(row.cells):
            set_cell_text(row.cells[ci], val)


# ---------------------------------------------------------------------------
# 实质差异（相对第一/二/四套）：
# - 无线开发底板 + WM2 + J1(IDC) + H3/H4/H5 UART 跳线 + H6 半成品
# - 电源：+5V→LDO，无虚构保护件
# - 非 CAN / 非 TB端子 / 非 TJA / 非 LM75 / 非 BLE 排针 / 非 WM1 多功能功率板
# ---------------------------------------------------------------------------

PARA_BY_INDEX: dict[int, str] = {
    5: "竞赛样题（第三套）",
    38: (
        "某物联网公司开发一款无线开发底板（核心板扩展形态）：以一颗 2.4GHz 短距无线模组（本卷器件名 WM2）为无线核心，"
        "通过板载 24P 核心座对外引出电源、复位与 GPIO，并提供 IDC 扩展口及 UART 外部/板载跳线选择，便于联调与二次开发。"
        "硬件工程师已完成板框、24P 核心座（H6）及主干电源/地扇出，但 LDO 供电、状态灯与按键、UART 跳线网络、"
        "IDC 扩展座等子电路尚未补全；无线模组 WM2 须选手按赛场附图自建符号与 PCB 封装后接入绿色高亮区。"
    ),
    39: (
        "作为 PCB Layout 工程师，您需要完成：自建无线模组 WM2 封装 → 按功能说明补全红色高亮子电路 → "
        "设置设计规则 → PCB 布局（模组天线朝向与净空、扩展口靠边、跳线区可操作）→ 布线与敷铜 → "
        "导出生产 Gerber 与 2D 预览，满足可制造性与开发底板布局要求。"
    ),
    45: "B-0  自建无线模组封装（WM2）",
    46: (
        "本设计核心无线器件为本卷第 2.4GHz 短距无线模组，器件命名为 WM2。"
        "赛场半成品工程库中不提供该模组封装，须自行绘制器件符号与 PCB 封装；"
        "封装名建议：WM2_Footprint。外形与焊盘以赛场下发的《无线模组封装尺寸附图（WM2）》为准，"
        "不得凭名称到元件库检索套用，不得使用其它射频/Wi‑Fi 模组封装顶替。"
    ),
    47: "1.器件符号设计",
    48: (
        "（1）新建器件 WM2；符号须至少包含本卷用到的电气引脚：3V3、GND、RESET（或 NRST）、"
        "UART 相关脚（如 TX/RX，网络名以半成品绿色区标签为准，可能写作 P1.5/TX、P1.4/RX 等），"
        "以及题面要求连接的 GPIO 子集；"
    ),
    49: (
        "（2）引脚电气类型合理设置；电源与地分开；在符号上标注第 1 脚或关键电源脚标识。"
        "不必画出模组全部引脚名称，但不得遗漏本卷要求连接的网络。"
        "注意：网络名形如 P1.0、P1.4/RX 表示模组口线，与扩展座位号 J1 的针脚号不是同一概念。"
    ),
    50: "2. PCB封装设计",
    51: (
        "（1）严格按照赛场《无线模组封装尺寸附图（WM2）》绘制：模组外形轮廓、焊盘阵列、间距与焊盘尺寸；"
        "天线端须在丝印层明确标识（如 ANT）；"
    ),
    52: (
        "（2）丝印外框清晰，第 1 脚或定位角加标识；禁止只画电源少数焊盘的简化假封装；"
        "禁止套用库内其它模组封装；"
    ),
    53: "（3）封装完成后导入本工程库；",
    54: (
        "（4）将器件放置到原理图绿色高亮的「无线模组」区，按半成品网络标签完成 3V3、GND、RESET 及题面要求的信号连接；"
        "布局阶段天线端朝板外并保留净空（净空区禁止敷铜、走线与过孔，以附图为准）。"
    ),
    55: "B-1  原理图补全",
    56: (
        "打开赛场下发的立创 EDA 半成品工程，仅补全红色高亮子电路。"
        "元件型号与封装以半成品及库内推荐为准，本卷不另发完整 BOM。"
        "下列为功能要求，阻容取值与工程标注冲突时以工程为准。"
        "本卷供电入口按半成品为 +5V 直入 LDO，不要求额外串联保险/TVS/防反电路。"
    ),
    57: "子电路功能说明（共 7 块）",
    58: (
        "（1）+5V 入口与 LDO 稳压：外部 +5V（入口形态以半成品为准）送入 3.3V LDO（位号/封装以半成品为准，常见 SOT-223 档），"
        "得到 3V3 供 WM2 与数字电路。按工程完成输入/输出去耦（如 10μF 与 100nF 组合）。"
        "本块考核供电入口与 LDO 网络完整，不要求选手自行增加保险、TVS 或防反二极管。"
    ),
    59: (
        "（2）复位：RESET（或 NRST）经约 10kΩ 上拉到 3V3，并联去耦电容到 GND（取值以工程为准，常见 100nF）；"
        "轻触开关 RST 一端接 RESET、一端接 GND。须保证 WM2 复位脚、H6 对应脚及扩展相关复位网连通（以半成品标签为准）。"
    ),
    60: (
        "（3）LED 指示（满分必做 3 只）：电源灯 PWR——按工程接 3V3/GND 与限流电阻；"
        "运行灯 RUN、网络/链路灯 NWK——分别经限流电阻接至半成品指定 GPIO 网络，极性以工程为准（常见低电平或高电平点亮，以标注为准）。"
        "若半成品另有 USR 灯，不要求选手补画，不纳入本卷采分。"
    ),
    61: (
        "（4）按键（满分必做 2 只用户键）：两只独立轻触键，一端分别接半成品指定 GPIO，另一端接 GND；"
        "使用内部上拉或工程已给上拉，按下为有效电平（通常为低）。"
        "RST 已在第（2）块计入；若板尚有其它键位（如网络键），不要求补画，不纳入本卷采分。"
    ),
    62: (
        "（5）UART 与 EXT/INT 跳线网络：补全 H3、H4、H5 三座 3P 跳线相关连接，实现与半成品标签同构的拓扑——"
        "RX 路径：板载/外部选择脚与公共脚（如 P1.4/RX_EXT、P1.4/RX_INT、P1.4/RX）；"
        "TX 路径：同理（如 P1.5/TX_EXT、P1.5/TX_INT、P1.5/TX）；"
        "公共 RX/TX 脚须接到 WM2 与 H6 对应网络。不要求撰写跳线使用说明，以原理图网络连通正确为准。"
    ),
    63: (
        "（6）WM2 全连接：在完成 B-0 落位后，按红色高亮与网络标签完成模组 3V3、GND、RESET、UART 及本卷要求的 GPIO 连接；"
        "不得遗漏电源/地；信号名以半成品为准。本块与 B-0 共同构成无线核心考核。"
    ),
    64: (
        "（7）IDC 扩展座 J1：放置 10P IDC 排线座（题面位号 J1；封装选用库内 2.54mm 10P IDC 或半成品指定封装）。"
        "至少连通：GND、3V3、以及半成品标签中的两路 GPIO（源工程对应 P2.1/P2.2 一类网络）与 RESET；"
        "其余针脚允许空置（NC）。注意 J1 的针脚序号与网络名 P1.x 无关，不得混接。"
        "麦克、螺丝孔等半成品已有电路不要求选手改画。"
    ),
    65: "B-2  设计规则设置",
    66: "（1）按本卷第（三）节工艺表设置 DRC，mil/mm 单位一致；",
    67: "（2）信号线默认 0.15mm；3V3 ≥0.8mm；+5V ≥1.0mm；",
    68: "（3）过孔：内径 0.3mm、外径 0.6mm；",
    69: "（4）安全间距：线-线、线-盘、孔-线均为 0.15mm；",
    70: "（5）阻焊：焊盘 1:1 开窗；阻焊桥最小 0.1mm。",
    71: "B-3  PCB布局",
    72: "将原理图更新到 PCB，在半成品给定板框内完成布局（板框与倒角不得修改）。",
    73: "布局要求",
    74: "（1）板框以半成品为准；H6 核心座与安装孔若已有则保持，不得平移出原功能区；",
    75: (
        "（2）+5V 入口与 LDO 集中在电源区，靠近电源入口；去耦电容紧贴 LDO 与 WM2 电源脚；"
    ),
    76: "（3）WM2 放置后天线端朝板外；天线净空区内禁止敷铜、走线与过孔；",
    77: (
        "（4）J1 扩展座放置在板边，便于插拔排线；H3/H4/H5 跳线区预留操作空间，避免被高器件遮挡；"
    ),
    78: (
        "（5）功能分区建议：电源区 | 无线模组区 | 核心座 H6 区 | UART 跳线区 | 扩展口 J1 区 | 人机区（键灯）；"
    ),
    79: "（6）丝印完整；WM2 天线标识与 J1 位号可见；位号不压焊盘。",
    80: "参考图：",
    81: "B-4 PCB布线",
    82: "（1）电源：3V3 ≥0.8mm，+5V ≥1.0mm；模组供电短而宽；",
    83: "（2）地：GND 完整连通，底层作地平面；模组地焊盘就近多过孔；",
    84: "（3）UART 与跳线相关走线尽量短、清晰；复位线远离大电流与噪声源；",
    85: "（4）天线净空内无信号线；晶振类器件下方不走其它信号（若板含）；",
    86: "（5）优先 45° 走线，避免 90°；",
    87: "（6）过孔合理，避免在 J1/H6 焊盘内乱打孔；",
    88: "（7）顶/底敷 GND 铜（遵守天线净空）；",
    89: "（8）网络连通率 100%。",
    91: "B-5  DRC检查与Gerber导出",
    92: "（1）DRC 错误为 0（警告可忽略）；",
    93: "（2）导出顶/底层、顶/底阻焊、顶/底丝印、板框、钻孔等；",
    94: "（3）Gerber 打 zip，命名见提交要求；",
    95: "（4）导出顶层 2D 预览 PNG。",
    96: "B-6 注意事项",
    97: (
        "⚠ 注意：文件保存到 D:\\提交资料\\模块B\\ ，赛终拷贝到 U 盘。"
        "工位号/赛位号替换为实际赛位号；U 盘包示例：赛位号_模块B。"
    ),
    98: "1. 合理分配本模块时间；时长与评分以技术工作文件/竞赛平台为准；",
    99: "2. 及时保存；",
    100: "3. 命名错误影响评分；",
    101: "4. 禁止拷贝他人文件或使用外部存储，违者按作弊处理；",
    102: "5. 除 WM2 外，优先使用库内封装，不得擅自乱建其它封装；",
    103: "6. 软件或设备异常请举手示意裁判。",
    110: (
        "您是 PCB 板厂 CAM 工程师。客户提交了一套「无线开发底板」的投产 Gerber"
        "（客户自述由开发底板设计导出，与赛场 EDA 半成品非同一文件），准备开料。"
        "请先完成 DFM 审核，识别并记录缺陷，再填写简版制程工艺卡。"
    ),
}


def apply_para_overrides(doc: Document) -> None:
    for idx, text in PARA_BY_INDEX.items():
        if idx < len(doc.paragraphs):
            set_paragraph_text(doc.paragraphs[idx], text)


def apply_table_overrides(doc: Document) -> None:
    if len(doc.tables) > 4:
        t = doc.tables[4]
        set_table_row(t, 1, ["1", "立创EDA工程文件", "半成品（含板框/H6 核心座；无 WM2 模组封装；红区待补）"])
        set_table_row(t, 2, ["2", "子电路功能说明", "本卷第 B-1 节 7 块红色高亮电路说明"])
        set_table_row(t, 3, ["3", "无线模组封装尺寸附图（WM2）", "B-0 绘制依据（赛场下发，禁止按名称搜库）"])
        set_table_row(t, 4, ["4", "立创EDA专业版软件", "赛场已预装"])

    if len(doc.tables) > 6:
        t = doc.tables[6]
        set_table_row(t, 1, ["1", "+5V 入口与 LDO", "+5V→3.3V LDO 及去耦（无额外保护件要求）"])
        set_table_row(t, 2, ["2", "复位", "上拉+电容+RST；连通 WM2/H6/扩展复位网"])
        set_table_row(t, 3, ["3", "LED 指示", "PWR + RUN + NWK（USR 不考）"])
        set_table_row(t, 4, ["4", "按键", "2 只用户键（其它键不考）"])
        set_table_row(t, 5, ["5", "UART 跳线", "H3/H4/H5 实现 EXT/INT 拓扑"])
        set_table_row(t, 6, ["6", "WM2 全连接", "电源/地/复位/UART/要求 GPIO"])
        set_table_row(t, 7, ["7", "扩展座 J1", "IDC 10P：GND/3V3/GPIO×2/RESET"])

    if len(doc.tables) > 7:
        t = doc.tables[7]
        set_table_row(t, 2, ["2", "自建封装库", "B0_WM2_工位{工位号}.elibz2", "Elibz2"])

    if len(doc.tables) > 8:
        t = doc.tables[8]
        set_table_row(t, 1, ["1", "Gerber文件包", "无线开发底板客户投产 Gerber"])
        set_table_row(t, 2, ["2", "板厂工艺能力参数表", "缺陷判定依据"])
        set_table_row(t, 3, ["3", "缺陷记录表", "记录所在层/类型/位置/违反条款"])
        set_table_row(t, 4, ["4", "制程工艺卡", "能量参数与文件完整性"])
        set_table_row(t, 5, ["5", "GB/T 4588-2025标准摘要", "参考"])


def scrub_tokens(doc: Document) -> list[str]:
    blob = "\n".join(p.text for p in doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                blob += "\n" + cell.text
    bad = []
    for token in [
        "PogoPin",
        "LoRa",
        "SX1268",
        "AHT30",
        "NE555",
        "BH1750",
        "PinHeader_2x5",
        "BLE模组",
        "BLE传感器",
        "TB_5.08",
        "TJA1051",
        "LM75",
        "CANH",
        "CANL",
        "CAN 总线",
        "CAN总线",
        "工控节点",
        "第一套",
        "第二套",
        "第四套",
        "E18",
        "WIRELM",
        "亿佰特",
        "Hi-12F",
        "安信可",
        "WM1",
    ]:
        if token in blob:
            bad.append(token)
    if blob.count("第三套") != 1:
        bad.append(f"第三套出现{blob.count('第三套')}次")
    return bad


def build_full_exam() -> Path:
    src = S1_EXAM / "印制电路制作工赛项_竞赛样题（完整版）.docx"
    dst = S3_EXAM / "印制电路制作工赛项_竞赛样题（完整版）.docx"
    shutil.copy2(src, dst)
    doc = Document(str(dst))
    apply_para_overrides(doc)
    apply_table_overrides(doc)
    bad = scrub_tokens(doc)
    doc.save(str(dst))
    print("CHECK residual:", bad if bad else "ok")
    return dst


def copy_player_forms() -> list[Path]:
    out = []
    for name in [
        "模块C_缺陷记录表.docx",
        "模块C_PCB制程工艺卡.docx",
        "学生组_模块D_检测记录表.docx",
    ]:
        src, dst = S1_FORM / name, S3_FORM / name
        shutil.copy2(src, dst)
        out.append(dst)
    return out


def build_c_answer() -> Path:
    src = S1_ANS / "模块C_缺陷记录表-参考答案.docx"
    dst = S3_ANS / "模块C_缺陷记录表-参考答案.docx"
    shutil.copy2(src, dst)
    doc = Document(str(dst))
    answers = [
        ["1", "顶层线路层", "最小线宽不足", "无线模组WM2扇出信号（UART）局部线宽约0.08mm", "违反：最小线宽0.10mm"],
        ["2", "顶层线路层", "最小线距不足", "H3 UART跳线座附近RX相关两网局部间距约0.08mm", "违反：最小线距0.10mm"],
        ["3", "底层线路层", "走线开路/缺口", "+5V至LDO输入侧底层供电干线在过孔附近缺口", "开路致命"],
        ["4", "顶层线路层", "走线短路", "NWK指示灯限流电阻与相邻GND铜皮桥接", "短路致命"],
        ["5", "钻孔层+焊盘", "焊环不足（孔破盘）", "J1扩展座附近过孔偏孔，单侧焊环约0.08mm", "焊环极限0.18mm"],
        ["6", "钻孔层", "过孔焊盘过小", "WM2地过孔焊盘外径约0.20mm", "最小过孔焊盘0.25mm"],
        ["7", "钻孔层", "孔间距过小", "H6核心座内侧两过孔孔边距约0.15mm", "孔边距≥0.20mm"],
        ["8", "顶层阻焊层", "阻焊覆盖焊盘", "用户键焊盘被阻焊覆盖未开窗", "焊盘需开窗"],
        ["9", "顶层阻焊层", "阻焊桥缺失/不足", "WM2焊盘阵列区间无法形成有效阻焊桥", "阻焊桥要求"],
        ["10", "顶层丝印层", "丝印上焊盘", "位号WM2压在模组焊盘上", "丝印距露铜≥0.15mm"],
        ["11", "顶层丝印层", "字符尺寸过小", "位号J1字高约0.7mm", "字高≥1.0mm"],
        ["12", "文件完整性", "缺少底层丝印层文件", "压缩包缺底层丝印层", "文件不完整"],
    ]
    table = doc.tables[0]
    ncols = len(table.columns)
    for i, row_vals in enumerate(answers):
        ri = i + 1
        if ri >= len(table.rows):
            break
        set_table_row(table, ri, row_vals[:ncols])
    doc.save(str(dst))
    return dst


def main() -> None:
    paths = [build_full_exam()]
    paths.extend(copy_player_forms())
    paths.append(build_c_answer())

    d1 = Document(str(S1_EXAM / "印制电路制作工赛项_竞赛样题（完整版）.docx"))
    d3 = Document(str(paths[0]))
    same = diff = 0
    for p1, p2 in zip(d1.paragraphs, d3.paragraphs):
        if not p1.text.strip() and not p2.text.strip():
            continue
        if p1.text == p2.text:
            same += 1
        else:
            diff += 1
    print("PARA same", same, "diff", diff)
    for p in paths:
        print("OK", p.name, p.stat().st_size)

    # 源 netlist 不得被本脚本删除（命题规格目录）
    tel = S3_SPEC / "Netlist_PCB6_2026-07-25.tel"
    print("source tel present:", tel.exists(), tel.stat().st_size if tel.exists() else 0)


if __name__ == "__main__":
    main()
