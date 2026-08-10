# -*- coding: utf-8 -*-
"""
第三套模块D参考答案：以第四套参考答案为母版。

坐标变换（用户操作：纵轴 ABCD ↔ 横轴；横轴 1234 ↔ 纵轴）：
  第四套约定：字母=行(纵)，数字=列(横)  →  标签如 B2 = 行B列2
  第三套约定：字母=列(横)，数字=行(纵)  →  同物理格新标签 = transpose
  公式：旧(R,C) → 新( letter(C), number(R) )
  例：C1→A3，A2→B1，B3→C2，B2→B2
"""
from __future__ import annotations

import re
import shutil
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parents[2]
SRC = ROOT / "02_样题" / "第四套" / "03_结果呈现_参考答案" / "模块D_检测记录表_参考答案.docx"
DST = ROOT / "02_样题" / "第三套" / "03_结果呈现_参考答案" / "模块D_检测记录表_参考答案.docx"

LETTERS = "ABCD"


def transpose_token(token: str) -> str:
    """Map single cell like A2 / C1. Unchanged if not grid."""
    m = re.fullmatch(r"([A-Da-d])([1-4])", token.strip())
    if not m:
        return token
    row_l = m.group(1).upper()
    col_n = int(m.group(2))
    new_letter = LETTERS[col_n - 1]
    new_num = LETTERS.index(row_l) + 1
    return f"{new_letter}{new_num}"


def transpose_grid_text(text: str) -> str:
    """Replace all A1–D4 tokens in a cell/paragraph (supports B3-C3 ranges)."""

    def repl(m: re.Match[str]) -> str:
        return transpose_token(m.group(0))

    return re.sub(r"[A-Da-d][1-4]", repl, text)


def set_paragraph_text(paragraph, text: str) -> None:
    if not paragraph.runs:
        paragraph.add_run(text)
        return
    paragraph.runs[0].text = text
    for r in paragraph.runs[1:]:
        r.text = ""


def set_cell_text(cell, text: str) -> None:
    if not cell.paragraphs:
        cell.text = text
        return
    set_paragraph_text(cell.paragraphs[0], text)
    for p in cell.paragraphs[1:]:
        set_paragraph_text(p, "")


# 第四套答案表 → 第三套（转置后网格 + 方位微调）
# 来源表：V01 B2 / V02 C1 / V03 A1 / V04 B3-C3 / V05 A2
V_ROWS = [
    # 编号, 所在面, 网格/特征, 缺陷类型, 现象简述
    [
        "V01",
        "顶层",
        "B2 / 电源走线",
        "走线颈缩",
        "走线一段突然变为不足4mil",
    ],
    [
        "V02",
        "钻孔层",
        "A3 / 过孔区（原第四套 C1，轴对调后）",
        "过孔内径过小",
        "内径过小",
    ],
    [
        "V03",
        "钻孔层",
        "A1 / 定位孔区（角部）",
        "漏钻孔",
        "有丝印无钻孔",
    ],
    [
        "V04",
        "顶层",
        "C2–C3 / U1 右侧阻焊异常（原 B3–C3 转置）",
        "无法形成有效阻焊桥",
        "焊盘阻焊引脚过小",
    ],
    [
        "V05",
        "顶层",
        "B1 / TP1 丝印（原 A2 转置）",
        "丝印缺失或压焊盘",
        "位号丝印不完整/缺失",
    ],
]


def main() -> None:
    if not SRC.exists():
        raise SystemExit(f"missing source: {SRC}")
    DST.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(SRC, DST)
    doc = Document(str(DST))

    # 题头说明：标明第三套 + 变换规则
    for p in doc.paragraphs:
        t = p.text
        if "黄色底纹为参考答案" in t:
            set_paragraph_text(
                p,
                "【黄色底纹为参考答案·第三套】工程母版=第四套模块D，网格轴已对调："
                "横轴为 A–D、纵轴为 1–4（相对第四套「纵 A–D / 横 1–4」做标签转置）。"
                "正式评分允许位置同义表述与合理测量误差；以实物三测标定为准，"
                "下列数值为设计/标定骨架（孔径按本套板 3.50 mm）。本文件不得发给选手。",
            )
        elif "误把正常线当 V02" in t and "3.50" in t:
            set_paragraph_text(
                p,
                "• 误把正常线当 V02、或 M05 与 V02 同点、或定位孔写成 2.00 与本套 3.50 冲突 → 按错扣。"
                " 网格须按第三套丝印（横 A–D / 纵 1–4）描述，勿直接照抄第四套坐标。",
            )
        elif t.strip().startswith("签注："):
            set_paragraph_text(
                p,
                "签注：命题/裁判长 ________　　日期 ________　　"
                "本答案随投板三测修订后锁版（第三套·自第四套轴对调派生）",
            )

    # 外观表
    if doc.tables:
        t0 = doc.tables[0]
        for i, row_vals in enumerate(V_ROWS):
            ri = i + 1
            if ri >= len(t0.rows):
                break
            for ci, val in enumerate(row_vals):
                if ci < len(t0.rows[ri].cells):
                    set_cell_text(t0.rows[ri].cells[ci], val)

    # 其余段落/表中若残留第四套网格码，做通用转置（保险）
    # 仅处理明确的 A1-D4，避免误伤「3.50」等
    for table in doc.tables[1:]:
        for row in table.rows:
            for cell in row.cells:
                new_t = transpose_grid_text(cell.text)
                if new_t != cell.text:
                    set_cell_text(cell, new_t)

    doc.save(str(DST))
    print("OK", DST)
    print("mapping check: C1->", transpose_token("C1"), "A2->", transpose_token("A2"), "B3->", transpose_token("B3"))
    for r in V_ROWS:
        print(" ", r[0], r[2])


if __name__ == "__main__":
    main()
