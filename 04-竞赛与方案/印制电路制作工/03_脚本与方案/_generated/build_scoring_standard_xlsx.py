# -*- coding: utf-8 -*-
"""按《评分标准demo物联网.xlsx》结构生成印制电路制作工第一套评分标准。"""
from copy import copy
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[2]
# demo 已归档；若需重生成 xlsx，请从归档路径复制 demo 到本机临时位置再改此路径
DEMO = (
    ROOT
    / "05_归档备份"
    / "sample-prep-cleanup-20260729_112827"
    / "第一套"
    / "评分标准demo物联网.xlsx"
)
OUT = ROOT / "02_样题" / "第一套" / "02_评分标准" / "印制电路制作工赛项_评分标准（第一套·0721口径）.xlsx"

# styles aligned with demo
FONT = Font(name="Arial", size=10)
FONT_YAHEI = Font(name="Microsoft YaHei", size=10)
FONT_BOLD = Font(name="Arial", size=10, bold=True)
FONT_HDR = Font(name="Microsoft YaHei", size=10)
ALIGN_C = Alignment(horizontal="center", vertical="center", wrap_text=True)
ALIGN_L = Alignment(horizontal="left", vertical="center", wrap_text=True)
THIN = Side(style="thin", color="000000")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
FILL_NONE = PatternFill(fill_type=None)
FILL_YELLOW = PatternFill("solid", fgColor="FFFF00")
FILL_WHITE = PatternFill("solid", fgColor="FFFFFF")


def style_range(ws, row, cols, *, fill=None, bold=False, center=True):
    for c in range(1, cols + 1):
        cell = ws.cell(row, c)
        cell.border = BORDER
        cell.alignment = ALIGN_C if center else ALIGN_L
        if c == 1 or c == 2:
            cell.font = FONT_HDR if not bold else Font(name="Microsoft YaHei", size=10, bold=True)
        else:
            cell.font = FONT_BOLD if bold else FONT
        if fill is not None:
            cell.fill = fill


def write_header(ws):
    headers = [
        "模块",
        "模块内容",
        "M=测量分J=评价分",
        "测量或评价内容",
        "测量依据",
        "最高分",
        "得分",
    ]
    for i, h in enumerate(headers, 1):
        cell = ws.cell(1, i, h)
        cell.font = FONT_HDR
        cell.alignment = ALIGN_C
        cell.border = BORDER
        cell.fill = FILL_WHITE
    ws.row_dimensions[1].height = 18


def module_header(ws, row, code, title, total):
    """模块小计行（仿 demo 黄底）"""
    ws.cell(row, 1, code)
    ws.cell(row, 2, title)
    for c in range(3, 6):
        ws.cell(row, c, None)
    ws.cell(row, 6, total)
    ws.cell(row, 7, None)
    # merge B like demo B2:E2 for title space? demo merges B2:E2 on some - check
    # demo: B2 only has title, C-E empty, merge B2:E2
    ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=5)
    for c in range(1, 8):
        cell = ws.cell(row, c)
        cell.border = BORDER
        cell.alignment = ALIGN_C
        cell.fill = FILL_YELLOW
        cell.font = FONT_BOLD if c == 6 else FONT_HDR
    ws.row_dimensions[row].height = 22
    return row


def detail(ws, row, *, sub=None, mj, content, basis, score, code=None):
    ws.cell(row, 1, code)
    ws.cell(row, 2, sub)
    ws.cell(row, 3, mj)
    ws.cell(row, 4, content)
    ws.cell(row, 5, basis)
    ws.cell(row, 6, score)
    ws.cell(row, 7, None)
    for c in range(1, 8):
        cell = ws.cell(row, c)
        cell.border = BORDER
        cell.fill = FILL_WHITE
        if c in (1, 3, 5, 6, 7):
            cell.alignment = ALIGN_C
            cell.font = FONT_BOLD if c == 6 else FONT
        elif c == 2:
            cell.alignment = ALIGN_C
            cell.font = FONT_YAHEI
        else:
            cell.alignment = ALIGN_L
            cell.font = FONT_YAHEI
    # height by content length
    h = 18 + min(80, (len(content) // 28) * 12)
    ws.row_dimensions[row].height = h
    return row


def build():
    wb = Workbook()
    ws = wb.active
    ws.title = "评分表"

    # column widths ~ demo
    widths = {
        "A": 8,
        "B": 22,
        "C": 16,
        "D": 58,
        "E": 18,
        "F": 8,
        "G": 9,
    }
    for k, v in widths.items():
        ws.column_dimensions[k].width = v

    write_header(ws)
    r = 2
    module_total_rows = []  # for SUM formula

    # ========== A 理论 30 ==========
    module_total_rows.append(r)
    module_header(ws, r, "A", "理论测试（竞赛平台）", 30)
    r += 1
    detail(
        ws,
        r,
        sub="平台测试成绩折算",
        mj="M",
        content="竞赛平台理论测试卷面满分100分，按技术文件折算计入总成绩：模块A得分＝平台卷面得分×0.30（保留两位小数，四舍五入）。题型题量以平台及赛前公布为准（0721参考：单选40+多选30+判断30）。",
        basis="平台导出成绩",
        score=30,
    )
    r += 1

    # ========== B EDA 25 ==========
    module_total_rows.append(r)
    module_header(ws, r, "B", "EDA工程设计", 25)
    r += 1

    # 评价 20
    r = detail(
        ws,
        r,
        sub="评价分（合计20）",
        mj="J",
        content="布局合理性：0～3档（无分组/基本分组/模块化清晰/专业美观EMC）。三名裁判独立给档后取平均，换算为满分10分：得分＝(平均档÷3)×10。",
        basis="现场/提交文件评审",
        score=10,
    )
    r += 1
    r = detail(
        ws,
        r,
        sub=None,
        mj="J",
        content="布线策略：0～3档（杂乱/基本整齐/策略明确/专业高速回流）。换算为满分10分：得分＝(平均档÷3)×10。",
        basis="现场/提交文件评审",
        score=10,
    )
    r += 1
    # 测量 5
    r = detail(
        ws,
        r,
        sub="测量分（合计5）",
        mj="M",
        content="DRC检查：无错误或违规数在允许范围内；有致命错误不得分。",
        basis="EDA软件DRC",
        score=1.5,
    )
    r += 1
    r = detail(
        ws,
        r,
        mj="M",
        content="网络连通/网表一致性：关键网络连通，原理图与PCB一致；严重开路或错连不得分。",
        basis="工程文件检查",
        score=1.5,
    )
    r += 1
    r = detail(
        ws,
        r,
        mj="M",
        content="Gerber完整性：规定层齐全、可正常打开；缺关键层按比例扣完为止。",
        basis="Gerber包",
        score=1.0,
    )
    r += 1
    r = detail(
        ws,
        r,
        mj="M",
        content="提交规范：文件命名、目录、格式符合赛题要求。",
        basis="提交物检查",
        score=1.0,
    )
    r += 1

    # ========== C CAM 20 ==========
    module_total_rows.append(r)
    module_header(ws, r, "C", "CAM审核与工艺文件编制", 20)
    r += 1

    r = detail(
        ws,
        r,
        sub="缺陷识别（合计14）",
        mj="M",
        content="标准缺陷共12处（锁版《模块C_缺陷记录表-参考答案》）。每处最高1.0分：类型对+层/面对+位置可对应＝1.0；类型对+层对但位置含糊＝0.5；类型错/层错/指错＝0。",
        basis="缺陷记录表+标准答案",
        score=12,
    )
    r += 1
    r = detail(
        ws,
        r,
        mj="M",
        content="无严重误报（2分）：不在标准清单且非同义描述的计误报；每处误报扣1.0分，扣完为止；不倒扣命中分。多写干扰项不另加分。",
        basis="缺陷记录表",
        score=2,
    )
    r += 1

    # expand 12 optional summary already above - add one row listing types for referee
    r = detail(
        ws,
        r,
        sub="工艺卡（合计6）",
        mj="M",
        content="层数正确（双面/2层）。",
        basis="工艺卡+Gerbv",
        score=1,
    )
    r += 1
    r = detail(
        ws,
        r,
        mj="M",
        content="外形尺寸（长×宽）与板框实测一致（允许合理误差，如±0.5mm或等价mil）。",
        basis="工艺卡+Gerbv",
        score=1,
    )
    r += 1
    r = detail(
        ws,
        r,
        mj="M",
        content="最小孔径与钻孔最小工具一致（如约12mil/0.30mm）。",
        basis="工艺卡+钻孔文件",
        score=1,
    )
    r += 1
    # 工艺卡满分6：层数1+外形1+孔径1+线宽线距1+开窗1+文件1
    r = detail(
        ws,
        r,
        mj="M",
        content="最小线宽与最小线距（正常区）合理；不得将故意缺陷极值填为规格。",
        basis="工艺卡+Gerbv",
        score=1,
    )
    r += 1
    r = detail(
        ws,
        r,
        mj="M",
        content="阻焊开窗方式表述正确（如1:1开窗）。",
        basis="工艺卡+阻焊/铜层对照",
        score=1,
    )
    r += 1
    r = detail(
        ws,
        r,
        mj="M",
        content="文件完整性与事实一致（本套须指出缺少底层丝印GBO等）。",
        basis="工艺卡+文件包",
        score=1,
    )
    r += 1

    # ========== D 质检 25 ==========
    module_total_rows.append(r)
    module_header(ws, r, "D", "成品板质量检测与缺陷分析", 25)
    r += 1

    # V01-V05
    visuals = [
        ("V01", "顶层线路缺口", "2", "顶层+缺口/断线类+位置大致正确；建议全对2/部分1/错0"),
        ("V02", "底层线宽变窄（颈缩）", "2", "底层+变窄/颈缩且不断开+位置大致正确；不得与M05同点"),
        ("V03", "孔破盘", "2", "孔破盘/环宽异常+面与位置大致正确；非3.50mm定位孔误报"),
        ("V04", "阻焊未开窗或开窗偏移", "2", "阻焊类缺陷+位置大致正确"),
        ("V05", "丝印缺失或压焊盘", "2", "丝印类缺陷+位置大致正确"),
    ]
    first_v = True
    for code, typ, sc, tip in visuals:
        r = detail(
            ws,
            r,
            sub="外观缺陷（合计10）" if first_v else None,
            mj="M",
            content=f"{code} {typ}。{tip}",
            basis="检测记录表+实物+标准答案",
            score=float(sc),
        )
        r += 1
        first_v = False

    measures = [
        ("M01", "板长L", "1", "卡尺；名义80.00±0.20"),
        ("M02", "板宽W", "1", "卡尺；名义60.00±0.20"),
        ("M03", "板厚T", "1", "厚度规/千分尺；名义1.60±0.15"),
        ("M04", "定位孔径", "1", "卡尺/孔规；名义3.50±0.10（本套板）"),
        ("M05", "指定点线宽或线距", "3", "测量显微镜；以板面M05标识为准；±0.05；禁止仅用卡尺测细线"),
    ]
    first_m = True
    for code, name, sc, tip in measures:
        r = detail(
            ws,
            r,
            sub="尺寸测量（合计7）" if first_m else None,
            mj="M",
            content=f"{code} {name}。{tip}",
            basis="检测记录表+量具+锁版真值",
            score=float(sc),
        )
        r += 1
        first_m = False

    r = detail(
        ws,
        r,
        sub="基础电气（合计5）",
        mj="M",
        content="E01 TP1－TP2开路：判定开路且与预期一致（如≥1MΩ或OL）得2.5分，否则0分。",
        basis="万用表+标准答案",
        score=2.5,
    )
    r += 1
    r = detail(
        ws,
        r,
        mj="M",
        content="E02 TP3－TP4短路：判定短路且与预期一致（如≤1Ω）得2.5分，否则0分。",
        basis="万用表+标准答案",
        score=2.5,
    )
    r += 1

    r = detail(
        ws,
        r,
        sub="综合判定（合计2）",
        mj="J",
        content="分类汇总基本合理得1分；质量处置与依据正确得1分。存在线路缺口、孔破盘等致命缺陷时应判报废，不得判接收。",
        basis="检测记录表D-5",
        score=2,
    )
    r += 1
    r = detail(
        ws,
        r,
        sub="记录规范（合计1）",
        mj="J",
        content="版本丝印与实物一致、单位mm、表头完整、无姓名单位等身份信息得1分；明显缺项得0分。",
        basis="检测记录表",
        score=1,
    )
    r += 1

    # ========== 总分 ==========
    last_detail = r - 1
    total_row = r
    ws.cell(total_row, 1, "总分")
    for c in range(2, 6):
        ws.cell(total_row, c, None)
    # 最高分：各模块小计行之和
    f_refs = ",".join(f"F{i}" for i in module_total_rows)
    ws.cell(total_row, 6, f"=SUM({f_refs})")
    # 得分：若填了G列明细可改；演示用模块小计G之和——G模块行一般空，改为SUM各明细G
    # demo: =SUM(F80,F72,...) for 得分 side - actually G88 = sum of module F headers for verification of max, wait
    # demo R88: F = =SUM(F2:F87)-100  weird; G = sum of module header F cells
    # Looking again: G88 = SUM(F80,F72,... module headers) - that's max total not score
    # F88 = SUM(all F)-100 
    # For our sheet: F total = sum module headers = 100
    # G total = SUM of all detail G scores - user fills G on detail rows
    # Better: G_total = SUM of G on all detail rows excluding module headers
    # Or leave G for judge to fill module subtotals
    detail_g_parts = []
    # simpler: G总分 = SUM(G2:G{last}) but module header G empty so ok if we only sum details
    # Module headers have empty G - SUM(G2:Glast) works when judges fill detail G
    ws.cell(total_row, 7, f"=SUM(G2:G{last_detail})")

    for c in range(1, 8):
        cell = ws.cell(total_row, c)
        cell.border = BORDER
        cell.alignment = ALIGN_C
        cell.font = FONT_BOLD
        cell.fill = FILL_WHITE
    ws.row_dimensions[total_row].height = 22

    # 说明行
    note_row = total_row + 2
    ws.cell(
        note_row,
        1,
        "说明：1.本表与技术工作文件0721定稿及第一套样题配套；总成绩A30+B25+C20+D25=100。"
        "2.M=测量分，J=评价分。3.各模块得分及总分保留两位小数，四舍五入。"
        "4.并列：先比操作技能(B+C+D)，再比D→C→B。"
        "5.模块D正式满分25分（不以40分计）。"
        "6.标准答案与尺寸/电气真值以赛前锁版及投板三测为准。"
        "7.结构参考《评分标准demo物联网.xlsx》。",
    )
    ws.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=7)
    ws.cell(note_row, 1).font = FONT_YAHEI
    ws.cell(note_row, 1).alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    ws.row_dimensions[note_row].height = 60

    # freeze header
    ws.freeze_panes = "A2"
    ws.print_title_rows = "1:1"

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print("wrote", OUT)

    # verify sums of detail max under each module roughly
    print("module header rows", module_total_rows)
    print("total_row", total_row)


if __name__ == "__main__":
    build()
