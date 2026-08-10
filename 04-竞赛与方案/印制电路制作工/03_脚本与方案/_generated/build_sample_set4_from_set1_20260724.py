# -*- coding: utf-8 -*-
"""
第四套选手材料：第一套 docx 母版 + PCB15 实质换题。
- 不写入第二套/第三套目录
- B-0 = 自制 2.4GHz 无线模组封装（题面代号 WM1；禁止写真实型号以免选手搜库）
- 仅标题「竞赛样题（第四套）」
"""
from __future__ import annotations

import shutil
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parents[2]
SET1 = ROOT / "02_样题" / "第一套"
SET4 = ROOT / "02_样题" / "第四套"
S1_EXAM = SET1 / "01_样题"
S1_ANS = SET1 / "03_结果呈现_参考答案"
S1_FORM = SET1 / "04_需填写内容"
S4_EXAM = SET4 / "01_样题"
S4_ANS = SET4 / "03_结果呈现_参考答案"
S4_FORM = SET4 / "04_需填写内容"
S4_USB = SET4 / "05_U盘资料"
for d in (S4_EXAM, S4_ANS, S4_FORM, SET4 / "02_评分标准", S4_USB, SET4 / "99_命题规格"):
    d.mkdir(parents=True, exist_ok=True)


def set_paragraph_text(paragraph, text: str) -> None:
    if not paragraph.runs:
        paragraph.add_run(text)
        return
    paragraph.runs[0].text = text
    for r in paragraph.runs[1:]:
        r.text = ""


def set_cell_text(cell, text: str) -> None:
    if not cell.paragraphs:
        cell.text = text
        return
    set_paragraph_text(cell.paragraphs[0], text)
    for p in cell.paragraphs[1:]:
        set_paragraph_text(p, "")


def set_table_row(table, row_idx: int, values: list[str]) -> None:
    row = table.rows[row_idx]
    for ci, val in enumerate(values):
        if ci < len(row.cells):
            set_cell_text(row.cells[ci], val)


PARA_BY_INDEX: dict[int, str] = {
    5: "竞赛样题（第四套）",
    38: (
        "某物联网设备商开发一款 Wi‑Fi 多功能控制节点板（工程代号 PCB15）：以一颗 2.4GHz 无线模组（本卷器件名 WM1）为无线核心，"
        "集成 RS485 总线、直流电机驱动、继电器输出及本地键灯人机接口，并可扩展 NFC/麦克等（半成品已预置部分电路）。"
        "硬件工程师已完成板框、扩展座与部分子电路，但电源链、人机、RS485、电机与继电器等仍待补全；"
        "且无线模组 WM1 须选手按赛场附图自建符号与 PCB 封装后接入绿色高亮区。"
    ),
    39: (
        "作为 PCB Layout 工程师，您需要完成：自建无线模组 WM1 封装 → 按说明补全红色高亮子电路 → 设置设计规则 → "
        "PCB 布局（模组天线朝向与净空、大电流路径）→ 布线与敷铜 → 导出生产 Gerber 与 2D 预览。"
    ),
    45: "B-0  自建无线模组封装（WM1）",
    46: (
        "本设计核心无线器件为本卷第 2.4GHz Wi‑Fi 模组，器件命名为 WM1。"
        "赛场半成品工程库中不提供该模组封装，须自行绘制器件符号与 PCB 封装；"
        "封装名建议：WM1_Footprint。外形与焊盘以赛场下发的《无线模组封装尺寸附图》为准，不得凭名称到元件库检索套用。"
    ),
    47: "1.器件符号设计",
    48: (
        "（1）新建器件 WM1；符号须至少包含本卷用到的电气引脚：3V3、GND、PWRON（或 EN）、"
        "以及与底板/扩展相关的 UART 或 GPIO 子集（网络名以半成品绿色区标签为准，如 IO/TXD/RXD 等）；"
    ),
    49: (
        "（2）引脚电气类型合理设置；电源与地分开；在符号上标注第 1 脚或关键电源脚标识。"
        "不必在符号上画出模组全部引脚名称，但不得遗漏本卷要求连接的网络。"
    ),
    50: "2. PCB封装设计",
    51: (
        "（1）严格按照赛场《无线模组封装尺寸附图》绘制：模组外形轮廓、焊盘阵列、间距与焊盘尺寸；"
        "天线端须在丝印层明确标识（如 ANT）；"
    ),
    52: (
        "（2）丝印外框清晰，第 1 脚或定位角加标识；禁止使用库内其它射频/Wi‑Fi 模组封装顶替；"
        "禁止只画电源四焊盘的简化假封装；"
    ),
    53: "（3）封装完成后导入本工程库；",
    54: (
        "（4）将器件放置到原理图绿色高亮的「无线模组」区，按半成品网络标签完成 3V3、GND、PWRON 及题面要求的信号连接；"
        "布局阶段天线端朝板外并保留净空。"
    ),
    55: "B-1  原理图补全",
    56: (
        "打开赛场下发的立创 EDA 半成品工程，仅补全红色高亮子电路。"
        "元件型号与封装以半成品库内推荐为准，本卷不另发完整 BOM。"
        "下列为功能要求，阻容取值与工程标注冲突时以工程为准。"
    ),
    57: "子电路功能说明（共 7 块）",
    58: (
        "（1）+5V 输入保护：外部 +5V 经自恢复保险丝 F1 → 保护器件（肖特基/TVS 等，位号以半成品 D2/D3/D4 区为准）"
        "得到节点 +5V，供 LDO、电机驱动与继电器线圈等。保护路径须完整、位号不与 LED 冲突。"
    ),
    59: (
        "（2）LDO 稳压：采用半成品指定的 3.3V LDO（位号 LDO1，SOT-23-5 等），+5V→3V3，"
        "供 WM1、数字电路与接口芯片；输入/输出滤波电容按工程预留位号焊接网络（常见 10μF 与 1μF 组合）。"
    ),
    60: (
        "（3）复位与按键：完成 RST 复位电路及 SW1、SW2 用户键；"
        "KEY1、KEY2 网络接至半成品扩展/主控侧标签；按下为有效电平（以工程为准，通常低有效）。"
        "本卷用户键为 2 只（另有拨码 SW3 半成品已处理，不必重画）。"
    ),
    61: (
        "（4）LED 指示：按红色高亮完成电源/状态 LED（如 LED1、LED2 等）及限流电阻；"
        "极性与限流取值按工程标注；不得占用保护二极管位号。"
    ),
    62: (
        "（5）RS485 接口：补全收发器 U3 与端子 CN1 相关连接；"
        "RS485_A/B 至 CN1；端接电阻 R11（120Ω）跨接 A-B（可串 0Ω 便于断开）；"
        "收发控制与 TX/RX 网络按半成品标签（RS485_TX/RX 等）连接完整。"
    ),
    63: (
        "（6）电机驱动：补全驱动芯片 U1 与输出端子 P1；"
        "控制端 MOTOR_INA/MOTOR_INB、输出 MOTOR_OUTA/MOTOR_OUTB 网络连通；"
        "供电与使能脚按半成品推荐接法，大电流路径在 PCB 阶段加宽。"
    ),
    64: (
        "（7）继电器输出：补全继电器 RLY2、驱动三极管及续流二极管等外围，以及端子 P2 触点引出；"
        "控制网络 RELAY 与指示灯（若高亮要求）按工程连接。本卷以继电器功率输出为第 7 块考核，"
        "NFC/麦克/蜂鸣器半成品已预置，不要求选手补画。"
    ),
    # 注意：第四套完整版可能已被人工改过（如 B-1、B-2 截图提交）。
    # 全量重跑本脚本会覆盖人工修改；B-2 双单位请优先用定点补丁，勿轻易全量 regenerate。
    65: "B-2  设计规则设置",
    66: "（1）根据工艺要求表设置 DRC 设计规则，单位支持 mil/mm 切换；",
    67: "（2）设置线宽规则：信号线 0.15mm（6mil）；电源线分级：3V3 ≥ 0.8mm（32mil），+5V 及电机电源路径 ≥ 1.0mm（40mil）（或按工程电源规则加宽）；",
    68: "（3）设置过孔规则：内径 0.3mm（12mil），外径 0.6mm（24mil）；",
    69: "（4）设置安全间距规则：线线间距、线焊盘间距、孔线间距均为 0.15mm（6mil）；",
    70: "（5）设置阻焊规则：焊盘 1:1 开窗，阻焊桥最小宽度 0.1mm（4mil）。",
    71: "B-3  PCB布局",
    72: "将原理图更新到 PCB，在半成品给定板框内完成布局（板框与倒角不得修改）。",
    73: "布局要求",
    74: "（1）板框以半成品为准；安装孔/螺丝孔若已有则保持；",
    75: "（2）电源入口、F1、LDO 集中在电源区，输入输出路径短、流向清晰；",
    76: "（3）WM1 放置后天线端朝板外短边；天线净空区内禁止敷铜、走线与过孔；",
    77: "（4）CN1（RS485）、P1（电机）、P2（继电器）等端子靠边，便于现场接线；",
    78: "（5）功能分区建议：电源区 | 无线模组区 | 总线接口区 | 功率驱动区（电机/继电器）| 人机区 | 扩展座区；",
    79: "（6）丝印完整；WM1 天线标识可见；位号不压焊盘。",
    80: "参考图：",
    81: "B-4 PCB布线",
    82: "（1）电源：3V3 ≥0.8mm，+5V 与电机供电加宽；模组供电短而宽；",
    83: "（2）地：GND 完整，底层地平面；模组地焊盘多过孔；",
    84: "（3）RS485 差分尽量平行；电机大电流回路远离复位与射频；",
    85: "（4）天线净空内无信号线；",
    86: "（5）优先 45° 走线，避免 90°；",
    87: "（6）过孔合理，避免端子焊盘内乱孔；",
    88: "（7）顶/底敷 GND 铜（遵守天线净空）；",
    89: "（8）网络连通率 100%。",
    91: "B-5  DRC检查与Gerber导出",
    92: "（1）DRC 错误为 0（警告可忽略）；",
    93: "（2）导出顶/底层、顶/底阻焊、顶/底丝印、板框、钻孔等；",
    94: "（3）Gerber 打 zip，命名见提交要求；",
    95: "（4）导出顶层 2D 预览 PNG。",
    96: "B-6 注意事项",
    97: (
        "⚠ 注意：文件保存到 D:\\提交资料\\模块B\\ ，赛终拷贝到 U 盘。"
        "工位号/赛位号替换为实际赛位号；U 盘包示例：赛位号_模块B。"
    ),
    98: "1. 合理分配时间；时长与评分以技术工作文件/竞赛平台为准；",
    99: "2. 及时保存；",
    100: "3. 命名错误影响评分；",
    101: "4. 禁止拷贝他人文件或使用外部存储，违者按作弊处理；",
    102: "5. 除 WM1 外，优先使用库内封装，不得擅自乱建其它封装；",
    103: "6. 软件或设备异常请举手示意裁判。",
    110: (
        "您是 PCB 板厂 CAM 工程师。客户提交了一套「Wi‑Fi 多功能控制节点（PCB15）」的投产 Gerber"
        "（与赛场 EDA 半成品非同一文件），准备开料。请先完成 DFM 审核并记录缺陷，再填写简版制程工艺卡。"
    ),
}


def apply_para_overrides(doc: Document) -> None:
    for idx, text in PARA_BY_INDEX.items():
        if idx < len(doc.paragraphs):
            set_paragraph_text(doc.paragraphs[idx], text)


def apply_table_overrides(doc: Document) -> None:
    if len(doc.tables) > 4:
        t = doc.tables[4]
        set_table_row(t, 1, ["1", "立创EDA工程文件", "半成品（含板框/扩展座；无 WM1 模组封装；红区待补）"])
        set_table_row(t, 2, ["2", "子电路功能说明", "本卷第 B-1 节 7 块说明"])
        set_table_row(t, 3, ["3", "无线模组封装尺寸附图", "B-0 绘制依据（赛场下发，禁止按名称搜库）"])
        set_table_row(t, 4, ["4", "立创EDA专业版软件", "赛场已预装"])

    if len(doc.tables) > 6:
        t = doc.tables[6]
        set_table_row(t, 1, ["1", "+5V 输入保护", "F1 + 保护二极管 → +5V"])
        set_table_row(t, 2, ["2", "LDO 3.3V", "LDO1：+5V→3V3，滤波按工程"])
        set_table_row(t, 3, ["3", "复位与按键", "RST；SW1/SW2→KEY1/KEY2"])
        set_table_row(t, 4, ["4", "LED 指示", "题面指定 LED + 限流"])
        set_table_row(t, 5, ["5", "RS485", "U3 + CN1 + R11(120Ω)"])
        set_table_row(t, 6, ["6", "电机驱动", "U1 + P1 + MOTOR_* 网络"])
        set_table_row(t, 7, ["7", "继电器输出", "RLY2 + P2 + 驱动外围"])

    if len(doc.tables) > 7:
        t = doc.tables[7]
        set_table_row(t, 2, ["2", "自建封装库", "B0_WM1_工位{工位号}.elibz2", "Elibz2"])

    if len(doc.tables) > 8:
        t = doc.tables[8]
        set_table_row(t, 1, ["1", "Gerber文件包", "Wi-Fi多功能控制节点客户投产 Gerber"])
        set_table_row(t, 2, ["2", "板厂工艺能力参数表", "缺陷判定依据"])
        set_table_row(t, 3, ["3", "缺陷记录表", "记录所在层/类型/位置/违反条款"])
        set_table_row(t, 4, ["4", "制程工艺卡", "能量参数与文件完整性"])
        set_table_row(t, 5, ["5", "GB/T 4588-2025标准摘要", "参考"])


def scrub(doc: Document) -> list[str]:
    blob = "\n".join(p.text for p in doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                blob += "\n" + cell.text
    bad = []
    for token in [
        "PogoPin", "LoRa", "AHT30", "NE555", "BH1750", "PinHeader_2x5",
        "TB_5.08", "第一套", "第二套", "第三套", "BLE传感器",
        "Hi-12F", "HI-12F", "Hi12F", "安信可", "AI-Thinker", "AI_Thinker",
    ]:
        if token in blob:
            bad.append(token)
    if blob.count("第四套") != 1:
        bad.append(f"第四套出现{blob.count('第四套')}次")
    return bad


def build_full_exam() -> Path:
    src = S1_EXAM / "印制电路制作工赛项_竞赛样题（完整版）.docx"
    dst = S4_EXAM / "印制电路制作工赛项_竞赛样题（完整版）.docx"
    shutil.copy2(src, dst)
    doc = Document(str(dst))
    apply_para_overrides(doc)
    apply_table_overrides(doc)
    bad = scrub(doc)
    doc.save(str(dst))
    print("CHECK residual:", bad if bad else "ok")
    return dst


def copy_player_forms() -> list[Path]:
    out = []
    for name in [
        "模块C_缺陷记录表.docx",
        "模块C_PCB制程工艺卡.docx",
        "学生组_模块D_检测记录表.docx",
    ]:
        src, dst = S1_FORM / name, S4_FORM / name
        shutil.copy2(src, dst)
        out.append(dst)
    return out


def build_c_answer() -> Path:
    src = S1_ANS / "模块C_缺陷记录表-参考答案.docx"
    dst = S4_ANS / "模块C_缺陷记录表-参考答案.docx"
    shutil.copy2(src, dst)
    doc = Document(str(dst))
    answers = [
        ["1", "顶层线路层", "最小线宽不足", "无线模组WM1扇出信号局部线宽约0.08mm", "违反：最小线宽0.10mm"],
        ["2", "顶层线路层", "最小线距不足", "RS485_A与RS485_B局部间距约0.08mm", "违反：最小线距0.10mm"],
        ["3", "底层线路层", "走线开路/缺口", "+5V给电机驱动供电的底层干线缺口", "开路致命"],
        ["4", "顶层线路层", "走线短路", "LED限流电阻与相邻GND铜皮桥接", "短路致命"],
        ["5", "钻孔层+焊盘", "焊环不足（孔破盘）", "H1排针区过孔偏孔，单侧焊环约0.08mm", "焊环极限0.18mm"],
        ["6", "钻孔层", "过孔焊盘过小", "模组地过孔焊盘外径约0.20mm", "最小过孔焊盘0.25mm"],
        ["7", "钻孔层", "孔间距过小", "继电器座附近两孔边距约0.15mm", "孔边距≥0.20mm"],
        ["8", "顶层阻焊层", "阻焊覆盖焊盘", "SW1按键焊盘被阻焊覆盖未开窗", "焊盘需开窗"],
        ["9", "顶层阻焊层", "阻焊桥缺失/不足", "U3密脚区间无法形成有效阻焊桥", "阻焊桥要求"],
        ["10", "顶层丝印层", "丝印上焊盘", "位号L1压在模组焊盘上", "丝印距露铜≥0.15mm"],
        ["11", "顶层丝印层", "字符尺寸过小", "位号R11字高约0.7mm", "字高≥1.0mm"],
        ["12", "文件完整性", "缺少底层丝印层文件", "压缩包缺底层丝印层", "文件不完整"],
    ]
    table = doc.tables[0]
    ncols = len(table.columns)
    for i, row_vals in enumerate(answers):
        ri = i + 1
        if ri >= len(table.rows):
            break
        set_table_row(table, ri, row_vals[:ncols])
    doc.save(str(dst))
    return dst


def main() -> None:
    # 保护：不得删除用户源工程
    src_eprj = S4_USB / "模块B.eprj2"
    print("source eprj2 present:", src_eprj.exists(), src_eprj.stat().st_size if src_eprj.exists() else 0)

    paths = [build_full_exam()]
    paths.extend(copy_player_forms())
    paths.append(build_c_answer())

    d1 = Document(str(S1_EXAM / "印制电路制作工赛项_竞赛样题（完整版）.docx"))
    d4 = Document(str(paths[0]))
    same = diff = 0
    for p1, p2 in zip(d1.paragraphs, d4.paragraphs):
        if not p1.text.strip() and not p2.text.strip():
            continue
        if p1.text == p2.text:
            same += 1
        else:
            diff += 1
    print("PARA same", same, "diff", diff)
    for p in paths:
        print("OK", p.name, p.stat().st_size)

    # 确认第二套未被本脚本改动（脚本只写 SET4）
    print("SET4 only write target:", SET4)


if __name__ == "__main__":
    main()
