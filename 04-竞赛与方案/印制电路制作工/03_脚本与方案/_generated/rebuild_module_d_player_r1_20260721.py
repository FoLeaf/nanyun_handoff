# -*- coding: utf-8 -*-
"""
学生组模块D · 选手材料包重做（R1 / 45 min / 基础电气）
- 对齐 命题边界确认表 + 缺陷与测点编号约定 + 设计规格
- 仅重写 选手/ 与 任务书要点.md
- 不回迁故障定位旧表
"""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

ROOT = Path(__file__).resolve().parents[2]
PACK = ROOT / "02_样题" / "学生组_模块D"
PLAYER = PACK / "选手"


def set_run_font(run, name_cn="仿宋", name_en="Times New Roman", size_pt=10.5, bold=False):
    run.bold = bold
    run.font.size = Pt(size_pt)
    run.font.name = name_en
    r = run._element
    rPr = r.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:ascii"), name_en)
    rFonts.set(qn("w:hAnsi"), name_en)
    rFonts.set(qn("w:eastAsia"), name_cn)
    rFonts.set(qn("w:cs"), name_en)


def add_para(doc, text, *, cn="仿宋", size=10.5, bold=False, align=None, space_after=6, space_before=0):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(space_before)
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    run = p.add_run(text)
    set_run_font(run, name_cn=cn, size_pt=size, bold=bold)
    return p


def set_cell_text(cell, text, *, cn="仿宋", size=9, bold=False, center=False):
    cell.text = ""
    p = cell.paragraphs[0]
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    set_run_font(run, name_cn=cn, size_pt=size, bold=bold)


def fill_table(table, rows, header_bold=True):
    for ri, row_data in enumerate(rows):
        for ci, val in enumerate(row_data):
            is_header = ri == 0 and header_bold
            set_cell_text(
                table.cell(ri, ci),
                str(val),
                cn="黑体" if is_header else "仿宋",
                size=9,
                bold=is_header,
                center=is_header or ci == 0,
            )


def setup_page(doc):
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.left_margin = Cm(2.0)
    sec.right_margin = Cm(2.0)
    sec.top_margin = Cm(1.8)
    sec.bottom_margin = Cm(1.8)


def new_doc():
    doc = Document()
    setup_page(doc)
    return doc


# ---------- 1. 任务书要点 ----------
def write_task_points():
    path = PLAYER / "任务书要点.md"
    path.write_text(
        """# 任务书要点（学生组·模块D）

> 版本：2026-07-21 · R1 / **45 分钟** / 含基础电气
> 完整表述以《印制电路制作工赛项_竞赛样题（完整版）》模块D 及本包随卷材料为准。

## 1. 定位与对象

| 项 | 内容 |
|----|------|
| 模块名称 | **模块D · 成品PCB裸板质量检测与判定**（学生组） |
| 检验对象 | 双面**成品裸板**质检样件（专用检测板，无元器件） |
| 板型概要 | FR-4 双面、约 80×60 mm、标称 1.6 mm、绿阻焊白丝印、**无铅喷锡** |
| 版本丝印 | `MOD-D-S-A` 或 `MOD-D-S-B`（以实物为准） |
| 正式时长 | **45 分钟** |
| 模块分值 | 40 原始分（外观14 / 尺寸·微几何12 / 电气8 / 综合4 / 规范2） |

### 与模块 B/C 边界

> 本模块为独立质检任务。检验对象为赛场提供的**双面成品裸板质检样件**（专用检测板），与模块 B（EDA 工程设计）、模块 C（CAM 审核与工艺文件编制）相互独立；**仅做基础开短路点测，不做故障定位与网络分段，不上电，不量测 Gerber**。

## 2. 任务结构（R1 · 45 min）

| 步骤 | 名称 | 建议时长 | 选手做什么 |
|------|------|--------:|------------|
| **D-1** | 检验准备 | 3 min | 工具检查；识读验收/尺寸/电气/分区说明；确认板面方向与版本丝印 |
| **D-2** | 外观缺陷检测 | 12 min | 顶/底 **5 处必检（V01–V05，F1）**；分区引导记录；干扰项可不计 |
| **D-3** | 尺寸与指定点微几何 | 10 min | **M01–M04**（L/W/T + 定位孔径）+ **M05** 指定点线宽或线距（显微镜） |
| **D-4** | 基础电气 | 8 min | **E01** 开路（TP1–TP2）+ **E02** 短路（TP3–TP4）；记电阻与判定 |
| **D-5** | 分类判定与提交 | 12 min | 缺陷分类；接收/返工/报废；记录复核提交 |
| **合计** | | **45** | |

## 3. 必检内容编号（与记录表一致）

| 前缀 | 编号 | 内容 |
|------|------|------|
| 外观 V | V01 | 顶层线路缺口 |
| | V02 | 底层线宽变窄（定性） |
| | V03 | 孔破盘 |
| | V04 | 阻焊未开窗或开窗偏移 |
| | V05 | 丝印缺失或压焊盘 |
| 测量 M | M01–M04 | 板长 / 板宽 / 板厚 / 定位孔径 |
| | M05 | 指定点线宽**或**线距（与 V02 不同点） |
| 电气 E | E01 | TP1–TP2 开路 |
| | E02 | TP3–TP4 短路 |

## 4. 成果物（须提交/填写）

1. 《学生组_模块D_检测记录表》（含 D-1～D-5 全部字段）
2. 同步完成并留存对照：验收要求表、尺寸基准表、基础电气检测表、板面分区说明（随卷识读）
3. 按赛场要求：纸质签署 / 电子命名示例 `赛位号_模块D`

## 5. 允许工具

游标卡尺、厚度规（或千分尺）、测量显微镜（或带刻度放大镜）、放大镜、**数字万用表**、照明/侧光、防静电用品（以赛场清单为准）。

## 6. 明确不做

- 复杂网络分段、故障定位树、应通应断全表、高阻分析
- 上电功能测试、焊接/贴装质量、实际返修操作
- 测量 Gerber / 对照设计文件量数
- 普通卡尺测 0.30 mm 级小孔；无指定点的全板扫微距

## 7. 随卷材料索引（选手侧）

| 文件 | 用途 |
|------|------|
| `学生组_模块D_检测记录表.docx` | 主提交表，D-1～D-5 |
| `学生组_模块D_尺寸基准表.docx` | M01–M05 基准与公差 |
| `学生组_模块D_验收要求表.docx` | F1 外观判定 + 综合处置 |
| `学生组_模块D_基础电气检测表.docx` | E01/E02 点测与电阻规则 |
| `学生组_模块D_板面分区说明.docx` | A1–D4、特征名、TP、版本 |
| `学生组_模块D_选手材料清单.docx` | 材料与工具清单 |
| `学生组_模块D_仪器使用与安全要求.md` | 仪器使用与安全（不上电） |
""",
        encoding="utf-8",
    )
    print("[OK]", path.name)


# ---------- 2. 检测记录表 ----------
def build_record_docx(path: Path):
    doc = new_doc()
    add_para(
        doc,
        "PCB成品裸板检测记录表（学生组·模块D）",
        cn="黑体",
        size=12,
        bold=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )
    add_para(
        doc,
        "备注：结合《验收要求表》《尺寸基准表》《基础电气检测表》《板面分区说明》填写。"
        "时长建议 45 分钟（D-1～D-5）。命名：赛位号_模块D。不得标注姓名（按赛场规则）。"
        "电气仅为点对点开短路，非故障定位。",
        size=9,
    )

    t0 = doc.add_table(rows=3, cols=4)
    t0.style = "Table Grid"
    fill_table(
        t0,
        [
            ["赛位号", "", "日期", ""],
            ["样件编号", "", "版本（丝印）", "□A(MOD-D-S-A)  □B(MOD-D-S-B)"],
            ["板面方向确认", "顶/底已确认□", "记录格式提示", "版本-编号，如 A-V01"],
        ],
        header_bold=False,
    )
    for cell in t0.rows[0].cells:
        for p in cell.paragraphs:
            for r in p.runs:
                set_run_font(r, name_cn="黑体", size_pt=9, bold=True)

    add_para(doc, "一、D-1 检验准备（建议 3 min）", cn="黑体", size=11, bold=True, space_before=8)
    add_para(
        doc,
        "□ 文件齐全（记录表/验收表/尺寸基准/基础电气/分区说明）  "
        "□ 卡尺零位已检  □ 厚度规可用  □ 测量显微镜/放大镜可用  "
        "□ 万用表电阻档可用  □ 样件无异常运输损伤（如有：________）  "
        "□ 版本丝印已抄录",
        size=9.5,
    )

    add_para(doc, "二、D-2 外观缺陷检测记录（V01–V05 必检，建议 12 min）", cn="黑体", size=11, bold=True, space_before=6)
    add_para(
        doc,
        "须检顶底两面。必检类型：V01 顶层线路缺口；V02 底层线宽变窄（定性）；V03 孔破盘；"
        "V04 阻焊未开窗或开窗偏移；V05 丝印缺失或压焊盘。干扰项可不计必检。记录格式：版本-编号。",
        size=9,
    )
    t1 = doc.add_table(rows=9, cols=8)
    t1.style = "Table Grid"
    rows = [
        [
            "编号",
            "所在面",
            "网格区",
            "特征名",
            "缺陷类型（标准名/同义）",
            "现象简述",
            "严重程度",
            "是否必检",
        ]
    ]
    # 5 required + 3 extra rows
    presets = [
        ("V01", "顶", "", "走线区", "顶层线路缺口", "", "轻/中/重", "是"),
        ("V02", "底", "", "走线区", "底层线宽变窄", "", "轻/中/重", "是"),
        ("V03", "底", "", "孔阵", "孔破盘", "", "轻/中/重", "是"),
        ("V04", "顶", "", "阻焊对比区", "阻焊未开窗/开窗偏移", "", "轻/中/重", "是"),
        ("V05", "", "", "丝印条", "丝印缺失/压焊盘", "", "轻/中/重", "是"),
    ]
    for r in presets:
        rows.append(list(r))
    for _ in range(3):
        rows.append(["", "", "", "", "", "", "", "否/干扰"])
    fill_table(t1, rows)

    add_para(doc, "三、D-3 尺寸与指定点微几何（M01–M05，建议 10 min）", cn="黑体", size=11, bold=True, space_before=8)
    add_para(
        doc,
        "单位 mm。M01–M04 用卡尺/厚度规；M05 用测量显微镜测**指定点**（与 V02 不同点）。"
        "禁止普通卡尺测 0.30 mm 级小孔；禁止无指定点全板扫微距。",
        size=9,
    )
    t2 = doc.add_table(rows=6, cols=6)
    t2.style = "Table Grid"
    fill_table(
        t2,
        [
            ["编号", "测量项目", "工具", "基准/公差（见尺寸基准表）", "实测值(mm)", "合格/不合格"],
            ["M01", "板长 L", "游标卡尺", "80.00±0.20", "", ""],
            ["M02", "板宽 W", "游标卡尺", "60.00±0.20", "", ""],
            ["M03", "板厚 T", "厚度规/千分尺", "1.60±0.15", "", ""],
            ["M04", "定位孔径", "卡尺/孔规", "2.00±0.10", "", ""],
            ["M05", "指定点线宽或线距", "测量显微镜", "见尺寸基准表指定点", "", ""],
        ],
    )
    add_para(
        doc,
        "M05 指定点特征名/位置（抄自尺寸基准表或板面丝印）：____________________  "
        "测量对象：□线宽  □线距    读数方法：____________________",
        size=9.5,
        space_before=4,
    )

    add_para(doc, "四、D-4 基础电气（E01–E02，建议 8 min）", cn="黑体", size=11, bold=True, space_before=8)
    add_para(
        doc,
        "仅测指定测试点对。默认判据：开路 ≥1 MΩ 或 OL；短路 ≤1 Ω 或蜂鸣导通。"
        "详细规则见《基础电气检测表》。**不做故障定位、网络分段、上电。**",
        size=9,
    )
    t3 = doc.add_table(rows=3, cols=8)
    t3.style = "Table Grid"
    fill_table(
        t3,
        [
            ["编号", "类型", "测试点对", "档位", "电阻读数", "单位", "判定(开路/短路/正常)", "与预期符合"],
            ["E01", "开路", "TP1 – TP2", "", "", "", "", "□是 □否"],
            ["E02", "短路", "TP3 – TP4", "", "", "", "", "□是 □否"],
        ],
    )

    add_para(doc, "五、D-5 分类、综合判定与复核（建议 12 min）", cn="黑体", size=11, bold=True, space_before=8)
    add_para(
        doc,
        "1. 外观必检计数：V01____ V02____ V03____ V04____ V05____；已发现必检合计____/5；干扰项____。",
        size=10,
    )
    add_para(
        doc,
        "2. 分类汇总（可按类型）：线路类____；孔/盘类____；阻焊类____；丝印类____；其它____。",
        size=10,
    )
    add_para(doc, "3. 尺寸：□ M01–M05 全部合格    □ 有不合格项：________________", size=10)
    add_para(doc, "4. 电气：□ E01/E02 均与预期符合    □ 不符合项：________________", size=10)
    add_para(doc, "5. 综合质量处置：□接收    □返工    □报废", size=10)
    add_para(
        doc,
        "6. 依据（对照验收要求表条款/现象名）：________________________________________________",
        size=10,
    )
    add_para(doc, "7. 简要说明：________________________________________________", size=10)
    add_para(
        doc,
        "复核：□ 单位齐全  □ 编号与版本一致  □ 分类与明细一致  □ 处置结论已填  □ 页码/附件完整  "
        "签注（按赛场要求）：________",
        size=10,
        space_before=4,
    )
    add_para(
        doc,
        "（标准名称 GB/T 4588-2025 条款待正式文本核对；本卷以竞赛用验收表述为准。评分见技术文件。）",
        size=9,
        space_before=6,
    )
    doc.save(str(path))
    print("[OK]", path.name)


# ---------- 3. 尺寸基准表 ----------
def build_dimension_docx(path: Path):
    doc = new_doc()
    add_para(
        doc,
        "尺寸基准表（学生组·模块D）",
        cn="黑体",
        size=12,
        bold=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )
    add_para(
        doc,
        "选手用·随卷下发。按下列基准测量并在《检测记录表》D-3 填写实测值与合格判定。"
        "不含裁判标定真值。投板后若名义微调，以赛场最新尺寸基准表为准。",
        size=9,
    )

    add_para(doc, "一、测量项目、工具与基准（M01–M05）", cn="黑体", size=11, bold=True)
    t = doc.add_table(rows=6, cols=6)
    t.style = "Table Grid"
    fill_table(
        t,
        [
            ["编号", "测量项目", "工具与精度要求", "单位", "名义/合格范围（骨架）", "备注"],
            [
                "M01",
                "板长 L",
                "游标卡尺（建议分度 0.02 mm）",
                "mm",
                "80.00 ± 0.20\n（79.80～80.20）",
                "沿长边；避开毛刺/外形缺口",
            ],
            [
                "M02",
                "板宽 W",
                "游标卡尺（建议分度 0.02 mm）",
                "mm",
                "60.00 ± 0.20\n（59.80～60.20）",
                "沿短边；垂直于板长",
            ],
            [
                "M03",
                "板厚 T",
                "厚度规或千分尺",
                "mm",
                "1.60 ± 0.15\n（1.45～1.75）",
                "板中平坦区；避开铜瘤/丝印厚堆",
            ],
            [
                "M04",
                "定位孔径",
                "卡尺或孔规（孔径适配）",
                "mm",
                "2.00 ± 0.10\n（1.90～2.10）",
                "指定定位孔内径；全卷统一测此孔",
            ],
            [
                "M05",
                "指定点线宽或线距",
                "10～20× 测量显微镜\n或带刻度放大镜",
                "mm",
                "见下节指定点；\n允差 ±0.05 或对照阈值",
                "与 V02 **不同点**；仅 1 处",
            ],
        ],
    )

    add_para(doc, "二、M05 指定点说明（骨架 · 投板后可改坐标不改编号）", cn="黑体", size=11, bold=True, space_before=8)
    t2 = doc.add_table(rows=4, cols=2)
    t2.style = "Table Grid"
    fill_table(
        t2,
        [
            ["项", "内容"],
            ["特征标识", "板面丝印/特征名：**M05测点**（或赛场示意图标注点）；记录时抄写特征名"],
            ["测量对象", "二选一写死于赛场说明：□ 线宽  □ 线距（正式卷只保留一项；默认优先**线宽**）"],
            [
                "方法",
                "在显微镜下对准指定点；估读或刻度读数至 0.01 mm 量级；填写实测值并与合格范围比对。"
                "不做全板扫描。不得以 V02 底层线宽变窄点代替 M05。",
            ],
        ],
        header_bold=True,
    )
    add_para(
        doc,
        "M05 合格范围（骨架默认，投板标定后可修订）：标称示例 0.25 mm，合格带 0.20～0.30 mm（±0.05）；"
        "或以验收阈值「≥ / ≤ 某某 mm 为合格」。正式下发以本表最新印刷版为准。",
        size=9.5,
        space_before=4,
    )

    add_para(doc, "三、测量与记录规则", cn="黑体", size=11, bold=True, space_before=8)
    for s in [
        "1. 单位统一为 mm；建议保留至 0.01 mm。",
        "2. 每项须填：实测值 + 与基准比对判定（合格/不合格）。",
        "3. 工艺参数可按 mm（mil）理解；**填写实测以 mm 为主**。",
        "4. **禁止**：普通卡尺测量 0.30 mm 级小孔；无指定点的全板线宽/线距扫描；量 Gerber。",
        "5. M04 仅测尺寸基准表指定的定位孔（建议标称 2.00 mm）；不得自选微孔充数。",
        "6. 本表不提供裁判答案真值；以赛场样件与本表基准为准。",
        "7. 厚度测量避开缺陷堆高与异常镀层堆积区。",
    ]:
        add_para(doc, s, size=10.5)

    add_para(doc, "四、与外观项关系", cn="黑体", size=11, bold=True, space_before=8)
    add_para(
        doc,
        "V02「底层线宽变窄」为**外观定性**采分点；M05 为**指定点定量/半定量**采分点，二者不得使用同一几何点重复计分。"
        "尺寸不合格项须在综合判定中体现。",
        size=10.5,
    )
    doc.save(str(path))
    print("[OK]", path.name)


# ---------- 4. 验收要求表 ----------
def build_acceptance_docx(path: Path):
    doc = new_doc()
    add_para(
        doc,
        "验收要求表（学生组·模块D）",
        cn="黑体",
        size=12,
        bold=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )
    add_para(
        doc,
        "选手用·随卷下发。用于 F1 外观缺陷判定与单件质量处置。"
        "引用标准名称 GB/T 4588-2025：对应条款待正式文本核对；下表为**本卷竞赛用验收表述**。",
        size=9,
    )

    add_para(doc, "一、外观必检 F1（V01–V05）判定语句", cn="黑体", size=11, bold=True)
    t = doc.add_table(rows=6, cols=4)
    t.style = "Table Grid"
    fill_table(
        t,
        [
            ["编号", "类型（标准名）", "不合格（应判为缺陷）", "合格说明"],
            [
                "V01",
                "顶层线路缺口",
                "顶层导体可见断开/缺口，铜皮不连续，存在开路风险",
                "顶层该区域线路连续，无明显缺口",
            ],
            [
                "V02",
                "底层线宽变窄",
                "底层可见线宽局部明显变窄/颈缩（定性可辨），明显偏离相邻正常线宽",
                "底层线宽均匀，无明显颈缩（本项定性，不替代 M05）",
            ],
            [
                "V03",
                "孔破盘",
                "孔与焊盘关系异常导致焊盘破损/破盘，焊环严重缺失或盘体破裂可辨",
                "焊盘完整可辨，无破盘",
            ],
            [
                "V04",
                "阻焊未开窗或开窗偏移",
                "应开窗处阻焊覆盖焊盘（未开窗），或开窗明显偏移导致异常露铜/盖盘",
                "开窗与焊盘基本对准，无未开窗盖盘",
            ],
            [
                "V05",
                "丝印缺失或压焊盘",
                "关键字符缺失/残缺不可辨，或丝印明显压在焊盘上影响可焊/可识别",
                "字符完整可辨，丝印未压有效焊盘",
            ],
        ],
    )
    add_para(
        doc,
        "说明：V04、V05 卷面按「二选一现象」记录实际所见（未开窗 **或** 开窗偏移；缺失 **或** 压焊盘），"
        "类型名写标准名或允许同义即可。底面缺陷须检（含 V02 等）。",
        size=9,
        space_before=4,
    )

    add_para(doc, "二、尺寸与微几何", cn="黑体", size=11, bold=True, space_before=8)
    add_para(
        doc,
        "以《尺寸基准表》M01–M05 为准：任一项超出合格范围 → 该尺寸项**不合格**。"
        "M05 仅评指定点；不得因未扫描全板微距扣分。",
        size=10.5,
    )

    add_para(doc, "三、基础电气", cn="黑体", size=11, bold=True, space_before=8)
    t2 = doc.add_table(rows=3, cols=3)
    t2.style = "Table Grid"
    fill_table(
        t2,
        [
            ["编号", "预期", "竞赛判定"],
            ["E01（TP1–TP2）", "开路", "电阻 ≥ 1 MΩ 或表显 OL/开路 → 与预期符合；否则不符合"],
            ["E02（TP3–TP4）", "短路", "电阻 ≤ 1 Ω 或导通蜂鸣 → 与预期符合；否则不符合"],
        ],
    )
    add_para(
        doc,
        "电气不合格指：测点对结果与预期不符（该开路却低阻/导通，或该短路却高阻/开路）。"
        "不做网络分段与故障定位。",
        size=9.5,
        space_before=4,
    )

    add_para(doc, "四、单件质量处置（接收 / 返工 / 报废）", cn="黑体", size=11, bold=True, space_before=8)
    t3 = doc.add_table(rows=4, cols=2)
    t3.style = "Table Grid"
    fill_table(
        t3,
        [
            ["结论", "条件（竞赛表述）"],
            [
                "接收",
                "V01–V05 均未构成不合格（或仅干扰项误报不计），且 M01–M05 全部合格，"
                "且 E01/E02 均与预期符合。",
            ],
            [
                "返工",
                "存在可返工类缺陷（如 V05 丝印类、V04 局部阻焊类等，板仍具备返工价值），"
                "且无致命线路开路/严重破盘/电气致命不符/尺寸严重超差；或尺寸轻微超差且可返工。",
            ],
            [
                "报废",
                "存在致命缺陷导致板不可用，例如：V01 线路缺口/开路、V03 严重孔破盘、"
                "外形严重破损（若发现）、尺寸严重超差、E01/E02 与预期严重不符且表明网络失效等。",
            ],
        ],
    )
    add_para(
        doc,
        "选手须在记录表勾选唯一结论并写简要依据（对照上表条款/编号）。"
        "正式等级与扣分细则以技术工作文件为准；样题采用上表可操作表述。"
        "干扰项（临界合格特征）不计入必检 5，误报不强制按缺陷处置。",
        size=9.5,
        space_before=6,
    )
    doc.save(str(path))
    print("[OK]", path.name)


# ---------- 5. 基础电气检测表 ----------
def build_electrical_docx(path: Path):
    doc = new_doc()
    add_para(
        doc,
        "基础电气检测表（学生组·模块D）",
        cn="黑体",
        size=12,
        bold=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )
    add_para(
        doc,
        "选手用·随卷下发。**仅**指定测试点对的开路/短路基础检测与电阻记录。"
        "本表**不是**故障定位表，不含网络分段、应通应断全表、高阻分析。",
        size=9,
    )

    add_para(doc, "一、检测项目（E01–E02）", cn="黑体", size=11, bold=True)
    t = doc.add_table(rows=3, cols=5)
    t.style = "Table Grid"
    fill_table(
        t,
        [
            ["编号", "类型", "测试点对（丝印）", "预期结果", "电阻判据（默认）"],
            ["E01", "开路", "TP1 – TP2", "开路", "≥ 1 MΩ，或表显 OL / 开路"],
            ["E02", "短路", "TP3 – TP4", "短路", "≤ 1 Ω，或连续档/蜂鸣导通"],
        ],
    )

    add_para(doc, "二、仪器与操作", cn="黑体", size=11, bold=True, space_before=8)
    for s in [
        "1. 仪器：数字万用表；优先电阻档，短路项可用蜂鸣通断档辅助，但须记录电阻读数或明确导通结论。",
        "2. 表笔接触 TP 焊盘金属面，避免只碰到阻焊；保持稳定读数后再记录。",
        "3. 记录字段：测试点对、档位、电阻读数、单位、判定（开路/短路/正常）、与预期是否符合。",
        "4. 单位：Ω / kΩ / MΩ 须写清；开路可用「OL」并在判定栏写「开路」。",
        "5. 每项独立测量；不得凭经验不测直接填结论。",
    ]:
        add_para(doc, s, size=10.5)

    add_para(doc, "三、电阻判定规则（竞赛口径）", cn="黑体", size=11, bold=True, space_before=8)
    t2 = doc.add_table(rows=4, cols=3)
    t2.style = "Table Grid"
    fill_table(
        t2,
        [
            ["情形", "读数特征", "判定写法"],
            ["开路成立", "≥ 1 MΩ 或 OL", "判定：开路；若项目为 E01，则「与预期符合」"],
            ["短路成立", "≤ 1 Ω 或蜂鸣导通", "判定：短路；若项目为 E02，则「与预期符合」"],
            [
                "介于中间/异常",
                "例如数 kΩ～数百 kΩ 等",
                "如实记录读数；判定可写「异常/非预期」；「与预期符合」选否，并交由综合判定说明",
            ],
        ],
    )
    add_para(
        doc,
        "阈值说明：上表为命题默认骨架；赛场若公布微调阈值，以赛场书面说明为准。"
        "裁判评分以实物标定与同义采点说明为准。",
        size=9.5,
        space_before=4,
    )

    add_para(doc, "四、测试点说明", cn="黑体", size=11, bold=True, space_before=8)
    t3 = doc.add_table(rows=5, cols=3)
    t3.style = "Table Grid"
    fill_table(
        t3,
        [
            ["丝印", "用途", "注意"],
            ["TP1", "E01 一端", "仅服务开路项"],
            ["TP2", "E01 另一端", "与 TP1 组成开路网络两端"],
            ["TP3", "E02 一端", "仅服务短路项"],
            ["TP4", "E02 另一端", "与 TP3 为设计短路对"],
        ],
    )
    add_para(
        doc,
        "本模块仅 **TP1–TP4** 四个专用测试点（默认无 TP0）。"
        "禁止按旧版多网络故障定位方式扩展测点或自建故障树。",
        size=9.5,
        space_before=4,
    )

    add_para(doc, "五、禁止事项", cn="黑体", size=11, bold=True, space_before=8)
    for s in [
        "1. 不上电、不加外部电源做功能测试。",
        "2. 不做故障定位、网络分段、最短路径推断、高阻机理分析。",
        "3. 不要求填写「应通应断全表」或 NET 网络名表。",
        "4. 不得刮除阻焊扩大露铜、不得破坏样件。",
        "5. 不得将归档旧件中的「电气连通与故障定位表」当作本卷现行要求。",
    ]:
        add_para(doc, s, size=10.5)

    add_para(doc, "六、填表示例", cn="黑体", size=11, bold=True, space_before=8)
    t4 = doc.add_table(rows=3, cols=7)
    t4.style = "Table Grid"
    fill_table(
        t4,
        [
            ["编号", "测试点对", "档位", "读数", "单位", "判定", "与预期"],
            ["E01", "TP1–TP2", "20 MΩ", "OL", "—", "开路", "是"],
            ["E02", "TP3–TP4", "200 Ω", "0.3", "Ω", "短路", "是"],
        ],
    )
    doc.save(str(path))
    print("[OK]", path.name)


# ---------- 6. 板面分区说明 ----------
def build_partition_docx(path: Path):
    doc = new_doc()
    add_para(
        doc,
        "板面分区说明（学生组·模块D）",
        cn="黑体",
        size=12,
        bold=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )
    add_para(
        doc,
        "选手用·随卷下发。用于理解板面坐标、特征名、版本丝印与测试点。"
        "本说明不是“一区一错”提示。测试点**仅**服务基础开短路，**非**故障定位网络。",
        size=9,
    )

    add_para(doc, "一、版本丝印", cn="黑体", size=11, bold=True)
    add_para(
        doc,
        "板面丝印版本：`MOD-D-S-A`（A 版）或 `MOD-D-S-B`（B 版）。"
        "请在记录表勾选并与袋装标签核对。A/B 缺陷类型与数量相同，位置与部分真值不同。",
        size=10.5,
    )

    add_para(doc, "二、A1–D4 网格", cn="黑体", size=11, bold=True, space_before=8)
    add_para(
        doc,
        "A1–D4 是印在成品裸板顶层丝印上的位置坐标网格，用来标明缺陷所在区域，方便书写与复核。"
        "它不是 16 个错误点，也不表示每个格子必有缺陷。",
        size=10.5,
    )
    add_para(
        doc,
        "约定：行用字母 A、B、C、D（自上而下）；列用数字 1、2、3、4（自左而右）。"
        "例如 B2 表示第 B 行第 2 列区域。以实物丝印为准。",
        size=10.5,
    )
    add_para(doc, "示意（顶层俯视）：", size=10.5)
    t = doc.add_table(rows=5, cols=5)
    t.style = "Table Grid"
    grid = [
        ["", "1", "2", "3", "4"],
        ["A", "A1", "A2", "A3", "A4"],
        ["B", "B1", "B2", "B3", "B4"],
        ["C", "C1", "C2", "C3", "C4"],
        ["D", "D1", "D2", "D3", "D4"],
    ]
    fill_table(t, grid)

    add_para(doc, "三、特征区名称（丝印）", cn="黑体", size=11, bold=True, space_before=8)
    add_para(doc, "记录时建议“所在面 + 网格 + 特征名”，例如：“顶层 B2 / 走线区”。", size=10.5)
    t2 = doc.add_table(rows=8, cols=2)
    t2.style = "Table Grid"
    fill_table(
        t2,
        [
            ["特征名", "含义（便于定位，非缺陷清单）"],
            ["走线区", "主要信号走线分布区域（V01/V02/M05 相关）"],
            ["孔阵", "多个通孔/焊盘排列区域（V03 相关）"],
            ["阻焊对比区", "阻焊开窗/覆盖特征对比区域（V04 相关）"],
            ["丝印条", "字符/标识集中区域（V05 相关）"],
            ["板框", "外形边缘区域"],
            ["测试点区", "TP1–TP4 专用测试点区域"],
            ["M05测点", "指定线宽/线距测量点标识（与 V02 不同点）"],
        ],
    )

    add_para(doc, "四、测试点 TP1–TP4（基础电气）", cn="黑体", size=11, bold=True, space_before=8)
    t3 = doc.add_table(rows=5, cols=3)
    t3.style = "Table Grid"
    fill_table(
        t3,
        [
            ["丝印", "建议面", "用途"],
            ["TP1", "顶", "E01 开路一端"],
            ["TP2", "顶或底", "E01 开路另一端"],
            ["TP3", "顶", "E02 短路一端"],
            ["TP4", "顶", "E02 短路另一端"],
        ],
    )
    add_para(
        doc,
        "仅上述 4 点用于本模块电气得分。禁止扩展为旧版 TP1–TP8 故障定位网络测法。"
        "具体操作与判据见《基础电气检测表》。",
        size=9.5,
        space_before=4,
    )

    add_para(doc, "五、顶面与底面", cn="黑体", size=11, bold=True, space_before=8)
    add_para(
        doc,
        "须检查顶面与底面。底面至少存在不少于 2 处必检外观相关特征（含 V02 等）。"
        "底面网格可按“翻转后对应关系”记录为“底层 + 网格/特征”；以实物丝印与赛场说明为准。"
        "版本丝印一般在顶面。",
        size=10.5,
    )

    add_para(doc, "六、缺陷与测点记录写法示例", cn="黑体", size=11, bold=True, space_before=8)
    for s in [
        "示例1：A-V01 — 顶层 / B2 / 走线区 — 顶层线路缺口，铜皮局部断开。",
        "示例2：A-V02 — 底层 / C2 / 走线区 — 底层线宽变窄（定性）。",
        "示例3：A-M05 — 顶层 / 指定 M05测点 — 线宽实测 0.24 mm，合格。",
        "示例4：A-E01 — TP1–TP2 — 电阻 OL — 判定开路 — 与预期符合。",
    ]:
        add_para(doc, s, size=10.5)

    add_para(doc, "七、重要说明", cn="黑体", size=11, bold=True, space_before=8)
    for s in [
        "1. 网格用于定位，不是“一格一个错误”。",
        "2. 缺陷分布可能不均匀；同一网格可有多处，部分网格可以没有缺陷。",
        "3. 请检查顶层与底层；勿只查正面。",
        "4. 本模块电气仅为 TP 点对点开短路；不做故障定位与网络分段。",
        "5. 本模块不上电，不量测 Gerber，不进行焊接/贴装质量评价。",
        "6. 编号前缀：V 外观、M 尺寸、E 电气、TP 测试点；与检测记录表一致。",
    ]:
        add_para(doc, s, size=10.5)

    doc.save(str(path))
    print("[OK]", path.name)


# ---------- 7. 选手材料清单 ----------
def build_material_list_docx(path: Path):
    doc = new_doc()
    add_para(
        doc,
        "选手材料清单（学生组·模块D）",
        cn="黑体",
        size=12,
        bold=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )
    add_para(
        doc,
        "印制电路制作工赛项 · 成品PCB裸板质量检测与判定（学生组）· R1 / 45 min",
        size=10.5,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )

    add_para(doc, "一、选手可见材料与附件", cn="黑体", size=11, bold=True)
    t = doc.add_table(rows=10, cols=3)
    t.style = "Table Grid"
    fill_table(
        t,
        [
            ["序号", "材料名称", "说明"],
            [
                "1",
                "双面成品裸板质检样件",
                "FR-4 双面、无元器件；约 80×60 mm；无铅喷锡；含 F1 五类外观 + M05 + E01/E02；版本 A 或 B",
            ],
            [
                "2",
                "学生组_模块D_检测记录表",
                "提交用；覆盖 D-1～D-5（含外观/尺寸/电气/综合）",
            ],
            ["3", "学生组_模块D_验收要求表", "F1 判定语句 + 接收/返工/报废"],
            ["4", "学生组_模块D_尺寸基准表", "M01–M05 工具、单位、合格范围"],
            ["5", "学生组_模块D_基础电气检测表", "TP 开短路与电阻规则；非故障定位"],
            ["6", "学生组_模块D_板面分区说明", "A1–D4、特征名、TP1–TP4、版本丝印"],
            ["7", "学生组_模块D_仪器使用与安全要求", "仪器使用与安全；不上电"],
            ["8", "本清单", "可与任务书要点合并印发"],
            ["9", "任务书要点（可选印发）", "D-1～D-5 与 45 min 结构摘要"],
        ],
    )

    add_para(doc, "二、赛场工位工具（允许）", cn="黑体", size=11, bold=True, space_before=8)
    t2 = doc.add_table(rows=7, cols=3)
    t2.style = "Table Grid"
    fill_table(
        t2,
        [
            ["序号", "工具", "用途"],
            ["1", "游标卡尺", "M01/M02/M04（板长/宽/定位孔径）"],
            ["2", "厚度规或千分尺", "M03 板厚"],
            ["3", "测量显微镜（10～20×）或带刻度放大镜", "M05 指定点线宽/线距；外观辅助"],
            ["4", "放大镜", "外观 V01–V05"],
            ["5", "数字万用表", "E01/E02 开短路与电阻"],
            ["6", "照明/侧光、防静电用品等", "以赛场现场清单为准"],
        ],
    )

    add_para(doc, "三、裁判内部材料（不下发选手）", cn="黑体", size=11, bold=True, space_before=8)
    add_para(
        doc,
        "缺陷参考答案（A/B）、尺寸与线宽参考答案、电气参考答案、同义答案与采点说明、"
        "评分表骨架（40 分）、标准板/备用板、缺陷照片与坐标（投板后）。",
        size=10.5,
    )

    add_para(doc, "四、明确不提供 / 不作为现行考核", cn="黑体", size=11, bold=True, space_before=8)
    add_para(
        doc,
        "通断网络表、电气连通与故障定位表、应通应断全表、网络分段表、上电功能测试、"
        "Gerber 量测、实际返修操作；`_归档移出_电气旧件/` 内旧表不得作现行下发。",
        size=10.5,
    )
    doc.save(str(path))
    print("[OK]", path.name)


# ---------- 8. 仪器使用与安全要求 ----------
def write_instrument_safety():
    path = PLAYER / "学生组_模块D_仪器使用与安全要求.md"
    path.write_text(
        """# 仪器使用与安全要求（学生组·模块D）

> 版本：2026-07-21 · R1 / 45 min
> 选手用·随卷下发（可打印为 docx 等价文本）。
> **本模块不上电、不做功能通电测试。**

---

## 1. 适用范围

适用于学生组模块D「成品PCB裸板质量检测与判定」工位：外观检查、尺寸与指定点微几何测量、基础开短路点测。

---

## 2. 仪器与正确用途

| 仪器 | 用途 | 使用要点 |
|------|------|----------|
| 游标卡尺 | 板长、板宽、定位孔径 | 测量前对零；轻卡勿划伤焊盘；读数 mm，建议 0.01 mm |
| 厚度规/千分尺 | 板厚 | 测平坦区；避开丝印厚堆与缺陷隆起 |
| 测量显微镜（10～20×）或带刻度放大镜 | 指定点线宽/线距；外观辅助 | 仅测尺寸基准表指定点；估读稳定后记录 |
| 放大镜 | 外观缺陷 | 可配合侧光观察缺口、破盘、阻焊与丝印 |
| 数字万用表 | TP1–TP2、TP3–TP4 | 电阻档/蜂鸣档；表笔接触金属焊盘；记录档位与读数 |

---

## 3. 安全与样件保护

1. **禁止对样件上电**，禁止外接电源、信号源做功能测试。
2. 禁止刮削阻焊、刀挑线路、钻孔扩孔等破坏性操作。
3. 表笔与卡尺测量时用力适度，避免撬起焊盘或划伤喷锡面。
4. 遵守防静电要求（若赛场提供腕带/台垫，按规定佩戴使用）。
5. 爱护量具：卡尺轻拿轻放；显微镜调焦避免镜头撞击板面。
6. 发现样件运输损伤或仪器故障，立即举手向裁判示意，不得自行拆修仪器。

---

## 4. 万用表专项

1. 先确认档位再测量；禁止在电阻档误用电流档硬测。
2. 开路项（E01）：高阻档观察是否 OL/≥1 MΩ。
3. 短路项（E02）：低阻档或蜂鸣档，确认 ≤1 Ω 或导通。
4. 不得将表笔短接后的读数当作样件结果。
5. 不得扩展测点做故障定位树。

---

## 5. 卡尺与显微镜专项

1. 卡尺：**禁止**用于 0.30 mm 级小孔作为得分项。
2. 显微镜：只测 **M05 指定点**；禁止无指定点全板扫线宽/线距。
3. 读数与单位必须写入记录表；无单位可能影响规范分。

---

## 6. 环境与提交

1. 保持工位整洁；样件用后放回防静电袋/托盘。
2. 记录表字迹清楚；版本（A/B）与样件丝印一致。
3. 按赛场要求提交纸质/电子成果；电子命名示例：`赛位号_模块D`。

---

## 7. 与考核边界的关系

| 允许 | 禁止 |
|------|------|
| 外观目检 + 放大观察 | 失效机理长文分析 |
| 尺寸与 1 处微几何 | 量 Gerber |
| 4 个 TP 基础开短路 | 故障定位 / 网络分段 / 上电 |

详见《命题边界确认表》与《任务书要点》。
""",
        encoding="utf-8",
    )
    print("[OK]", path.name)


def main():
    PLAYER.mkdir(parents=True, exist_ok=True)
    write_task_points()
    build_record_docx(PLAYER / "学生组_模块D_检测记录表.docx")
    build_dimension_docx(PLAYER / "学生组_模块D_尺寸基准表.docx")
    build_acceptance_docx(PLAYER / "学生组_模块D_验收要求表.docx")
    build_electrical_docx(PLAYER / "学生组_模块D_基础电气检测表.docx")
    build_partition_docx(PLAYER / "学生组_模块D_板面分区说明.docx")
    build_material_list_docx(PLAYER / "学生组_模块D_选手材料清单.docx")
    write_instrument_safety()
    print("[DONE] player materials R1")


if __name__ == "__main__":
    main()
