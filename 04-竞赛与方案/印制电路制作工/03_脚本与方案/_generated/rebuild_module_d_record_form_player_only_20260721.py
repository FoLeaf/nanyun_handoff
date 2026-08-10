# -*- coding: utf-8 -*-
"""
模块D 检测记录表 · 选手-only 优化
对齐完整版模块D：无时长、无命题黑话、字段可填即可
"""
from __future__ import annotations

import shutil
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "02_样题" / "学生组_模块D" / "选手" / "学生组_模块D_检测记录表.docx"
BACKUP = ROOT / "05_归档备份" / f"module-d-record-form-{datetime.now().strftime('%Y%m%d_%H%M%S')}"

C_NAVY = "1F3864"
C_BLUE = "2E75B6"
C_GRAY = "595959"
C_RED = "C00000"


def set_run(run, cn="微软雅黑", en=None, size=10.5, bold=False, color=None):
    en = en or cn
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = en
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    for k in ("w:ascii", "w:hAnsi", "w:cs"):
        rFonts.set(qn(k), en)
    rFonts.set(qn("w:eastAsia"), cn)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_p(doc, text, *, cn="微软雅黑", size=10.5, bold=False, color=None, align=None, before=0, after=4):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    run = p.add_run(text)
    set_run(run, cn=cn, en=cn if cn != "黑体" else "黑体", size=size, bold=bold, color=color)
    return p


def shade(cell, fill):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def cell_text(cell, text, *, header=False, size=9, bold=False, center=False):
    cell.text = ""
    p = cell.paragraphs[0]
    if center or header:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(str(text))
    set_run(
        run,
        cn="微软雅黑",
        size=size,
        bold=bold or header,
        color="FFFFFF" if header else None,
    )
    if header:
        shade(cell, C_BLUE)


def fill_table(table, rows, header=True):
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            is_h = header and ri == 0
            cell_text(table.cell(ri, ci), val, header=is_h, size=9 if is_h else 9.5, bold=is_h, center=is_h)


def setup(doc):
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.left_margin = Cm(1.8)
    sec.right_margin = Cm(1.8)
    sec.top_margin = Cm(1.5)
    sec.bottom_margin = Cm(1.5)


def build():
    if OUT.exists():
        BACKUP.mkdir(parents=True, exist_ok=True)
        shutil.copy2(OUT, BACKUP / OUT.name)
        print("[BACKUP]", BACKUP)

    doc = Document()
    setup(doc)

    # header like B/C
    add_p(doc, '江西省"振兴杯"职业技能大赛', cn="黑体", size=18, bold=True, color=C_NAVY, align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
    add_p(doc, '"印制电路制作工"赛项', cn="黑体", size=16, bold=True, color=C_NAVY, align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
    add_p(doc, "━━━━━━━━━━━━━━━━━━━━", size=11, color=C_BLUE, align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
    add_p(doc, "模块D：成品PCB裸板质量检测与判定", cn="黑体", size=20, bold=True, color=C_RED, align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
    add_p(doc, "检测记录表（学生组）", cn="微软雅黑", size=14, bold=True, color=C_GRAY, align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
    add_p(doc, "━━━━━━━━━━━━━━━━━━━━", size=11, color=C_BLUE, align=WD_ALIGN_PARAGRAPH.CENTER, after=6)

    add_p(
        doc,
        "填写说明：判定依据见《竞赛样题》模块D第（三）节「检测依据」。本表填写后提交。"
        "命名示例：赛位号_模块D。不得标注姓名（按赛场规则）。",
        size=10,
        bold=True,
        color=C_RED,
        after=8,
    )

    # basic info
    t0 = doc.add_table(rows=2, cols=4)
    t0.style = "Table Grid"
    fill_table(
        t0,
        [
            ["赛位号", "", "日期", ""],
            ["样件编号", "", "版本（丝印）", "□ MOD-D-S-A    □ MOD-D-S-B"],
        ],
        header=False,
    )
    for row in t0.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for r in p.runs:
                    set_run(r, size=9.5)

    add_p(doc, "一、D-1 检验准备", cn="黑体", size=12, bold=True, color=C_NAVY, before=10, after=4)
    add_p(
        doc,
        "□ 已阅读赛题模块D第（三）节检测依据    □ 顶面/底面方向已确认    "
        "□ 卡尺已对零    □ 厚度规可用    □ 显微镜/放大镜可用    □ 万用表可用    "
        "□ 版本丝印已抄录    □ 样件外观运输情况：□正常  □异常：__________",
        size=10,
        after=6,
    )

    add_p(doc, "二、D-2 外观缺陷检测（V01～V05）", cn="黑体", size=12, bold=True, color=C_NAVY, before=6, after=3)
    add_p(
        doc,
        "须检查顶面与底面。对照赛题第（三）节外观判定表填写。位置写清面、网格或特征名及方位即可。",
        size=10,
        after=4,
    )
    t1 = doc.add_table(rows=8, cols=6)
    t1.style = "Table Grid"
    rows = [
        ["编号", "所在面", "网格/特征", "缺陷类型", "现象简述", "判定（合格/不合格）"],
        ["V01", "顶", "", "顶层线路缺口", "", ""],
        ["V02", "底", "", "底层线宽变窄", "", ""],
        ["V03", "", "", "孔破盘", "", ""],
        ["V04", "", "", "阻焊未开窗或开窗偏移", "", ""],
        ["V05", "", "", "丝印缺失或压焊盘", "", ""],
        ["", "", "", "", "", ""],
        ["", "", "", "", "", ""],
    ]
    fill_table(t1, rows)

    add_p(doc, "三、D-3 尺寸与指定点测量（M01～M05）", cn="黑体", size=12, bold=True, color=C_NAVY, before=10, after=3)
    add_p(
        doc,
        "单位 mm。对照赛题第（三）节尺寸基准表。M05 只测指定点，不得与 V02 为同一处。",
        size=10,
        after=3,
    )
    add_p(
        doc,
        "M05 指定点位置/标识：____________________    测量对象：□线宽  □线距",
        size=10,
        after=4,
    )
    t2 = doc.add_table(rows=6, cols=5)
    t2.style = "Table Grid"
    fill_table(
        t2,
        [
            ["编号", "项目", "工具", "实测值 (mm)", "合格/不合格"],
            ["M01", "板长 L", "游标卡尺", "", ""],
            ["M02", "板宽 W", "游标卡尺", "", ""],
            ["M03", "板厚 T", "厚度规/千分尺", "", ""],
            ["M04", "定位孔径", "卡尺/孔规", "", ""],
            ["M05", "指定点线宽或线距", "测量显微镜", "", ""],
        ],
    )

    add_p(doc, "四、D-4 基础电气检测（E01～E02）", cn="黑体", size=12, bold=True, color=C_NAVY, before=10, after=3)
    add_p(
        doc,
        "仅测指定测试点。对照赛题第（三）节电气判据表。",
        size=10,
        after=4,
    )
    t3 = doc.add_table(rows=3, cols=7)
    t3.style = "Table Grid"
    fill_table(
        t3,
        [
            ["编号", "类型", "测试点", "档位", "读数", "判定（开路/短路/正常）", "与预期一致"],
            ["E01", "开路", "TP1–TP2", "", "", "", "□是  □否"],
            ["E02", "短路", "TP3–TP4", "", "", "", "□是  □否"],
        ],
    )

    add_p(doc, "五、D-5 分类判定与提交", cn="黑体", size=12, bold=True, color=C_NAVY, before=10, after=4)
    add_p(doc, "1. 外观必检：V01____  V02____  V03____  V04____  V05____  （已检查□）", size=10, after=3)
    add_p(doc, "2. 分类汇总：线路____  孔/盘____  阻焊____  丝印____  其它____", size=10, after=3)
    add_p(doc, "3. 尺寸：□ M01～M05 全部合格    □ 有不合格：________________", size=10, after=3)
    add_p(doc, "4. 电气：□ E01、E02 均与预期一致    □ 不符合：________________", size=10, after=3)
    add_p(doc, "5. 质量处置：□ 接收    □ 返工    □ 报废", size=10, after=3)
    add_p(
        doc,
        "6. 依据（对照赛题检测依据，简要说明）：________________________________________________",
        size=10,
        after=3,
    )
    add_p(doc, "7. 其他说明：________________________________________________", size=10, after=6)
    add_p(
        doc,
        "提交前检查：□ 表头信息完整  □ 单位齐全  □ 版本与实物一致  □ 结论已填",
        size=10,
        after=4,
    )
    add_p(doc, "签注（按赛场要求）：____________________", size=10, after=2)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUT))
    print("[OK]", OUT)


def verify():
    d = Document(str(OUT))
    text = "\n".join(p.text for p in d.paragraphs)
    for t in d.tables:
        for row in t.rows:
            for c in row.cells:
                text += "\n" + c.text
    checks = {
        "无45分钟": "45" not in text and "分钟" not in text,
        "无R1": "R1" not in text,
        "无干扰项话术": "干扰项" not in text,
        "无分册附件清单": "验收表" not in text and "分区说明" not in text,
        "有V01": "V01" in text,
        "有M05": "M05" in text,
        "有E01": "E01" in text,
        "有接收": "接收" in text,
        "指向第（三）节": "第（三）节" in text or "检测依据" in text,
    }
    for k, v in checks.items():
        print(("OK" if v else "FAIL"), k)
    if not all(checks.values()):
        raise SystemExit(1)
    print("VERIFY PASS", "paras", len(d.paragraphs), "tables", len(d.tables))


if __name__ == "__main__":
    build()
    verify()
