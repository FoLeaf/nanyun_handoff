# -*- coding: utf-8 -*-
"""
完整回写：竞赛样题（完整版）· 学生组模块D

口径冻结：
- R1 / 45 min / 基础电气 E01·E02
- 题干第（三）节内嵌全部文字表格素材
- 唯一单独提交附件：检测记录表
- 规范 F：黑体标题 + 仿宋正文 12

同时：删除模块D区域残留旧表；同步现行包说明。
"""
from __future__ import annotations

import shutil
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

ROOT = Path(__file__).resolve().parents[2]
SAMPLE = ROOT / "02_样题"
PACK = SAMPLE / "学生组_模块D"
FULL = SAMPLE / "印制电路制作工赛项_竞赛样题（完整版）.docx"
BACKUP = ROOT / "05_归档备份" / f"module-d-full-rewrite-{datetime.now().strftime('%Y%m%d_%H%M%S')}"

C_BLUE = "2E75B6"


def set_run_font(run, name_cn="仿宋", name_en="Times New Roman", size_pt=12, bold=False, color_hex=None):
    run.bold = bold
    run.font.size = Pt(size_pt) if size_pt else None
    run.font.name = name_en
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:ascii"), name_en)
    rFonts.set(qn("w:hAnsi"), name_en)
    rFonts.set(qn("w:eastAsia"), name_cn)
    rFonts.set(qn("w:cs"), name_en)
    if color_hex:
        run.font.color.rgb = RGBColor.from_string(color_hex)


def add_para(doc, text, *, cn="仿宋", size=12, bold=False, space_after=4, space_before=0):
    p = doc.add_paragraph()
    try:
        p.style = doc.styles["Normal"]
    except Exception:
        pass
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(space_before)
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    run = p.add_run(text)
    en = "黑体" if cn == "黑体" else "Times New Roman"
    set_run_font(run, name_cn=cn, name_en=en, size_pt=size, bold=bold)
    return p


def shade_cell(cell, fill_hex):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill_hex)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def set_cell(cell, text, *, header=False, size=9.5):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if header else WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(str(text))
    set_run_font(
        run,
        name_cn="黑体" if header else "仿宋",
        name_en="Times New Roman",
        size_pt=9 if header else size,
        bold=header,
        color_hex="FFFFFF" if header else None,
    )
    if header:
        shade_cell(cell, C_BLUE)


def add_table(doc, rows):
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            set_cell(table.cell(ri, ci), val, header=(ri == 0))
    # spacer para after table for readability
    add_para(doc, "", space_after=2)
    return table


def el_text(el) -> str:
    return "".join(t.text or "" for t in el.iter(qn("w:t"))).strip()


def find_module_d_body_start(body) -> int:
    """Return body child index where Module D region starts (include orphan pre-title tables)."""
    kids = list(body.iterchildren())
    title_idx = None
    for i, child in enumerate(kids):
        tag = child.tag.split("}")[-1]
        if tag != "p":
            continue
        t = el_text(child)
        if t.startswith("模块D：") and ("成品" in t or "质量检测" in t or "学生组" in t):
            title_idx = i
            break
    if title_idx is None:
        raise RuntimeError("未找到模块D标题段落")

    start = title_idx
    # 向前收起紧邻的模块D旧材料表（标题前游离表）
    j = title_idx - 1
    while j >= 0:
        child = kids[j]
        tag = child.tag.split("}")[-1]
        if tag == "tbl":
            tx = el_text(child)
            if any(
                k in tx
                for k in (
                    "检测记录表",
                    "尺寸基准表",
                    "验收要求表",
                    "基础电气",
                    "板面分区",
                    "成品裸板质检样件",
                    "双面成品裸板",
                )
            ):
                start = j
                j -= 1
                continue
        break
    return start


def remove_body_from(body, start_idx: int):
    kids = list(body.iterchildren())
    # 保留 sectPr
    to_remove = []
    for i, child in enumerate(kids):
        if i < start_idx:
            continue
        tag = child.tag.split("}")[-1]
        if tag == "sectPr":
            continue
        to_remove.append(child)
    for el in to_remove:
        body.remove(el)
    print(f"[FULL] removed {len(to_remove)} body elements from index {start_idx}")


def write_module_d(doc: Document):
    H = lambda t, sz=14: add_para(doc, t, cn="黑体", size=sz, bold=True, space_before=6, space_after=4)
    B = lambda t, **kw: add_para(doc, t, cn="仿宋", size=12, bold=False, **kw)
    L = lambda t: add_para(doc, t, cn="黑体", size=12, bold=True, space_before=4, space_after=2)

    H("模块D：成品PCB裸板质量检测与判定（学生组）", 16)

    H("（一）模块考核点")
    B(
        "本模块（学生组）考核选手对双面成品PCB裸板（无元器件，已完成线路、阻焊、丝印和表面处理）进行"
        "外观缺陷检测、关键尺寸与指定点微几何测量、基础电气开短路点测、缺陷分类与单件质量处置（接收/返工/报废）的能力。"
        "正式考核时长45分钟（D-1～D-5）；模块内原始分40分（外观14／尺寸·微几何12／电气8／综合4／规范2；折算见技术工作文件）。"
        "不考核复杂网络分段、故障定位树、应通应断全表、上电功能调试、实际返修、贴装/焊接质量及Gerber量测。"
    )

    H("（二）模块简介")
    L("【任务背景】")
    B(
        "本模块为独立质检任务。检验对象为赛场提供的双面成品裸板质检样件（专用检测板），"
        "与模块B（EDA工程设计）、模块C（CAM审核与工艺文件编制）相互独立；"
        "仅做基础开短路点测，不做故障定位与网络分段，不上电，不量测Gerber。"
        "某批次双面成品裸板完成后进入抽检环节，你作为质检人员，需完成规范检验、准确测量、基础点测，"
        "对照本卷题干中的验收要求、尺寸基准与电气判据进行分类判定，并给出接收／返工／报废结论及简要依据。"
    )
    L("【检验对象】")
    B(
        "双面成品裸板质检样件1块（无元器件）。外形约80×60 mm，标称板厚1.6 mm；FR-4双面、绿阻焊白丝印、无铅喷锡。"
        "板面丝印含版本标识（MOD-D-S-A或MOD-D-S-B）、A1–D4分区网格、特征名及专用测试点（TP1–TP4）。"
        "须完成：外观必检5处（V01–V05）、尺寸M01–M04与指定点线宽/线距M05、基础电气E01开路与E02短路；"
        "可能含1～2处临界合格干扰特征（不计必检5）。须检查顶面与底面（底面缺陷不少于2处）。"
    )
    L("【提供材料】")
    B("1. 双面成品裸板质检样件1块；")
    B("2. 《学生组_模块D_检测记录表》（空白，选手填写并提交的唯一表格附件）；")
    B("3. 本卷模块D第（三）节题干内嵌素材（板面分区与测试点、尺寸基准、外观验收要求、基础电气判据、仪器与安全、材料与工具清单）。")
    L("【使用工具】")
    B(
        "游标卡尺、厚度规（或千分尺）、测量显微镜（10～20×或带刻度放大镜）、放大镜、数字万用表、侧光/照明、防静电用品等（以赛场清单为准）。"
    )

    H("（三）题干素材（直接使用，无需另附分册）")
    B(
        "下列表格与说明为本模块试题组成部分。选手依据本卷填写《检测记录表》并提交；"
        "除检测记录表外，不再另发验收表/尺寸表/电气表/分区说明等文字表格附件。"
    )

    L("1. 板面分区与测试点说明")
    B("版本丝印：MOD-D-S-A / MOD-D-S-B（以实物为准）。答题版本栏须与丝印一致。")
    B("分区网格：顶层参考A1–D4（列1–4，行A–D）。记录缺陷时填写所在面＋网格/特征名＋方位。")
    B("特征区名称（示例，以板面丝印为准）：走线区、孔阵、阻焊对比区、丝印条、板框、测试点区。")
    B("专用测试点仅服务基础电气E01/E02，不得理解为多网络故障定位表：")
    add_table(
        doc,
        [
            ["丝印", "用途", "说明"],
            ["TP1", "E01一端", "开路网络一端"],
            ["TP2", "E01另一端", "开路网络另一端"],
            ["TP3", "E02一端", "短路网络一端"],
            ["TP4", "E02另一端", "短路网络另一端"],
        ],
    )

    L("2. 尺寸基准要求（M01–M05）")
    B(
        "单位mm。M01–M04用卡尺/厚度规；M05用测量显微镜测指定点（与V02不同点）。"
        "禁止普通卡尺测0.30 mm级小孔；禁止无指定点全板扫微距。"
        "名义与公差为竞赛骨架，实物以赛场样件为准；判定时实测与下表比对。"
    )
    add_table(
        doc,
        [
            ["编号", "项目", "工具", "名义值", "允许误差", "判定"],
            ["M01", "板长L", "游标卡尺", "80.00", "±0.20", "合格/不合格"],
            ["M02", "板宽W", "游标卡尺", "60.00", "±0.20", "合格/不合格"],
            ["M03", "板厚T", "厚度规/千分尺", "1.60", "±0.15", "合格/不合格"],
            ["M04", "定位孔径", "卡尺/孔规", "2.00", "±0.10", "合格/不合格"],
            ["M05", "指定点线宽或线距", "测量显微镜", "见板面指定点丝印/特征名", "±0.05或对照阈值", "合格/不合格"],
        ],
    )
    B(
        "M05指定点：以板面丝印或特征名唯一标识的测量位为准（与V02底层线宽变窄外观点不得为同一采分点）。"
        "测量对象勾选线宽或线距其一；读数保留至0.01 mm（或赛场规定）。"
    )

    L("3. 外观缺陷验收要求（F1 · V01–V05）")
    B("下列为竞赛用不合格判定表述。须检顶底两面；底面缺陷不少于2处。")
    add_table(
        doc,
        [
            ["编号", "类型", "不合格判定（竞赛用表述）", "备注"],
            ["V01", "顶层线路缺口", "顶层导体出现可见断开/缺口，存在连通风险", "须定位顶面"],
            ["V02", "底层线宽变窄", "底层局部线宽明显变窄（外观定性，不要求在此处读数）", "与M05不同点"],
            ["V03", "孔破盘", "孔与焊盘关系异常导致破盘/焊盘缺损", "顶或底以实物为准"],
            ["V04", "阻焊未开窗或开窗偏移", "应开窗处被阻焊覆盖，或开窗相对焊盘明显偏移", "二选一现象"],
            ["V05", "丝印缺失或压焊盘", "字符缺失/残缺，或丝印压覆焊盘", "二选一现象"],
        ],
    )
    B("干扰项：可出现1～2处接近合格临界的特征，不计入必检5；误报不强制重扣外观分（评分细则另定）。")

    L("4. 基础电气检测要求（E01–E02）")
    B(
        "仅对指定测试点做点对点开路/短路判定。不做网络分段、故障定位树、应通应断全表、高阻分析、上电功能测试。"
        "电阻判据骨架如下（赛场书面说明或实物标定优先）："
    )
    add_table(
        doc,
        [
            ["编号", "类型", "测试点对", "预期", "默认电阻判据"],
            ["E01", "开路", "TP1–TP2", "开路", "≥1 MΩ 或表显OL/开路"],
            ["E02", "短路", "TP3–TP4", "短路", "≤1 Ω 或蜂鸣导通"],
        ],
    )
    B("记录要求：测试点对、档位、电阻读数、单位、判定（开路/短路/正常）、与预期是否符合。")

    L("5. 综合质量处置规则（接收／返工／报废）")
    B("接收：无必检致命缺陷，尺寸M01–M05与电气E01/E02均符合本卷要求。")
    B("返工：存在可返工类缺陷（如丝印、部分阻焊类）且尺寸/电气主体可接受——按本卷验收表述与裁判细则。")
    B("报废：存在线路缺口、孔破盘等严重影响功能/可靠性，或电气与预期严重不符且不可接受。")
    B("选手须在检测记录表给出结论＋简要依据（对照本卷第（三）节条款/现象名）。不要求深层工艺根因分析。")

    L("6. 仪器使用与安全要点")
    B("（1）卡尺使用前对零；测量板厚避开铜瘤/丝印厚堆；定位孔测量方法全卷统一。")
    B("（2）显微镜仅用于M05指定点及必要时外观确认；爱护光学部件与样件。")
    B("（3）万用表仅用于指定TP点测；注意档位，禁止对样件进行上电功能测试。")
    B("（4）防静电、轻拿轻放；如遇设备故障举手示意裁判。")

    L("7. 本模块材料与工具清单（摘要）")
    add_table(
        doc,
        [
            ["类别", "名称", "备注"],
            ["样件", "双面成品裸板质检样件", "A或B版；无元器件"],
            ["提交表", "学生组_模块D_检测记录表", "唯一须提交的表格附件"],
            ["题干素材", "本卷第（三）节全部表格与说明", "已嵌入，不另附分册"],
            ["工具", "卡尺/厚度规/显微镜/放大镜/万用表等", "以赛场清单为准"],
        ],
    )

    H("（四）模块任务")
    B(
        "请在45分钟内完成D-1～D-5，并将结果填写在《学生组_模块D_检测记录表》中提交。"
        "判定一律以本卷第（三）节题干素材为准。"
    )

    L("D-1  检验准备（建议3分钟）")
    B("核对样件编号与版本丝印、确认顶/底方向与分区约定；检查卡尺零位、厚度规、放大镜/显微镜、照明及万用表；在检测记录表填写基本信息并勾选准备项。")
    B("（1）核对样件编号、版本丝印（MOD-D-S-A / MOD-D-S-B）与记录表；")
    B("（2）确认顶/底方向与本卷分区说明一致；")
    B("（3）检查量具与万用表可用；")
    B("（4）勾选准备项；如有运输损伤先记录。")

    L("D-2  外观缺陷检测（建议12分钟）")
    B(
        "对样件线路、焊盘/孔、阻焊、丝印等进行外观检查（须检顶面与底面），完成5处必检（V01–V05），"
        "按本卷第（三）节第3款判定，将结果记入检测记录表。可结合分区网格定位；干扰项可不计入必检5。"
    )
    B("（1）每发现一处缺陷填写一行：所在面、类型、网格/特征、现象、严重程度等；")
    B("（2）位置描述应便于复核（如“顶层B2/走线区/中部”），无需精确坐标；")
    B("（3）类型应具体（如“线路缺口”“孔破盘”），避免仅写“线路有问题”。")

    L("D-3  尺寸与指定点微几何测量（建议10分钟）")
    B("按本卷第（三）节第2款测量M01–M05，填写实测值并判定合格/不合格。")
    B("（1）M01板长；（2）M02板宽；（3）M03板厚；（4）M04定位孔径；（5）M05指定点线宽或线距（显微镜）。")
    B("不要求测量普通卡尺无法可靠完成的项目（如0.30 mm级微孔）。记录以mm为主。")

    L("D-4  基础电气检测（建议8分钟）")
    B("按本卷第（三）节第4款，在TP1–TP4完成E01开路、E02短路点测，记录电阻与判定。")
    B("仅做指定测试点基础开短路；不做故障定位、网络分段、应通应断全表、高阻分析与上电功能测试。")

    L("D-5  分类判定与提交（建议12分钟）")
    B("（1）对已记录外观缺陷按类型分类计数；")
    B("（2）汇总尺寸M01–M05是否全部满足基准；")
    B("（3）汇总电气E01/E02结果；")
    B("（4）对照本卷第（三）节第3、5款，给出接收/返工/报废结论及简要依据；")
    B("（5）不要求深层制造工艺根因分析；")
    B(
        "（6）检查记录完整性后提交《学生组_模块D_检测记录表》；如需电子版，示例目录D:\\提交资料\\模块D\\，命名示例：赛位号_模块D。"
    )

    L("注意事项")
    B("1. 请合理分配D-1～D-5时间（建议合计45分钟）；时长与评分以技术工作文件/竞赛平台为准；")
    B("2. 本模块与模块B/C独立；仅做基础开短路点测，不做故障定位与网络分段，不上电，不量测Gerber；")
    B("3. 判定依据以本卷题干素材为准；检测记录表为唯一提交表格附件；")
    B("4. 爱护样件与仪器；按安全规范使用防静电措施与万用表；")
    B("5. 成果上不得标注姓名等身份信息；如遇设备故障请举手示意裁判。")


def sync_front_tables(doc: Document):
    """确保模块映射与成果物表为现行口径（若已正确则跳过）。"""
    for table in doc.tables:
        # 模块映射表
        if len(table.rows) >= 5 and "模块编号" in table.rows[0].cells[0].text:
            for row in table.rows:
                if row.cells[0].text.strip() == "模块D":
                    if "成品PCB裸板" not in row.cells[1].text:
                        row.cells[1].text = ""
                        p = row.cells[1].paragraphs[0]
                        run = p.add_run("成品PCB裸板质量检测与判定（学生组）")
                        set_run_font(run, name_cn="仿宋", size_pt=12)
                    break
        # 成果物表
        if len(table.columns) >= 3 and "提交内容" in table.rows[0].cells[1].text:
            for row in table.rows:
                if row.cells[0].text.strip() == "模块D":
                    want = "学生组_模块D_检测记录表（唯一提交表格；外观/尺寸/电气/综合判定）"
                    if want not in row.cells[1].text:
                        row.cells[1].text = ""
                        p = row.cells[1].paragraphs[0]
                        run = p.add_run(want)
                        set_run_font(run, name_cn="仿宋", size_pt=12)
                    break


def write_package_docs():
    (SAMPLE / "00_学生组模块D现行包说明.md").write_text(
        """# 学生组模块D现行包说明（2026-07-21 · R1 · 题干嵌入）

## 现行口径

| 项 | 内容 |
|----|------|
| 定位 | **成品PCB裸板质量检测与判定**（学生组） |
| 时长 | **45 分钟**（D-1～D-5：3 / 12 / 10 / 8 / 12） |
| 减负包 | **R1** |
| 电气 | **基础开短路 2 项**（TP1–TP4；非故障定位） |
| 外观 | F1 必检 **5**（V01–V05） |
| 尺寸·微几何 | M01–M04 + **M05** 指定点 |
| 分值 | 40 原始分（14+12+8+4+2） |
| 完整版正文 | `印制电路制作工赛项_竞赛样题（完整版）.docx` **模块D**（第（三）节内嵌全部文字/表格素材） |
| 独立包 | `02_样题/学生组_模块D/` |

## 对选手正式下发

| 材料 | 形式 |
|------|------|
| 赛题模块D全文 | 完整版样题内（含题干素材表） |
| **检测记录表** | `学生组_模块D/选手/学生组_模块D_检测记录表.docx`（**唯一**单独表格附件） |
| 样件与量具 | 赛场提供 |

> 验收/尺寸/电气/分区/安全/清单等：**只在题干第（三）节**，不另附分册。

## 命题包索引

- 边界/规格：`命题边界确认表.md`、`检测板/*`、`技术文件待同步清单.md`、`打样验证与验收标准.md`
- 选手提交表：`选手/学生组_模块D_检测记录表.docx`
- 选手目录其他 docx：命题底稿/备份，**默认不另发**
- 裁判：`裁判/*`（不向选手下发）

## 禁止混发

| 旧口径/文件 | 处理 |
|-------------|------|
| 90 min / 无电气 / 8 缺陷 | 废止 |
| 故障定位/通断网络表 | 不得现行下发 |
| 人工复审完整版 | 非现行 |
| 分册式验收/尺寸/电气/分区表作选手附件 | **废止**（已嵌入题干） |

## 备份

- `05_归档备份/module-d-full-rewrite-*`（本次回写前）
- 更早：`module-d-embed-stem-*`、`module-d-r1-45min-*`、`module-d-format-align-*`
""",
        encoding="utf-8",
    )
    print("[OK] 现行包说明")

    (PACK / "完整版学生组模块D正文稿_R1.md").write_text(
        """# 完整版 · 学生组模块D（权威以 docx 为准）

> 文件：`02_样题/印制电路制作工赛项_竞赛样题（完整版）.docx`
> 口径：R1 / 45 min / 题干嵌入 / 唯一提交《检测记录表》

## 结构

1. （一）模块考核点
2. （二）模块简介（背景／对象／材料／工具）
3. **（三）题干素材**
   1. 分区与 TP1–TP4
   2. 尺寸基准 M01–M05
   3. 外观验收 V01–V05
   4. 基础电气 E01–E02
   5. 接收／返工／报废
   6. 仪器与安全
   7. 材料工具清单摘要
4. （四）模块任务 D-1～D-5
5. 注意事项

## 回写脚本

`03_脚本与方案/_generated/rewrite_full_exam_module_d_20260721.py`
""",
        encoding="utf-8",
    )
    print("[OK] 正文稿镜像")


def verify(doc: Document):
    texts = [p.text for p in doc.paragraphs]
    full = "\n".join(texts)
    # no orphan booklet materials table as separate player list of 6 items in front of title
    body = doc.element.body
    kids = list(body.iterchildren())
    title_i = None
    for i, ch in enumerate(kids):
        if ch.tag.split("}")[-1] == "p" and el_text(ch).startswith("模块D："):
            title_i = i
            break
    assert title_i is not None
    # previous non-sectPr should not be old 6-row materials if it lists 尺寸基准表 as separate handout
    prev = kids[title_i - 1] if title_i > 0 else None
    if prev is not None and prev.tag.split("}")[-1] == "tbl":
        tx = el_text(prev)
        if "尺寸基准表" in tx and "检测记录表" in tx and "验收要求表" in tx:
            raise AssertionError("模块D标题前仍残留分册材料表")

    checks = {
        "标题": "模块D：成品PCB裸板质量检测与判定" in full,
        "45分钟": "45分钟" in full,
        "题干素材": "题干素材" in full,
        "不再另发": "不再另发" in full,
        "唯一": "唯一" in full,
        "V01": "V01" in full,
        "M05": "M05" in full,
        "E01": "E01" in full,
        "TP1": "TP1" in full,
        "模块B保留": any("模块B：" in t for t in texts),
        "模块C保留": any("模块C：" in t for t in texts),
    }
    # tables should include TP / M / V / E
    all_tbl = "\n".join(el_text(t._tbl) for t in doc.tables)
    checks["表TP"] = "TP1" in all_tbl
    checks["表M01"] = "M01" in all_tbl
    checks["表V01"] = "V01" in all_tbl
    checks["表E01"] = "E01" in all_tbl
    # must NOT have separate handout list of 尺寸基准表 as player material in module D section tables after title
    # (清单摘要表 only has 提交表/题干素材)
    bad = 0
    for t in doc.tables:
        tx = el_text(t._tbl)
        if "模块D_尺寸基准表" in tx or ("尺寸基准表" in tx and "验收要求表" in tx and "基础电气检测表" in tx):
            bad += 1
    checks["无分册材料表"] = bad == 0

    for k, v in checks.items():
        print(("OK" if v else "FAIL"), k)
    if not all(checks.values()):
        raise SystemExit(1)
    print("VERIFY PASS; tables=", len(doc.tables), "paras=", len(doc.paragraphs))


def main():
    BACKUP.mkdir(parents=True, exist_ok=True)
    shutil.copy2(FULL, BACKUP / FULL.name)
    print("[BACKUP]", BACKUP)

    doc = Document(str(FULL))
    body = doc.element.body
    start = find_module_d_body_start(body)
    print("[FULL] Module D body start index", start, "->", el_text(list(body.iterchildren())[start])[:60])
    remove_body_from(body, start)
    write_module_d(doc)
    sync_front_tables(doc)
    doc.save(str(FULL))
    print("[OK] saved", FULL.name)

    write_package_docs()

    doc2 = Document(str(FULL))
    verify(doc2)
    print("DONE")


if __name__ == "__main__":
    main()
