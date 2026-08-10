# -*- coding: utf-8 -*-
# 注意：本脚本为历史/辅助生成器，权威提交文件以 01_技术文件 与 02_样题 中现有 docx 为准。
# 请勿用本脚本输出直接覆盖提交集权威文件（尤其 V1 技术文件与完整版样题）。
# 默认输出到 03_脚本与方案/_generated/（可复跑草稿目录）。
"""
生成额外内容文件：作品提交要求和附件
这些内容在省级参考格式中没有对应章节，单独保存
"""
import os
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml

def set_cell_shading(cell, color):
    shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color}" w:val="clear"/>')
    cell._tc.get_or_add_tcPr().append(shading)

def add_formatted_paragraph(doc, text, font_name='宋体', font_size=Pt(12),
                            bold=False, alignment=WD_ALIGN_PARAGRAPH.JUSTIFY,
                            space_before=Pt(0), space_after=Pt(0),
                            line_spacing=Pt(28), first_line_indent=None,
                            indent=True):
    p = doc.add_paragraph()
    p.alignment = alignment
    pf = p.paragraph_format
    pf.space_before = space_before
    pf.space_after = space_after
    pf.line_spacing = line_spacing
    if first_line_indent:
        pf.first_line_indent = first_line_indent
    elif indent and not first_line_indent:
        pf.first_line_indent = Cm(0.74)
    run = p.add_run(text)
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn('w:eastAsia'), font_name)
    run.font.size = font_size
    run.font.bold = bold
    return p

def add_heading_custom(doc, text, level=1):
    sizes = {1: Pt(18), 2: Pt(15), 3: Pt(12)}
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf = p.paragraph_format
    pf.space_before = Pt(12) if level == 1 else Pt(6)
    pf.space_after = Pt(6) if level == 1 else Pt(3)
    pf.line_spacing = Pt(28)
    run = p.add_run(text)
    run.font.name = '黑体'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
    run.font.size = sizes.get(level, Pt(12))
    run.font.bold = True
    return p

def add_body_text(doc, text, indent=True):
    return add_formatted_paragraph(
        doc, text, font_name='宋体', font_size=Pt(12),
        first_line_indent=Cm(0.74) if indent else None
    )

def add_table_with_style(doc, headers, rows, col_widths=None):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = 'Table Grid'
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ''
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(header)
        run.font.name = '黑体'
        run._element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
        run.font.size = Pt(10.5)
        run.font.bold = True
        set_cell_shading(cell, 'D9E2F3')
    for r_idx, row_data in enumerate(rows):
        for c_idx, cell_text in enumerate(row_data):
            if c_idx < len(headers):
                cell = table.rows[r_idx + 1].cells[c_idx]
                cell.text = ''
                p = cell.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = p.add_run(str(cell_text))
                run.font.name = '宋体'
                run._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
                run.font.size = Pt(10.5)
    if col_widths and len(col_widths) == len(headers):
        for i, width in enumerate(col_widths):
            for row in table.rows:
                row.cells[i].width = width
    return table

# 读取原文件
_SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
_ROOT = os.path.dirname(_SCRIPT_DIR)
_GEN_DIR = os.path.join(_SCRIPT_DIR, '_generated')
os.makedirs(_GEN_DIR, exist_ok=True)
# 优先读草稿生成物；若不存在则尝试归档旧版（只读输入）
_src_candidates = [
    os.path.join(_GEN_DIR, '印制电路制作工职业技能竞赛技术文件_draft.docx'),
    os.path.join(_ROOT, '05_归档备份', 'cross-review-2026-07-16', '01_技术文件', '印制电路制作工职业技能竞赛技术文件.docx'),
    os.path.join(_ROOT, '01_技术文件', '印制电路制作工竞赛技术工作文件V1.docx'),
]
_src = next((p for p in _src_candidates if os.path.exists(p)), _src_candidates[-1])
src = Document(_src)

def get_table_data(idx):
    if idx < len(src.tables):
        t = src.tables[idx]
        return [[cell.text.strip() for cell in row.cells] for row in t.rows]
    return []

# 创建额外内容文档
doc = Document()
section = doc.sections[0]
section.page_width = Cm(21)
section.page_height = Cm(29.7)
section.top_margin = Cm(2.54)
section.bottom_margin = Cm(2.54)
section.left_margin = Cm(3.17)
section.right_margin = Cm(3.17)

# ===== 封面 =====
for _ in range(6):
    doc.add_paragraph()

add_formatted_paragraph(doc, '印制电路制作工职业技能竞赛',
                        font_name='黑体', font_size=Pt(26),
                        bold=True, alignment=WD_ALIGN_PARAGRAPH.CENTER)
add_formatted_paragraph(doc, '附加文件',
                        font_name='黑体', font_size=Pt(26),
                        bold=True, alignment=WD_ALIGN_PARAGRAPH.CENTER)
doc.add_paragraph()
add_formatted_paragraph(doc, '作品提交要求与附件表格',
                        font_name='宋体', font_size=Pt(16),
                        alignment=WD_ALIGN_PARAGRAPH.CENTER)
for _ in range(4):
    doc.add_paragraph()
add_formatted_paragraph(doc, '（本文件为技术工作文件的补充材料）',
                        font_name='宋体', font_size=Pt(12),
                        alignment=WD_ALIGN_PARAGRAPH.CENTER)

# ===== 正文 =====
doc.add_page_break()

# ============================================================
# 作品提交要求
# ============================================================
add_heading_custom(doc, '作品提交要求', level=1)

add_heading_custom(doc, '（一）文件夹结构', level=2)
add_body_text(doc, '选手须按以下统一文件夹结构提交作品：')

folder_lines = [
    '选手编号_作品提交/',
    '├── 01_工程源文件/',
    '├── 02_Gerber文件/',
    '├── 03_NC_Drill文件/',
    '├── 04_DRC报告/',
    '├── 05_DFM审查表/',
    '├── 06_缺陷检测记录表/',
    '└── 07_作品说明.pdf',
]
for line in folder_lines:
    add_formatted_paragraph(doc, line, font_name='Consolas', font_size=Pt(11),
                            indent=False, line_spacing=Pt(20),
                            first_line_indent=None)

add_heading_custom(doc, '（二）提交物清单', level=2)

submit_rows = get_table_data(3)  # Table 3: 提交物清单
add_table_with_style(doc, submit_rows[0], submit_rows[1:],
                     col_widths=[Cm(1.5), Cm(4.5), Cm(3.5), Cm(4.5)])

add_body_text(doc, '注意：提交截止后，选手不得再修改或补充任何文件。'
    '未按要求提交的作品，相关模块按零分处理。')

# ============================================================
# 附件1：选手报名表
# ============================================================
doc.add_page_break()
add_heading_custom(doc, '附件1：选手报名表', level=1)

reg_rows = get_table_data(8)  # Table 8: 报名表
add_table_with_style(doc, reg_rows[0], reg_rows[1:],
                     col_widths=[Cm(3), Cm(11)])

# ============================================================
# 附件2：竞赛日程表
# ============================================================
doc.add_page_break()
add_heading_custom(doc, '附件2：竞赛日程表', level=1)

add_body_text(doc, '竞赛总时长为4小时，各阶段时间安排由组委会在赛前统一公布。竞赛流程如下：')

schedule_rows = get_table_data(9)  # Table 9: 日程表
add_table_with_style(doc, schedule_rows[0], schedule_rows[1:],
                     col_widths=[Cm(1.5), Cm(3), Cm(9)])

# ============================================================
# 附件3：设备与软件清单
# ============================================================
doc.add_page_break()
add_heading_custom(doc, '附件3：设备与软件清单', level=1)

add_heading_custom(doc, '一、组委会统一提供：', level=2)
hw_rows = get_table_data(10)  # Table 10: 设备清单
add_table_with_style(doc, hw_rows[0], hw_rows[1:],
                     col_widths=[Cm(1), Cm(3.5), Cm(3.5), Cm(2.5), Cm(3)])

add_heading_custom(doc, '二、选手自带工具（缺陷检测环节）：', level=2)
tool_rows = get_table_data(11)  # Table 11: 自带工具
add_table_with_style(doc, tool_rows[0], tool_rows[1:],
                     col_widths=[Cm(1.5), Cm(3.5), Cm(3.5), Cm(5)])

# ============================================================
# 附件4：PCB制造规则表
# ============================================================
doc.add_page_break()
add_heading_custom(doc, '附件4：PCB制造规则表', level=1)

rule_rows = get_table_data(12)  # Table 12: 制造规则
add_table_with_style(doc, rule_rows[0], rule_rows[1:],
                     col_widths=[Cm(1.5), Cm(3.5), Cm(5), Cm(4)])

# ============================================================
# 附件5：作品提交清单
# ============================================================
doc.add_page_break()
add_heading_custom(doc, '附件5：作品提交清单', level=1)

submit_rows2 = get_table_data(13)  # Table 13: 提交清单
add_table_with_style(doc, submit_rows2[0], submit_rows2[1:],
                     col_widths=[Cm(1.5), Cm(4.5), Cm(3.5), Cm(4.5)])

add_formatted_paragraph(doc, '')
add_formatted_paragraph(doc, '选手签名：____________    日期：____________',
                        font_name='宋体', font_size=Pt(12), indent=False)

# ============================================================
# 附件6：DFM审查表
# ============================================================
doc.add_page_break()
add_heading_custom(doc, '附件6：DFM审查表', level=1)

add_formatted_paragraph(doc, '选手编号：____________    审查日期：____________',
                        font_name='宋体', font_size=Pt(12), indent=False)

dfm_rows = get_table_data(14)  # Table 14: DFM审查表
add_table_with_style(doc, dfm_rows[0], dfm_rows[1:],
                     col_widths=[Cm(1.5), Cm(3), Cm(3), Cm(3), Cm(3.5)])

add_formatted_paragraph(doc, '')
add_formatted_paragraph(doc, '审查结论：□ 全部合格    □ 存在不合格项（需整改）',
                        font_name='宋体', font_size=Pt(12), indent=False)
add_formatted_paragraph(doc, '审查人签名：____________',
                        font_name='宋体', font_size=Pt(12), indent=False)

# ============================================================
# 附件7：缺陷检测记录表
# ============================================================
doc.add_page_break()
add_heading_custom(doc, '附件7：缺陷检测记录表', level=1)

add_formatted_paragraph(doc, '选手编号：____________    检测日期：____________',
                        font_name='宋体', font_size=Pt(12), indent=False)

defect_rows = get_table_data(15)  # Table 15: 缺陷检测记录
add_table_with_style(doc, defect_rows[0], defect_rows[1:],
                     col_widths=[Cm(1.5), Cm(2), Cm(2.5), Cm(3), Cm(2.5), Cm(2.5)])

add_formatted_paragraph(doc, '')
add_formatted_paragraph(doc, '检测人签名：____________    日期：____________',
                        font_name='宋体', font_size=Pt(12), indent=False)

# ============================================================
# 附件8：裁判评分表
# ============================================================
doc.add_page_break()
add_heading_custom(doc, '附件8：裁判评分表', level=1)

add_formatted_paragraph(doc, '选手编号：____________    裁判签名：____________',
                        font_name='宋体', font_size=Pt(12), indent=False)

score_rows = get_table_data(16)  # Table 16: 裁判评分表
add_table_with_style(doc, score_rows[0], score_rows[1:],
                     col_widths=[Cm(2.5), Cm(3.5), Cm(2), Cm(2), Cm(4)])

add_formatted_paragraph(doc, '')
add_formatted_paragraph(doc, '裁判长签名：____________    日期：____________',
                        font_name='宋体', font_size=Pt(12), indent=False)

# ============================================================
# 附件9：安全与纪律承诺书
# ============================================================
doc.add_page_break()
add_heading_custom(doc, '附件9：安全与纪律承诺书', level=1)

add_body_text(doc, '本人自愿参加本次印制电路制作工职业技能竞赛，郑重承诺如下：')
add_body_text(doc, '1. 本人已认真阅读并理解竞赛技术文件和各项规则。')
add_body_text(doc, '2. 本人保证所提交的个人信息真实有效。')
add_body_text(doc, '3. 本人承诺遵守竞赛纪律，不携带违禁物品进入赛场。')
add_body_text(doc, '4. 本人承诺独立完成竞赛任务，不抄袭、不协助他人。')
add_body_text(doc, '5. 本人承诺遵守赛场安全管理规定，注意人身安全。')
add_body_text(doc, '6. 本人承诺服从裁判和工作人员的管理和指挥。')
add_body_text(doc, '7. 如违反上述承诺，本人愿意接受取消竞赛成绩等处理。')

for _ in range(2):
    doc.add_paragraph()

add_formatted_paragraph(doc, '承诺人签名：____________    日期：____________',
                        font_name='宋体', font_size=Pt(12), indent=False)
add_formatted_paragraph(doc, '身份证号：____________',
                        font_name='宋体', font_size=Pt(12), indent=False)
add_formatted_paragraph(doc, '联系电话：____________',
                        font_name='宋体', font_size=Pt(12), indent=False)

# ============================================================
# 附件10：赛场异常情况记录表
# ============================================================
doc.add_page_break()
add_heading_custom(doc, '附件10：赛场异常情况记录表', level=1)

abnormal_rows = get_table_data(17)  # Table 17: 异常记录
add_table_with_style(doc, abnormal_rows[0], abnormal_rows[1:],
                     col_widths=[Cm(1.5), Cm(3.5), Cm(9)])

# 保存
output_path = os.path.join(_GEN_DIR, '印制电路制作工_附加文件_draft.docx')
doc.save(output_path)
print(f'附加文件已保存: {output_path}')
