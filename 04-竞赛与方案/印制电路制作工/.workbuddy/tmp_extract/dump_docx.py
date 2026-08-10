import sys
from docx import Document

path = sys.argv[1]
out = sys.argv[2]
doc = Document(path)
lines = []
lines.append(f"FILE: {path}")
lines.append(f"PARAGRAPHS: {len(doc.paragraphs)}  TABLES: {len(doc.tables)}")
lines.append("")

# Paragraphs
for i, p in enumerate(doc.paragraphs):
    t = p.text.strip()
    if t:
        style = p.style.name if p.style else ""
        lines.append(f"[P{i}|{style}] {t}")

# Tables
for ti, tbl in enumerate(doc.tables):
    lines.append("")
    lines.append(f"===== TABLE {ti} ({len(tbl.rows)}x{len(tbl.columns)}) =====")
    for r in tbl.rows:
        cells = [c.text.replace("\n", " / ").strip() for c in r.cells]
        lines.append(" | ".join(cells))

text = "\n".join(lines)
with open(out, "w", encoding="utf-8") as f:
    f.write(text)
print(f"wrote {len(lines)} lines -> {out}")
