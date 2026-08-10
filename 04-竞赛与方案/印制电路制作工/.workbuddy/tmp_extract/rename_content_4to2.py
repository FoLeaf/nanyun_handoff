import docx, openpyxl, re
from docx.oxml.ns import qn

DOCX = r"D:/nanyun/04-竞赛与方案/印制电路制作工/02_样题/第二套/01竞赛要求/印制电路制作工赛项_竞赛样题（完整版）.docx"
XLSX = r"D:/nanyun/04-竞赛与方案/印制电路制作工/02_样题/第二套/04评分标准/印制电路制作工赛项_评分标准（第二套）.xlsx"

# --- DOCX: replace '第四套' -> '第二套' across all XML parts (document/headers/footers) ---
doc = docx.Document(DOCX)
dcount = 0
for part in doc.part.package.iter_parts():
    if not str(part.partname).endswith('.xml'):
        continue
    el = getattr(part, 'element', None)
    if el is None:
        continue
    for t in el.iter(qn('w:t')):
        if t.text and '第四套' in t.text:
            t.text = t.text.replace('第四套', '第二套')
            dcount += 1
doc.save(DOCX)
print(f"DOCX: 替换 '第四套'->'第二套' 共 {dcount} 处文本节点")

# --- XLSX: replace '第四套' -> '第二套' across all cells ---
wb = openpyxl.load_workbook(XLSX)
xcount = 0
for ws in wb.worksheets:
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and '第四套' in c.value:
                c.value = c.value.replace('第四套', '第二套')
                xcount += 1
wb.save(XLSX)
print(f"XLSX: 替换 '第四套'->'第二套' 共 {xcount} 个单元格")

# --- verify no '第四套' remains ---
import zipfile
for label, f in [("DOCX", DOCX), ("XLSX", XLSX)]:
    z = zipfile.ZipFile(f)
    rem = 0
    for n in z.namelist():
        if n.endswith('.xml'):
            txt = re.sub(r'<[^>]+>', '', z.read(n).decode('utf-8', 'ignore'))
            rem += txt.count('第四套')
    print(f"{label} 残留 '第四套' 数: {rem}")
