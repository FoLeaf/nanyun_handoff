# -*- coding: utf-8 -*-
"""【已废弃】请改用 build_sample_set2_from_set1_20260723.py

本脚本会自建版式，与第一套字体/页边距不一致，且正文过多强调「第二套」。
保留仅作历史参考，禁止再运行覆盖选手材料。
"""
raise SystemExit("DEPRECATED: use build_sample_set2_from_set1_20260723.py")
"""第二套样题阶段1：完整版初稿 + C/D 空白表 + C 参考答案初稿。

输出目录：02_样题/第二套/
不修改 01_技术文件。
"""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "02_样题" / "第二套"
OUT.mkdir(parents=True, exist_ok=True)

FONT = "仿宋"
SZ_TITLE = 18
SZ_H1 = 16
SZ_BODY = 12
SZ_TABLE = 10


def set_run(run, size=SZ_BODY, bold=False, font=FONT):
    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    run.font.size = Pt(size)
    run.bold = bold


def pfmt(p, *, center=False, first=False, before=0, after=6):
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.JUSTIFY
    if first:
        p.paragraph_format.first_line_indent = Cm(0.74)


def add_text(doc, text, *, size=SZ_BODY, bold=False, center=False, first=False, before=0, after=6):
    p = doc.add_paragraph()
    pfmt(p, center=center, first=first, before=before, after=after)
    r = p.add_run(text)
    set_run(r, size=size, bold=bold)
    return p


def add_h(doc, text, size=SZ_H1):
    return add_text(doc, text, size=size, bold=True, before=12, after=8)


def set_cell(cell, text, *, bold=False, size=SZ_TABLE, center=False):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(1)
    r = p.add_run(text)
    set_run(r, size=size, bold=bold)


def add_table(doc, headers, rows, col_widths=None):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        set_cell(table.rows[0].cells[i], h, bold=True, center=True)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            set_cell(table.rows[ri + 1].cells[ci], str(val))
    if col_widths:
        for row in table.rows:
            for i, w in enumerate(col_widths):
                if i < len(row.cells):
                    row.cells[i].width = Cm(w)
    return table


def setup_doc():
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = FONT
    style.font.size = Pt(SZ_BODY)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    for sec in doc.sections:
        sec.top_margin = Cm(2.5)
        sec.bottom_margin = Cm(2.5)
        sec.left_margin = Cm(2.8)
        sec.right_margin = Cm(2.5)
    return doc


# ---------------------------------------------------------------------------
# 完整版样题
# ---------------------------------------------------------------------------

def build_full_exam():
    doc = setup_doc()
    add_text(doc, '2026年江西省"振兴杯"职业技能大赛', size=SZ_TITLE, bold=True, center=True, after=2)
    add_text(doc, '"印制电路制作工"赛项', size=SZ_TITLE, bold=True, center=True, after=2)
    add_text(doc, "竞赛样题（第二套）", size=SZ_TITLE, bold=True, center=True, after=2)
    add_text(doc, "（阶段1初稿 · 工程附件待回填）", size=11, center=True, after=12)

    add_h(doc, "竞赛说明")
    add_text(
        doc,
        "本赛项为“印制电路制作工”职业技能竞赛，考核选手在PCB设计、CAM审核、工艺文件编制及成品质量检测等方面的综合职业能力。"
        "本套为公开样题第二套，与第一套考核结构、难度与模块权重平行，仅更换任务场景与预设缺陷内容。"
        "时长与评分以技术工作文件及竞赛平台为准（技术工作文件本身不因本套样题而修改）。",
        first=True,
    )
    add_text(doc, "一、竞赛内容", bold=True, before=8)
    add_table(
        doc,
        ["模块编号", "模块名称"],
        [
            ["模块A", "理论测试"],
            ["模块B", "EDA工程设计"],
            ["模块C", "CAM审核与工艺文件编制"],
            ["模块D", "成品板质量检测与缺陷分析"],
        ],
    )
    add_text(doc, "二、竞赛规则与注意事项", bold=True, before=8)
    for t in [
        "选手不得携带任何书籍、资料、存储设备及通讯工具进入赛场，已带入的应交监考人员统一保管。",
        "竞赛开始前15分钟，选手进入赛场，检查赛位设备、工具及材料是否齐全，如有问题立即向监考人员报告。",
        "竞赛信号发出后方可开始答题或操作；竞赛结束信号发出后应立即停止操作，按要求提交成果物。",
        "选手在试卷或成果物上不得以任何方式标注姓名、单位等个人信息，违者按作弊处理。",
        "操作技能模块中，选手应严格遵守安全操作规程，防止设备和人身安全事故发生。",
    ]:
        add_text(doc, t, first=True)

    add_text(doc, "三、成果物提交要求", bold=True, before=8)
    add_text(
        doc,
        "所有电子文件请统一提交到赛场指定U盘根目录下，文件命名中“赛位号”以实际赛位号为准。",
        first=True,
    )
    add_table(
        doc,
        ["模块", "提交内容", "命名规则"],
        [
            ["模块B", "立创EDA工程文件 + Gerber压缩包 + 2D预览图等", "赛位号_模块B（本地：D:\\提交资料\\模块B\\）"],
            ["模块C", "缺陷记录表 + 简版制程工艺卡", "作答簿提交"],
            ["模块D", "模块D检测记录表", "作答簿提交"],
        ],
    )

    # ---- A ----
    add_h(doc, "模块A：理论测试")
    add_text(
        doc,
        "模块A为理论测试（机考）：试题由竞赛平台统一下发与作答。正式题型、题量与评分细则以竞赛平台及技术工作文件为准。"
        "本竞赛样题仅对模块A作简要说明，不提供理论试题或赛题展示；操作技能赛题见模块B/C/D。",
        first=True,
    )

    # ---- B ----
    add_h(doc, "模块B：EDA工程设计")
    add_text(doc, "（一）模块考核点", bold=True)
    add_text(
        doc,
        "本模块考核选手使用立创EDA专业版进行PCB布局布线设计的能力，包括设计规则设置、元件布局、信号走线、电源地处理、敷铜及Gerber导出等。",
        first=True,
    )
    add_text(doc, "（二）模块简介", bold=True)
    add_text(doc, "【任务背景】", bold=True)
    add_text(
        doc,
        "某物联网公司开发一款BLE（低功耗蓝牙）传感器节点板，用于近距离环境数据采集与无线上报。"
        "核心板基于STM32G070系列主控芯片，集成BLE通信模组、I2C温湿度传感器、USB/排针供电接口等功能。"
        "硬件工程师已完成核心电路的原理图设计，但部分子电路尚未完成，且有一颗底板连接座需要自建封装。",
        first=True,
    )
    add_text(
        doc,
        "作为PCB Layout工程师，您需要完成：自建连接器封装 → 补全原理图 → 设置设计规则 → PCB布局 → PCB布线 → 输出生产Gerber文件。",
        first=True,
    )
    add_text(doc, "【提供材料】", bold=True)
    add_text(doc, "赛场下发立创EDA半成品工程文件（模块B）。元件型号与封装以半成品工程及库内推荐封装为准，本赛题不单独下发完整BOM。", first=True)
    add_text(doc, "（三）设计工艺要求", bold=True)
    add_text(doc, "以下参数基于板厂常规工艺能力（FR-4双面板），设计必须满足：", first=True)
    add_table(
        doc,
        ["序号", "项目", "要求"],
        [
            ["1", "板材/层数", "FR-4，双面板"],
            ["2", "信号线宽/线距", "≥0.15mm（6mil）"],
            ["3", "电源线宽", "3V3≥0.8mm；5V≥1.0mm"],
            ["4", "过孔", "内径0.3mm，外径0.6mm"],
            ["5", "安全间距", "线线/线盘/孔线 0.15mm"],
            ["6", "阻焊", "焊盘1:1开窗；阻焊桥最小0.1mm"],
        ],
    )

    add_text(doc, "（四）模块任务", bold=True, before=8)
    add_text(doc, "B-0  自建连接器封装", bold=True)
    add_text(
        doc,
        "本设计中有一颗底板连接座，立创EDA库中没有对应封装，需要自行绘制器件符号和PCB封装。",
        first=True,
    )
    add_text(doc, "1.器件符号设计", bold=True)
    add_text(doc, "（1）在立创EDA中新建器件，器件命名为 PinHeader_2x5；", first=True)
    add_text(doc, "（2）双列共10引脚：左列1–5（上到下），右列6–10（上到下）；1脚为电源相关（以半成品网络标注为准）。", first=True)
    add_text(doc, "2.PCB封装设计", bold=True)
    add_text(doc, "（1）按2.54mm间距双排针常规尺寸绘制焊盘与丝印外框，第1脚附近加圆点标识；", first=True)
    add_text(doc, "（2）封装完成后导入本工程库，并放置到原理图底板连接座部分（绿色高亮），按半成品网络补全。", first=True)

    add_text(doc, "B-1  原理图补全", bold=True, before=6)
    add_text(
        doc,
        "打开提供的立创EDA工程文件，根据以下子电路功能描述补全（原理图中已红色高亮）。",
        first=True,
    )
    subs = [
        "（1）电源输入保护电路：5V输入经自恢复保险丝F1 + 单向TVS（D2）+ 肖特基防反（D1），输出节点+5V。",
        "（2）LDO稳压：AMS1117-3.3，输入+5V、输出3V3，输入输出各并联约1μF陶瓷电容至GND。",
        "（3）复位与BOOT0：NRST经10kΩ上拉至3V3并可并联约100nF至GND；复位键一端NRST一端GND；BOOT0经10kΩ下拉GND（或按工程网络）。",
        "（4）LED指示：电源灯LED-PWR经限流电阻接3V3与GND；用户LED1/LED2分别接MCU指定GPIO（以半成品标注为准），低电平点亮，限流约1kΩ。注意二极管位号不得与D1/D2混用。",
        "（5）按键：SW1~SW4一端接MCU GPIO，一端GND；使用内部上拉，无需外加上拉。",
        "（6）BLE模组接口：按半成品给定的UART（或SPI）引脚连接模组；供电3V3/GND；复位/使能脚按工程标注处理；天线端朝板边布局在PCB阶段完成。",
        "（7）I2C传感器（如AHT30）：VDD=3V3，GND=GND；SDA/SCL接MCU I2C脚，各4.7kΩ上拉至3V3。",
        "（8）可选：NE555无稳态测频输出至MCU定时器捕获脚，或有源蜂鸣器驱动电路（以半成品高亮区为准，二选一已预置其一）。",
    ]
    for s in subs:
        add_text(doc, s, first=True)

    add_text(doc, "B-2  设计规则设置", bold=True, before=6)
    for s in [
        "（1）根据工艺要求表设置DRC，单位支持mil/mm切换；",
        "（2）信号线0.15mm；3V3≥0.8mm；5V≥1.0mm；",
        "（3）过孔内径0.3mm、外径0.6mm；",
        "（4）线线/线盘/孔线间距0.15mm；",
        "（5）焊盘1:1开窗，阻焊桥最小0.1mm。",
    ]:
        add_text(doc, s, first=True)

    add_text(doc, "B-3  PCB布局", bold=True, before=6)
    for s in [
        "（1）板框以半成品工程为准，不得改动板框尺寸；",
        "（2）电源模块集中在电源输入（PinHeader_2x5）附近；",
        "（3）晶振靠近MCU晶振脚；",
        "（4）BLE模组靠板边缘，天线端靠近板边；",
        "（5）按电源区、主控区、射频/模组区、传感器区、接口区功能分区；",
        "（6）丝印清晰，不压焊盘。",
    ]:
        add_text(doc, s, first=True)

    add_text(doc, "B-4  PCB布线", bold=True, before=6)
    for s in [
        "（1）电源线加宽符合B-2；信号线0.15mm；",
        "（2）优先保证GND连通，底层主要作地平面；",
        "（3）晶振下方不走其他信号，周围敷地；",
        "（4）优先45°走线，避免90°；",
        "（5）顶/底层敷GND铜；网络连通率100%。",
    ]:
        add_text(doc, s, first=True)

    add_text(doc, "B-5  DRC检查与Gerber导出", bold=True, before=6)
    for s in [
        "（1）DRC 0错误（警告可忽略）；",
        "（2）导出顶/底线路、顶/底阻焊、顶/底丝印、板框、钻孔等生产层；",
        "（3）Gerber打zip；导出顶层2D预览PNG；",
        "（4）保存至 D:\\提交资料\\模块B\\ ，U盘包命名：赛位号_模块B。",
    ]:
        add_text(doc, s, first=True)

    add_text(doc, "B-6 注意事项", bold=True, before=6)
    for s in [
        "请合理分配时间；时长与评分以技术工作文件/平台为准。",
        "随时保存；命名错误影响评分。",
        "禁止拷贝他人文件或使用外部存储；违者按作弊处理。",
        "除PinHeader_2x5外，优先使用库内封装，不得擅自乱建封装。",
        "软件故障举手示意裁判。",
    ]:
        add_text(doc, s, first=True)

    # ---- C ----
    add_h(doc, "模块C：CAM审核与工艺文件编制")
    add_text(doc, "（一）模块考核点", bold=True)
    add_text(doc, "考核使用Gerbv进行Gerber文件DFM审核及简版制程工艺卡编制的能力。", first=True)
    add_text(doc, "（二）模块简介", bold=True)
    add_text(doc, "【任务背景】", bold=True)
    add_text(
        doc,
        "您是PCB板厂CAM工程师。客户提交了一套BLE传感器节点板的Gerber文件准备投产。"
        "开料前需进行DFM审核，识别缺陷并编制简版制程工艺卡。",
        first=True,
    )
    add_text(doc, "【提供材料】模块C Gerber压缩包；缺陷记录表；制程工艺卡模板。", first=True)
    add_text(doc, "【工具】Gerbv（赛场预装）。", first=True)
    add_text(doc, "（三）板厂工艺能力参数表", bold=True)
    add_text(doc, "以下为板厂常规FR-4双面板工艺能力参数，缺陷判定以此为标准（数值沿用公开样题第一套口径）：", first=True)
    add_table(
        doc,
        ["项目", "能力要求"],
        [
            ["最小线宽", "≥0.10mm"],
            ["最小线距", "≥0.10mm"],
            ["有铜插件孔焊环极限", "≥0.18mm（单侧）"],
            ["最小过孔焊盘外径", "≥0.25mm"],
            ["过孔孔边到孔边", "≥0.20mm"],
            ["阻焊/开窗", "焊盘需开窗；阻焊桥满足最小要求"],
            ["丝印", "字符高度≥1.0mm；字符到露铜焊盘间距≥0.15mm"],
            ["文件", "生产层文件齐全（含顶/底丝印等约定层）"],
        ],
    )
    add_text(doc, "C-1 Gerber缺陷识别", bold=True, before=8)
    add_text(
        doc,
        "使用Gerbv打开Gerber包，对照工艺能力表检查各层缺陷，填入缺陷记录表。"
        "检测范围含线路、钻孔、阻焊、丝印、文件完整性。每处一行：所在层、缺陷类型、缺陷类别、位置描述。",
        first=True,
    )
    add_text(doc, "C-2 制程工艺卡编制", bold=True)
    add_text(doc, "根据Gerber信息与工艺能力，填写简版制程工艺卡基本信息（产品名称、板材、层数、板厚、铜厚、表面处理、阻焊色、最小线宽线距孔径等）。", first=True)
    add_text(doc, "C-3 注意事项", bold=True)
    add_text(doc, "合理分配时间；仅用批准工具；注意mil/mm单位。", first=True)

    # ---- D ----
    add_h(doc, "模块D：成品板质量检测与缺陷分析")
    add_text(doc, "（一）考核要求", bold=True)
    add_text(
        doc,
        "对赛场提供的双面成品PCB裸板完成：外观缺陷检测；关键尺寸与指定点线宽/线距；指定测试点开路/短路；"
        "缺陷分类与质量处置结论（接收/返工/报废）及简要依据。"
        "仅检验本模块样件；只做指定测试点开短路；不上电；不测量Gerber。",
        first=True,
    )
    add_text(
        doc,
        "本模块为独立质检任务。检验对象为赛场提供的双面成品裸板质检样件（专用检测板），"
        "与模块B、模块C相互独立；仅做基础开短路点测，不做故障定位与网络分段，不上电，不量测Gerber。",
        first=True,
    )
    add_text(doc, "（二）任务说明", bold=True)
    add_text(doc, "【检验对象】双面成品裸板1块。外形约80mm×60mm，板厚1.6mm；FR-4双面，绿阻焊，白丝印，无铅喷锡。", first=True)
    add_text(doc, "丝印含版本标识（MOD-D-S2-A 或 MOD-D-S2-B）、分区参考及测试点TP1～TP4。", first=True)
    add_text(doc, "【提供材料】裸板样件1块；《模块D检测记录表》空白表。", first=True)
    add_text(doc, "【工具】游标卡尺、厚度规/千分尺、测量显微镜或带刻度放大镜、放大镜、数字万用表、防静电用品等。", first=True)

    add_text(doc, "（三）检测依据", bold=True)
    add_text(doc, "1. 板面分区与测试点", bold=True)
    add_text(doc, "版本以板面丝印为准（MOD-D-S2-A / MOD-D-S2-B），记录表版本栏须与实物一致。", first=True)
    add_text(doc, "分区网格：顶层参考A1～D4。记录缺陷：所在面+网格或特征名+方位。", first=True)
    add_table(
        doc,
        ["测试点", "用途说明"],
        [
            ["TP1–TP2", "用于E01开路检测（按记录表指定网络）"],
            ["TP3–TP4", "用于E02短路检测（按记录表指定网络）"],
        ],
    )
    add_text(doc, "2. 尺寸基准（M01～M05）", bold=True)
    add_table(
        doc,
        ["编号", "项目", "工具", "标称/公差（mm）"],
        [
            ["M01", "板长 L", "卡尺", "80.00±0.20（以板面/答案键为准）"],
            ["M02", "板宽 W", "卡尺", "60.00±0.20"],
            ["M03", "板厚 T", "厚度规/千分尺", "1.60±0.15"],
            ["M04", "定位孔径", "卡尺/孔规", "2.00±0.10（禁止卡尺测≤0.30小孔）"],
            ["M05", "指定点线宽或线距", "测量显微镜", "按板面标识；保留0.01；不得与V02同点"],
        ],
    )
    add_text(doc, "3. 外观缺陷判定（V01～V05）", bold=True)
    add_text(doc, "须检查顶面与底面。下列5类为必检；记录类型、位置与现象：", first=True)
    add_table(
        doc,
        ["编号", "必检类型"],
        [
            ["V01", "线路缺口/颈缩类"],
            ["V02", "线宽或线距异常（定性）"],
            ["V03", "孔破盘/焊环不足"],
            ["V04", "阻焊未开窗或开窗偏移/桥异常"],
            ["V05", "丝印缺失、残缺或压焊盘"],
        ],
    )
    add_text(doc, "4. 基础电气（E01～E02）", bold=True)
    add_text(doc, "仅在指定测试点间用电阻档/通断档测量；记录读数、判定（开路/短路/正常）、是否与预期一致。", first=True)

    add_text(doc, "（四）模块任务", bold=True)
    for s in [
        "D-1 检验准备：核对版本丝印；确认顶底与分区；工具对零；运输损伤先注明。",
        "D-2 外观：完成V01～V05，位置描述清楚即可。",
        "D-3 尺寸：测M01～M05并判定合格/不合格。",
        "D-4 电气：TP间完成E01、E02。",
        "D-5 分类判定：汇总后给出接收/返工/报废及简要依据；提交检测记录表。",
    ]:
        add_text(doc, s, first=True)
    add_text(doc, "注意事项：合理安排时间；爱护样件与仪器；成果物不得标注身份信息。", first=True)

    path = OUT / "印制电路制作工赛项_竞赛样题（第二套）.docx"
    doc.save(path)
    return path


# ---------------------------------------------------------------------------
# 空白表 / 参考答案
# ---------------------------------------------------------------------------

def build_c_defect_blank():
    doc = setup_doc()
    add_text(doc, "缺陷记录表（模块C · 第二套）", size=SZ_H1, bold=True, center=True)
    add_text(doc, "赛位号：__________    日期：__________", center=True)
    add_text(doc, "产品：BLE传感器节点板（客户投产Gerber）", after=8)
    rows = [[str(i), "", "", "", ""] for i in range(1, 16)]
    add_table(doc, ["序号", "所在层", "缺陷类型", "缺陷类别", "位置/描述"], rows)
    add_text(doc, "说明：缺陷类别可选 线路 / 钻孔焊盘 / 阻焊 / 丝印 / 文件完整性 / 其他。", before=8, size=10)
    path = OUT / "模块C_缺陷记录表.docx"
    doc.save(path)
    return path


def build_c_defect_answer():
    doc = setup_doc()
    add_text(doc, "缺陷记录表（参考答案·模块C·第二套）", size=SZ_H1, bold=True, center=True)
    add_text(doc, "【阶段1初稿】共12处；Gerber植入后校核位置描述。", after=8)
    rows = [
        ["1", "顶层线路", "最小线宽不足", "线路", "MCU附近信号线局部约0.08mm，小于0.10mm"],
        ["2", "顶层线路", "最小线距不足", "线路", "BLE模组焊盘引出两线间距约0.08mm"],
        ["3", "底层线路", "走线开路/缺口", "线路", "3V3干线近LDO输出处缺口"],
        ["4", "顶层线路", "走线短路", "线路", "用户LED限流电阻附近SIG与GND桥接"],
        ["5", "钻孔+焊盘", "焊环不足（孔破盘）", "钻孔/焊盘", "过孔偏出焊盘，单侧焊环约0.08mm"],
        ["6", "钻孔", "过孔焊盘过小", "钻孔/焊盘", "过孔焊盘外径约0.20mm，小于0.25mm"],
        ["7", "钻孔", "孔间距过小", "钻孔/焊盘", "传感器附近两过孔孔边距约0.15mm"],
        ["8", "顶层阻焊", "阻焊覆盖焊盘", "阻焊", "某0603电阻焊盘未开窗"],
        ["9", "顶层阻焊", "阻焊桥缺失/不足", "阻焊", "细间距IC引脚间无法形成有效阻焊桥"],
        ["10", "顶层丝印", "丝印上焊盘", "丝印", "位号压在电容焊盘上"],
        ["11", "顶层丝印", "字符尺寸过小", "丝印", "某位号字高约0.7mm，小于1.0mm"],
        ["12", "文件包", "缺少底层丝印层文件", "文件完整性", "缺少底层丝印层文件"],
    ]
    add_table(doc, ["序号", "所在层", "缺陷类型", "缺陷类别", "位置/描述"], rows)
    path = OUT / "模块C_缺陷记录表-参考答案.docx"
    doc.save(path)
    return path


def build_process_card():
    doc = setup_doc()
    add_text(doc, "印制电路板（PCB）制程工艺卡（简版）", size=SZ_H1, bold=True, center=True)
    add_text(doc, "第二套样题 · 选手填写", center=True, after=8)
    add_text(doc, "注：本工艺卡依据GB/T 4588-2025及IPC相关标准编制。", size=10)
    add_table(
        doc,
        ["项目", "填写", "项目", "填写"],
        [
            ["产品名称", "BLE传感器节点板（示例可改）", "图号/版本", ""],
            ["板材类型", "", "板厚", ""],
            ["铜箔厚度", "", "层数", ""],
            ["表面处理", "", "阻焊颜色", ""],
            ["最小线宽", "", "最小线距", ""],
            ["最小孔径", "", "页次", ""],
        ],
    )
    add_text(doc, "", after=8)
    add_table(
        doc,
        ["编制", "审核", "日期"],
        [["", "", ""]],
    )
    path = OUT / "模块C_PCB制程工艺卡.docx"
    doc.save(path)
    return path


def build_d_record():
    doc = setup_doc()
    add_text(doc, "模块D 检测记录表（第二套 · 学生组）", size=SZ_H1, bold=True, center=True)
    add_table(
        doc,
        ["字段", "填写"],
        [
            ["赛位号", ""],
            ["样件编号", ""],
            ["版本丝印", "MOD-D-S2-A / MOD-D-S2-B（圈选）"],
            ["日期", ""],
        ],
    )
    add_text(doc, "一、外观缺陷（V01～V05及补充）", bold=True, before=10)
    add_table(
        doc,
        ["序号", "所在面", "类型", "网格/特征", "现象描述"],
        [[str(i), "", "", "", ""] for i in range(1, 9)],
    )
    add_text(doc, "二、尺寸测量（M01～M05）", bold=True, before=10)
    add_table(
        doc,
        ["编号", "项目", "标称/公差", "实测(mm)", "判定"],
        [
            ["M01", "板长L", "80.00±0.20", "", "合格/不合格"],
            ["M02", "板宽W", "60.00±0.20", "", "合格/不合格"],
            ["M03", "板厚T", "1.60±0.15", "", "合格/不合格"],
            ["M04", "定位孔径", "2.00±0.10", "", "合格/不合格"],
            ["M05", "指定点线宽/线距", "板面标识", "", "合格/不合格"],
        ],
    )
    add_text(doc, "三、基础电气（E01～E02）", bold=True, before=10)
    add_table(
        doc,
        ["编号", "测试点", "档位", "读数", "判定", "与预期一致"],
        [
            ["E01", "TP1–TP2", "", "", "开路/短路/正常", "是/否"],
            ["E02", "TP3–TP4", "", "", "开路/短路/正常", "是/否"],
        ],
    )
    add_text(doc, "四、综合判定", bold=True, before=10)
    add_text(doc, "处置结论：□接收  □返工  □报废", before=4)
    add_text(doc, "简要依据：", before=4)
    add_text(doc, "________________________________________________________________")
    add_text(doc, "________________________________________________________________")
    add_text(doc, "选手确认：__________    裁判：__________", before=12)
    path = OUT / "学生组_模块D_检测记录表.docx"
    doc.save(path)
    return path


def main():
    paths = [
        build_full_exam(),
        build_c_defect_blank(),
        build_c_defect_answer(),
        build_process_card(),
        build_d_record(),
    ]
    print("OUT DIR:", OUT)
    for p in paths:
        print("OK", p.relative_to(ROOT), p.stat().st_size)


if __name__ == "__main__":
    main()
