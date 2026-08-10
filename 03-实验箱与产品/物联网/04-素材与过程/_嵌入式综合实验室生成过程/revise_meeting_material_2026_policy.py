from pathlib import Path
import shutil

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


SOURCE = Path(r"D:\nanyun\物联网\6-2026.6.16学院提交会议讨论材料-嵌入式系统综合实训室.docx")
OUTPUT = Path(r"D:\nanyun\物联网\6-2026.6.16学院提交会议讨论材料-嵌入式系统综合实训室-评审修订版.docx")
PROJECT_NAME = "嵌入式综合实验室（畜禽养殖环境智能监测系统开发）建设项目"


def set_font(run, size=11.5, bold=False, name="仿宋_GB2312"):
    run.font.name = name
    run.font.size = Pt(size)
    run.bold = bold
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for key in ("w:ascii", "w:hAnsi", "w:eastAsia"):
        rfonts.set(qn(key), name)


def clear_cell(cell):
    tc = cell._tc
    for child in list(tc):
        if child.tag != qn("w:tcPr"):
            tc.remove(child)
    tc.append(OxmlElement("w:p"))


def set_cell(cell, text, size=12, bold=False, align=WD_ALIGN_PARAGRAPH.CENTER):
    clear_cell(cell)
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(text)
    set_font(run, size=size, bold=bold)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_para(cell, text, *, heading=False, indent=True, size=11.5, after=3):
    p = cell.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT if heading else WD_ALIGN_PARAGRAPH.JUSTIFY
    fmt = p.paragraph_format
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(after)
    fmt.line_spacing = 1.12
    if indent:
        fmt.first_line_indent = Pt(size * 2)
    run = p.add_run(text)
    set_font(run, size=12 if heading else size, bold=heading)
    return p


def auto_height(row):
    row.height = None
    row.height_rule = WD_ROW_HEIGHT_RULE.AUTO
    trpr = row._tr.get_or_add_trPr()
    for child in list(trpr):
        if child.tag == qn("w:trHeight"):
            trpr.remove(child)


def main():
    shutil.copyfile(SOURCE, OUTPUT)
    doc = Document(OUTPUT)
    table = doc.tables[0]

    set_cell(table.rows[0].cells[1], PROJECT_NAME, size=12, bold=True)

    background = table.rows[5].cells[0]
    clear_cell(background)
    add_para(
        background,
        "面向江西及学校服务区域畜禽养殖产业智能化升级需求，现代养殖现场正加快应用环境多参数感知、通风补光与精准环控、设备状态监测、异常告警和数据化运维。专业教学需要由单一传感器或单片机验证，转向以猪舍、鸡舍典型生产任务为载体，贯通“感知采集—工业通信—实时决策—设备执行—数据反馈”的嵌入式系统综合实训。",
    )
    add_para(
        background,
        "2026年2月，教育部印发《教育部关于深化职业教育教学关键要素改革的意见》（教职成〔2026〕1号），明确推进专业、课程、教材、教师、实习实训“五要素”联动改革，要求坚持需求牵引，以产定教、以产引教、以产改教、以产促教，建设产教融合实习实训基地和虚拟仿真实训基地，将产业最新技术、标准、装备和真实职业场景转化为教学内容。项目以该文件作为主要政策依据，推动畜禽养殖生产任务、嵌入式课程、教学资源和实训条件协同建设。",
    )
    add_para(
        background,
        "项目拟建设约120㎡嵌入式综合实验室，配置13套嵌入式综合实训设备，每套支持4名学生协同操作，可满足单班52名学生同时开展实训；同步建设生猪、蛋鸡智能养殖实训系统、畜禽智能养殖综合实训沙盘和嵌入式智能综合实训云平台，补齐真实设备接入、工业总线联调、虚拟仿真、过程评价和教学资源配套等条件。现提请会议审议项目建设必要性、采购内容和预算安排。",
    )
    auto_height(table.rows[5])

    overview = table.rows[7].cells[0]
    clear_cell(overview)
    add_para(overview, "一、解决问题的主要思路、方法、措施", heading=True, indent=False, after=2)
    add_para(
        overview,
        "坚持“产业场景牵引、嵌入式技术支撑、软硬件协同、虚实结合、资源同步交付”的建设思路。先以猪舍精准环控、鸡舍温感联动、饲喂饮水状态、设备巡检和故障诊断等行业任务定义实训项目，再配置ESP32无线感知、STM32实时控制、RT-Thread多任务系统、RS485/Modbus工业通信、边缘网关和教学云平台，避免脱离养殖场景堆叠技术。",
    )

    add_para(overview, "二、项目建设具体内容", heading=True, indent=False, after=2)
    add_para(
        overview,
        "1. 行业应用与综合沙盘。建设生猪智能养殖环境监测与精准环控实训系统、蛋鸡智能养殖状态监测与禽舍联动控制实训系统各1套，以及畜禽智能养殖嵌入式感知与环控综合实训沙盘1套。沙盘直观标注猪舍、鸡舍、传感点位、控制节点、RS485总线和执行设备，重点演示“鸡舍温度采集—RS485/Modbus传输—RT-Thread阈值与回差判断—风机/卷帘联动—状态反馈与告警”闭环。方案阶段配套CAD风格初步效果图，明确为非施工图。",
        indent=False,
    )
    add_para(
        overview,
        "2. 嵌入式实训设备。建设13套嵌入式综合实验箱，配置主控、通信、传感器、RFID识别和执行器开发包各13套。每套支持4人按硬件接线与传感器调试、嵌入式程序开发、通信及平台配置、数据分析与系统测试等角色协同操作。ESP32用于低成本、低功耗和Wi-Fi/BLE无线养殖感知节点；STM32用于GPIO、ADC、PWM、定时器、中断、执行机构和RS485控制；RT-Thread部署在STM32平台，用于线程调度、设备驱动、线程通信和复杂环控任务组织。",
        indent=False,
    )
    add_para(
        overview,
        "3. 虚拟仿真与教学云平台。建设嵌入式智能综合实训云平台1套，支持学生、教师、管理员三类角色和“教师创建任务—学生完成仿真或真实设备实验—本地保存数据—选择性同步至学校本地平台—AI辅助评分—教师复核—数据归档”班级任务闭环。虚拟仿真支持引脚悬停提示、错误接线自动告警、积木式I²C初始化和温度读取，以及TFT屏和0.96英寸OLED等常用外设适配；系统支持部署至学校本地服务器。",
        indent=False,
    )
    add_para(
        overview,
        "4. 教学资源与应用保障。同步交付ESP32、STM32、RT-Thread和RS485/Modbus相关实训指导书，硬件接线、虚拟仿真和平台使用教学视频，生猪环境监测、鸡舍温感联动等典型项目案例，全套项目源代码、接线图、点位表、通信协议、评分规则和验收测试表，做到硬件、软件和教学资源同步培训、同步试运行、同步验收。",
        indent=False,
    )

    add_para(overview, "三、初步预算及预算依据", heading=True, indent=False, after=2)
    add_para(
        overview,
        "项目预算总额69.7230万元。其中：嵌入式综合实验箱13套29.2500万元；五类开发包各13套合计4.1730万元；嵌入式智能综合实训云平台1套19.5000万元；生猪实训系统1套2.6500万元；蛋鸡实训系统1套2.6500万元；畜禽智能养殖综合实训沙盘1套11.5000万元。按政府采购品目分类，教学专用仪器50.2230万元，行业应用软件19.5000万元，合计69.7230万元。",
    )
    add_para(
        overview,
        "预算依据为项目采购清单的数量、技术参数和单价测算，并结合市场询价、同类项目配置和学校资产配置要求综合确定。技术路线调整不改变设备数量、采购分类和项目总额，资金拟纳入2027年度预算。",
    )

    add_para(overview, "四、提请会议审议事项", heading=True, indent=False, after=2)
    add_para(
        overview,
        "提请会议审议是否同意项目按照教职成〔2026〕1号文件和本次评审意见完善建设方案，按69.7230万元纳入2027年度预算，并按学校规定启动政府采购、合同签订、资产登记、安装联调、教师培训和建设验收程序。",
    )
    auto_height(table.rows[7])

    # 保留负责人、意见、签名、陈述人、会议时间和会议决定等现场栏目为空。
    for row, cells in [(1, [3]), (2, [1, 3]), (3, [1, 3]), (8, [3]), (10, [0])]:
        for col in cells:
            if row == 10:
                clear_cell(table.rows[row].cells[col])
                table.rows[row].cells[col].paragraphs[0].paragraph_format.space_after = Pt(54)
            else:
                set_cell(table.rows[row].cells[col], "")
    auto_height(table.rows[10])
    for row in (5, 7, 10):
        table.rows[row].cells[0].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
