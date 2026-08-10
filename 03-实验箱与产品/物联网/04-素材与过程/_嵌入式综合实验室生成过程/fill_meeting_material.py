from pathlib import Path
import shutil
import os
import stat

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


SOURCE = Path(r"D:\nanyun\物联网\6-2026.6.16-学院提交会议讨论材料-嵌入式系统综合实训室.docx")
OUTPUT = Path(r"D:\nanyun\物联网\6-2026.6.16-学院提交会议讨论材料-嵌入式系统综合实训室-填写版.docx")

PROJECT_NAME = "嵌入式综合实验室（畜禽养殖环境智能监测系统开发）建设项目"


def set_run_font(run, size=12, bold=False, font="仿宋_GB2312"):
    run.font.name = font
    run.font.size = Pt(size)
    run.bold = bold
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), font)
    rfonts.set(qn("w:hAnsi"), font)
    rfonts.set(qn("w:eastAsia"), font)


def clear_cell(cell):
    tc = cell._tc
    for child in list(tc):
        if child.tag != qn("w:tcPr"):
            tc.remove(child)
    tc.append(OxmlElement("w:p"))


def add_para(cell, text="", *, bold=False, first_indent=True, space_after=3,
             align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=11.5):
    p = cell.add_paragraph()
    p.alignment = align
    pf = p.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(space_after)
    pf.line_spacing = 1.15
    if first_indent:
        pf.first_line_indent = Pt(size * 2)
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold)
    return p


def set_cell_text(cell, text, *, size=12, bold=False, align=WD_ALIGN_PARAGRAPH.CENTER):
    clear_cell(cell)
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def remove_fixed_height(row):
    row.height = None
    row.height_rule = WD_ROW_HEIGHT_RULE.AUTO
    trpr = row._tr.get_or_add_trPr()
    for child in list(trpr):
        if child.tag == qn("w:trHeight"):
            trpr.remove(child)


def main():
    if OUTPUT.exists():
        os.chmod(OUTPUT, stat.S_IREAD | stat.S_IWRITE)
    shutil.copy2(SOURCE, OUTPUT)
    os.chmod(OUTPUT, stat.S_IREAD | stat.S_IWRITE)
    doc = Document(OUTPUT)
    table = doc.tables[0]

    # 项目名称；不填写需由学院现场确认或签署的栏目。
    set_cell_text(table.rows[0].cells[1], PROJECT_NAME, size=12, bold=True)

    background = table.rows[5].cells[0]
    clear_cell(background)
    add_para(
        background,
        "随着畜禽养殖向环境精准调控、装备智能控制和数据化运维发展，专业教学需要从单一传感器或单片机验证，升级为覆盖“感知采集—通信传输—边缘决策—执行联动—平台评价”的嵌入式系统综合实训。现有条件在多控制器协同、协议接入、边缘计算、故障诊断以及猪舍、禽舍真实业务任务方面仍有不足，难以充分支撑畜禽智能化养殖专业群复合型技术技能人才培养。",
    )
    add_para(
        background,
        "项目拟建设约120㎡嵌入式综合实验室，按13个实训工位、每组2人配置，可同时满足26名学生开展实训。建设内容以嵌入式综合实验箱、主控/通信/传感器/RFID识别/执行器开发包和嵌入式智能综合实训云平台为基础，引入生猪、蛋鸡智能养殖实训系统及畜禽智能养殖综合实训沙盘，形成“底层驱动—联网通信—边缘控制—行业应用”贯通的教学环境。",
    )
    add_para(
        background,
        "项目依据《国家职业教育改革实施方案》、职业教育专业教学标准、学校专业群建设与实训条件改善任务，以及2027年度预算申报要求提出。建设后将服务嵌入式系统开发、智能硬件应用、畜禽养殖环境监测与环控、技能竞赛训练、教师科研和社会培训。现提请会议审议项目建设必要性、采购内容和预算安排。",
    )
    remove_fixed_height(table.rows[5])

    overview = table.rows[7].cells[0]
    clear_cell(overview)

    add_para(overview, "一、解决问题的主要思路、方法、措施", bold=True,
             first_indent=False, space_after=2, align=WD_ALIGN_PARAGRAPH.LEFT, size=12)
    add_para(
        overview,
        "坚持“嵌入式技术主线、畜禽养殖场景牵引、软硬件协同、虚实结合、教学赛研共用”的建设思路。以STM32、OpenHarmony、ZigBee和ARM边缘网关为核心，组织GPIO、ADC、PWM、UART、I²C、SPI、RS485/Modbus、MQTT等递进式实训；通过猪舍、禽舍和综合沙盘完成环境感知、边缘规则、执行机构控制、数据可视化及故障诊断闭环；通过云平台统一设备接入、任务发布、过程记录和教学评价。",
    )

    add_para(overview, "二、项目建设具体内容", bold=True,
             first_indent=False, space_after=2, align=WD_ALIGN_PARAGRAPH.LEFT, size=12)
    add_para(
        overview,
        "1. 建设13套嵌入式综合实验箱，配置主控开发包、通信开发包、传感器开发包、RFID识别开发包、执行器开发包各13套，形成13个双人实训工位，可同时容纳26名学生。",
        first_indent=False,
    )
    add_para(
        overview,
        "2. 建设1套嵌入式智能综合实训云平台，支持设备建模、协议接入、在线仿真、数据可视化、规则联动、AI教学辅助、实验任务和过程评价。",
        first_indent=False,
    )
    add_para(
        overview,
        "3. 建设生猪智能养殖环境监测与精准环控实训系统、蛋鸡智能养殖状态监测与禽舍联动控制实训系统各1套，开展温湿度、光照、空气质量、饲喂饮水状态、产蛋计数、巡检识别及通风补光等嵌入式综合项目。",
        first_indent=False,
    )
    add_para(
        overview,
        "4. 建设1套畜禽智能养殖嵌入式感知与环控综合实训沙盘，集成猪舍、禽舍、传感器、低压执行机构、嵌入式节点、边缘网关和本地触控终端，用于系统联调、控制策略验证及故障诊断。",
        first_indent=False,
    )

    add_para(overview, "三、初步预算及预算依据", bold=True,
             first_indent=False, space_after=2, align=WD_ALIGN_PARAGRAPH.LEFT, size=12)
    add_para(
        overview,
        "项目预算总额69.7230万元。其中：嵌入式综合实验箱13套29.2500万元；五类开发包各13套合计4.1730万元；嵌入式智能综合实训云平台1套19.5000万元；生猪实训系统1套2.6500万元；蛋鸡实训系统1套2.6500万元；畜禽智能养殖综合实训沙盘1套11.5000万元。按政府采购品目分类，教学专用仪器50.2230万元，行业应用软件19.5000万元，合计69.7230万元。",
    )
    add_para(
        overview,
        "预算依据为最新项目采购清单的数量、技术参数和单价测算，并结合市场询价、同类项目配置和学校资产配置要求综合确定；资金拟纳入2027年度预算，后续按学校政府采购、合同管理、资产登记和项目验收制度执行。",
    )

    add_para(overview, "四、提请会议审议事项", bold=True,
             first_indent=False, space_after=2, align=WD_ALIGN_PARAGRAPH.LEFT, size=12)
    add_para(
        overview,
        "提请会议审议是否同意该项目按69.7230万元纳入2027年度预算，并按学校规定启动政府采购、合同签订、资产登记和建设验收程序。随附项目采购清单、建设方案、2027年预算支出申报表、项目经费申报书、政府采购预算表、资产配置表及采购可行性论证报告。",
    )
    remove_fixed_height(table.rows[7])

    # 会议决定留白，但保证留有书写空间。
    decision = table.rows[10].cells[0]
    clear_cell(decision)
    p = decision.paragraphs[0]
    p.paragraph_format.space_after = Pt(54)
    remove_fixed_height(table.rows[10])

    # 防止长内容被拆成多行时出现单元格顶部拥挤。
    for row_idx in (5, 7, 10):
        table.rows[row_idx].cells[0].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
