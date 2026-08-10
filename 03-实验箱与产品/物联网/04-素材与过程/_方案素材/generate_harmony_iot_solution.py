from pathlib import Path
import math
import re
import shutil
from zipfile import ZipFile

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Cm, Pt, RGBColor, Inches
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parent.parent
ASSET = ROOT / "_方案素材"
ORIG = ASSET / "鸿蒙原图"
DIAG = ASSET / "原创图示"
DIAG.mkdir(parents=True, exist_ok=True)
OUT = ROOT / "鸿蒙物联网综合实训室建设方案-图文并茂版.docx"


def font_path(*names):
    for name in names:
        p = Path("C:/Windows/Fonts") / name
        if p.exists():
            return str(p)
    return None


FONT_MAIN = font_path("msyh.ttc", "simhei.ttf", "simsun.ttc")
FONT_SONG = font_path("simsun.ttc", "msyh.ttc")
FONT_BOLD = font_path("msyhbd.ttc", "msyh.ttc", "simhei.ttf")


def f(size, bold=False):
    return ImageFont.truetype(FONT_BOLD if bold and FONT_BOLD else FONT_MAIN, size)


def wrap(draw, text, font, max_width):
    lines = []
    for part in str(text).split("\n"):
        cur = ""
        for ch in part:
            test = cur + ch
            if draw.textlength(test, font=font) <= max_width:
                cur = test
            else:
                if cur:
                    lines.append(cur)
                cur = ch
        if cur:
            lines.append(cur)
    return lines


def rounded(draw, box, fill, outline="#CBD5E1", radius=18, width=2):
    draw.rounded_rectangle(box, radius=radius, fill=fill, outline=outline, width=width)


def center_text(draw, box, text, font, fill="#0F172A", spacing=8):
    x1, y1, x2, y2 = box
    lines = wrap(draw, text, font, x2 - x1 - 28)
    heights = [draw.textbbox((0, 0), line, font=font)[3] for line in lines]
    total = sum(heights) + spacing * (len(lines) - 1)
    y = y1 + (y2 - y1 - total) / 2
    for line, h in zip(lines, heights):
        w = draw.textlength(line, font=font)
        draw.text((x1 + (x2 - x1 - w) / 2, y), line, font=font, fill=fill)
        y += h + spacing


def arrow(draw, start, end, fill="#475569", width=4):
    draw.line([start, end], fill=fill, width=width)
    sx, sy = start
    ex, ey = end
    ang = math.atan2(ey - sy, ex - sx)
    size = 12
    pts = [
        (ex, ey),
        (ex - size * math.cos(ang - math.pi / 6), ey - size * math.sin(ang - math.pi / 6)),
        (ex - size * math.cos(ang + math.pi / 6), ey - size * math.sin(ang + math.pi / 6)),
    ]
    draw.polygon(pts, fill=fill)


def title(draw, text, subtitle=None):
    draw.text((60, 38), text, font=f(38, True), fill="#0F172A")
    draw.line((60, 92, 1540, 92), fill="#3B82F6", width=4)
    if subtitle:
        draw.text((60, 108), subtitle, font=f(21), fill="#475569")


def save_canvas(name, draw_fn, size=(1600, 980)):
    img = Image.new("RGB", size, "#F8FAFC")
    draw = ImageDraw.Draw(img)
    draw_fn(draw, size)
    path = DIAG / f"{name}.png"
    img.save(path, quality=95)
    return path


def diagram_policy():
    def draw_fn(d, size):
        title(d, "政策-产业-专业建设逻辑图", "以职业教育数字化与鸿蒙物联网生态支撑专业实训条件升级")
        colors = ["#DBEAFE", "#D1FAE5", "#FEF3C7", "#FCE7F3"]
        boxes = [
            (100, 210, 420, 410, "政策牵引\n职业教育数字化\n现代职教体系\n专业教学标准"),
            (470, 210, 790, 410, "产业需求\nAIoT终端接入\n边缘网关集成\n数据可视化运维"),
            (840, 210, 1160, 410, "技术底座\nOpenHarmony\nSTM32/ZigBee\nMQTT/Modbus"),
            (1210, 210, 1530, 410, "专业建设\n课程实训平台\n云平台\n50人项目化教学"),
        ]
        for i, (x1, y1, x2, y2, txt) in enumerate(boxes):
            rounded(d, (x1, y1, x2, y2), colors[i], "#93C5FD")
            center_text(d, (x1, y1, x2, y2), txt, f(27, True))
            if i < len(boxes) - 1:
                arrow(d, (x2 + 16, 310), (boxes[i + 1][0] - 16, 310), "#2563EB")
        rounded(d, (180, 560, 1420, 820), "#FFFFFF", "#94A3B8", 22)
        center_text(
            d,
            (220, 590, 1380, 790),
            "建设结果：形成“硬件可操作、软件可管理、数据可沉淀、项目可评价、场景可迁移”的鸿蒙物联网综合实训环境",
            f(32, True),
            "#0F172A",
        )
    return save_canvas("01_policy_logic", draw_fn)


def diagram_teaching_loop():
    def draw_fn(d, size):
        title(d, "教学做一体化闭环图", "课堂讲授、设备操作、平台评价和项目复盘贯通")
        cx, cy, r = 800, 510, 310
        labels = [
            ("理论导入\n知识点讲解", 800, 175, "#DBEAFE"),
            ("硬件操作\n实验箱与模块", 1135, 510, "#D1FAE5"),
            ("平台联调\n数据上报与控制", 800, 845, "#FEF3C7"),
            ("项目评价\n报告与复盘", 465, 510, "#FCE7F3"),
        ]
        points = [(800, 310), (1000, 510), (800, 710), (600, 510)]
        for i in range(4):
            arrow(d, points[i], points[(i + 1) % 4], "#64748B", 5)
        for txt, x, y, color in labels:
            rounded(d, (x - 160, y - 70, x + 160, y + 70), color, "#94A3B8")
            center_text(d, (x - 150, y - 60, x + 150, y + 60), txt, f(25, True))
        rounded(d, (620, 405, 980, 615), "#FFFFFF", "#60A5FA", 28, 3)
        center_text(d, (650, 430, 950, 590), "教师发布任务\n学生分组实训\n数据自动沉淀", f(27, True), "#1E3A8A")
    return save_canvas("02_teaching_loop", draw_fn)


def diagram_ability_route():
    def draw_fn(d, size):
        title(d, "学生能力递进路线图", "由单点验证逐步进入端、边、云、用完整工程链路")
        stages = [
            ("基础验证", "GPIO/UART/I2C\nADC/PWM/中断\nOpenHarmony工程创建", "#DBEAFE"),
            ("模块组合", "温湿度/光照/CO2\nZigBee组网\nWiFi与MQTT上报", "#D1FAE5"),
            ("综合项目", "网关协议转换\n规则引擎\n平台大屏与报告", "#FEF3C7"),
            ("行业创新", "智慧农业/养殖\n微生物发酵监测\n数字孪生与AI问答", "#FCE7F3"),
        ]
        x = 120
        for i, (h, body, color) in enumerate(stages):
            rounded(d, (x, 260, x + 310, 650), color, "#94A3B8", 22)
            center_text(d, (x + 20, 300, x + 290, 390), h, f(30, True), "#0F172A")
            center_text(d, (x + 25, 410, x + 285, 620), body, f(22), "#334155")
            if i < 3:
                arrow(d, (x + 325, 455), (x + 385, 455), "#2563EB", 5)
            x += 370
        d.text((190, 760), "能力输出：硬件调试、协议通信、平台建模、数据分析、规则控制、项目文档与现场表达", font=f(27, True), fill="#0F172A")
    return save_canvas("03_ability_route", draw_fn)


def diagram_space():
    def draw_fn(d, size):
        title(d, "120㎡空间布局示意图", "25组学生工位、50人同步实训，兼顾教学演示、平台运维和设备收纳")
        room = (120, 160, 1480, 860)
        rounded(d, room, "#FFFFFF", "#334155", 12, 4)
        rounded(d, (160, 190, 1440, 285), "#DBEAFE", "#60A5FA", 14)
        center_text(d, (160, 190, 1440, 285), "教师演示区：讲台 / 大屏投影 / 平台管理终端 / 实物展示台", f(25, True), "#1E3A8A")
        start_x, start_y = 190, 335
        cell_w, cell_h, gap_x, gap_y = 210, 78, 42, 28
        idx = 1
        for row in range(5):
            for col in range(5):
                x1 = start_x + col * (cell_w + gap_x)
                y1 = start_y + row * (cell_h + gap_y)
                rounded(d, (x1, y1, x1 + cell_w, y1 + cell_h), "#ECFDF5", "#34D399", 10)
                center_text(d, (x1, y1, x1 + cell_w, y1 + cell_h), f"{idx:02d}组\n2人/实验箱", f(18, True), "#064E3B", 3)
                idx += 1
        rounded(d, (1240, 335, 1440, 525), "#FEF3C7", "#F59E0B", 12)
        center_text(d, (1240, 335, 1440, 525), "网络与\n服务器柜", f(22, True), "#78350F")
        rounded(d, (1240, 555, 1440, 760), "#FCE7F3", "#F472B6", 12)
        center_text(d, (1240, 555, 1440, 760), "设备收纳\n耗材备件\n充电维护", f(21, True), "#831843")
        d.text((170, 805), "建议：低压供电优先、强弱电分离、桌面预留USB/网口/安全插座，教师端可集中展示网关日志、数据大屏和学生过程记录。", font=f(21), fill="#334155")
    return save_canvas("04_space_layout", draw_fn)


def diagram_architecture():
    def draw_fn(d, size):
        title(d, "端-边-云总体架构图", "从鸿蒙终端采集到云平台展示、规则联动和教学评价")
        layers = [
            ("感知与执行层", ["温湿度", "光照", "CO2", "重量", "RFID/NFC", "风扇/水泵/舵机"], "#DBEAFE"),
            ("终端控制层", ["OpenHarmony模块", "STM32模块", "ZigBee节点", "通信开发包"], "#D1FAE5"),
            ("边缘网关层", ["协议转换", "规则判定", "本地缓存", "触控看板"], "#FEF3C7"),
            ("AIoT平台层", ["设备建模", "数据可视化", "任务发布", "过程评价", "AI助教"], "#FCE7F3"),
            ("场景应用层", ["智能家居", "环境监测", "智慧农业", "智慧养殖", "发酵监测"], "#EDE9FE"),
        ]
        y = 175
        for i, (name, items, color) in enumerate(layers):
            rounded(d, (120, y, 1480, y + 120), color, "#94A3B8", 18)
            d.text((155, y + 38), name, font=f(28, True), fill="#0F172A")
            x = 430
            for item in items:
                rounded(d, (x, y + 25, x + 160, y + 95), "#FFFFFF", "#CBD5E1", 12)
                center_text(d, (x + 5, y + 25, x + 155, y + 95), item, f(18, True), "#334155")
                x += 175
            if i < len(layers) - 1:
                arrow(d, (800, y + 130), (800, y + 160), "#475569", 5)
            y += 150
    return save_canvas("05_architecture", draw_fn)


def diagram_network():
    def draw_fn(d, size):
        title(d, "实训室网络拓扑图", "千兆有线、无线接入、边缘网关与平台服务器统一部署")
        nodes = [
            ("互联网/校内网", (700, 150, 900, 230), "#DBEAFE"),
            ("防火墙/路由器", (680, 300, 920, 385), "#D1FAE5"),
            ("核心交换机", (680, 455, 920, 540), "#FEF3C7"),
            ("AIoT平台服务器\n数据备份存储", (1080, 455, 1380, 560), "#FCE7F3"),
            ("教师管理终端\n投屏/大屏", (230, 455, 520, 560), "#EDE9FE"),
            ("无线AP", (245, 675, 455, 760), "#DBEAFE"),
            ("25组学生工位\n实验箱/PC", (610, 670, 990, 780), "#D1FAE5"),
            ("边缘网关/触控看板", (1110, 675, 1370, 770), "#FEF3C7"),
        ]
        for txt, box, color in nodes:
            rounded(d, box, color, "#94A3B8", 16)
            center_text(d, box, txt, f(22, True), "#0F172A")
        arrow(d, (800, 230), (800, 300), "#2563EB")
        arrow(d, (800, 385), (800, 455), "#2563EB")
        arrow(d, (680, 500), (520, 500), "#475569")
        arrow(d, (920, 500), (1080, 500), "#475569")
        arrow(d, (800, 540), (800, 670), "#475569")
        arrow(d, (720, 540), (455, 675), "#475569")
        arrow(d, (880, 540), (1110, 675), "#475569")
        d.text((170, 845), "部署要点：学生端设备、虚拟仿真、边缘网关和平台服务器统一接入；教学数据按班级、任务、设备和时间维度沉淀。", font=f(23), fill="#334155")
    return save_canvas("06_network", draw_fn)


def diagram_equipment():
    def draw_fn(d, size):
        title(d, "核心设备组成图", "以鸿蒙物联网实验箱为中心，形成可扩展的分组实训套件")
        center = (610, 360, 990, 580)
        rounded(d, center, "#FFFFFF", "#2563EB", 24, 4)
        center_text(d, center, "鸿蒙物联网实验箱\n便携一体化 / 模块快接\n9个通用接口位", f(27, True), "#1E3A8A")
        items = [
            ("主控开发包\nSTM32/鸿蒙/ZigBee", (120, 180, 470, 300), "#DBEAFE"),
            ("通信开发包\nWiFi/4G/BLE/LoRa/NB-IoT", (1130, 180, 1480, 300), "#D1FAE5"),
            ("传感器开发包\n环境/安防/农业/气体", (120, 660, 470, 780), "#FEF3C7"),
            ("RFID识别开发包\nHF/UHF/NFC/二维码", (1130, 660, 1480, 780), "#FCE7F3"),
            ("执行器开发包\n风扇/水泵/舵机/灯光", (625, 760, 975, 880), "#EDE9FE"),
            ("AIoT教学云平台\n设备/任务/评价/仿真", (625, 130, 975, 250), "#CCFBF1"),
        ]
        for txt, box, color in items:
            rounded(d, box, color, "#94A3B8", 16)
            center_text(d, box, txt, f(21, True), "#0F172A")
            arrow(d, ((box[0]+box[2])//2, (box[1]+box[3])//2), ((center[0]+center[2])//2, (center[1]+center[3])//2), "#64748B", 3)
    return save_canvas("07_equipment", draw_fn)


def diagram_course():
    def draw_fn(d, size):
        title(d, "课程体系设计图", "基础课程、平台课程、综合项目、行业场景分层支撑")
        cols = [
            ("基础开发课程", ["OpenHarmony工程", "STM32外设", "ZigBee组网", "传感器采集"], "#DBEAFE"),
            ("平台集成课程", ["MQTT/HTTP", "Modbus接入", "设备建模", "数据可视化"], "#D1FAE5"),
            ("综合项目课程", ["规则联动", "本地看板", "虚实仿真", "项目报告"], "#FEF3C7"),
            ("行业场景课程", ["智慧农业", "智慧养殖", "发酵监测", "数字孪生"], "#FCE7F3"),
        ]
        x = 95
        for name, items, color in cols:
            rounded(d, (x, 190, x + 330, 780), color, "#94A3B8", 20)
            center_text(d, (x + 10, 220, x + 320, 295), name, f(27, True))
            y = 340
            for item in items:
                rounded(d, (x + 40, y, x + 290, y + 70), "#FFFFFF", "#CBD5E1", 12)
                center_text(d, (x + 40, y, x + 290, y + 70), item, f(21, True), "#334155")
                y += 95
            x += 370
        d.text((160, 850), "课程资源建议配套：实验指导书、源码、接线图、平台配置截图、报告模板、评分量规、教师培训材料。", font=f(23), fill="#334155")
    return save_canvas("08_course_system", draw_fn)


def diagram_extension():
    def draw_fn(d, size):
        title(d, "拓展应用与前瞻发展路线图", "以能力预留方式支撑后续校企协同、课程共建和创新项目")
        stages = [
            ("近期：建成即用", "教师培训\n课程导入\n设备联调\n试运行验收", "#DBEAFE"),
            ("中期：资源共建", "项目库迭代\n教材素材沉淀\n教师课题孵化", "#D1FAE5"),
            ("远期：生态拓展", "校企协同课程\n开放接口二开\n行业应用样板", "#FEF3C7"),
            ("持续：质量提升", "竞赛训练预留\n优秀作品库\n数据化评价", "#FCE7F3"),
        ]
        x = 120
        for i, (h, body, color) in enumerate(stages):
            rounded(d, (x, 260, x + 310, 650), color, "#94A3B8", 24)
            center_text(d, (x + 25, 300, x + 285, 390), h, f(28, True))
            center_text(d, (x + 25, 415, x + 285, 620), body, f(22))
            if i < len(stages) - 1:
                arrow(d, (x + 325, 455), (x + 385, 455), "#2563EB", 5)
            x += 370
        d.text((135, 780), "说明：本章只采用“可支撑、可拓展、建议后续开展”的前瞻表述，不写既有产教融合业绩或赛事成绩。", font=f(24, True), fill="#7C2D12")
    return save_canvas("09_extension_route", draw_fn)


def diagram_competition():
    def draw_fn(d, size):
        title(d, "竞赛训练能力映射图", "面向物联网、嵌入式、人工智能与鸿蒙生态方向预留训练条件")
        center = (600, 380, 1000, 585)
        rounded(d, center, "#FFFFFF", "#2563EB", 28, 4)
        center_text(d, center, "鸿蒙物联网实训室\n竞赛训练能力预留", f(30, True), "#1E3A8A")
        items = [
            ("物联网应用开发\n设备接入/协议/平台", (110, 170, 465, 310), "#DBEAFE"),
            ("嵌入式系统\nMCU外设/驱动/调试", (1135, 170, 1490, 310), "#D1FAE5"),
            ("人工智能应用\nAI助教/视觉/语音接口", (110, 660, 465, 800), "#FEF3C7"),
            ("鸿蒙生态创新\nOpenHarmony轻量设备", (1135, 660, 1490, 800), "#FCE7F3"),
        ]
        for txt, box, color in items:
            rounded(d, box, color, "#94A3B8", 18)
            center_text(d, box, txt, f(23, True))
            arrow(d, ((box[0]+box[2])//2, (box[1]+box[3])//2), (800, 480), "#64748B", 3)
        d.text((240, 865), "定位：作为后续训练条件与能力支撑，不承诺现有参赛案例、奖项或既有合作成果。", font=f(24, True), fill="#7C2D12")
    return save_canvas("10_competition_map", draw_fn)


def diagram_implementation():
    def draw_fn(d, size):
        title(d, "实施进度流程图", "从现场确认到验收归档的五阶段实施路径")
        steps = [
            ("1 现场确认", "尺寸/电力/网络\n工位与清单深化"),
            ("2 到货部署", "设备清点\n资产编号\n环境布置"),
            ("3 平台联调", "服务器部署\n设备接入\n数据采集"),
            ("4 资源导入", "课程包\n账号权限\n教师培训"),
            ("5 试运行验收", "基础实验\n综合项目\n文档归档"),
        ]
        x = 80
        for i, (h, body) in enumerate(steps):
            color = ["#DBEAFE", "#D1FAE5", "#FEF3C7", "#FCE7F3", "#EDE9FE"][i]
            rounded(d, (x, 290, x + 260, 640), color, "#94A3B8", 20)
            center_text(d, (x + 20, 330, x + 240, 420), h, f(25, True))
            center_text(d, (x + 25, 455, x + 235, 600), body, f(21))
            if i < 4:
                arrow(d, (x + 275, 465), (x + 330, 465), "#2563EB", 5)
            x += 300
        d.text((145, 760), "交付资料：深化设计、设备清单、平台账号、课程资源、培训记录、试运行记录、验收报告、运维制度。", font=f(24), fill="#334155")
    return save_canvas("11_implementation", draw_fn)


def generate_diagrams():
    return {
        "policy": diagram_policy(),
        "teaching": diagram_teaching_loop(),
        "ability": diagram_ability_route(),
        "space": diagram_space(),
        "architecture": diagram_architecture(),
        "network": diagram_network(),
        "equipment": diagram_equipment(),
        "course": diagram_course(),
        "extension": diagram_extension(),
        "competition": diagram_competition(),
        "implementation": diagram_implementation(),
    }


def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def set_cell_text(cell, text, bold=False, color="000000", size=10.5):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if bold else WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(str(text))
    run.font.name = "宋体"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("第 ")
    fldChar1 = OxmlElement("w:fldChar")
    fldChar1.set(qn("w:fldCharType"), "begin")
    instrText = OxmlElement("w:instrText")
    instrText.set(qn("xml:space"), "preserve")
    instrText.text = "PAGE"
    fldChar2 = OxmlElement("w:fldChar")
    fldChar2.set(qn("w:fldCharType"), "end")
    run._r.append(fldChar1)
    run._r.append(instrText)
    run._r.append(fldChar2)
    paragraph.add_run(" 页")


def set_run_font(run, name="宋体", size=11, bold=False, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def paragraph(doc, text="", style=None, first_line=True, align=None):
    p = doc.add_paragraph(style=style)
    if text:
        run = p.add_run(text)
        set_run_font(run, "宋体", 11)
    p.paragraph_format.line_spacing = 1.35
    p.paragraph_format.space_after = Pt(6)
    if first_line:
        p.paragraph_format.first_line_indent = Pt(22)
    if align:
        p.alignment = align
    return p


def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    run = p.add_run(text)
    set_run_font(run, "宋体", 10.5)
    p.paragraph_format.line_spacing = 1.25
    p.paragraph_format.space_after = Pt(3)
    return p


def heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        run.font.name = "微软雅黑"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        run.font.color.rgb = RGBColor.from_string("0F172A")
        if level == 1:
            run.font.size = Pt(18)
        elif level == 2:
            run.font.size = Pt(14)
        else:
            run.font.size = Pt(12)
    p.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(8)
    return p


def caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    set_run_font(run, "宋体", 9.5, False, "475569")
    p.paragraph_format.space_after = Pt(8)
    return p


def add_image(doc, path, cap, width_cm=15.7):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(path), width=Cm(width_cm))
    caption(doc, cap)


def add_table(doc, title_text, headers, rows, widths=None):
    caption(doc, title_text)
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        set_cell_shading(hdr[i], "1E3A8A")
        set_cell_text(hdr[i], h, True, "FFFFFF", 10)
    for row in rows:
        cells = table.add_row().cells
        for i, v in enumerate(row):
            set_cell_text(cells[i], v, False, "111827", 9.5)
            if i == 0:
                set_cell_shading(cells[i], "EFF6FF")
    if widths:
        for row in table.rows:
            for idx, width in enumerate(widths):
                row.cells[idx].width = Cm(width)
    doc.add_paragraph()
    return table


def create_doc(diagrams):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Cm(21)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(2.2)
    sec.bottom_margin = Cm(2.0)
    sec.left_margin = Cm(2.2)
    sec.right_margin = Cm(2.2)
    sec.header_distance = Cm(1.2)
    sec.footer_distance = Cm(1.0)

    styles = doc.styles
    styles["Normal"].font.name = "宋体"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    styles["Normal"].font.size = Pt(11)
    for style_name in ["Heading 1", "Heading 2", "Heading 3"]:
        st = styles[style_name]
        st.font.name = "微软雅黑"
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        st.font.color.rgb = RGBColor.from_string("0F172A")

    header_p = sec.header.paragraphs[0]
    header_p.text = "鸿蒙物联网综合实训室建设方案"
    header_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header_p.runs:
        set_run_font(run, "微软雅黑", 9, False, "64748B")
    add_page_number(sec.footer.paragraphs[0])

    cover = doc.add_paragraph()
    cover.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cover.paragraph_format.space_before = Pt(90)
    r = cover.add_run("鸿蒙物联网综合实训室建设方案")
    set_run_font(r, "微软雅黑", 28, True, "0F172A")
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub.add_run("面向 AIoT 项目化教学、虚实结合实训与行业场景应用的图文并茂版方案")
    set_run_font(r, "微软雅黑", 14, False, "334155")
    info = doc.add_paragraph()
    info.alignment = WD_ALIGN_PARAGRAPH.CENTER
    info.paragraph_format.space_before = Pt(36)
    r = info.add_run("项目名称：信息工程系物联网综合实训室建设项目\n建设面积：约120㎡    同步容量：50人（25组）    预算控制：约86万元\n方案版本：V1.0    生成日期：2026年7月")
    set_run_font(r, "宋体", 12, False, "334155")
    add_image(doc, diagrams["policy"], "图 0-1 方案总体建设逻辑", 15.0)
    doc.add_page_break()

    heading(doc, "目录", 1)
    toc_items = [
        "一、项目概述",
        "二、教学实验总体设计",
        "三、鸿蒙物联网综合实训室",
        "四、拓展应用与前瞻发展",
        "五、竞赛训练与能力提升展望",
        "六、建设清单与实施保障",
        "七、资料依据",
    ]
    for item in toc_items:
        paragraph(doc, item, first_line=False)
    doc.add_page_break()

    heading(doc, "一、项目概述", 1)
    heading(doc, "1.1 政策背景", 2)
    paragraph(doc, "信息工程系物联网综合实训室建设项目以职业教育数字化转型、现代职业教育体系建设和物联网应用技术专业人才培养为牵引，面向课程教学、实验实训、技能训练和创新项目研发等多类教学任务，建设一间可支撑50人同步实训的综合性实训空间。")
    paragraph(doc, "从专业建设角度看，物联网应用技术已从单一传感器采集、单片机验证，逐步发展为端侧设备开发、边缘网关集成、云端平台管理、数据分析展示和行业场景应用的完整工程链路。实训室建设需要避免只采购零散器材，而应形成“设备、平台、课程、项目、评价、运维”一体化条件。")
    add_image(doc, diagrams["policy"], "图 1-1 政策、产业与专业建设逻辑图")
    add_table(
        doc,
        "表 1-1 项目约束与建设定位",
        ["项目维度", "公告要求/建设口径", "方案响应"],
        [
            ["建设面积", "规划建设面积约120㎡", "按教师演示区、25组学生工位、网络服务器区、设备收纳区、文化展示区组织空间。"],
            ["同步容量", "可容纳50人同时练习和教学实训", "采用25组分组工位，建议2人一组，每组可开展设备接入、数据采集、平台联调和项目报告。"],
            ["预算控制", "项目预算约86万元", "设备、平台、课程资源、场景系统、文化建设和装修改造统筹控制。"],
            ["建设内容", "课程教学实训平台、物联网云平台、物联网应用实训设备套件、课程资源、应用套件和耗材", "以鸿蒙物联网实验箱和AIoT教学云平台为核心，配套主控、通信、传感、RFID、执行器及行业场景。"],
            ["建设风格", "完整、详尽、可行的实施方案", "图文并茂呈现总体架构、空间布局、网络拓扑、设备组成、课程体系、实施进度和验收指标。"],
        ],
        [3.0, 5.4, 7.2],
    )

    heading(doc, "1.2 专业背景", 2)
    paragraph(doc, "物联网技术是电子信息、软件技术、人工智能、现代农业、生物工程等专业群共同需要的基础能力。面向真实岗位，学生不仅要理解传感器和控制器的单点原理，还要掌握多协议通信、设备上云、数据可视化、规则联动、边缘网关与平台运维等综合能力。")
    paragraph(doc, "鸿蒙/OpenHarmony轻量设备开发为物联网教学提供了新的技术切入点。通过鸿蒙物联网模块、STM32控制模块、ZigBee无线节点、中心网关和AIoT云平台协同，学生能够在统一实训环境中完成“采集、通信、平台、控制、展示、评价”的完整闭环。")
    paragraph(doc, "本项目建议建设一间兼具课程教学、项目实训、行业场景模拟和后续能力拓展的鸿蒙物联网综合实训室，使专业教学从传统验证型实验升级为工程任务型、数据驱动型和场景融合型实训。")

    heading(doc, "1.3 人才需求", 2)
    paragraph(doc, "物联网相关岗位对学生的要求正在从“会接线、会烧录、会读数”转向“能接入、会建模、能联调、懂数据、会交付”。建设鸿蒙物联网综合实训室，可以面向以下岗位能力开展递进式训练：")
    for t in [
        "物联网设备开发与调试：掌握MCU外设、鸿蒙轻量设备开发、串口调试、传感器采集和执行器控制。",
        "物联网系统集成与运维：掌握ZigBee、WiFi、4G、BLE、LoRa、NB-IoT/Cat.1等通信方式，以及MQTT、HTTP、Modbus等协议接入。",
        "AIoT平台应用与数据可视化：掌握设备建模、数据上报、规则联动、报警管理、历史曲线、项目大屏和实验报告引用。",
        "行业应用项目实施：能够围绕智慧农业、智慧养殖、微生物发酵环境监测、智能家居等场景完成需求分析、方案设计、联调测试和成果展示。",
    ]:
        bullet(doc, t)

    heading(doc, "1.4 知识体系", 2)
    add_table(
        doc,
        "表 1-2 鸿蒙物联网实训室知识体系",
        ["知识模块", "核心内容", "对应能力"],
        [
            ["基础理论", "传感器原理、嵌入式系统、计算机网络、数据库基础、信息安全基础", "理解设备采集、边缘处理和数据传输的基础逻辑。"],
            ["端侧开发", "OpenHarmony工程创建、GPIO/UART/I2C/SPI/ADC/PWM、STM32外设、ZigBee节点", "完成终端驱动、数据采集、设备控制和本地调试。"],
            ["通信协议", "WiFi、ZigBee、BLE、LoRa、4G、NB-IoT/Cat.1、MQTT、HTTP、Modbus", "完成多协议接入、网关汇聚和云端数据上报。"],
            ["平台应用", "设备建模、产品分类、命令下发、规则联动、报警记录、数据大屏、过程评价", "完成平台侧配置、课堂任务发布、数据沉淀和实验报告。"],
            ["行业项目", "智慧农业、智慧养殖、微生物发酵、环境监测、智能家居、资产识别", "形成跨专业项目开发、展示答辩和持续迭代能力。"],
        ],
        [3.0, 6.0, 6.6],
    )

    heading(doc, "1.5 专业方向与就业方向", 2)
    paragraph(doc, "实训室可服务物联网应用技术、电子信息、人工智能应用、软件技术、生物工程和现代农业等专业群。围绕专业方向，可形成鸿蒙物联网设备开发、AIoT平台应用、智能传感与检测、多协议通信、行业场景系统集成等教学方向。")
    add_table(
        doc,
        "表 1-3 专业方向与就业岗位映射",
        ["专业方向", "典型课程/实训", "岗位能力指向"],
        [
            ["鸿蒙物联网设备开发", "OpenHarmony轻量设备、GPIO与通信接口、设备上云", "物联网终端开发、嵌入式助理工程师、硬件调试。"],
            ["AIoT平台应用", "设备接入、数据可视化、规则引擎、实验过程评价", "平台实施、物联网运维、数据看板配置。"],
            ["智能传感与检测", "环境采集、气体检测、重量检测、CO2趋势分析", "传感器应用、检测系统维护、现场数据采集。"],
            ["行业场景集成", "智慧农业、智慧养殖、发酵监测、智能家居", "系统集成、项目实施、售前技术支持、应用开发。"],
        ],
        [3.3, 5.6, 6.7],
    )

    heading(doc, "二、教学实验总体设计", 1)
    heading(doc, "2.1 建设需求及目标", 2)
    paragraph(doc, "本项目建设目标是形成一套面向真实工程链路的鸿蒙物联网综合实训环境。实训室建成后，应能够支撑课堂演示、学生分组训练、虚拟仿真预习、真实设备联调、平台过程评价、行业场景项目和后续能力拓展。")
    add_image(doc, diagrams["teaching"], "图 2-1 教学做一体化闭环图")
    add_table(
        doc,
        "表 2-1 建设目标分解",
        ["目标类别", "目标描述", "验收关注点"],
        [
            ["教学目标", "支撑物联网基础、鸿蒙开发、通信协议、平台应用、行业项目等课程。", "课程资源可导入、任务可发布、学生可提交报告。"],
            ["实训目标", "25组学生工位同步开展实验，完成从采集到控制的闭环训练。", "50人同时实训，设备可稳定接入，数据可采集。"],
            ["平台目标", "建设AIoT教学云平台，实现设备、课程、任务、资源、成绩和报告统一管理。", "教师、学生、管理员角色可用，设备和实验数据可查询。"],
            ["场景目标", "形成智慧农业、智慧养殖、微生物发酵、智能家居等可迁移项目。", "至少形成基础实验、综合项目和行业项目的分层案例。"],
            ["持续目标", "预留产教协同、创新课题、竞赛训练和二次开发能力。", "接口开放、资源可迭代、数据可沉淀、过程可评价。"],
        ],
        [3.0, 6.5, 5.9],
    )
    heading(doc, "2.2 建设规划", 2)
    paragraph(doc, "实训室按照“一平台、两类终端、三层能力、四类场景”的思路进行规划。一平台指AIoT教学云平台；两类终端指真实硬件终端与在线仿真终端；三层能力指感知控制基础、协议网关集成、行业应用创新；四类场景指智能家居、环境监测、智慧农业/发酵、智慧养殖与数字孪生。")
    add_image(doc, diagrams["ability"], "图 2-2 学生能力递进路线图")
    add_table(
        doc,
        "表 2-2 实训体系规划",
        ["层级", "典型内容", "能力目标"],
        [
            ["基础验证", "GPIO、UART、I2C、SPI、ADC、PWM、定时器、中断、OpenHarmony工程创建", "验证硬件驱动、接口调用和基础调试能力。"],
            ["模块组合", "温湿度采集、光照采集、继电器控制、ZigBee组网、WiFi联网", "完成多模块协同和数据上报。"],
            ["综合项目", "网关协议转换、MQTT主题设计、规则引擎、平台大屏、本地触控看板", "形成端、边、云、控闭环项目。"],
            ["行业创新", "发酵监测、智慧养殖、智慧农业、资产识别、数字孪生、AI问答", "支撑跨专业项目、毕业设计和后续竞赛训练。"],
        ],
        [3.0, 7.0, 5.4],
    )

    heading(doc, "三、鸿蒙物联网综合实训室", 1)
    heading(doc, "3.1 总体规划", 2)
    paragraph(doc, "鸿蒙物联网综合实训室建议以120㎡空间为载体，设置教师演示区、25组学生实训工位、网络与服务器柜、设备收纳区和文化展示区。空间组织兼顾课堂讲授、分组实操、设备维护、成果展示与项目答辩。")
    add_image(doc, diagrams["space"], "图 3-1 120㎡空间布局示意图")
    if (ORIG / "02.png").exists():
        add_image(doc, ORIG / "02.png", "图 3-2 原始鸿蒙方案空间与设备部署示意图（复用本地素材）")
    paragraph(doc, "每组工位建议按2名学生组织，配置电脑终端、鸿蒙物联网实验箱、基础线材和低压安全电源。扩展开发包可按课程模块配置到工位，也可采用共享轮换方式提升预算利用率。教师端可集中展示代码运行、网关日志、平台数据大屏和学生任务进度。")

    heading(doc, "3.2 基本思路", 2)
    paragraph(doc, "第一，强调真实场景再现。通过智能家居、环境监测、智慧农业、智慧养殖和微生物发酵监测等场景，让学生理解传感器采集、边缘判断、平台展示和控制执行在实际项目中的作用。")
    paragraph(doc, "第二，强调虚实结合。真实设备承担动手操作和工程联调，在线仿真承担课前预习、课后练习和混合考核，两类终端统一接入平台，便于教师查看过程数据和实验结果。")
    paragraph(doc, "第三，强调模块化和可扩展。实验箱采用统一快接接口、统一供电和通信触点，支持主控、通信、传感、RFID、执行器等模块按课程需要组合，降低课堂准备难度，提高设备复用率。")
    paragraph(doc, "第四，强调数据沉淀和评价闭环。平台记录学生设备连接、代码运行、仿真日志、数据上报、报告提交和教师批注，为过程性评价、项目复盘和课程资源迭代提供依据。")
    add_image(doc, diagrams["architecture"], "图 3-3 端-边-云总体架构图")
    if (ORIG / "01.png").exists():
        add_image(doc, ORIG / "01.png", "图 3-4 原始鸿蒙方案总体架构图（复用本地素材）")

    heading(doc, "3.3 基础设施构建", 2)
    paragraph(doc, "基础设施建设应满足安全、稳定、易维护和可展示四项要求。教室应完成强弱电整理、网络接入、教师演示设备、服务器与交换设备、设备收纳、文化展示和必要的装修改造。")
    add_image(doc, diagrams["network"], "图 3-5 实训室网络拓扑图")
    add_table(
        doc,
        "表 3-1 基础设施建设内容",
        ["建设项", "建设内容", "建设要求"],
        [
            ["教师演示区", "讲台、教师终端、大屏或投影、实物展示、平台管理入口", "满足教师发布任务、展示数据大屏、演示代码运行和组织答辩。"],
            ["学生工位区", "25组双人工位、电脑端、实验箱、线材、低压电源、网口/无线覆盖", "支持50人同步实训，布线清晰，设备取放安全。"],
            ["网络服务器区", "千兆交换机、路由器/防火墙、平台服务器或边缘计算设备、备份存储", "支持真实设备、仿真设备和平台稳定接入。"],
            ["收纳维护区", "实验箱收纳、模块分类、耗材备件、充电维护、资产编号", "建立设备借还、损坏登记和备件替换流程。"],
            ["文化展示区", "专业建设展示、鸿蒙生态介绍、优秀作品、课程路径、实训安全规范", "体现物联网专业特色和实训室育人氛围。"],
        ],
        [3.0, 6.7, 5.7],
    )

    heading(doc, "3.4 课程体系设计", 2)
    paragraph(doc, "课程体系以岗位能力为导向，按照基础开发、平台集成、综合项目和行业场景四层递进。基础开发解决单点原理和驱动能力，平台集成解决协议与设备建模能力，综合项目解决端边云协同能力，行业场景解决跨专业应用和项目交付能力。")
    add_image(doc, diagrams["course"], "图 3-6 课程体系设计图")
    add_table(
        doc,
        "表 3-2 课程资源配置建议",
        ["课程模块", "实验/项目内容", "资源形态"],
        [
            ["MCU与鸿蒙基础", "GPIO、UART、I2C、SPI、ADC、PWM、OpenHarmony工程创建、WiFi配网", "实验指导书、源码、接线图、运行截图、报告模板。"],
            ["无线通信与组网", "ZigBee协调器与节点、WiFi联网、BLE近场、4G/LoRa/NB-IoT数据上传", "通信案例、AT指令样例、平台配置截图、故障排查清单。"],
            ["传感与执行控制", "温湿度、光照、CO2、气体、重量、RFID、风扇、水泵、舵机、灯光联动", "模块说明、实验步骤、阈值策略、数据曲线样例。"],
            ["AIoT平台应用", "设备建模、MQTT主题、HTTP接口、Modbus接入、规则联动、报警记录、大屏配置", "平台账号、任务模板、评价量规、实验数据导出模板。"],
            ["行业综合项目", "智慧农业、智慧养殖、微生物发酵环境监测、智能家居、资产识别", "项目任务书、答辩PPT模板、验收清单、优秀作品库。"],
        ],
        [3.2, 7.5, 4.7],
    )

    heading(doc, "3.5 核心实验设备与平台", 2)
    paragraph(doc, "核心设备以鸿蒙物联网实验箱为学生端实训终端，配套主控、通信、传感器、RFID识别、执行器开发包和AIoT教学云平台。设备体系覆盖嵌入式控制、鸿蒙开发、ZigBee无线传感网、中心网关多协议接入、环境采集、光照采集、开关量输出和本地触控可视化等实验。")
    add_image(doc, diagrams["equipment"], "图 3-7 核心设备组成图")
    add_table(
        doc,
        "表 3-3 核心设备配置表",
        ["类别", "主要配置", "教学用途"],
        [
            ["鸿蒙物联网实验箱", "通用实验平台本体、9个平台标配模块、电源适配器、USB连接线、配套线材和标配教学资源", "作为学生核心实训终端，承载鸿蒙、STM32、ZigBee、网关、传感与执行闭环实验。"],
            ["主控开发包", "STM32扩展控制模块、鸿蒙扩展模块、ZigBee单片机模块、接口拓展板", "用于多控制器协同、接口扩展、程序下载、串口调试和创新项目开发。"],
            ["通信开发包", "WiFi、ZigBee、4G、BLE、LoRa、NB-IoT/Cat.1等通信模块", "支撑局域网、广域网、低功耗和远距采集等多制式通信训练。"],
            ["传感器开发包", "霍尔、人体红外、气体、红外对射、重量、火焰、PM2.5、土壤湿度、超声波、CO2等模块", "覆盖安防、环境、农业、仓储、气体和距离等典型感知对象。"],
            ["RFID识别开发包", "HF高频、UHF超高频、NFC、条码/二维码识别模块", "面向门禁、考勤、资产盘点、仓储出入库和标签数据追踪实验。"],
            ["执行器开发包", "舵机、风扇、步进电机、LED灯光、水泵、喷雾器、窗帘执行器等", "支撑智能家居、灌溉、通风、告警和机械动作联动实验。"],
            ["AIoT教学云平台", "B/S架构、设备接入、数据可视化、规则联动、硬件在线仿真、教学评价和AI辅助", "作为教学组织中心、设备管理中心和数据沉淀中心。"],
        ],
        [3.1, 6.2, 6.1],
    )
    if (ORIG / "03.png").exists():
        add_image(doc, ORIG / "03.png", "图 3-8 采集、通信、平台、控制教学闭环图（复用本地素材）")
    add_table(
        doc,
        "表 3-4 典型实验项目表",
        ["项目类别", "项目名称", "项目成果"],
        [
            ["基础实验", "GPIO控制、串口通信、I2C传感器读取、PWM调光、ADC采集", "完成单点驱动、代码运行截图和实验报告。"],
            ["联网实验", "WiFi配网、MQTT数据上报、HTTP接口调用、Modbus设备接入", "完成设备建模、数据上报和命令下发。"],
            ["组网实验", "ZigBee协调器配置、节点入网退网、数据透传、多节点汇聚", "完成无线传感网络拓扑和网关汇聚。"],
            ["平台实验", "规则引擎、报警记录、历史曲线、项目大屏、虚拟设备接入", "完成可视化看板和平台配置文档。"],
            ["行业项目", "发酵监测、智慧养殖、智慧农业、智能家居、资产识别", "完成需求分析、系统联调、项目展示和答辩。"],
        ],
        [3.0, 6.7, 5.7],
    )

    heading(doc, "3.6 实验室特色与优势", 2)
    for t in [
        "鸿蒙特色：引入OpenHarmony轻量设备开发，结合鸿蒙模块完成配网、设备状态上传、云端控制和场景联动。",
        "工程闭环：围绕“采集-通信-平台-控制-展示-评价”组织课程，避免碎片化实验。",
        "虚实结合：真实硬件与在线仿真统一接入平台，降低课堂准备成本，增强课前预习和课后复盘能力。",
        "行业融合：面向智慧农业、智慧养殖、微生物发酵等学院相关专业群，增强跨专业服务能力。",
        "持续迭代：预留RESTful API、MQTT、HTTP、Modbus、WebSocket等接口，支持后续二次开发和资源更新。",
    ]:
        bullet(doc, t)

    heading(doc, "四、拓展应用与前瞻发展", 1)
    heading(doc, "4.1 前瞻定位", 2)
    paragraph(doc, "考虑到当前产教融合、科研创新和竞赛训练尚处于能力预留阶段，本方案不将相关内容表述为既有业绩或已发生成果，而是定位为实训室建成后的拓展方向。通过平台接口、课程资源、项目案例和数据沉淀，后续可逐步支撑校企协同课程共建、教师教学研究、学生创新项目和行业应用样板。")
    add_image(doc, diagrams["extension"], "图 4-1 拓展应用与前瞻发展路线图")
    add_table(
        doc,
        "表 4-1 前瞻拓展方向",
        ["拓展方向", "建议开展方式", "预期价值"],
        [
            ["校企协同课程共建", "围绕鸿蒙设备接入、AIoT平台运维、行业场景实施等模块共建任务书和案例库。", "使课程内容更贴近工程链路，便于后续持续更新。"],
            ["教师培训与教研", "开展平台使用、设备联调、课程组织、项目评价和故障排查培训。", "帮助教师独立完成任务发布、过程监控、报告批阅和课程迭代。"],
            ["学生创新项目", "基于开放接口和模块化设备，支持学生围绕智慧农业、养殖、发酵等场景进行项目开发。", "形成作品库、报告库、数据样例和展示素材。"],
            ["行业场景拓展", "将环境监测、资产识别、仓储管理、智能家居等场景作为后续扩展包。", "提升实训室跨专业服务和持续建设空间。"],
            ["资源持续运营", "每学期复盘课程资源，每年更新通信模块、AI能力或行业案例。", "避免设备一次性建设后资源固化。"],
        ],
        [3.4, 7.0, 5.0],
    )
    heading(doc, "4.2 行业场景应用建议", 2)
    paragraph(doc, "行业场景不作为既有项目成果呈现，而作为实训室建成后可开展的教学应用方向。建议优先围绕学院专业群特征选择微生物发酵环境监测、智慧养殖监测和智慧农业管理等方向。")
    if (ORIG / "04.png").exists():
        add_image(doc, ORIG / "04.png", "图 4-2 发酵监测与智慧养殖场景拓展示意图（复用本地素材）")
    add_table(
        doc,
        "表 4-2 行业场景应用建议",
        ["场景", "关键参数/对象", "教学任务"],
        [
            ["微生物发酵环境监测", "温度、湿度、CO2、重量、气体、光照、搅拌、补料", "完成多参数采集、PID/阈值控制、网关上报、趋势分析和异常告警。"],
            ["智慧养殖监测", "环境温湿度、氨气/CO2、设备状态、猪舍档案、数据看板", "完成环境监测、生产流程看板、数字孪生大屏和AI问答接口探索。"],
            ["智慧农业", "土壤湿度、光照、温湿度、水泵、喷雾、窗帘/遮阳", "完成自动灌溉、阈值联动、远距通信和农业数据可视化。"],
            ["智能家居", "人体红外、门磁、灯光、窗帘、风扇、安防告警", "完成情景模式、规则联动、远程控制和本地看板。"],
            ["资产识别与仓储", "HF/UHF/NFC、二维码、出入库记录、标签追踪", "完成资产盘点、权限识别、仓储流程模拟和数据追踪。"],
        ],
        [3.0, 5.4, 7.0],
    )

    heading(doc, "五、竞赛训练与能力提升展望", 1)
    heading(doc, "5.1 能力预留原则", 2)
    paragraph(doc, "实训室建成后可面向物联网应用开发、嵌入式系统、人工智能应用、鸿蒙生态创新等方向预留训练条件。本方案不写现有参赛案例、相关成绩或合作院校案例，仅从课程和设备能力角度说明后续可支撑的训练方向。")
    add_image(doc, diagrams["competition"], "图 5-1 竞赛训练能力映射图")
    add_table(
        doc,
        "表 5-1 竞赛训练能力映射表",
        ["能力方向", "可训练内容", "实训室支撑条件"],
        [
            ["物联网应用开发", "设备接入、协议通信、规则联动、数据看板、项目交付", "鸿蒙实验箱、通信开发包、AIoT平台、行业场景项目。"],
            ["嵌入式系统", "MCU外设、传感器驱动、串口调试、执行器控制、故障排查", "STM32模块、传感器开发包、执行器开发包、实验指导书。"],
            ["人工智能应用", "AI助教、语音/视觉接口、边缘AI探索、知识库答疑", "AIoT平台高级能力、开放接口、项目化课程资源。"],
            ["鸿蒙生态创新", "OpenHarmony轻量设备开发、WiFi配网、设备上云、鸿蒙节点联动", "鸿蒙物联网模块、DevEco开发流程、云平台联动案例。"],
            ["工程表达与答辩", "需求分析、方案设计、联调记录、演示汇报、报告撰写", "任务书、报告模板、过程数据、项目大屏和优秀作品库。"],
        ],
        [3.0, 6.2, 6.2],
    )
    heading(doc, "5.2 能力提升路径", 2)
    paragraph(doc, "建议在实训室投入使用后，按照“课程训练—综合项目—校内展示—专项训练—外部赛项对接”的路径逐步推进。初期以课程达标和教师熟练使用为主，中期沉淀优秀作品和训练题库，后期再根据学校专业建设安排选择适配赛项进行专项训练。")

    heading(doc, "六、建设清单与实施保障", 1)
    heading(doc, "6.1 建设清单", 2)
    add_table(
        doc,
        "表 6-1 建设内容清单",
        ["序号", "建设内容", "建议配置", "用途说明"],
        [
            ["1", "鸿蒙物联网实验箱", "按25组学生工位配置，配套电源、线材和基础教学资源", "学生核心实训终端，支撑鸿蒙、STM32、ZigBee、网关、传感与执行实验。"],
            ["2", "主控/通信/传感/RFID/执行器开发包", "按课程需要配置到工位或共享轮换使用", "扩展实验边界，支撑多协议通信、资产识别、环境检测和执行联动。"],
            ["3", "AIoT教学云平台", "1套，支持管理员、教师、学生角色", "承担设备接入、任务发布、数据可视化、规则联动、在线仿真和评价。"],
            ["4", "课程教学包", "基础实验、综合项目、行业项目、源码、指导书、报告模板", "支撑课堂教学、课后练习、项目答辩和持续资源建设。"],
            ["5", "行业场景系统", "智慧农业、智慧养殖、微生物发酵环境监测等按预算深化", "增强跨专业应用能力和项目化教学吸引力。"],
            ["6", "网络与基础设施", "交换机、路由器、服务器/边缘计算、收纳柜、教师演示设备", "保障平台稳定运行、设备接入和课堂展示。"],
            ["7", "文化建设与装修改造", "专业展示、制度上墙、成果展示、布线整理、安全标识", "营造物联网专业氛围并保障实训安全。"],
        ],
        [1.3, 3.6, 5.2, 6.2],
    )
    heading(doc, "6.2 预算控制建议", 2)
    paragraph(doc, "以下预算为围绕86万元总额的方案测算口径，便于论证阶段把握建设重心。正式采购可根据学校最终清单、品牌型号、服务范围和现场深化结果调整。")
    add_table(
        doc,
        "表 6-2 预算分配建议表",
        ["类别", "建设内容", "预算建议（万元）", "说明"],
        [
            ["学生端实训设备", "25组鸿蒙物联网实验箱及基础配套", "43", "优先保障50人同步实训的核心条件。"],
            ["扩展开发包", "主控、通信、传感、RFID、执行器扩展模块", "15", "可按课程模块共享轮换，提高预算利用率。"],
            ["平台与仿真", "AIoT教学云平台、设备接入、在线仿真、数据可视化", "9", "支撑教学管理、过程评价和数据沉淀。"],
            ["行业场景系统", "智慧农业、智慧养殖、微生物发酵监测等场景", "6", "突出学院专业群特色，形成项目化教学样板。"],
            ["网络与基础设施", "服务器/边缘设备、交换机、路由、收纳、教师演示配套", "5", "保障设备稳定接入和课堂展示。"],
            ["课程资源与培训", "实验指导书、源码、模板、教师培训、耗材备件", "4", "保障建成后可教学、可运维、可迭代。"],
            ["文化建设与装修改造", "文化墙、安全标识、布线整理、局部环境优化", "4", "满足公告中实训室文化建设及部分装修改造要求。"],
            ["合计", "约86万元预算控制", "86", "最终以深化清单和正式报价为准。"],
        ],
        [3.0, 5.5, 2.4, 5.2],
    )
    heading(doc, "6.3 实施进度", 2)
    add_image(doc, diagrams["implementation"], "图 6-1 实施进度流程图")
    add_table(
        doc,
        "表 6-3 实施阶段与交付物",
        ["阶段", "工作项", "主要内容", "交付物"],
        [
            ["第1阶段", "现场确认与深化设计", "确认教室尺寸、电力网络、工位数量、服务器部署方式和课程优先级。", "深化设计与施工准备清单。"],
            ["第2阶段", "设备到货与环境部署", "完成实验箱、开发包、平台服务器、网络设备、演示设备和收纳设施到位。", "设备清点、资产编号、网络与供电可用。"],
            ["第3阶段", "平台安装与联调", "部署AIoT平台，接入真实设备和虚拟设备，配置账号、课程、项目和数据大屏。", "平台可登录、设备可接入、数据可采集。"],
            ["第4阶段", "课程资源导入与教师培训", "导入实验指导书、源码、PPT、任务模板和报告模板，完成教师实操培训。", "教师可独立发布任务、查看过程、批阅报告。"],
            ["第5阶段", "试运行与验收", "按基础实验、综合项目、行业场景开展试运行，完成问题整改和验收归档。", "验收报告、运维制度和交付资料。"],
        ],
        [2.0, 3.0, 6.2, 4.7],
    )
    heading(doc, "6.4 培训与运维保障", 2)
    for t in [
        "教师培训：覆盖平台登录、课程发布、设备接入、实验数据查看、报告批阅、常见故障排查和课程资源维护。",
        "学生培训：覆盖设备安全使用、模块取放、接线规范、账号登录、报告提交和数据引用规范。",
        "设备运维：建立资产编号、借还记录、损坏登记、耗材补充、备件替换和定期巡检机制。",
        "平台运维：定期备份数据库、实验报告、项目源码、平台配置和账号权限，确保教学数据安全。",
        "安全管理：执行低压安全用电、强弱电分离、网络访问控制、账号分级权限和课堂设备归位检查。",
    ]:
        bullet(doc, t)
    add_table(
        doc,
        "表 6-4 验收指标表",
        ["验收类别", "验收指标", "建议验收方式"],
        [
            ["容量验收", "25组工位可支撑50人同步开展物联网实训。", "现场抽查工位、设备和网络接入情况。"],
            ["设备验收", "实验箱、开发包、平台服务器、网络设备、场景系统和耗材与清单一致。", "清点资产编号、规格参数和配套资料。"],
            ["平台验收", "教师、学生、管理员角色可用，真实设备和虚拟设备可接入，数据可展示。", "登录平台并完成设备建模、数据上报和命令下发。"],
            ["课程验收", "基础实验、综合项目、行业项目、源码、指导书、报告模板可使用。", "抽取样例实验完成全流程演示。"],
            ["图文资料验收", "深化方案、设备清单、网络拓扑、空间布局、培训记录、运维制度完整。", "查验交付文档和培训签到/记录。"],
            ["安全验收", "强弱电布置、安全标识、设备收纳、低压供电和账号权限满足实训管理要求。", "现场检查和试运行记录确认。"],
        ],
        [3.0, 7.0, 5.4],
    )

    heading(doc, "七、资料依据", 1)
    paragraph(doc, "本方案综合使用本地项目资料和公开资料进行论证。公开资料仅用于政策、专业建设和技术趋势说明，不作为供应商业绩、合作案例或奖项证明。")
    add_table(
        doc,
        "表 7-1 主要资料来源",
        ["资料类别", "来源", "用途"],
        [
            ["本地公告", "《信息工程学院.txt》", "确定项目名称、预算约86万元、面积约120㎡、50人容量和建设内容。"],
            ["结构底稿", "《传感器与检测技术创新实验室解决方案(1)(1).docx》", "借鉴项目概述、总体设计、建设内容、拓展应用、建设清单等章节结构。"],
            ["设备参数", "《新南云鸿蒙物联网实验箱技术参数-7.2.xlsx》", "提炼鸿蒙实验箱、开发包、AIoT平台、课程包和行业系统配置。"],
            ["鸿蒙底稿", "《鸿蒙物联网实训室建设方案-图片全优化版(1).docx》", "复用总体架构、空间部署、教学闭环和行业场景等本地图片素材。"],
            ["教育政策", "教育部现代职业教育体系建设改革、职业教育数字化相关公开资料", "支撑建设必要性和专业升级方向。"],
            ["专业标准", "物联网应用技术专业教学标准及职业院校技能训练公开资料", "支撑课程体系、岗位能力和竞赛训练能力映射。"],
            ["技术资料", "OpenHarmony官方文档、HarmonyOS设备开发官网", "支撑OpenHarmony轻量设备开发、设备接入和鸿蒙物联网技术路线。"],
        ],
        [3.0, 6.7, 5.7],
    )
    paragraph(doc, "参考网址：教育部官网 https://www.moe.gov.cn/；OpenHarmony文档 https://docs.openharmony.cn/；HarmonyOS设备开发 https://device.harmonyos.com/；全国职业院校技能大赛公开资料 https://www.chinaskills-jsw.org/。", first_line=False)

    return doc


def main():
    diagrams = generate_diagrams()
    doc = create_doc(diagrams)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
