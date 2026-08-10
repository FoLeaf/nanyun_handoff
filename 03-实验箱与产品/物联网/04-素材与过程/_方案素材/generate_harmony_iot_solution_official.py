from pathlib import Path
import subprocess
import textwrap
from urllib.parse import quote

from PIL import Image, ImageEnhance, ImageOps
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parent.parent
ASSET = ROOT / "_方案素材"
CHARTS = ASSET / "frontend_charts"
CHARTS.mkdir(parents=True, exist_ok=True)
PROCESSED_IMAGES = ASSET / "processed_images"
PROCESSED_IMAGES.mkdir(parents=True, exist_ok=True)
OUT = ROOT / "鸿蒙物联网综合实训室建设方案-公文优化版-产品实景增强版.docx"
PDF = ROOT / "鸿蒙物联网综合实训室建设方案-公文优化版-产品实景增强版.pdf"
LOGO = Path("D:/nanyun/lgo/常规.png")
LOGO_DOC = CHARTS / "logo_doc.png"
USER_IMAGE_SOURCES = {
    "overview": Path(r"C:/Users/19y/Downloads/ChatGPT Image 2026年7月3日 15_06_50.png"),
    "space": Path(r"C:/Users/19y/Downloads/ChatGPT Image 2026年7月3日 15_11_34.png"),
    "product": Path(r"C:/Users/19y/Downloads/ChatGPT Image 2026年7月3日 14_47_33.png"),
    "scenario": Path(r"C:/Users/19y/Downloads/ChatGPT Image 2026年7月3日 14_57_12.png"),
}
EDGE = Path("C:/Program Files (x86)/Microsoft/Edge/Application/msedge.exe")
CHROME = Path("C:/Program Files/Google/Chrome/Application/chrome.exe")
BROWSER = EDGE if EDGE.exists() else CHROME

BLACK = "000000"
MUTED = "666666"
LINE = "BFBFBF"
LIGHT = "F7F7F7"
BLUE = "075FBF"
ORANGE = "D98200"


def ensure_logo():
    if not LOGO.exists():
        return None
    im = Image.open(LOGO).convert("RGBA")
    im.thumbnail((1400, 680))
    bg = Image.new("RGBA", im.size, "WHITE")
    bg.alpha_composite(im)
    bg.convert("RGB").save(LOGO_DOC, quality=95)
    return LOGO_DOC


def prepare_user_images():
    processed = {}
    for key, src in USER_IMAGE_SOURCES.items():
        if not src.exists():
            continue
        im = Image.open(src).convert("RGB")
        im = ImageOps.exif_transpose(im)
        # Keep the provided product/effect images truthful; only normalize brightness,
        # contrast, and file weight for stable DOCX rendering.
        im = ImageEnhance.Brightness(im).enhance(1.03)
        im = ImageEnhance.Contrast(im).enhance(1.04)
        max_width = 1800 if key != "product" else 1500
        if im.width > max_width:
            ratio = max_width / im.width
            im = im.resize((max_width, int(im.height * ratio)), Image.Resampling.LANCZOS)
        im = ImageOps.expand(im, border=8, fill="#d9d9d9")
        out = PROCESSED_IMAGES / f"user_{key}.jpg"
        im.save(out, quality=88, optimize=True)
        processed[key] = out
    return processed


def html_shell(title, subtitle, body):
    return f"""<!doctype html>
<html>
<head>
<meta charset="utf-8">
<style>
  * {{ box-sizing: border-box; }}
  body {{
    margin: 0;
    width: 1600px;
    height: 940px;
    background: #ffffff;
    color: #111111;
    font-family: "Microsoft YaHei", "SimHei", "SimSun", sans-serif;
  }}
  .canvas {{
    width: 1600px;
    height: 940px;
    padding: 54px 70px;
    background: #fff;
    border: 1px solid #e5e5e5;
  }}
  .chart-title {{
    font-family: "SimHei", "Microsoft YaHei", sans-serif;
    font-size: 34px;
    font-weight: 700;
    color: #000;
    margin: 0;
    letter-spacing: 0;
  }}
  .chart-subtitle {{
    font-size: 18px;
    color: #555;
    margin-top: 8px;
    padding-bottom: 14px;
    border-bottom: 3px solid #222;
  }}
  .accent-line {{
    height: 4px;
    width: 90px;
    background: linear-gradient(90deg, {css_color(BLUE)}, {css_color(ORANGE)});
    margin-top: 12px;
  }}
  .body {{ position: relative; height: 760px; margin-top: 34px; }}
  .box {{
    position: absolute;
    border: 2px solid #777;
    background: #fff;
    color: #111;
    padding: 18px 20px;
    display: flex;
    flex-direction: column;
    justify-content: center;
    align-items: center;
    text-align: center;
    line-height: 1.45;
  }}
  .box .k {{ font-size: 23px; font-weight: 700; color: #000; }}
  .box .v {{ font-size: 18px; color: #222; margin-top: 7px; }}
  .box.light {{ background: #f8f8f8; }}
  .num {{
    display: inline-flex;
    align-items: center;
    justify-content: center;
    width: 34px;
    height: 34px;
    border: 2px solid #111;
    border-radius: 50%;
    font-weight: 700;
    font-size: 18px;
    margin-bottom: 10px;
    background: #fff;
  }}
  .blue-rule {{ border-top: 5px solid {css_color(BLUE)}; }}
  .orange-rule {{ border-top: 5px solid {css_color(ORANGE)}; }}
  .thin {{ border-color: #aaa; }}
  svg {{ position: absolute; left: 0; top: 0; width: 100%; height: 100%; pointer-events: none; }}
  .note {{
    position: absolute;
    left: 60px;
    right: 60px;
    bottom: 22px;
    padding: 18px 24px;
    border-left: 5px solid #111;
    background: #f8f8f8;
    font-size: 20px;
    line-height: 1.5;
    color: #111;
  }}
  .mini {{
    position: absolute;
    border: 1.6px solid #999;
    background: #fff;
    padding: 10px 12px;
    text-align: center;
    font-size: 17px;
    line-height: 1.35;
  }}
  .grid-label {{ font-weight: 700; font-size: 20px; }}
  .room {{ position:absolute; border: 3px solid #111; background:#fff; }}
  .desk {{ position:absolute; border:1.5px solid #777; background:#f9f9f9; text-align:center; font-size:15px; line-height:1.2; display:flex; align-items:center; justify-content:center; }}
  .lane {{
    position: absolute;
    border: 1.8px solid #8a8a8a;
    background: #fff;
    padding: 16px 18px;
    font-size: 18px;
    line-height: 1.42;
  }}
  .lane-title {{
    font-family: "SimHei", "Microsoft YaHei", sans-serif;
    font-size: 22px;
    font-weight: 700;
    color: #000;
    padding-bottom: 8px;
    margin-bottom: 10px;
    border-bottom: 2px solid #cfcfcf;
  }}
  .chip {{
    display: inline-block;
    margin: 5px 4px;
    padding: 6px 10px;
    border: 1.4px solid #999;
    background: #fafafa;
    font-size: 16px;
    line-height: 1.2;
  }}
  .port {{
    position: absolute;
    width: 74px;
    height: 48px;
    border: 1.6px solid #777;
    background: #f9f9f9;
    display: flex;
    align-items: center;
    justify-content: center;
    font-size: 15px;
    font-weight: 700;
    color: #000;
  }}
</style>
</head>
<body>
  <div class="canvas">
    <div class="chart-title">{title}</div>
    <div class="chart-subtitle">{subtitle}</div>
    <div class="accent-line"></div>
    <div class="body">{body}</div>
  </div>
</body>
</html>"""


def css_color(hex_value):
    return "#" + hex_value


def box(x, y, w, h, k, v="", cls=""):
    return f'<div class="box {cls}" style="left:{x}px;top:{y}px;width:{w}px;height:{h}px;"><div class="k">{k}</div><div class="v">{v}</div></div>'


def arrow(x1, y1, x2, y2):
    return f'<line x1="{x1}" y1="{y1}" x2="{x2}" y2="{y2}" stroke="#333" stroke-width="3" marker-end="url(#arrow)" />'


def svg_defs(lines):
    return f'''<svg viewBox="0 0 1460 760">
      <defs><marker id="arrow" markerWidth="10" markerHeight="8" refX="9" refY="4" orient="auto"><path d="M0,0 L10,4 L0,8 Z" fill="#333"/></marker></defs>
      {lines}
    </svg>'''


def write_chart(name, title, subtitle, body):
    html = html_shell(title, subtitle, body)
    html_path = CHARTS / f"{name}.html"
    png_path = CHARTS / f"{name}.png"
    html_path.write_text(html, encoding="utf-8")
    cmd = [
        str(BROWSER),
        "--headless=new",
        "--disable-gpu",
        "--hide-scrollbars",
        "--force-device-scale-factor=2",
        "--window-size=1600,940",
        f"--screenshot={png_path}",
        html_path.resolve().as_uri(),
    ]
    subprocess.run(cmd, check=True, stdout=subprocess.PIPE, stderr=subprocess.PIPE)
    return png_path


def render_frontend_charts():
    charts = {}
    charts["policy"] = write_chart(
        "01_policy_logic",
        "政策-产业-专业建设逻辑图",
        "以职业教育数字化与鸿蒙物联网技术栈支撑实训条件升级",
        svg_defs(
            arrow(312, 228, 392, 228)
            + arrow(650, 228, 730, 228)
            + arrow(988, 228, 1068, 228)
        )
        + box(40, 110, 270, 235, "政策牵引", "职业教育数字化<br>现代职教体系<br>专业教学标准", "blue-rule")
        + box(390, 110, 270, 235, "产业需求", "AIoT终端接入<br>边缘网关集成<br>数据可视化运维", "orange-rule")
        + box(730, 110, 270, 235, "技术底座", "OpenHarmony<br>STM32 / ZigBee<br>MQTT / Modbus", "blue-rule")
        + box(1070, 110, 300, 235, "专业建设", "课程实训平台<br>云平台<br>50人项目化教学", "orange-rule")
        + '<div class="note">建设结果：形成“硬件可操作、软件可管理、数据可沉淀、项目可评价、场景可迁移”的鸿蒙物联网综合实训环境。</div>',
    )

    charts["teaching"] = write_chart(
        "02_teaching_loop",
        "教学做一体化闭环图",
        "课堂讲授、设备操作、平台评价和项目复盘贯通",
        svg_defs(
            arrow(730, 180, 1010, 320)
            + arrow(1080, 420, 820, 610)
            + arrow(640, 610, 380, 420)
            + arrow(450, 320, 730, 180)
        )
        + box(580, 70, 300, 120, "理论导入", "知识点讲解 / 任务发布", "blue-rule")
        + box(1010, 300, 300, 120, "硬件操作", "实验箱 / 模块 / 接线调试", "orange-rule")
        + box(580, 610, 300, 120, "平台联调", "数据上报 / 控制执行", "blue-rule")
        + box(150, 300, 300, 120, "项目评价", "报告 / 数据 / 过程复盘", "orange-rule")
        + box(555, 310, 350, 150, "教学组织中心", "教师发布任务<br>学生分组实训<br>数据自动沉淀", "light")
        + '<div class="note">闭环重点：每个知识点都落到“采集-通信-平台-控制-报告”的工程任务，减少孤立验证型实验。</div>',
    )

    body = ""
    xs = [50, 395, 740, 1085]
    heads = ["基础验证", "模块组合", "综合项目", "行业创新"]
    vals = [
        "GPIO / UART / I2C<br>ADC / PWM / 中断<br>OpenHarmony工程",
        "温湿度 / 光照 / CO2<br>ZigBee组网<br>WiFi与MQTT上报",
        "网关协议转换<br>规则引擎<br>平台大屏与报告",
        "智慧农业 / 养殖<br>发酵监测<br>数字孪生与AI问答",
    ]
    lines = ""
    for i, x in enumerate(xs):
        body += f'<div class="box {"blue-rule" if i%2==0 else "orange-rule"}" style="left:{x}px;top:135px;width:285px;height:330px;"><span class="num">{i+1}</span><div class="k">{heads[i]}</div><div class="v">{vals[i]}</div></div>'
        if i < 3:
            lines += arrow(x + 285, 300, xs[i + 1], 300)
    charts["ability"] = write_chart(
        "03_ability_route",
        "学生能力递进路线图",
        "由单点验证逐步进入端、边、云、用完整工程链路",
        svg_defs(lines)
        + body
        + '<div class="note">能力输出：硬件调试、协议通信、平台建模、数据分析、规则控制、项目文档与现场表达。</div>',
    )

    desks = ""
    idx = 1
    for r in range(5):
        for c in range(5):
            desks += f'<div class="desk" style="left:{150+c*180}px;top:{210+r*88}px;width:135px;height:55px;">{idx:02d}组<br>2人</div>'
            idx += 1
    charts["space"] = write_chart(
        "04_space_layout",
        "120㎡空间布局示意图",
        "25组学生工位、50人同步实训，兼顾演示、运维和设备收纳",
        '<div class="room" style="left:80px;top:60px;width:1300px;height:590px;"></div>'
        + '<div class="mini grid-label" style="left:120px;top:90px;width:1220px;height:70px;border:2px solid #111;">教师演示区：讲台 / 大屏投影 / 平台管理终端 / 实物展示台</div>'
        + desks
        + '<div class="mini grid-label" style="left:1120px;top:210px;width:190px;height:150px;border-top:5px solid #075FBF;">网络与<br>服务器柜</div>'
        + '<div class="mini grid-label" style="left:1120px;top:400px;width:190px;height:170px;border-top:5px solid #D98200;">设备收纳<br>耗材备件<br>充电维护</div>'
        + '<div class="note">布置建议：低压供电优先、强弱电分离、桌面预留USB/网口/安全插座，教师端集中展示网关日志和数据大屏。</div>',
    )

    y0 = 70
    layer_body = ""
    layers = [
        ("感知与执行层", "温湿度、光照、CO2、重量、RFID、风扇、水泵、舵机"),
        ("终端控制层", "OpenHarmony模块、STM32模块、ZigBee节点、通信开发包"),
        ("边缘网关层", "协议转换、规则判定、本地缓存、触控看板"),
        ("AIoT平台层", "设备建模、数据可视化、任务发布、过程评价、AI助教"),
        ("场景应用层", "智能家居、环境监测、智慧农业、智慧养殖、发酵监测"),
    ]
    svg_lines = ""
    for i, (k, v) in enumerate(layers):
        layer_body += box(90, y0 + i * 112, 1270, 78, k, v, "blue-rule" if i % 2 == 0 else "orange-rule")
        if i < len(layers) - 1:
            svg_lines += arrow(725, y0 + i * 112 + 82, 725, y0 + (i + 1) * 112 - 4)
    charts["architecture"] = write_chart(
        "05_architecture",
        "端-边-云总体架构图",
        "从鸿蒙终端采集到云平台展示、规则联动和教学评价",
        svg_defs(svg_lines) + layer_body,
    )

    charts["network"] = write_chart(
        "06_network",
        "实训室网络拓扑图",
        "千兆有线、无线接入、边缘网关与平台服务器统一部署",
        svg_defs(
            arrow(725, 125, 725, 190)
            + arrow(725, 280, 725, 345)
            + arrow(560, 390, 360, 390)
            + arrow(890, 390, 1090, 390)
            + arrow(650, 430, 470, 560)
            + arrow(725, 430, 725, 560)
            + arrow(800, 430, 1050, 560)
        )
        + box(610, 55, 230, 70, "互联网 / 校内网", "", "blue-rule")
        + box(590, 190, 270, 90, "防火墙 / 路由器", "", "orange-rule")
        + box(590, 345, 270, 90, "核心交换机", "", "blue-rule")
        + box(1070, 345, 320, 95, "AIoT平台服务器", "数据备份存储", "orange-rule")
        + box(80, 345, 320, 95, "教师管理终端", "投屏 / 大屏", "orange-rule")
        + box(235, 560, 270, 95, "无线AP", "", "blue-rule")
        + box(575, 560, 300, 95, "25组学生工位", "实验箱 / PC", "orange-rule")
        + box(1010, 560, 330, 95, "边缘网关", "触控看板", "blue-rule"),
    )

    charts["equipment"] = write_chart(
        "07_equipment",
        "核心设备组成图",
        "以鸿蒙物联网实验箱为中心，形成可扩展的分组实训套件",
        svg_defs(
            arrow(410, 170, 650, 340)
            + arrow(1040, 170, 810, 340)
            + arrow(410, 565, 650, 455)
            + arrow(1040, 565, 810, 455)
            + arrow(725, 650, 725, 500)
            + arrow(725, 210, 725, 340)
        )
        + box(580, 330, 290, 155, "鸿蒙物联网实验箱", "便携一体化 / 模块快接<br>9个通用接口位", "blue-rule")
        + box(90, 90, 320, 120, "主控开发包", "STM32 / 鸿蒙 / ZigBee", "orange-rule")
        + box(1040, 90, 340, 120, "通信开发包", "WiFi / 4G / BLE / LoRa / NB-IoT", "blue-rule")
        + box(90, 525, 320, 120, "传感器开发包", "环境 / 安防 / 农业 / 气体", "blue-rule")
        + box(1040, 525, 340, 120, "RFID识别开发包", "HF / UHF / NFC / 二维码", "orange-rule")
        + box(565, 640, 320, 100, "执行器开发包", "风扇 / 水泵 / 舵机 / 灯光", "blue-rule")
        + box(565, 90, 320, 100, "AIoT教学云平台", "设备 / 任务 / 评价 / 仿真", "orange-rule"),
    )

    ports = ""
    for r in range(3):
        for c in range(3):
            ports += f'<div class="port" style="left:{618+c*84}px;top:{282+r*58}px;">接口{r*3+c+1}</div>'
    charts["modular_box"] = write_chart(
        "08_modular_box",
        "模块化实验箱结构与教学组合图",
        "9个同构通用接口位支撑模块快接、层叠互联和任务化组合",
        svg_defs(
            arrow(410, 160, 595, 300)
            + arrow(1040, 160, 915, 300)
            + arrow(410, 570, 595, 425)
            + arrow(1040, 570, 915, 425)
            + arrow(725, 132, 725, 252)
            + arrow(725, 620, 725, 470)
        )
        + box(575, 250, 370, 245, "鸿蒙物联网实验箱", "", "blue-rule")
        + ports
        + '<div class="mini" style="left:585px;top:456px;width:350px;height:34px;border:0;background:#fff;font-size:15px;">9个同构通用接口位 / 模块快接 / 免跳线</div>'
        + box(75, 90, 335, 130, "主控组合", "STM32 / OpenHarmony / ZigBee单片机<br>程序下载、外设驱动、串口调试", "orange-rule")
        + box(1040, 90, 345, 130, "通信组合", "WiFi / 4G / BLE / LoRa / NB-IoT<br>MQTT、HTTP、TCP/UDP数据接入", "blue-rule")
        + box(75, 525, 335, 130, "感知识别组合", "传感器 / RFID / NFC / 二维码<br>环境采集、身份识别、资产盘点", "blue-rule")
        + box(1040, 525, 345, 130, "执行控制组合", "舵机 / 风扇 / 水泵 / 灯光 / 喷雾<br>阈值联动、远程控制、场景闭环", "orange-rule")
        + box(565, 70, 320, 80, "网关与平台组合", "协议转换 / 本地看板 / AIoT云平台", "orange-rule")
        + box(565, 630, 320, 80, "行业项目组合", "智慧农业 / 智慧养殖 / 发酵监测", "blue-rule")
        + '<div class="note">模块化价值：同一实验箱可按课程任务快速重组，减少重复接线和固定硬件绑定，让学生把注意力放在接口、协议、数据和系统联调上。</div>',
    )

    charts["course_support"] = write_chart(
        "09_course_support_matrix",
        "高职物联网专业课程支撑矩阵图",
        "课程群、设备模块和能力目标三层对应，回答每门课能做什么、训练什么",
        svg_defs(
            arrow(425, 345, 535, 345)
            + arrow(925, 345, 1035, 345)
            + '<line x1="480" y1="180" x2="480" y2="600" stroke="#d0d0d0" stroke-width="2" stroke-dasharray="8 8"/>'
            + '<line x1="980" y1="180" x2="980" y2="600" stroke="#d0d0d0" stroke-width="2" stroke-dasharray="8 8"/>'
        )
        + '<div class="lane" style="left:35px;top:75px;width:390px;height:540px;border-top:5px solid #075FBF;"><div class="lane-title">课程群</div>'
        + '<span class="chip">物联网技术基础</span><span class="chip">电工电子技术</span><span class="chip">C语言/嵌入式C</span><span class="chip">单片机应用</span><span class="chip">传感器与检测</span><span class="chip">RFID技术</span><span class="chip">无线传感网络</span><span class="chip">物联网通信</span><span class="chip">设备安装与调试</span><span class="chip">云平台应用</span><span class="chip">系统集成实践</span></div>'
        + '<div class="lane" style="left:535px;top:75px;width:390px;height:540px;border-top:5px solid #D98200;"><div class="lane-title">模块与平台</div>'
        + '<span class="chip">9接口实验箱</span><span class="chip">STM32主控</span><span class="chip">OpenHarmony模块</span><span class="chip">ZigBee节点</span><span class="chip">WiFi/4G/BLE/LoRa/NB</span><span class="chip">传感器包</span><span class="chip">RFID/NFC/二维码</span><span class="chip">执行器包</span><span class="chip">边缘网关</span><span class="chip">AIoT教学云平台</span><span class="chip">行业场景套件</span></div>'
        + '<div class="lane" style="left:1035px;top:75px;width:390px;height:540px;border-top:5px solid #075FBF;"><div class="lane-title">能力目标</div>'
        + '<span class="chip">接口识别</span><span class="chip">硬件调试</span><span class="chip">程序开发</span><span class="chip">传感检测</span><span class="chip">无线组网</span><span class="chip">协议接入</span><span class="chip">平台建模</span><span class="chip">数据可视化</span><span class="chip">规则联动</span><span class="chip">系统集成</span><span class="chip">项目交付</span></div>'
        + '<div class="note">课程组织建议：同一模块可跨多门课程复用，同一课程可由“基础验证-模块组合-平台联调-综合项目”逐级展开。</div>',
    )

    cols = ""
    for i, (k, items) in enumerate(
        [
            ("基础开发课程", "OpenHarmony工程<br>STM32外设<br>ZigBee组网<br>传感器采集"),
            ("平台集成课程", "MQTT/HTTP<br>Modbus接入<br>设备建模<br>数据可视化"),
            ("综合项目课程", "规则联动<br>本地看板<br>虚实仿真<br>项目报告"),
            ("行业场景课程", "智慧农业<br>智慧养殖<br>发酵监测<br>数字孪生"),
        ]
    ):
        cols += f'<div class="box {"blue-rule" if i%2==0 else "orange-rule"}" style="left:{60+i*350}px;top:100px;width:300px;height:480px;"><span class="num">{i+1}</span><div class="k">{k}</div><div class="v">{items}</div></div>'
    charts["course"] = write_chart(
        "10_course_system",
        "课程体系设计图",
        "基础课程、平台课程、综合项目、行业场景分层支撑",
        cols + '<div class="note">课程资源建议配套：实验指导书、源码、接线图、平台配置截图、报告模板、评分量规、教师培训材料。</div>',
    )

    body = ""
    xs = [65, 400, 735, 1070]
    for i, (k, v) in enumerate(
        [
            ("近期：建成即用", "教师培训<br>课程导入<br>设备联调<br>试运行验收"),
            ("中期：资源共建", "项目库迭代<br>教材素材沉淀<br>教师课题孵化"),
            ("远期：生态拓展", "校企协同课程<br>开放接口二开<br>行业应用样板"),
            ("持续：质量提升", "竞赛训练预留<br>优秀作品库<br>数据化评价"),
        ]
    ):
        body += f'<div class="box {"blue-rule" if i%2==0 else "orange-rule"}" style="left:{xs[i]}px;top:150px;width:285px;height:315px;"><span class="num">{i+1}</span><div class="k">{k}</div><div class="v">{v}</div></div>'
    charts["extension"] = write_chart(
        "11_extension_route",
        "拓展应用与前瞻发展路线图",
        "以能力预留方式支撑后续校企协同、课程共建和创新项目",
        svg_defs(arrow(350, 308, 400, 308) + arrow(685, 308, 735, 308) + arrow(1020, 308, 1070, 308))
        + body
        + '<div class="note">说明：本章采用“可支撑、可拓展、建议后续开展”的前瞻表述，不写既有产教融合业绩或赛事成果。</div>',
    )

    charts["competition"] = write_chart(
        "12_competition_map",
        "竞赛训练能力映射图",
        "面向物联网、嵌入式、人工智能与鸿蒙生态方向预留训练条件",
        svg_defs(
            arrow(400, 200, 640, 355)
            + arrow(1050, 200, 820, 355)
            + arrow(400, 590, 640, 475)
            + arrow(1050, 590, 820, 475)
        )
        + box(580, 345, 300, 150, "竞赛训练能力预留", "鸿蒙物联网实训室", "blue-rule")
        + box(70, 120, 330, 130, "物联网应用开发", "设备接入 / 协议 / 平台", "orange-rule")
        + box(1050, 120, 330, 130, "嵌入式系统", "MCU外设 / 驱动 / 调试", "blue-rule")
        + box(70, 550, 330, 130, "人工智能应用", "AI助教 / 视觉 / 语音接口", "blue-rule")
        + box(1050, 550, 330, 130, "鸿蒙生态创新", "OpenHarmony轻量设备", "orange-rule"),
    )

    steps = ""
    for i, (k, v) in enumerate(
        [
            ("现场确认", "尺寸 / 电力 / 网络<br>工位与清单深化"),
            ("到货部署", "设备清点<br>资产编号<br>环境布置"),
            ("平台联调", "服务器部署<br>设备接入<br>数据采集"),
            ("资源导入", "课程包<br>账号权限<br>教师培训"),
            ("试运行验收", "基础实验<br>综合项目<br>文档归档"),
        ]
    ):
        steps += f'<div class="box {"blue-rule" if i%2==0 else "orange-rule"}" style="left:{30+i*285}px;top:150px;width:235px;height:300px;"><span class="num">{i+1}</span><div class="k">{k}</div><div class="v">{v}</div></div>'
    charts["implementation"] = write_chart(
        "13_implementation",
        "实施进度流程图",
        "从现场确认到验收归档的五阶段实施路径",
        svg_defs(arrow(265, 300, 315, 300) + arrow(550, 300, 600, 300) + arrow(835, 300, 885, 300) + arrow(1120, 300, 1170, 300))
        + steps
        + '<div class="note">交付资料：深化设计、设备清单、平台账号、课程资源、培训记录、试运行记录、验收报告、运维制度。</div>',
    )
    return charts


def set_run_font(run, name="仿宋", size=12, bold=False, color=BLACK):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def set_para_spacing(p, before=0, after=6, line=1.35, first=True):
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = line
    if first:
        p.paragraph_format.first_line_indent = Pt(24)


def add_para(doc, text, first=True, align=None, size=12, bold=False, font="仿宋", after=6):
    p = doc.add_paragraph()
    set_para_spacing(p, after=after, first=first)
    if align is not None:
        p.alignment = align
    r = p.add_run(text)
    set_run_font(r, font, size, bold, BLACK)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.style = f"Heading {level}" if level <= 3 else "Normal"
    p.paragraph_format.space_before = Pt(14 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(8 if level == 1 else 5)
    p.paragraph_format.line_spacing = 1.25
    r = p.add_run(text)
    set_run_font(r, "黑体", 16 if level == 1 else (14 if level == 2 else 12), True, BLACK)
    return p


def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def set_cell_margins(cell, top=100, bottom=100, start=140, end=140):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for m, v in [("top", top), ("bottom", bottom), ("start", start), ("end", end)]:
        node = tcMar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tcMar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_cell_text(cell, text, header=False, center=False, font_size=10):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if (header or center) else WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.line_spacing = 1.25
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(str(text))
    set_run_font(r, "黑体" if header else "仿宋", 10.5 if header else font_size, header, BLACK)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)


def set_table_width(table, widths_cm):
    table.autofit = False
    for row in table.rows:
        for idx, width in enumerate(widths_cm):
            row.cells[idx].width = Cm(width)


def set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    trPr.append(tbl_header)


def set_row_cant_split(row):
    trPr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    trPr.append(cant_split)


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(text)
    set_run_font(r, "仿宋", 10, False, MUTED)


def add_image(doc, path, caption, width=15.2):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(str(path), width=Cm(width))
    add_caption(doc, caption)


def add_table(doc, caption, headers, rows, widths, font_size=10):
    add_caption(doc, caption)
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    set_repeat_table_header(table.rows[0])
    set_row_cant_split(table.rows[0])
    for i, h in enumerate(headers):
        set_cell_shading(hdr[i], LIGHT)
        set_cell_text(hdr[i], h, header=True, center=True)
    for row in rows:
        data_row = table.add_row()
        set_row_cant_split(data_row)
        cells = data_row.cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], value, center=(i == 0 or len(str(value)) <= 6), font_size=font_size)
    set_table_width(table, widths)
    doc.add_paragraph()
    return table


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = paragraph.add_run("第 ")
    set_run_font(r, "仿宋", 10, False, BLACK)
    run = paragraph.add_run()
    fld1 = OxmlElement("w:fldChar")
    fld1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld2 = OxmlElement("w:fldChar")
    fld2.set(qn("w:fldCharType"), "end")
    run._r.append(fld1)
    run._r.append(instr)
    run._r.append(fld2)
    r2 = paragraph.add_run(" 页")
    set_run_font(r2, "仿宋", 10, False, BLACK)


def build_doc(charts):
    logo = ensure_logo()
    user_images = prepare_user_images()
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Cm(21)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(3.0)
    sec.bottom_margin = Cm(2.5)
    sec.left_margin = Cm(2.7)
    sec.right_margin = Cm(2.7)
    sec.header_distance = Cm(1.3)
    sec.footer_distance = Cm(1.5)
    for style_name in ["Normal", "Heading 1", "Heading 2", "Heading 3"]:
        st = doc.styles[style_name]
        st.font.name = "仿宋" if style_name == "Normal" else "黑体"
        st._element.rPr.rFonts.set(qn("w:eastAsia"), st.font.name)
        st.font.color.rgb = RGBColor.from_string(BLACK)
    header = sec.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hr = header.add_run("鸿蒙物联网综合实训室建设方案")
    set_run_font(hr, "仿宋", 9, False, MUTED)
    add_page_number(sec.footer.paragraphs[0])

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(95)
    p.paragraph_format.space_after = Pt(18)
    r = p.add_run("鸿蒙物联网综合实训室\n建设方案")
    set_run_font(r, "华文中宋", 28, True, BLACK)
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_after = Pt(32)
    r = p2.add_run("面向 AIoT 项目化教学、虚实结合实训与行业场景应用")
    set_run_font(r, "仿宋", 15, False, BLACK)
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.paragraph_format.space_before = Pt(12)
    meta.paragraph_format.space_after = Pt(80)
    r = meta.add_run("项目名称：信息工程系物联网综合实训室建设项目\n建设面积：约 120 ㎡    同步容量：50 人（25 组）\n方案定位：建设解决方案    版本说明：产品实景增强版")
    set_run_font(r, "仿宋", 13, False, BLACK)
    if logo:
        lp = doc.add_paragraph()
        lp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        lp.add_run().add_picture(str(logo), width=Cm(9.5))
    doc.add_page_break()

    add_heading(doc, "目录", 1)
    for item in ["一、项目概述", "二、教学实验总体设计", "三、鸿蒙物联网综合实训室建设内容", "四、拓展应用与场景延伸", "五、竞赛训练与能力提升展望", "六、实施计划、培训运维与验收", "七、资料依据"]:
        add_para(doc, item, first=False, size=12.5, after=4)
    doc.add_page_break()

    add_heading(doc, "一、项目概述", 1)
    add_heading(doc, "1.1 建设背景", 2)
    add_para(doc, "信息工程系物联网综合实训室建设项目以职业教育数字化转型、现代职业教育体系建设和物联网应用技术专业人才培养为牵引，面向课程教学、实验实训、技能训练和创新项目研发等多类任务，建设一间可支撑50人同步实训的综合性实训空间。")
    add_para(doc, "从专业建设角度看，物联网应用技术已从单一传感器采集、单片机验证，逐步发展为端侧设备开发、边缘网关集成、云端平台管理、数据分析展示和行业场景应用的完整工程链路。实训室建设应形成设备、平台、课程、项目、评价、运维一体化条件。")
    add_image(doc, charts["policy"], "图 1-1 政策、产业与专业建设逻辑图")
    add_table(
        doc,
        "表 1-1 项目约束与建设定位",
        ["项目维度", "建设口径", "方案响应"],
        [
            ["建设面积", "规划建设面积约120㎡", "按教师演示区、25组学生工位、网络服务器区、设备收纳区、文化展示区组织空间。"],
            ["同步容量", "可容纳50人同时练习和教学实训", "采用25组分组工位，建议2人一组，每组可开展设备接入、数据采集、平台联调和项目报告。"],
            ["建设内容", "课程教学实训平台、物联网云平台、物联网应用实训设备套件、课程资源、应用套件和耗材", "以鸿蒙物联网实验箱和AIoT教学云平台为核心，配套主控、通信、传感、RFID、执行器及行业场景。"],
            ["建设风格", "完整、详尽、可行的实施方案", "图文并茂呈现总体架构、空间布局、网络拓扑、设备组成、课程体系、实施进度和验收指标。"],
        ],
        [3.0, 5.1, 7.5],
    )
    add_heading(doc, "1.2 专业背景与人才需求", 2)
    add_para(doc, "物联网相关岗位对学生的要求正在从会接线、会烧录、会读数，转向能接入、会建模、能联调、懂数据、会交付。建设鸿蒙物联网综合实训室，可以面向物联网设备开发、系统集成与运维、AIoT平台应用、行业场景项目实施等能力开展递进式训练。")
    add_table(
        doc,
        "表 1-2 鸿蒙物联网实训室知识体系",
        ["知识模块", "核心内容", "对应能力"],
        [
            ["基础理论", "传感器原理、嵌入式系统、计算机网络、数据库基础、信息安全基础", "理解设备采集、边缘处理和数据传输的基础逻辑。"],
            ["端侧开发", "OpenHarmony工程创建、GPIO/UART/I2C/SPI/ADC/PWM、STM32外设、ZigBee节点", "完成终端驱动、数据采集、设备控制和本地调试。"],
            ["通信协议", "WiFi、ZigBee、BLE、LoRa、4G、NB-IoT/Cat.1、MQTT、HTTP、Modbus", "完成多协议接入、网关汇聚和云端数据上报。"],
            ["平台应用", "设备建模、命令下发、规则联动、报警记录、数据大屏、过程评价", "完成平台侧配置、课堂任务发布、数据沉淀和实验报告。"],
            ["行业项目", "智慧农业、智慧养殖、微生物发酵、环境监测、智能家居、资产识别", "形成跨专业项目开发、展示答辩和持续迭代能力。"],
        ],
        [3.0, 6.4, 6.2],
    )
    add_heading(doc, "1.3 建设效果总览", 2)
    add_para(doc, "从建设效果看，实训室应呈现“实训空间、设备套件、教学平台、课程资源、行业场景”一体化形态。教师端可集中展示物联网系统架构、平台数据和课堂任务，学生端以双人工位开展模块化实验箱操作、代码开发、设备联网和平台联调。")
    if "overview" in user_images:
        add_image(doc, user_images["overview"], "图 1-2 鸿蒙物联网综合实训室整体效果图")

    add_heading(doc, "二、教学实验总体设计", 1)
    add_heading(doc, "2.1 建设目标", 2)
    add_para(doc, "本项目建设目标是形成一套面向真实工程链路的鸿蒙物联网综合实训环境。实训室建成后，应能够支撑课堂演示、学生分组训练、虚拟仿真预习、真实设备联调、平台过程评价、行业场景项目和后续能力拓展。")
    add_image(doc, charts["teaching"], "图 2-1 教学做一体化闭环图")
    add_image(doc, charts["ability"], "图 2-2 学生能力递进路线图")
    add_table(
        doc,
        "表 2-1 实训体系规划",
        ["层级", "典型内容", "能力目标"],
        [
            ["基础验证", "GPIO、UART、I2C、SPI、ADC、PWM、定时器、中断、OpenHarmony工程创建", "验证硬件驱动、接口调用和基础调试能力。"],
            ["模块组合", "温湿度采集、光照采集、继电器控制、ZigBee组网、WiFi联网", "完成多模块协同和数据上报。"],
            ["综合项目", "网关协议转换、MQTT主题设计、规则引擎、平台大屏、本地触控看板", "形成端、边、云、控闭环项目。"],
            ["行业创新", "发酵监测、智慧养殖、智慧农业、资产识别、数字孪生、AI问答", "支撑跨专业项目、毕业设计和后续训练。"],
        ],
        [2.7, 6.9, 6.0],
    )

    add_heading(doc, "三、鸿蒙物联网综合实训室建设内容", 1)
    add_heading(doc, "3.1 实训空间与教学场景", 2)
    add_para(doc, "鸿蒙物联网综合实训室建议以120㎡空间为载体，设置教师演示区、25组学生实训工位、网络与服务器柜、设备收纳区和文化展示区。空间组织兼顾课堂讲授、分组实操、设备维护、成果展示与项目答辩。")
    add_para(doc, "实训场景采用双人工位组织方式，每组配套计算机、鸿蒙物联网实验箱和必要通信/传感/执行模块，便于在同一课堂内完成“设备安装、代码开发、数据上报、平台展示、报告提交”的完整任务。")
    if "space" in user_images:
        add_image(doc, user_images["space"], "图 3-1 50人分组实训场景效果图")
    add_image(doc, charts["space"], "图 3-2 120㎡空间布局示意图")
    add_image(doc, charts["network"], "图 3-3 实训室网络拓扑图")
    add_table(
        doc,
        "表 3-1 基础设施建设内容",
        ["建设项", "建设内容", "建设要求"],
        [
            ["教师演示区", "讲台、教师终端、大屏或投影、实物展示、平台管理入口", "满足教师发布任务、展示数据大屏、演示代码运行和组织答辩。"],
            ["学生工位区", "25组双人工位、电脑端、实验箱、线材、低压电源、网口/无线覆盖", "支持50人同步实训，布线清晰，设备取放安全。"],
            ["网络服务器区", "千兆交换机、路由器/防火墙、平台服务器或边缘计算设备、备份存储", "支持真实设备、仿真设备和平台稳定接入。"],
            ["收纳维护区", "实验箱收纳、模块分类、耗材备件、充电维护、资产编号", "建立设备借还、损坏登记和备件替换流程。"],
            ["文化展示区", "专业建设展示、鸿蒙生态介绍、优秀作品、课程路径、实训安全规范", "体现物联网专业特色和实训室育人氛围。"],
        ],
        [3.0, 6.6, 6.0],
    )
    add_heading(doc, "3.2 技术架构与设备体系", 2)
    add_para(doc, "核心设备以鸿蒙物联网实验箱为学生端实训终端，配套主控、通信、传感器、RFID识别、执行器开发包和AIoT教学云平台。设备体系覆盖嵌入式控制、鸿蒙开发、ZigBee无线传感网、中心网关多协议接入、环境采集、开关量输出和本地触控可视化等实验。")
    add_image(doc, charts["architecture"], "图 3-4 端-边-云总体架构图")
    if "product" in user_images:
        add_image(doc, user_images["product"], "图 3-5 鸿蒙物联网实验箱产品示意图", width=12.8)
        doc.add_page_break()
    add_table(
        doc,
        "表 3-2 核心设备配置表",
        ["类别", "主要配置", "教学用途"],
        [
            ["鸿蒙物联网实验箱", "通用实验平台本体、平台标配模块、电源适配器、USB连接线、配套线材和教学资源", "作为学生核心实训终端，承载鸿蒙、STM32、ZigBee、网关、传感与执行闭环实验。"],
            ["主控开发包", "STM32扩展控制模块、鸿蒙扩展模块、ZigBee单片机模块、接口拓展板", "用于多控制器协同、接口扩展、程序下载、串口调试和创新项目开发。"],
            ["通信开发包", "WiFi、ZigBee、4G、BLE、LoRa、NB-IoT/Cat.1等通信模块", "支撑局域网、广域网、低功耗和远距采集等多制式通信训练。"],
            ["传感器开发包", "霍尔、人体红外、气体、红外对射、重量、火焰、PM2.5、土壤湿度、超声波、CO2等模块", "覆盖安防、环境、农业、仓储、气体和距离等典型感知对象。"],
            ["RFID识别开发包", "HF高频、UHF超高频、NFC、条码/二维码识别模块", "面向门禁、考勤、资产盘点、仓储出入库和标签数据追踪实验。"],
            ["执行器开发包", "舵机、风扇、步进电机、LED灯光、水泵、喷雾器、窗帘执行器等", "支撑智能家居、灌溉、通风、告警和机械动作联动实验。"],
            ["AIoT教学云平台", "B/S架构、设备接入、数据可视化、规则联动、硬件在线仿真、教学评价和AI辅助", "作为教学组织中心、设备管理中心和数据沉淀中心。"],
        ],
        [3.0, 6.0, 6.6],
    )
    add_heading(doc, "3.3 实验箱模块化设计", 2)
    add_para(doc, "鸿蒙物联网实验箱采用便携式一体化结构，面向课堂高频使用和学生分组训练设计。平台支持模块化快接、层叠互联和免跳线实验教学，能够将主控、通信、传感、识别、执行、网关和平台任务组合到同一教学载体中。")
    add_para(doc, "实验箱设置9个同构通用模块接口位，各接口位的电气接口、供电触点和通信触点保持一致。教学组织时，平台模块序号仅用于说明产品构成和课程安排，不代表固定硬件绑定关系，学生可根据实验任务将不同模块放置到任意接口位开展验证和联调。")
    add_image(doc, charts["modular_box"], "图 3-6 模块化实验箱结构与教学组合图")
    add_table(
        doc,
        "表 3-3 实验箱模块化教学组合表",
        ["模块类型", "典型模块", "组合方式", "支撑课程", "典型任务"],
        [
            ["主控模块", "STM32扩展、OpenHarmony扩展、ZigBee单片机、接口拓展板", "插入任意同构接口位，与传感、通信或执行模块组合", "C语言程序设计、单片机应用技术、嵌入式系统开发、OpenHarmony轻量设备开发", "GPIO控制、串口调试、I2C/SPI外设读取、PWM输出、鸿蒙节点控制。"],
            ["通信模块", "WiFi、ZigBee、4G、BLE、LoRa、NB-IoT/Cat.1", "与主控模块、网关模块或云平台组合，形成局域、广域和低功耗通信链路", "无线传感网络技术、物联网通信技术、物联网设备安装与调试", "WiFi配网、MQTT/HTTP上报、ZigBee组网、LoRa远距采集、4G远程接入。"],
            ["传感模块", "霍尔、人体红外、气体、红外对射、重量、火焰、PM2.5、土壤湿度、超声波、CO2", "与主控、通信、平台组合，完成采集、判断、上传和曲线展示", "传感器与检测技术、数据采集与可视化、智慧农业综合实训", "环境检测、安防触发、距离测量、重量采集、土壤湿度监测、CO2通风联动。"],
            ["识别模块", "HF、UHF、NFC、条码/二维码识别", "与主控、网关、平台组合，形成身份识别和资产流转实验", "RFID技术与应用、物联网系统集成与项目实践", "门禁验证、资产盘点、仓储出入库、设备编号扫码、标签数据读取。"],
            ["执行模块", "舵机、风扇、步进电机、LED、水泵、喷雾器、窗帘执行器", "与传感阈值、规则引擎和平台命令组合，完成感知到控制闭环", "电工电子技术、单片机应用技术、AIoT云平台应用、系统集成实践", "风扇调速、水泵灌溉、灯光报警、喷雾补湿、窗帘联动、舵机门锁模拟。"],
            ["网关与平台", "中心网关、七寸触控终端、AIoT教学云平台、虚实结合仿真", "汇聚多协议设备，完成本地看板、云端建模、规则联动和教学评价", "边缘网关与协议转换、AIoT云平台应用、数据采集与可视化", "Modbus接入、ZigBee数据汇聚、MQTT主题设计、平台大屏、在线仿真、实验评价。"],
            ["行业套件", "智慧农业、智慧养殖、微生物发酵环境监测等场景", "按项目任务选择传感、通信、执行、网关和平台模块组合", "物联网系统集成与项目实践、毕业设计、岗位综合实训", "需求分析、系统搭建、联调记录、项目演示、报告撰写和答辩展示。"],
        ],
        [2.2, 3.6, 3.4, 3.2, 3.2],
        font_size=9,
    )
    add_heading(doc, "3.4 高职物联网专业课程支撑体系", 2)
    add_para(doc, "实训室面向高职物联网应用技术专业课程群建设，覆盖专业基础课、专业核心课、综合实践课和岗位综合训练。课程支撑不是把设备简单分配给单门课程，而是围绕“端侧开发、网络通信、边缘网关、云平台应用、行业项目交付”构建递进式训练链路。")
    add_para(doc, "在教学实施中，教师可根据课程进度选择不同模块组合：低年级侧重认识物联网系统、电工电子基础和C语言程序控制；中年级侧重传感检测、嵌入式开发、RFID、无线组网和通信协议；高年级侧重边缘网关、AIoT平台、系统集成、行业场景项目和岗位综合实训。")
    add_image(doc, charts["course_support"], "图 3-7 高职物联网专业课程支撑矩阵图")
    add_table(
        doc,
        "表 3-4 高职物联网专业课程支撑矩阵表",
        ["课程名称", "对应模块/平台", "可开展实验项目", "支撑能力目标"],
        [
            ["物联网技术基础", "鸿蒙物联网实验箱、端-边-云架构图、AIoT教学云平台", "认识感知层、网络层、平台层和应用层；观察真实设备上报、平台看板和命令下发流程。", "建立物联网系统整体认知，理解典型岗位任务和工程链路。"],
            ["电工电子技术", "电源接口、GPIO、继电器、LED、风扇、ADC/PWM相关模块", "低压供电识别、开关量输入输出、模拟量采集、PWM调光调速、继电器控制。", "掌握基础电路、安全用电、信号类型和执行控制基础。"],
            ["C语言程序设计/嵌入式C", "STM32主控、OpenHarmony扩展、接口拓展板、串口调试工具", "GPIO控制、按键采集、UART通信、数组与结构体存储传感数据、函数封装和任务流程控制。", "形成面向硬件控制的程序设计、调试和模块化编码能力。"],
            ["单片机应用技术/嵌入式系统开发", "STM32、ZigBee单片机、传感器包、执行器包", "I2C/SPI外设读取、ADC采集、定时器与中断、PWM输出、传感器驱动和执行器联动。", "掌握MCU外设驱动、硬件调试、接口扩展和故障定位能力。"],
            ["传感器与检测技术", "温湿度、光照、CO2、气体、PM2.5、土壤湿度、重量、超声波、人体红外等模块", "环境参数采集、气体异常检测、距离测量、重量检测、人体入侵触发、土壤湿度监测。", "理解传感器选型、信号采集、阈值判断、数据校准和检测应用。"],
            ["RFID技术与应用", "HF、UHF、NFC、条码/二维码识别模块、AIoT平台", "门禁刷卡、NFC标签读取、UHF资产盘点、二维码设备编号、仓储出入库模拟。", "掌握近场识别、标签数据读写、资产追踪和业务流程建模能力。"],
            ["无线传感网络技术/ZigBee技术应用", "ZigBee协调器、路由器、终端节点、中心网关", "节点入网退网、点播/组播/广播、网络自愈、多节点数据汇聚、低功耗终端实验。", "掌握无线传感网络拓扑、节点角色、数据汇聚和组网调试能力。"],
            ["物联网通信技术", "WiFi、4G、BLE、LoRa、NB-IoT/Cat.1、MQTT/HTTP/TCP/UDP/Modbus", "WiFi配网、4G远程接入、BLE近场交互、LoRa远距采集、NB-IoT数据上传、MQTT主题设计。", "理解多制式通信适用场景，完成联网配置、协议接入和数据传输。"],
            ["物联网设备安装与调试", "模块化实验箱、25组工位、线材、电源、网络设备、资产编号体系", "模块安装与拆换、接口识别、串口枚举、网络配置、设备清点、常见故障排查。", "训练现场安装、规范接线、联调记录、安全操作和交付文档能力。"],
            ["边缘网关与协议转换", "中心网关、七寸触控终端、ZigBee节点、Modbus设备、MQTT/HTTP接口", "ZigBee数据汇聚、Modbus RTU/TCP接入、协议转换、本地缓存、断网续传、本地触控看板。", "掌握边缘侧设备汇聚、协议适配、规则判定和现场运维能力。"],
            ["AIoT云平台应用", "AIoT教学云平台、真实设备、虚拟设备、规则引擎、教学评价模块", "产品建模、设备密钥配置、属性上报、事件记录、命令下发、报警规则、班级任务与报告批阅。", "掌握平台配置、设备管理、数据可视化、规则联动和教学过程评价。"],
            ["数据采集与可视化", "传感器包、边缘网关、AIoT平台、项目大屏、历史数据查询", "实时数据卡片、历史曲线、仪表盘、报警列表、项目大屏、实验数据导出与报告分析。", "形成数据采集、清洗观察、图表表达和工程报告撰写能力。"],
            ["OpenHarmony轻量设备开发", "OpenHarmony扩展模块、DevEco工程环境、WiFi模块、云平台接口", "工程创建、GPIO控制、WiFi配网、设备状态上传、云端命令下发、鸿蒙节点联动。", "掌握鸿蒙轻量设备开发流程、外设调用、联网接入和生态创新基础。"],
            ["物联网系统集成与项目实践", "实验箱、主控/通信/传感/RFID/执行器包、网关、AIoT平台、行业套件", "智能家居、智慧农业、智慧养殖、微生物发酵环境监测、资产识别等综合项目。", "完成需求分析、方案设计、软硬件联调、平台配置、项目展示和验收交付。"],
            ["毕业设计/岗位综合实训", "完整实训室环境、开放接口、课程资源、项目模板、报告与答辩模板", "自选行业场景方案设计、原型搭建、故障排查、数据看板、成果演示和毕业答辩。", "提升岗位迁移、综合应用、团队协作、文档表达和持续迭代能力。"],
        ],
        [3.0, 4.1, 4.5, 4.0],
        font_size=8.5,
    )
    add_heading(doc, "3.5 课程资源与典型实验项目", 2)
    add_image(doc, charts["course"], "图 3-8 课程体系设计图")
    add_table(
        doc,
        "表 3-5 课程资源配置建议",
        ["课程模块", "实验/项目内容", "资源形态"],
        [
            ["MCU与鸿蒙基础", "GPIO、UART、I2C、SPI、ADC、PWM、OpenHarmony工程创建、WiFi配网", "实验指导书、源码、接线图、运行截图、报告模板。"],
            ["无线通信与组网", "ZigBee协调器与节点、WiFi联网、BLE近场、4G/LoRa/NB-IoT数据上传", "通信案例、AT指令样例、平台配置截图、故障排查清单。"],
            ["传感与执行控制", "温湿度、光照、CO2、气体、重量、RFID、风扇、水泵、舵机、灯光联动", "模块说明、实验步骤、阈值策略、数据曲线样例。"],
            ["AIoT平台应用", "设备建模、MQTT主题、HTTP接口、Modbus接入、规则联动、报警记录、大屏配置", "平台账号、任务模板、评价量规、实验数据导出模板。"],
            ["行业综合项目", "智慧农业、智慧养殖、微生物发酵环境监测、智能家居、资产识别", "项目任务书、答辩PPT模板、验收清单、优秀作品库。"],
        ],
        [3.0, 7.0, 5.6],
    )
    doc.add_page_break()
    add_table(
        doc,
        "表 3-6 典型实验项目表",
        ["项目类别", "项目名称", "项目成果"],
        [
            ["基础实验", "GPIO控制、串口通信、I2C传感器读取、PWM调光、ADC采集", "完成单点驱动、代码运行截图和实验报告。"],
            ["联网实验", "WiFi配网、MQTT数据上报、HTTP接口调用、Modbus设备接入", "完成设备建模、数据上报和命令下发。"],
            ["组网实验", "ZigBee协调器配置、节点入网退网、数据透传、多节点汇聚", "完成无线传感网络拓扑和网关汇聚。"],
            ["平台实验", "规则引擎、报警记录、历史曲线、项目大屏、虚拟设备接入", "完成可视化看板和平台配置文档。"],
            ["行业项目", "发酵监测、智慧养殖、智慧农业、智能家居、资产识别", "完成需求分析、系统联调、项目展示和答辩。"],
        ],
        [3.0, 6.8, 5.8],
    )

    add_heading(doc, "四、拓展应用与场景延伸", 1)
    add_heading(doc, "4.1 行业场景应用展示", 2)
    add_para(doc, "在基础课程和综合实训之外，实训室可围绕智慧农业、智能家居、工业物联网、环境监测等方向组织项目化教学。学生通过实验箱、网关和AIoT平台完成数据采集、规则联动、平台可视化和应用场景解释，形成从设备到应用的完整表达能力。")
    if "scenario" in user_images:
        add_image(doc, user_images["scenario"], "图 4-1 AIoT平台与行业应用展示效果图")
    add_heading(doc, "4.2 前瞻发展方向", 2)
    add_para(doc, "产教融合、科研创新和竞赛训练在本方案中定位为实训室建成后的拓展方向。通过平台接口、课程资源、项目案例和数据沉淀，后续可逐步支撑校企协同课程共建、教师教学研究、学生创新项目和行业应用样板。")
    add_image(doc, charts["extension"], "图 4-2 拓展应用与前瞻发展路线图")
    add_table(
        doc,
        "表 4-1 前瞻拓展方向",
        ["拓展方向", "建议开展方式", "预期价值"],
        [
            ["校企协同课程共建", "围绕鸿蒙设备接入、AIoT平台运维、行业场景实施等模块共建任务书和案例库。", "使课程内容更贴近工程链路，便于后续持续更新。"],
            ["教师培训与教研", "开展平台使用、设备联调、课程组织、项目评价和故障排查培训。", "帮助教师独立完成任务发布、过程监控、报告批阅和课程迭代。"],
            ["学生创新项目", "基于开放接口和模块化设备，支持学生围绕智慧农业、养殖、发酵等场景进行项目开发。", "形成作品库、报告库、数据样例和展示素材。"],
            ["行业场景拓展", "将环境监测、资产识别、仓储管理、智能家居等场景作为后续扩展包。", "提升实训室跨专业服务和持续建设空间。"],
        ],
        [3.0, 7.2, 5.4],
    )

    add_heading(doc, "五、竞赛训练与能力提升展望", 1)
    add_para(doc, "实训室建成后可面向物联网应用开发、嵌入式系统、人工智能应用、鸿蒙生态创新等方向预留训练条件。本方案不写现有参赛案例、相关成绩或合作院校案例，仅从课程和设备能力角度说明后续可支撑的训练方向。")
    add_image(doc, charts["competition"], "图 5-1 竞赛训练能力映射图")
    add_table(
        doc,
        "表 5-1 竞赛训练能力映射表",
        ["能力方向", "可训练内容", "实训室支撑条件"],
        [
            ["物联网应用开发", "设备接入、协议通信、规则联动、数据看板、项目交付", "鸿蒙实验箱、通信开发包、AIoT平台、行业场景项目。"],
            ["嵌入式系统", "MCU外设、传感器驱动、串口调试、执行器控制、故障排查", "STM32模块、传感器开发包、执行器开发包、实验指导书。"],
            ["人工智能应用", "AI助教、语音/视觉接口、边缘AI探索、知识库答疑", "AIoT平台高级能力、开放接口、项目化课程资源。"],
            ["鸿蒙生态创新", "OpenHarmony轻量设备开发、WiFi配网、设备上云、鸿蒙节点联动", "鸿蒙物联网模块、DevEco开发流程、云平台联动案例。"],
            ["工程表达与答辩", "需求分析、方案设计、联调记录、演示汇报、报告撰写", "任务书、报告模板、过程数据、项目大屏和优秀作品库。"],
        ],
        [3.0, 6.2, 6.4],
    )

    doc.add_page_break()
    add_heading(doc, "六、实施计划、培训运维与验收", 1)
    add_heading(doc, "6.1 建设内容清单", 2)
    add_table(
        doc,
        "表 6-1 建设内容清单",
        ["序号", "建设内容", "建议配置", "用途说明"],
        [
            ["1", "鸿蒙物联网实验箱", "按25组学生工位配置，配套电源、线材和基础教学资源", "学生核心实训终端，支撑鸿蒙、STM32、ZigBee、网关、传感与执行实验。"],
            ["2", "主控/通信/传感/RFID/执行器开发包", "按课程需要配置到工位或共享轮换使用", "扩展实验边界，支撑多协议通信、资产识别、环境检测和执行联动。"],
            ["3", "AIoT教学云平台", "支持管理员、教师、学生角色", "承担设备接入、任务发布、数据可视化、规则联动、在线仿真和评价。"],
            ["4", "课程教学包", "基础实验、综合项目、行业项目、源码、指导书、报告模板", "支撑课堂教学、课后练习、项目答辩和持续资源建设。"],
            ["5", "行业场景系统", "智慧农业、智慧养殖、微生物发酵环境监测等按建设需求深化", "增强跨专业应用能力和项目化教学吸引力。"],
            ["6", "网络与基础设施", "交换机、路由器、服务器/边缘计算、收纳柜、教师演示设备", "保障平台稳定运行、设备接入和课堂展示。"],
            ["7", "文化建设与装修改造", "文化墙、安全标识、布线整理、局部环境优化", "营造物联网专业氛围并保障实训安全。"],
        ],
        [1.4, 4.0, 5.0, 5.2],
    )
    add_heading(doc, "6.2 实施进度", 2)
    add_image(doc, charts["implementation"], "图 6-1 实施进度流程图")
    add_table(
        doc,
        "表 6-2 实施阶段与交付物",
        ["阶段", "工作项", "主要内容", "交付物"],
        [
            ["第1阶段", "现场确认与深化设计", "确认教室尺寸、电力网络、工位数量、服务器部署方式和课程优先级。", "深化设计与施工准备清单。"],
            ["第2阶段", "设备到货与环境部署", "完成实验箱、开发包、平台服务器、网络设备、演示设备和收纳设施到位。", "设备清点、资产编号、网络与供电可用。"],
            ["第3阶段", "平台安装与联调", "部署AIoT平台，接入真实设备和虚拟设备，配置账号、课程、项目和数据大屏。", "平台可登录、设备可接入、数据可采集。"],
            ["第4阶段", "课程资源导入与教师培训", "导入实验指导书、源码、PPT、任务模板和报告模板，完成教师实操培训。", "教师可独立发布任务、查看过程、批阅报告。"],
            ["第5阶段", "试运行与验收", "按基础实验、综合项目、行业场景开展试运行，完成问题整改和验收归档。", "验收报告、运维制度和交付资料。"],
        ],
        [2.0, 3.0, 6.3, 4.3],
    )
    add_heading(doc, "6.3 培训与运维保障", 2)
    add_para(doc, "实训室建成后应建立设备资产台账、课程资源版本、平台账号权限、数据备份、耗材备件、故障工单和教师共建等运维机制。建议每学期开展一次课程资源复盘和设备健康巡检，每年围绕新通信模块、新行业案例或新AI能力进行一次迭代升级。")
    add_table(
        doc,
        "表 6-3 验收指标表",
        ["验收类别", "验收指标", "建议验收方式"],
        [
            ["容量验收", "25组工位可支撑50人同步开展物联网实训。", "现场抽查工位、设备和网络接入情况。"],
            ["设备验收", "实验箱、开发包、平台服务器、网络设备、场景系统和耗材与清单一致。", "清点资产编号、规格参数和配套资料。"],
            ["平台验收", "教师、学生、管理员角色可用，真实设备和虚拟设备可接入，数据可展示。", "登录平台并完成设备建模、数据上报和命令下发。"],
            ["课程验收", "基础实验、综合项目、行业项目、源码、指导书、报告模板可使用。", "抽取样例实验完成全流程演示。"],
            ["资料验收", "深化方案、设备清单、网络拓扑、空间布局、培训记录、运维制度完整。", "查验交付文档和培训签到/记录。"],
            ["安全验收", "强弱电布置、安全标识、设备收纳、低压供电和账号权限满足实训管理要求。", "现场检查和试运行记录确认。"],
        ],
        [3.0, 6.8, 5.7],
    )

    add_heading(doc, "七、资料依据", 1)
    add_para(doc, "本方案综合使用本地项目资料和公开资料进行论证。公开资料仅用于政策、专业建设和技术趋势说明，不作为供应商业绩、合作案例或奖项证明。")
    add_table(
        doc,
        "表 7-1 主要资料来源",
        ["资料类别", "来源", "用途"],
        [
            ["本地公告", "《信息工程学院.txt》", "确定项目名称、面积约120㎡、50人容量和建设内容。"],
            ["结构底稿", "《传感器与检测技术创新实验室解决方案(1)(1).docx》", "借鉴项目概述、总体设计、建设内容、拓展应用、建设清单等章节结构。"],
            ["设备参数", "《新南云鸿蒙物联网实验箱技术参数-7.2.xlsx》", "提炼鸿蒙实验箱、开发包、AIoT平台、课程包和行业系统配置。"],
            ["专业标准", "教育部职业教育专业教学标准（2025年修（制）订）公开资料", "用于对齐高职物联网专业课程建设、实践教学和人才培养要求。"],
            ["专业简介", "教育部职业教育专业简介（2022年修订）公开资料", "用于补充专业面向、课程群和岗位能力表述。"],
            ["技术资料", "OpenHarmony官方文档、HarmonyOS设备开发官网", "支撑OpenHarmony轻量设备开发、设备接入和鸿蒙物联网技术路线。"],
        ],
        [3.0, 7.0, 5.5],
    )
    if logo:
        add_para(doc, "", first=False, after=24)
        lp = doc.add_paragraph()
        lp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        lp.add_run().add_picture(str(logo), width=Cm(8.0))
    return doc


def main():
    charts = render_frontend_charts()
    doc = build_doc(charts)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
